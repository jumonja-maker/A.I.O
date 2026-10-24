"""Opt-in bounded structural evidence; never route admission or visual fidelity.

Native files remain available when evidence is missing. No renderer, network,
source edits, formula evaluation or external relationship resolution is used.
"""
from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
import re
import struct
import time
from xml.etree import ElementTree as ET
import zipfile

BUDGET_BYTES = 64*1024*1024
BUDGET_SECONDS = 8
WORD = re.compile(r'\S+')


class _Limited(Exception):
    pass


class _Budget:
    def __init__(self, cancel, seconds, max_bytes):
        self.cancel, self.until, self.left = cancel, time.monotonic()+seconds, max_bytes
    def check(self, count=0):
        self.left -= count
        if self.left < 0 or time.monotonic()>self.until:
            raise _Limited('inspection_budget')
        if self.cancel is not None and self.cancel.is_set():
            raise _Limited('cancelled')
    def wrap(self, stream):
        budget = self
        class Reader:
            def read(self, size=-1):
                budget.check()
                block = stream.read(min(65536, size) if size >= 0 else 65536)
                budget.check(len(block))
                # Reject entities/DTDs including declarations split by reads.
                # Inspection-only normalization catches UTF-16/32 declarations
                # as well as ASCII, including markers split across reads.
                test=(self.tail+block).replace(b'\0', b'').upper()
                self.tail=test[-32:]
                if b'<!DOCTYPE' in test or b'<!ENTITY' in test:
                    raise _Limited('xml_entities_not_inspected')
                return block
            tail=b''
        return Reader()


def _add_text(result, digest, text):
    for match in WORD.finditer(text):
        word = match.group()
        digest.update(word.encode('utf-8')+b'\0')
        result['counts']['words'] += 1


def _xml(stream, result, digest, budget, family):
    paragraph = {'p','h'} if family=='odf' else {'p'} if family=='ooxml' else {'p','h1','h2','h3','li','text'}
    depth = 0
    text_depth = 0
    parents = []
    for event, element in ET.iterparse(budget.wrap(stream),events=('start','end')):
        budget.check()
        local=element.tag.rsplit('}',1)[-1]
        if event=='start':
            parents.append(element)
            depth += 1
            if depth > 256:
                raise _Limited('xml_depth')
            if local in paragraph and not text_depth:
                text_depth=depth
            counts=result['counts']
            if local in ('tbl','table'):counts['tables']+=1
            if local in ('f','oMath','math') or any(k.rsplit('}',1)[-1]=='formula' for k in element.attrib):
                counts['formulas']+=1
            if local in ('chart',):counts['charts']+=1
            if local=='s' and family=='odf' and text_depth:
                count=element.get('{urn:oasis:names:tc:opendocument:xmlns:text:1.0}c','1')
                if not count.isdigit() or int(count)>4096:raise _Limited('expanded_spaces')
                element.text=' '*int(count)
            continue
        if text_depth==depth:
            text=''.join(element.itertext())
            if text.strip():
                result['counts']['paragraphs']+=1
                _add_text(result,digest,text)
            text_depth=0
        if not text_depth:
            element.clear()
            if len(parents)>1:
                parents[-2].remove(element)
        parents.pop()
        depth -= 1


def _xlsx_values(source, budget):
    from openpyxl import load_workbook
    budget.check()
    # read_only limits worksheet rows, but openpyxl eagerly retains shared
    # strings and styles. Account for these parts before invoking its parser;
    # a small compressed XLSX must not bypass the optional inspection budget.
    with zipfile.ZipFile(source) as archive:
        # OPC content types may legally name a nonstandard shared-strings part.
        # Read that declaration under the same budget before resolving names.
        declared_strings = set()
        if '[Content_Types].xml' in archive.namelist():
            with archive.open('[Content_Types].xml') as stream:
                for _, item in ET.iterparse(budget.wrap(stream), events=('end',)):
                    if item.get('ContentType') == ('application/vnd.openxmlformats-officedocument.'
                                                    'spreadsheetml.sharedStrings+xml'):
                        declared_strings.add(item.get('PartName', '')[1:])
                    item.clear()
        # In read-only mode worksheets stream, whereas styles, strings, theme,
        # properties and other package metadata may still be read eagerly.
        metadata = [item for item in archive.infolist()
                    if item.filename != '[Content_Types].xml'
                    and (item.filename in declared_strings or
                         item.filename.lower().endswith(('.xml', '.rels'))
                         and not item.filename.startswith('xl/worksheets/'))]
        if sum(item.file_size for item in metadata) > budget.left:
            raise _Limited('xlsx_metadata_budget')
        for item in metadata:
            with archive.open(item) as stream:
                reader = budget.wrap(stream)
                while reader.read(65536):
                    pass
    budget.check()
    book=load_workbook(source,read_only=True,data_only=False,keep_links=False)
    digest=hashlib.sha256()
    result={'rows':0,'columns':0,'valuesDigest':None,'sheets':len(book.sheetnames)}
    try:
        if len(book.worksheets)!=1:
            # A flat table digest is not a workbook-wide comparison.
            return result
        sheet=book.worksheets[0]
        for row in sheet.iter_rows():
            budget.check()
            result['rows']+=1
            values=[]
            for cell in row:
                value=cell.value
                values.append('' if value is None else str(value))
            result['columns']=max(result['columns'],len(values))
            data=json.dumps(values,ensure_ascii=True,separators=(',',':')).encode('ascii')+b'\n'
            budget.check(len(data));digest.update(data)
        result['valuesDigest']=digest.hexdigest()
        return result
    finally:
        book.close()


def inspect_document(source, *, cancel=None, seconds=BUDGET_SECONDS, max_bytes=BUDGET_BYTES):
    """Return facts and explicit unknowns; budgets limit inspection, not export."""
    source=Path(source)
    budget=_Budget(cancel,seconds,max_bytes)
    result={'schema':1,'status':'not_checked','complete':False,'kind':'unknown',
            'counts':dict.fromkeys(('words','paragraphs','tables','images','formulas','charts'),0),
            'textDigest':None,'unknown':['visual_layout','fonts','rendered_pixels','formula_results']}
    digest=hashlib.sha256()
    try:
        with source.open('rb') as stream:
            magic=stream.read(32)
        if magic.startswith(b'PK\x03\x04'):
            with zipfile.ZipFile(source) as archive:
                infos=archive.infolist()
                if len(infos)>10000:raise _Limited('package_entry_budget')
                names=[i.filename for i in infos]
                if len(set(names))!=len(names):raise _Limited('duplicate_package_entries')
                odf='mimetype' in names and 'content.xml' in names
                if odf:
                    mime=archive.getinfo('mimetype')
                    if mime.file_size>256:raise _Limited('unknown_package_type')
                    kind=archive.read(mime).decode('ascii')
                    if 'opendocument' not in kind:raise _Limited('unknown_package_type')
                    result['kind']='odf'
                    result['documentClass']=kind.rsplit('.',1)[-1]
                    documents=['content.xml']
                    images=[n for n in names if n.startswith('Pictures/') and not n.endswith('/')]
                    result['counts']['embeddedObjects']=sum(n.endswith('/content.xml') for n in names)
                elif 'word/document.xml' in names:
                    result['kind']='ooxml';result['documentClass']='text'
                    documents=['word/document.xml']
                    images=[n for n in names if n.startswith('word/media/') and not n.endswith('/')]
                elif 'xl/workbook.xml' in names:
                    result['kind']='ooxml';result['documentClass']='spreadsheet'
                    documents=sorted(n for n in names if re.fullmatch(r'xl/worksheets/[^/]+\.xml',n))
                    result['counts']['sheets']=len(documents)
                    images=[n for n in names if n.startswith('xl/media/') and not n.endswith('/')]
                    result['unknown'].append('spreadsheet_values')
                elif 'ppt/presentation.xml' in names:
                    result['kind']='ooxml';result['documentClass']='presentation'
                    documents=sorted(n for n in names if re.fullmatch(r'ppt/slides/slide[0-9]+\.xml',n))
                    result['counts']['slides']=len(documents)
                    images=[n for n in names if n.startswith('ppt/media/') and not n.endswith('/')]
                elif 'META-INF/container.xml' in names:
                    result['kind']='epub';result['documentClass']='text'
                    documents=sorted(n for n in names if n.lower().endswith(('.xhtml','.html','.htm')))
                    images=[n for n in names if n.lower().endswith(('.png','.jpg','.jpeg','.gif','.svg','.webp'))]
                    result['counts']['chapters']=len(documents)
                    result['unknown'].append('epub_reading_order')
                else:
                    raise _Limited('unsupported_package')
                result['counts']['images']=len(images)
                for name in documents:
                    budget.check()
                    with archive.open(name) as stream:
                        _xml(stream,result,digest,budget,result['kind'])
                if result.get('documentClass')!='spreadsheet':
                    result['textDigest']=digest.hexdigest()
                if result.get('documentClass')=='spreadsheet' and result['kind']=='ooxml':
                    result['cellValues']=_xlsx_values(source,budget)
                    if result['cellValues'].get('valuesDigest'):
                        result['counts'].update(rows=result['cellValues']['rows'],columns=result['cellValues']['columns'])
                        result['unknown'].remove('spreadsheet_values')
                result.update(status='checked_subset',complete=True)
        elif magic.startswith(b'%PDF-'):
            result['kind']='pdf';result['documentClass']='rendered'
            if source.stat().st_size>max_bytes:raise _Limited('inspection_budget')
            # Optional small reader; missing support is explicit. PDF text
            # extraction discrepancies are evidence, never proof of lost ink.
            try:
                from pypdf import PdfReader
            except ImportError:
                raise _Limited('pdf_reader_unavailable')
            reader=PdfReader(str(source),strict=False)
            try:
                if reader.is_encrypted:raise _Limited('encrypted_pdf')
                result['counts']['pages']=len(reader.pages)
                # Page counts do not prove table/formula/image preservation.
                result['counts']={k:v for k,v in result['counts'].items() if k=='pages'}
                result.update(status='checked_subset',complete=True)
                result['unknown'].extend(['text','tables','images','formulas'])
            finally:
                reader.close()
        elif magic.startswith(b'BM') and len(magic)>=26 and struct.unpack('<I',magic[14:18])[0]>=40:
            result.update(kind='raster',documentClass='image',complete=True,status='checked_subset')
            result['counts']={'images':1,'width':abs(struct.unpack('<i',magic[18:22])[0]),
                              'height':abs(struct.unpack('<i',magic[22:26])[0])}
        elif magic.startswith(b'\x89PNG\r\n\x1a\n') and len(magic)>=24:
            result.update(kind='raster',documentClass='image',complete=True,status='checked_subset')
            result['counts']={'images':1,'width':struct.unpack('>I',magic[16:20])[0],
                              'height':struct.unpack('>I',magic[20:24])[0]}
        elif source.suffix.lower()=='.svg':
            result.update(kind='svg',documentClass='image')
            with source.open('rb') as stream:_xml(stream,result,digest,budget,'svg')
            result.update(status='checked_subset',complete=True,textDigest=digest.hexdigest())
        elif source.suffix.lower() in ('.csv','.tsv','.tab'):
            from .csv_import import inspect_csv_stream, DELIMITERS
            # Choose from a small prefix, then inspect the complete stream
            # once. Never rescan a large file for all four possible delimiters.
            with source.open('r',encoding='utf-8-sig',newline='') as stream:
                sample=stream.read(65536)
            choices=[]
            import io
            for delimiter in DELIMITERS:
                try:
                    rows=[]
                    for row in csv.reader(io.StringIO(sample),delimiter=delimiter,strict=True):
                        if row:rows.append(row)
                        if len(rows)>=8:break
                    if len(rows)>=2 and len(rows[0])>1 and all(len(r)==len(rows[0]) for r in rows):
                        choices.append(delimiter)
                except csv.Error:pass
            if len(choices)!=1:raise _Limited('csv_dialect_undetermined')
            value=inspect_csv_stream(source,delimiter=choices[0],cancel=cancel,
                seconds=max(.01,budget.until-time.monotonic()),max_bytes=max_bytes)
            if not value.get('complete'):raise _Limited(value.get('reason','csv_incomplete'))
            result.update(kind='csv',documentClass='spreadsheet',complete=True,status='checked_subset',csv=value)
            result['counts']={'rows':value['rows'],'columns':value['columns']}
            result['cellValues']={k:value[k] for k in ('rows','columns','valuesDigest')}
        else:
            result['reason']='unsupported_format'
    except _Limited as error:
        result.update(status='not_checked',complete=False,reason=str(error))
    except Exception:
        result.update(status='not_checked',complete=False,reason='inspection_failed')
    return result


def compare_documents(before, after):
    evidence={'schema':1,'status':'not_checked','qualityCompared':False,
              'scope':'structural_subset','before':before,'after':after,
              'risks':[],'formatLimitations':[],
              'unknown':['visual_layout','fonts','rendered_pixels','formula_results']}
    if not before.get('complete') or not after.get('complete'):
        return evidence
    risks=evidence['risks']
    same_class=before.get('documentClass')==after.get('documentClass')
    if same_class:
        for feature in ('tables','images','formulas','slides','sheets','pages','width','height','rows','columns'):
            a,b=before['counts'].get(feature),after['counts'].get(feature)
            if a is not None and b is not None and b<a:
                risks.append({'code':'STRUCTURE_COUNT_DECREASED','feature':feature,'before':a,'after':b,
                              'certainty':'potential_loss'})
        if before.get('textDigest') and after.get('textDigest') and before['textDigest']!=after['textDigest']:
            risks.append({'code':'TEXT_DIFFERS','feature':'text','certainty':'requires_review'})
        if (before.get('cellValues',{}).get('valuesDigest') and after.get('cellValues',{}).get('valuesDigest')
                and before['cellValues']['valuesDigest']!=after['cellValues']['valuesDigest']):
            risks.append({'code':'CELL_VALUES_DIFFER','feature':'cells','certainty':'requires_review'})
    if before['counts'].get('formulas',0) and after.get('kind') in ('pdf','raster','svg','epub','csv'):
        evidence['formatLimitations'].append('FORMULAS_NOT_EDITABLE_IN_TARGET')
    if before.get('documentClass')=='spreadsheet' and after.get('kind')=='csv':
        evidence['formatLimitations'].append('CSV_HAS_ONE_TABLE_AND_NO_WORKBOOK_STYLES')
    comparable=same_class and (bool(before.get('textDigest') and after.get('textDigest')) or
        bool(set(before['counts']).intersection(after['counts'])-{'words','paragraphs'}))
    if comparable:
        evidence.update(status='risks_found' if risks else 'checked_subset',qualityCompared=True)
    else:
        evidence['unknown'].append('no_comparable_structure')
    return evidence


def inspect_isolated(source, workspace, *, cancel=None, seconds=BUDGET_SECONDS):
    """Keep XML/PDF parsers off the host and enforce a killable deadline."""
    import os
    import subprocess
    import sys
    from ..native_process import ProcessTree
    workspace=Path(workspace)
    workspace.mkdir(parents=True,exist_ok=False)
    request=workspace/'quality-request.json'
    request.write_text(json.dumps({'source':str(Path(source).resolve()),'seconds':seconds}),encoding='utf-8')
    command=([sys.executable,'--aio-guard-quality',str(request)] if getattr(sys,'frozen',False) else
             [sys.executable,'-m','app_backend.converter_guard.quality_worker',str(request)])
    unknown={'schema':1,'status':'not_checked','complete':False,'kind':'unknown','counts':{},
             'reason':'inspection_worker_unavailable'}
    until=time.monotonic()+seconds+2
    try:
        with ProcessTree(command,cwd=Path(__file__).resolve().parents[2]) as tree:
            while tree.process.poll() is None:
                if cancel is not None and cancel.is_set():
                    return {**unknown,'reason':'cancelled'}
                if time.monotonic()>=until:
                    return {**unknown,'reason':'inspection_deadline'}
                try:tree.process.wait(timeout=.05)
                except subprocess.TimeoutExpired:pass
            output=workspace/'quality-result.json'
            if tree.process.returncode or not output.is_file() or output.stat().st_size>65536:
                return unknown
            result=json.loads(output.read_text(encoding='ascii'))
            return result if isinstance(result,dict) and result.get('schema')==1 else unknown
    except (OSError,ValueError):
        return unknown
