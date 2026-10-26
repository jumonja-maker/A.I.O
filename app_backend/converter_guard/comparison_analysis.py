"""Local, conservative before/after evidence; no AI or conversion engine.

Public results contain bounded descriptions, not source paths or complete
document text. A checked subset never certifies visual/semantic fidelity.
The caller owns the file leases and the killable job-level deadline.
"""
from __future__ import annotations

import csv
from decimal import Decimal, InvalidOperation
import hashlib
import io
import json
from pathlib import Path
import posixpath
import re
import stat
import threading
import time
import unicodedata
from urllib.parse import unquote, urlsplit
import xml.etree.ElementTree as ET
import zipfile

from .errors import GuardError
from .media_fidelity import ALIASES, ANIMATED_TARGETS, RASTER, display_dimensions, display_page_dimensions, inspect_media

MAX_READ = 32 * 1024 * 1024
MAX_XML = 16 * 1024 * 1024
MAX_TEXT = 512000
MAX_CELLS = 50000
MAX_NODES = 200000
MAX_FINDINGS = 64
MAX_SECONDS = 12
MAX_PDF_PAGES = 256
OFFICE = 'urn:oasis:names:tc:opendocument:xmlns:office:1.0'
TABLE = 'urn:oasis:names:tc:opendocument:xmlns:table:1.0'
TEXT = 'urn:oasis:names:tc:opendocument:xmlns:text:1.0'
WORD = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
STRICT_WORD = 'http://purl.oclc.org/ooxml/wordprocessingml/main'
DRAWING = 'http://schemas.openxmlformats.org/drawingml/2006/main'
STRICT_DRAWING = 'http://purl.oclc.org/ooxml/drawingml/main'
PRESENTATION = 'http://schemas.openxmlformats.org/presentationml/2006/main'
STRICT_PRESENTATION = 'http://purl.oclc.org/ooxml/presentationml/main'
UNCHECKED = ('visual_layout', 'fonts', 'colour_accuracy', 'formula_evaluation', 'semantic_meaning')
_SKIP_TEXT_SUBTREES = {
    '{%s}%s' % (namespace, local)
    for namespace, locals_ in ((OFFICE, ('annotation',)), (TEXT, ('note', 'tracked-changes')),
        (WORD, ('del', 'moveFrom', 'instrText', 'delText')),
        (STRICT_WORD, ('del', 'moveFrom', 'instrText', 'delText')))
    for local in locals_
}


class _Limited(ValueError):
    pass


class _Budget:
    def __init__(self, cancel):
        self.cancel = cancel
        self.until = time.monotonic() + MAX_SECONDS
        self.used = 0

    def check(self, count=0):
        if self.cancel is not None and self.cancel.is_set():
            raise GuardError('CANCELLED', 'Анализ до/после отменён.', stage='comparison')
        self.used += count
        if self.used > MAX_READ or time.monotonic() > self.until:
            raise _Limited('inspection_budget')

    def read(self, stream, maximum=MAX_XML):
        data = bytearray()
        while True:
            self.check()
            block = stream.read(min(65536, maximum + 1 - len(data)))
            if not block:
                return bytes(data)
            self.check(len(block))
            data.extend(block)
            if len(data) > maximum:
                raise _Limited('inspection_byte_budget')


def _normal(value):
    return ' '.join(unicodedata.normalize('NFC', str(value)).split())


def _identity(path):
    path = Path(path)
    if path.is_symlink():
        raise _Limited('source_link_not_inspected')
    stamp = path.stat()
    if not stat.S_ISREG(stamp.st_mode):
        raise _Limited('non_regular_file_not_inspected')
    return (str(path.resolve(strict=True)), stamp.st_dev, stamp.st_ino,
            stamp.st_size, stamp.st_mtime_ns, stamp.st_ctime_ns)


def _xml(data, budget):
    declaration = data.replace(b'\0', b'').upper()
    if b'<!DOCTYPE' in declaration or b'<!ENTITY' in declaration:
        raise _Limited('xml_declarations_not_inspected')
    root = ET.fromstring(data)
    pending, count = [(root, 1)], 0
    while pending:
        item, depth = pending.pop()
        budget.check()
        count += 1
        if depth > 128 or count > MAX_NODES:
            raise _Limited('xml_structure_budget')
        pending.extend((child, depth + 1) for child in item)
    return root


def _element_text(item):
    if item.tag in _SKIP_TEXT_SUBTREES:
        return ''
    parts = [item.text or '']
    for child in item:
        local = child.tag.rsplit('}', 1)[-1]
        if child.tag == '{%s}s' % TEXT:
            count = int(child.get('{%s}c' % TEXT, '1'))
            if not 1 <= count <= 4096:
                raise _Limited('expanded_spaces_budget')
            parts.append(' ' * count)
        elif child.tag in ('{%s}tab' % TEXT, '{%s}tab' % WORD, '{%s}tab' % STRICT_WORD):
            parts.append('\t')
        elif child.tag in ('{%s}line-break' % TEXT, '{%s}br' % WORD, '{%s}br' % STRICT_WORD):
            parts.append('\n')
        elif child.tag in _SKIP_TEXT_SUBTREES:
            pass
        else:
            parts.append(_element_text(child))
        if local in ('p', 'h'):
            parts.append(' ')
        parts.append(child.tail or '')
    return ''.join(parts)


def _paragraphs(root):
    """Visible text stories only; never revisit skipped or nested paragraphs."""
    tags = {'{%s}%s' % (namespace, local) for namespace, locals_ in (
        (TEXT, ('p', 'h')), (WORD, ('p',)), (STRICT_WORD, ('p',)),
        (DRAWING, ('p',)), (STRICT_DRAWING, ('p',))) for local in locals_}
    pending = [root]
    while pending:
        item = pending.pop()
        if item.tag in _SKIP_TEXT_SUBTREES:
            continue
        if item.tag in tags:
            yield item
        else:
            pending.extend(reversed(list(item)))


def _cell_text(cell):
    return '\n'.join(_element_text(item) for item in _paragraphs(cell))


def _paragraph_text(root, budget):
    paragraphs = []
    size = 0
    for item in _paragraphs(root):
        budget.check()
        value = _normal(_element_text(item))
        size += len(value)
        if size > MAX_TEXT:
            raise _Limited('text_budget')
        if value:
            paragraphs.append(value)
    return '\n'.join(paragraphs)


def _value(value, kind='text', *, formula=None):
    # Cell text is data: leading/trailing/repeated spaces must survive. Only
    # document-layout text is whitespace-normalized by its own extractor.
    return {'value': unicodedata.normalize('NFC', str(value)), 'type': kind,
            **({'formula': formula} if formula else {})}


def _trim(rows):
    result = []
    for row in rows:
        row = list(row)
        while row and not row[-1]['value'] and not row[-1].get('formula'):
            row.pop()
        result.append(row)
    while result and not result[-1]:
        result.pop()
    return result


def _decode_text(data, kind):
    # A BOM identifies Unicode unambiguously. Never guess legacy codepages or
    # silently replace characters: replacement would erase comparison evidence.
    encoding = ('utf-32' if data.startswith((b'\xff\xfe\0\0', b'\0\0\xfe\xff')) else
                'utf-16' if data.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig')
    try:
        text = data.decode(encoding, 'strict')
    except UnicodeError as error:
        raise _Limited(kind + '_encoding_not_inspected') from error
    if '\0' in text:
        raise _Limited(kind + '_encoding_not_inspected')
    return text


def _odf_grids(root, budget):
    sheets, cells, total_rows = [], 0, 0
    for table in root.findall('.//{%s}table' % TABLE):
        if table.findall('.//{%s}table' % TABLE):
            raise _Limited('nested_table_mapping_not_inspected')
        rows = []
        for row in table.iter('{%s}table-row' % TABLE):
            repeat = int(row.get('{%s}number-rows-repeated' % TABLE, '1'))
            if not 1 <= repeat <= MAX_CELLS:
                raise _Limited('repeated_rows_budget')
            total_rows += repeat
            if total_rows > MAX_CELLS:
                raise _Limited('row_budget')
            items = []
            for cell in row:
                if cell.tag not in ('{%s}table-cell' % TABLE, '{%s}covered-table-cell' % TABLE):
                    continue
                if cell.tag.endswith('}covered-table-cell') or any(
                    cell.get('{%s}%s' % (TABLE, key), '1') != '1'
                    for key in ('number-columns-spanned', 'number-rows-spanned')):
                    raise _Limited('merged_cells_mapping_not_inspected')
                columns = int(cell.get('{%s}number-columns-repeated' % TABLE, '1'))
                if not 1 <= columns <= MAX_CELLS:
                    raise _Limited('repeated_columns_budget')
                kind = cell.get('{%s}value-type' % OFFICE, 'string')
                text = _cell_text(cell)
                if kind in ('float', 'percentage', 'currency'):
                    value = _value(cell.get('{%s}value' % OFFICE, text), 'number')
                elif kind in ('date', 'time', 'boolean'):
                    value = _value(cell.get('{%s}%s-value' % (OFFICE, kind), text), kind)
                else:
                    value = _value(text)
                value['display'] = text
                formula = cell.get('{%s}formula' % TABLE)
                if formula:
                    value['formula'] = formula
                cells += columns * repeat
                if cells > MAX_CELLS:
                    raise _Limited('cell_budget')
                items.extend([value] * columns)
                budget.check()
            rows.extend([items] * repeat)
        sheets.append({'name': table.get('{%s}name' % TABLE, ''), 'rows': _trim(rows)})
    return sheets


def _word_grids(root, budget):
    result, cells = [], 0
    for namespace in (WORD, STRICT_WORD):
        for table in root.findall('.//{%s}tbl' % namespace):
            if table.findall('.//{%s}tbl' % namespace):
                raise _Limited('nested_table_mapping_not_inspected')
            rows = []
            for row in table.findall('{%s}tr' % namespace):
                values = []
                for cell in row.findall('{%s}tc' % namespace):
                    if cell.find('.//{%s}gridSpan' % namespace) is not None or cell.find('.//{%s}vMerge' % namespace) is not None:
                        raise _Limited('merged_cells_mapping_not_inspected')
                    cells += 1
                    if cells > MAX_CELLS:
                        raise _Limited('cell_budget')
                    values.append(_value(_cell_text(cell)))
                    budget.check()
                rows.append(values)
            result.append({'name': '', 'rows': _trim(rows)})
    return result


def _package_xml(archive, name, budget):
    with archive.open(name) as stream:
        return _xml(budget.read(stream), budget)


def _relationships(archive, base, budget):
    relname = posixpath.join(posixpath.dirname(base), '_rels', posixpath.basename(base) + '.rels')
    if relname not in archive.namelist():
        return {}
    result = {}
    for relation in _package_xml(archive, relname, budget):
        if relation.get('TargetMode', '').lower() == 'external':
            continue
        target = unquote(relation.get('Target', ''))
        parsed = urlsplit(target)
        if parsed.scheme or parsed.netloc or parsed.query or parsed.fragment or '\\' in target:
            continue
        name = posixpath.normpath(target.lstrip('/') if target.startswith('/') else
                                posixpath.join(posixpath.dirname(base), target))
        if name == '..' or name.startswith('../'):
            continue
        result[relation.get('Id', '')] = (relation.get('Type', '').rsplit('/', 1)[-1], name)
    return result


def _word_summary(archive, names, budget):
    root = _package_xml(archive, 'word/document.xml', budget)
    if root.tag not in ('{%s}document' % WORD, '{%s}document' % STRICT_WORD):
        raise _Limited('word_document_root_not_recognized')
    roots = [root]
    coverage = {'body'}
    stories = {name for type_, name in _relationships(archive, 'word/document.xml', budget).values()
               if type_ in ('header', 'footer', 'footnotes', 'endnotes')}
    # Include conventional story parts even for a recoverable package whose
    # relationships are absent. Detached comments and revisions stay unknown.
    stories.update(n for n in names if re.fullmatch(r'word/(?:header\d*|footer\d*|footnotes|endnotes)\.xml', n))
    for name in sorted(stories):
        if name not in names:
            raise _Limited('word_related_text_part_missing')
        story = _package_xml(archive, name, budget)
        local = story.tag.rsplit('}', 1)[-1]
        if local not in ('hdr', 'ftr', 'footnotes', 'endnotes'):
            raise _Limited('word_related_text_root_not_recognized')
        if local in ('footnotes', 'endnotes'):
            # Separator glyphs are layout machinery rather than note content.
            for note in list(story):
                if any(note.get('{%s}type' % namespace) in ('separator', 'continuationSeparator')
                       for namespace in (WORD, STRICT_WORD)):
                    story.remove(note)
        roots.append(story)
        coverage.add({'hdr': 'headers', 'ftr': 'footers'}.get(local, local))
    texts = [_paragraph_text(part, budget) for part in roots]
    if sum(map(len, texts)) > MAX_TEXT:
        raise _Limited('text_budget')
    return {'kind': 'ooxml', 'documentClass': 'text', 'complete': True,
        '_text': '\n'.join(texts), '_grids': _word_grids(root, budget),
        'textCoverage': sorted(coverage), 'hasStyles': 'word/styles.xml' in names,
        'unknown': ['comments_and_revisions', 'hidden_text', 'field_evaluation',
                    'textboxes_reading_order', 'word_table_structure'],
        'counts': {'tables': sum(1 for x in root.iter() if x.tag.rsplit('}', 1)[-1] == 'tbl'),
                   'textParts': len(roots),
                   'images': sum(n.startswith('word/media/') and not n.endswith('/') for n in names),
                   'formulas': sum(x.tag.rsplit('}', 1)[-1] == 'oMath' for part in roots for x in part.iter())}}


def _presentation_summary(archive, names, budget):
    root = _package_xml(archive, 'ppt/presentation.xml', budget)
    if root.tag not in ('{%s}presentation' % PRESENTATION, '{%s}presentation' % STRICT_PRESENTATION):
        raise _Limited('presentation_root_not_recognized')
    relations = _relationships(archive, 'ppt/presentation.xml', budget)
    slide_names = []
    for item in root.iter():
        if item.tag.rsplit('}', 1)[-1] != 'sldId':
            continue
        rid = next((value for key, value in item.attrib.items() if key.endswith('}id')), None)
        type_, name = relations.get(rid, ('', ''))
        if type_ != 'slide' or name not in names or name in slide_names:
            raise _Limited('presentation_slide_mapping_not_determined')
        slide_names.append(name)
        if len(slide_names) > 1000:
            raise _Limited('presentation_slide_budget')
    if not slide_names:
        raise _Limited('presentation_slide_mapping_not_determined')
    texts, grids, cells, text_size = [], [], 0, 0
    for name in slide_names:
        slide = _package_xml(archive, name, budget)
        if slide.tag not in ('{%s}sld' % PRESENTATION, '{%s}sld' % STRICT_PRESENTATION):
            raise _Limited('presentation_slide_root_not_recognized')
        texts.append(_paragraph_text(slide, budget))
        text_size += len(texts[-1])
        if text_size > MAX_TEXT:
            raise _Limited('text_budget')
        for namespace in (DRAWING, STRICT_DRAWING):
            for table in slide.iter('{%s}tbl' % namespace):
                rows = []
                for row in table.findall('{%s}tr' % namespace):
                    values = []
                    for cell in row.findall('{%s}tc' % namespace):
                        if any((key.rsplit('}', 1)[-1] in ('gridSpan', 'rowSpan') and value != '1') or
                               (key.rsplit('}', 1)[-1] in ('hMerge', 'vMerge') and value not in ('0', 'false'))
                               for key, value in cell.attrib.items()):
                            raise _Limited('merged_cells_mapping_not_inspected')
                        cells += 1
                        if cells > MAX_CELLS:
                            raise _Limited('cell_budget')
                        values.append(_value(_cell_text(cell)))
                    rows.append(values)
                grids.append({'name': '', 'rows': _trim(rows)})
    return {'kind': 'ooxml', 'documentClass': 'presentation', 'complete': True,
        '_text': '\n'.join(texts), '_grids': grids, 'textCoverage': ['slide_text', 'slide_table_text'],
        'hasStyles': True, 'counts': {'slides': len(slide_names), 'tables': len(grids),
            'cells': cells, 'images': sum(n.startswith('ppt/media/') and not n.endswith('/') for n in names)},
        'unknown': ['presentation_notes', 'presentation_masters_and_layouts',
                    'presentation_animations', 'charts_and_smartart', 'embedded_objects',
                    'table_to_slide_mapping', 'slide_object_positions']}


def _csv_summary(path, budget):
    with path.open('rb') as stream:
        data = budget.read(stream, MAX_READ)
    text = _decode_text(data, 'csv')
    choices = []
    for delimiter in (',', ';', '\t', '|'):
        reader = csv.reader(io.StringIO(text[:65536], newline=''), delimiter=delimiter, strict=True)
        sample = []
        try:
            for row in reader:
                if row:
                    sample.append(row)
                if len(sample) == 8:
                    break
        except csv.Error:
            continue
        if sample and len(sample[0]) > 1 and all(len(row) == len(sample[0]) for row in sample):
            choices.append(delimiter)
    if len(choices) != 1:
        raise _Limited('csv_dialect_not_determined')
    rows, count = [], 0
    for row in csv.reader(io.StringIO(text, newline=''), delimiter=choices[0], strict=True):
        budget.check()
        if len(rows) >= MAX_CELLS:
            raise _Limited('row_budget')
        count += len(row)
        if count > MAX_CELLS:
            raise _Limited('cell_budget')
        rows.append([_value(value) for value in row])
    return {'kind': 'csv', 'documentClass': 'spreadsheet', 'complete': True,
            '_text': text[:MAX_TEXT], '_grids': [{'name': '', 'rows': _trim(rows)}],
            '_text_truncated': len(text) > MAX_TEXT,
            'counts': {'sheets': 1, 'rows': len(rows), 'columns': max(map(len, rows), default=0), 'characters': len(text)},
            'unknown': ['cell_types', 'cell_formatting', 'sheet_names_and_identity'],
            'delimiter': choices[0]}


def _xlsx_summary(path, archive, budget):
    # Openpyxl eagerly loads metadata even in read_only mode. Inspect every
    # non-worksheet XML first, including nonstandard OPC shared string parts.
    declared_xml = set()
    if '[Content_Types].xml' in archive.namelist():
        with archive.open('[Content_Types].xml') as stream:
            declarations = _xml(budget.read(stream), budget)
        for entry in declarations:
            if entry.get('ContentType', '').endswith(('+xml', '/xml')):
                name = entry.get('PartName', '')
                if name.startswith('/'):
                    declared_xml.add(name[1:])
    parts = [item for item in archive.infolist() if item.filename in declared_xml or
             item.filename.lower().endswith(('.xml', '.rels'))]
    if sum(item.file_size for item in parts) > MAX_READ - budget.used:
        raise _Limited('xlsx_xml_budget')
    for item in parts:
        if item.filename != '[Content_Types].xml':
            with archive.open(item) as stream:
                _xml(budget.read(stream), budget)
    from openpyxl import load_workbook
    # We identify the container before the suffix. A spreadsheet incorrectly
    # named .docx should be reported as a model mismatch, not hidden by
    # openpyxl's extension whitelist.
    stream = path.open('rb')
    try:
        book = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
    except BaseException:
        stream.close()
        raise
    try:
        sheets, cells, formulas, styles = [], 0, 0, False
        for sheet in book:
            if sheet.max_row and sheet.max_column and sheet.max_row * sheet.max_column > MAX_CELLS:
                raise _Limited('cell_budget')
            rows = []
            for row in sheet.iter_rows():
                budget.check()
                values = []
                for cell in row:
                    cells += 1
                    if cells > MAX_CELLS:
                        raise _Limited('cell_budget')
                    # Read-only sparse sheets yield EmptyCell placeholders,
                    # which carry no style metadata or has_style property.
                    styles = styles or getattr(cell, 'has_style', False)
                    value = cell.value
                    if cell.data_type == 'f':
                        formulas += 1
                        values.append(_value('', 'unknown', formula=str(value)))
                    elif value is None:
                        values.append(_value(''))
                    elif isinstance(value, bool):
                        values.append(_value(str(value).lower(), 'boolean'))
                    elif isinstance(value, (int, float, Decimal)):
                        values.append(_value(str(value), 'number'))
                    elif hasattr(value, 'isoformat'):
                        values.append(_value(value.isoformat(), 'date'))
                    else:
                        values.append(_value(value))
                rows.append(values)
            sheets.append({'name': sheet.title, 'rows': _trim(rows)})
        return {'kind': 'ooxml', 'documentClass': 'spreadsheet', 'complete': True,
                '_grids': sheets, '_text': '', 'counts': {'sheets': len(sheets), 'cells': cells,
                'formulas': formulas,
                'images': sum(item.filename.startswith('xl/media/') and not item.is_dir() for item in archive.infolist())},
                'hasStyles': styles,
                'unknown': ['cell_types', 'cell_formatting', 'sheet_names_and_identity',
                            'merged_cells', 'cell_annotations', 'charts_and_embedded_objects']}
    finally:
        book.close()
        stream.close()


def _pdf_summary(path, budget):
    if path.stat().st_size > MAX_READ:
        raise _Limited('pdf_byte_budget')
    from pypdf import PdfReader
    reader = PdfReader(str(path), strict=False)
    try:
        if reader.is_encrypted:
            raise _Limited('encrypted_pdf_not_inspected')
        pages = len(reader.pages)
        if pages > MAX_PDF_PAGES:
            raise _Limited('pdf_page_budget')
        texts, text_size, images, seen, text_complete = [], 0, [], set(), True
        for page in reader.pages:
            budget.check()
            try:
                value = page.extract_text() or ''
                text_size += len(value)
                if text_size > MAX_TEXT:
                    raise _Limited('pdf_text_budget')
                texts.append(value)
            except _Limited:
                raise
            except Exception:
                text_complete = False
            resources = page.get('/Resources', {})
            resources = resources.get_object() if hasattr(resources, 'get_object') else resources
            objects = resources.get('/XObject', {})
            objects = objects.get_object() if hasattr(objects, 'get_object') else objects
            for reference in objects.values():
                budget.check()
                obj = reference.get_object()
                if obj.get('/Subtype') != '/Image':
                    continue
                key = getattr(reference, 'idnum', id(obj))
                if key in seen:
                    continue
                seen.add(key)
                if len(images) > 10000:
                    raise _Limited('pdf_image_budget')
                width, height = obj.get('/Width'), obj.get('/Height')
                if type(width) is int or isinstance(width, int):
                    images.append({'width': int(width), 'height': int(height or 0)})
        return {'kind': 'pdf', 'documentClass': 'rendered', 'complete': True,
                'counts': {'pages': pages, 'images': len(images)}, '_text': '\n'.join(texts),
                'pdfImages': images, '_text_complete': text_complete,
                'unknown': ['pdf_text_extraction', 'nested_or_inline_pdf_images']}
    finally:
        reader.close()


def collect_source_summary(source, *, cancel=None):
    """Return one private, bounded reusable summary for the current source."""
    source = Path(source)
    budget = _Budget(cancel)
    result = {'kind': 'unknown', 'documentClass': 'unknown', 'complete': False,
              'counts': {}, 'unknown': list(UNCHECKED), '_text': '', '_grids': []}
    try:
        identity = _identity(source)
        result['_identity'] = identity
        budget.check()
        with source.open('rb') as stream:
            magic = stream.read(32)
        if magic.startswith((b'\x89PNG\r\n\x1a\n', b'GIF87a', b'GIF89a', b'BM', b'\xff\xd8', b'II*\0', b'MM\0*')) or (
                magic.startswith(b'RIFF') and magic[8:12] == b'WEBP'):
            media = inspect_media(source, cancel=cancel, validate_container=True)
            result.update(kind='raster', documentClass='image', complete=True, media=media,
                          counts={k: media[k] for k in ('width', 'height', 'frames', 'pages')})
        elif magic.startswith(b'%PDF-'):
            result.update(_pdf_summary(source, budget))
        elif magic.startswith(b'PK\x03\x04'):
            with zipfile.ZipFile(source) as archive:
                names = archive.namelist()
                if len(names) > 10000 or len(names) != len(set(names)):
                    raise _Limited('package_entry_budget_or_duplicates')
                if 'xl/workbook.xml' in names:
                    result.update(_xlsx_summary(source, archive, budget))
                elif 'word/document.xml' in names:
                    result.update(_word_summary(archive, names, budget))
                elif 'ppt/presentation.xml' in names:
                    result.update(_presentation_summary(archive, names, budget))
                elif 'content.xml' in names and 'mimetype' in names:
                    with archive.open('mimetype') as stream:
                        mime = budget.read(stream, 256).decode('ascii')
                    if not mime.startswith('application/vnd.oasis.opendocument.'):
                        raise _Limited('unrecognized_package')
                    with archive.open('content.xml') as stream:
                        root = _xml(budget.read(stream), budget)
                    kind = mime.rsplit('.', 1)[-1]
                    result.update(kind='odf', documentClass={'spreadsheet': 'spreadsheet', 'presentation': 'presentation',
                        'graphics': 'drawing'}.get(kind, 'text'),
                        complete=True, _text=_paragraph_text(root, budget), _grids=_odf_grids(root, budget),
                        hasStyles=any(any(key.endswith('}style-name') for key in x.attrib) for x in root.iter()),
                        unknown=['notes_and_comments', 'tracked_changes', 'embedded_objects',
                                 'cell_types', 'sheet_names_and_identity', 'cell_formatting'],
                        counts={'tables': len(root.findall('.//{%s}table' % TABLE)),
                                'images': sum(n.startswith('Pictures/') and not n.endswith('/') for n in names),
                                'formulas': sum('{%s}formula' % TABLE in x.attrib for x in root.iter())})
                    if result['documentClass'] == 'spreadsheet':
                        result['counts']['sheets'] = len(result['_grids'])
                else:
                    result['reason'] = 'unsupported_package'
        elif source.suffix.lower() in ('.fods', '.fodt', '.fodp', '.fodg'):
            with source.open('rb') as stream:
                root = _xml(budget.read(stream), budget)
            kind = root.get('{%s}mimetype' % OFFICE, '')
            result.update(kind='odf', documentClass='spreadsheet' if kind.endswith('.spreadsheet') else 'text',
                          complete=True, _text=_paragraph_text(root, budget), _grids=_odf_grids(root, budget),
                          unknown=['notes_and_comments', 'tracked_changes', 'embedded_objects',
                                   'cell_types', 'sheet_names_and_identity', 'cell_formatting'],
                          counts={'tables': len(root.findall('.//{%s}table' % TABLE)),
                                  'formulas': sum('{%s}formula' % TABLE in x.attrib for x in root.iter())})
        elif source.suffix.lower() in ('.csv', '.tsv', '.tab'):
            result.update(_csv_summary(source, budget))
        elif source.suffix.lower() in ('.txt', '.text', '.md'):
            with source.open('rb') as stream:
                raw = budget.read(stream, MAX_TEXT * 4)
            value = _decode_text(raw, 'text')
            if len(value) > MAX_TEXT:
                raise _Limited('text_budget')
            result.update(kind='text', documentClass='text', complete=True, _text=value)
        else:
            result['reason'] = 'unsupported_format'
        budget.check()
        if _identity(source) != identity:
            raise _Limited('source_changed_during_analysis')
        text = result.get('_text', '')
        result['textDigest'] = hashlib.sha256(_normal(text).encode()).hexdigest() if text else None
        result['counts'].setdefault('characters', len(text))
    except GuardError as error:
        if error.code == 'CANCELLED':
            raise
        result.update(complete=False, reason=error.code)
    except (_Limited, UnicodeError, ET.ParseError, zipfile.BadZipFile, OSError, ValueError, KeyError) as error:
        result.update(complete=False, reason=str(error) if isinstance(error, _Limited) else 'inspection_failed')
    except Exception:
        result.update(complete=False, reason='inspection_failed')
    if not result.get('complete'):
        result['_text'], result['_grids'] = '', []
    return result


def _public(summary):
    return {key: summary[key] for key in ('kind', 'documentClass', 'complete', 'counts', 'reason',
                                        'textDigest', 'textCoverage') if key in summary}


def _number(value):
    try:
        result = Decimal(value)
        return result if result.is_finite() else None
    except (InvalidOperation, ValueError, TypeError):
        return None


def _compare_grids(before, after, findings, checked, unknown):
    from collections import Counter
    a, b = before.get('_grids', []), after.get('_grids', [])
    if not a or not b:
        return
    # Mapping multiple sheets/tables merely by order would invent losses when
    # exporters rename/reorder them. Only a single unambiguous grid is proven.
    if len(a) != 1 or len(b) != 1:
        unknown.add('multiple_table_or_sheet_mapping')
        return
    rows_a, rows_b = a[0]['rows'], b[0]['rows']
    changed, representation, additions, added_positions, unknown_formula = [], [], [], set(), False
    for ri, row in enumerate(rows_b):
        for ci, cell in enumerate(row):
            original = rows_a[ri][ci] if ri < len(rows_a) and ci < len(rows_a[ri]) else None
            if (cell['value'] or cell.get('formula')) and (
                    original is None or (not original['value'] and not original.get('formula'))):
                # Preservation of every original value is weaker than equality:
                # extra nonempty output cells must not certify the entire grid.
                value = cell['value'] or cell['formula']
                additions.append({'row': ri + 1, 'column': ci + 1, 'before': '', 'after': value[:120],
                                  **({'afterTruncated': True} if len(value) > 120 else {}),
                                  **({'afterKind': 'formula'} if cell.get('formula') else {})})
                added_positions.add((ri, ci))
    for ri, row in enumerate(rows_a):
        for ci, cell in enumerate(row):
            if (ri, ci) in added_positions:
                continue
            if cell.get('formula') and (cell['type'] == 'unknown' or not cell['value']):
                unknown_formula = True
                continue
            expected = cell['value']
            actual = rows_b[ri][ci]['value'] if ri < len(rows_b) and ci < len(rows_b[ri]) else ''
            if actual == expected:
                continue
            output_cell = rows_b[ri][ci] if ri < len(rows_b) and ci < len(rows_b[ri]) else None
            if output_cell and output_cell.get('formula') and (
                    output_cell['type'] == 'unknown' or not output_cell['value']):
                # A formula's absent cache is not an empty output value.
                unknown_formula = True
                continue
            special_type = cell['type'] in ('date', 'time', 'boolean') or (
                output_cell is not None and output_cell['type'] in ('date', 'time', 'boolean'))
            if special_type or (actual and cell['type'] == 'number' and _number(actual) is None):
                # ISO/native typed values and locale-formatted CSV display
                # strings cannot be called lost without a verified dialect.
                representation.append({'row': ri + 1, 'column': ci + 1,
                                       'before': expected[:120], 'after': actual[:120]})
                continue
            number_a, number_b = _number(expected), _number(actual)
            if number_a is not None and number_b is not None and number_a == number_b:
                if cell['type'] == 'text':
                    representation.append({'row': ri + 1, 'column': ci + 1,
                                           'before': expected[:120], 'after': actual[:120]})
                continue
            if not expected and not actual:
                continue
            changed.append({'row': ri + 1, 'column': ci + 1,
                            'before': expected[:120], 'after': actual[:120],
                            **({'beforeTruncated': True} if len(expected) > 120 else {}),
                            **({'afterTruncated': True} if len(actual) > 120 else {})})
    if changed:
        # Different coordinates alone are not lost content: an exporter may
        # reorder rows/columns. Confirm only missing value multiplicity in
        # the complete, unambiguous inspected grid; keep moved values uncertain.
        def token(value):
            number = _number(value)
            return ('number', number) if number is not None else ('text', value)
        source_values = Counter(token(cell['value']) for row in rows_a for cell in row
                                if cell['value'] and not (cell.get('formula') and cell['type'] == 'unknown'))
        result_values = Counter(token(cell['value']) for row in rows_b for cell in row if cell['value'])
        missing = source_values - result_values
        lost, moved = [], []
        for entry in changed:
            # Display snippets are bounded, but the proof must use the full
            # source value. A long cell must not become a fake "moved" value.
            full_value = rows_a[entry['row'] - 1][entry['column'] - 1]['value']
            key = token(full_value)
            if entry['before'] and missing[key]:
                lost.append(entry)
                missing[key] -= 1
            else:
                moved.append(entry)
        if lost:
            findings.append({'code': 'CELL_VALUES_CHANGED', 'category': 'confirmed_loss', 'feature': 'cell_values',
                'message': 'Исходные значения отсутствуют в проверенной таблице результата.', 'certainty': 'confirmed_subset',
                'before': len(lost), 'after': None, 'locations': lost[:25], 'totalAffected': len(lost)})
        if moved:
            findings.append({'code': 'CELL_POSITIONS_CHANGED', 'category': 'uncertain', 'feature': 'cell_relationships',
                'message': 'Расположение значений изменилось; потеря самих значений не доказана.',
                'certainty': 'requires_review', 'locations': moved[:25], 'totalAffected': len(moved)})
    elif not unknown_formula and not representation and not additions:
        checked.add('cell_values_single_grid')
    if additions:
        findings.append({'code': 'CELL_VALUES_ADDED', 'category': 'uncertain', 'feature': 'cell_values',
            'message': 'В результате появились непустые ячейки или формулы в исходно пустых позициях. Причина и допустимость добавления требуют проверки.',
            'certainty': 'requires_review', 'locations': additions[:25], 'totalAffected': len(additions)})
    if representation:
        findings.append({'code': 'CELL_REPRESENTATION_CHANGED', 'category': 'uncertain', 'feature': 'cell_text',
            # This bucket includes verified numeric equivalence of TEXT cells,
            # but also typed dates, booleans and numeric strings whose locale is
            # unknown. Never assert equality for the entire bucket: "1,23" may
            # be a rounded export of 1.23456789012345, not the same value.
            'message': 'Представление значений изменилось; совпадение данных с учётом типа, локализации и точности требует проверки.',
            'certainty': 'requires_review', 'locations': representation[:25], 'totalAffected': len(representation)})
    if unknown_formula:
        unknown.add('cached_formula_values')


def _compare_formulas(before, after, findings, checked, unknown):
    a, b = before.get('_grids', []), after.get('_grids', [])
    if before.get('documentClass') != 'spreadsheet' or after.get('documentClass') != 'spreadsheet':
        return
    if after.get('kind') not in ('ooxml', 'odf'):
        return
    if len(a) != 1 or len(b) != 1:
        unknown.add('editable_formula_mapping')
        return
    def formulas(grid):
        return {(ri + 1, ci + 1): cell['formula'] for ri, row in enumerate(grid['rows'])
                for ci, cell in enumerate(row) if cell.get('formula')}
    old, new = formulas(a[0]), formulas(b[0])
    if not old:
        return
    if old == new:
        checked.add('editable_formulas_single_grid')
    elif not new and after.get('kind') in ('ooxml', 'odf'):
        findings.append({'code': 'EDITABLE_FORMULAS_MISSING', 'category': 'confirmed_loss',
            'feature': 'editable_formulas',
            'message': 'В исходной таблице есть формулы, а в проверенной таблице результата формул нет.',
            'certainty': 'confirmed_subset', 'before': len(old), 'after': 0,
            'locations': [{'row': row, 'column': column, 'before': value[:120], 'after': ''}
                          for (row, column), value in list(old.items())[:25]], 'totalAffected': len(old)})
    else:
        findings.append({'code': 'FORMULA_STRUCTURE_CHANGED', 'category': 'uncertain',
            'feature': 'editable_formulas', 'message': 'Формулы или их расположение изменились; эквивалентность выражений не проверена.',
            'certainty': 'requires_review', 'before': len(old), 'after': len(new)})


def _compare_media_sequences(a, b, target, findings, checked, unknown):
    """Map static pages and timed frames before comparing their properties."""
    target = ALIASES.get(target, target)
    if target == 'tiff':
        expected, actual = max(a.get('pages', 1), a.get('frames', 1)), b.get('pages', 1)
        field = 'pages'
        if a.get('frames', 1) > 1:
            findings.append({'code': 'FORMAT_ANIMATION_TIMELINE_LIMITATION', 'category': 'format_limitation',
                'feature': 'animation_timeline', 'message': 'Кадры могут сохраниться как страницы TIFF, но формат не сохраняет временную шкалу анимации.',
                'certainty': 'format_contract'})
    elif target in ANIMATED_TARGETS:
        expected, actual = max(a.get('pages', 1), a.get('frames', 1)), b.get('frames', 1)
        field = 'frames'
        if a.get('pages', 1) > 1:
            findings.append({'code': 'FORMAT_PAGE_SEMANTICS_LIMITATION', 'category': 'format_limitation',
                'feature': 'page_semantics', 'message': 'Страницы исходника представлены кадрами анимации; их временная шкала в исходнике не задана.',
                'certainty': 'format_contract'})
            unknown.add('page_to_frame_timing')
    else:
        expected, actual, field = 1, 1, 'single_image'
        for kind in ('frames', 'pages'):
            if a.get(kind, 1) > 1:
                findings.append({'code': 'FORMAT_' + kind.upper() + '_LIMITATION', 'category': 'format_limitation',
                    'feature': kind, 'message': 'Выбранный формат не сохраняет ' +
                        ('анимацию.' if kind == 'frames' else 'несколько страниц.'),
                    'before': a[kind], 'after': b.get(kind, 1), 'certainty': 'format_contract'})
    if expected != actual:
        findings.append({'code': 'MEDIA_' + field.upper() + '_CHANGED', 'category': 'confirmed_loss',
            'feature': field, 'message': 'Количество ' + ('страниц' if field == 'pages' else 'кадров') + ' изменилось.',
            'before': expected, 'after': actual, 'certainty': 'confirmed_subset'})
    else:
        checked.add(field)
    if a.get('frames', 1) > 1 and target in ANIMATED_TARGETS and a['frames'] == b.get('frames'):
        if a.get('loopCount') is not None and b.get('loopCount') is not None:
            if a['loopCount'] != b['loopCount']:
                findings.append({'code': 'ANIMATION_LOOP_CHANGED', 'category': 'confirmed_loss',
                    'feature': 'animation_loop_count', 'message': 'Число повторов анимации изменилось.',
                    'before': a['loopCount'], 'after': b['loopCount'], 'certainty': 'confirmed_subset'})
            else:
                checked.add('animation_loop_count')
        else:
            unknown.add('animation_loop_count')
        old, new = a.get('durationsMs', []), b.get('durationsMs', [])
        if len(old) == len(new) == a['frames']:
            changes = [{'frame': i + 1, 'before': left, 'after': right}
                       for i, (left, right) in enumerate(zip(old, new))
                       if abs(left - right) > max(20, left * 0.05)]
            if changes:
                findings.append({'code': 'ANIMATION_TIMING_CHANGED', 'category': 'confirmed_loss',
                    'feature': 'frame_durations', 'message': 'Длительности кадров изменились сверх допуска 20 мс или 5%.',
                    'locations': changes[:25], 'totalAffected': len(changes), 'certainty': 'confirmed_subset'})
            else:
                checked.add('frame_durations_within_tolerance')
        else:
            unknown.add('frame_durations')
    elif a.get('frames', 1) > 1:
        unknown.add('frame_durations')
    if target == 'tiff' and b.get('pageDimensions'):
        old_dimensions = (display_page_dimensions(a) if a.get('pageDimensions') else
                          [list(display_dimensions(a))] * max(a.get('pages', 1), a.get('frames', 1)))
        new_dimensions = display_page_dimensions(b)
        if len(old_dimensions) == len(new_dimensions):
            changed = [{'page': i + 1, 'before': old, 'after': new}
                       for i, (old, new) in enumerate(zip(old_dimensions, new_dimensions)) if old != new]
            if changed:
                findings.append({'code': 'IMAGE_PAGE_DIMENSIONS_CHANGED', 'category': 'confirmed_loss',
                    'feature': 'page_dimensions', 'message': 'Размеры отдельных страниц изображения изменились.',
                    'locations': changed[:25], 'totalAffected': len(changed), 'certainty': 'confirmed_subset'})
            else:
                checked.add('page_dimensions')
    elif a.get('pages', 1) > 1:
        unknown.add('individual_page_to_frame_dimensions')
    if a.get('bitDepth') is not None and b.get('bitDepth') is not None:
        if a['bitDepth'] > b['bitDepth']:
            supported = target in ('png', 'apng', 'tiff')
            findings.append({'code': 'IMAGE_BIT_DEPTH_REDUCED' if supported else 'FORMAT_BIT_DEPTH_LIMITATION',
                'category': 'confirmed_loss' if supported else 'format_limitation', 'feature': 'bit_depth',
                'message': 'Глубина цвета изображения уменьшилась.', 'before': a['bitDepth'], 'after': b['bitDepth'],
                'certainty': 'confirmed_subset' if supported else 'format_contract'})
        else:
            checked.add('bit_depth_not_reduced')
    else:
        unknown.add('bit_depth')
    unknown.update(('image_pixel_values', 'colour_profile', 'exif_and_other_metadata'))


def _multi_csv_summary(paths, before, budget, findings, checked, unknown):
    """Prove all sheet VALUES by multiplicity without guessing sheet order.

    This does not certify names or map an ambiguous mismatch to a source row.
    Complete equal multiset evidence is useful for existing all-sheet exports.
    """
    from collections import Counter
    if len(paths) > 256:
        raise _Limited('sheet_output_budget')
    identities = [_identity(path) for path in paths]
    if len({identity[0] for identity in identities}) != len(paths):
        raise _Limited('duplicate_primary_outputs')
    expected, observed, cells = Counter(), Counter(), 0

    def signature(grid):
        nonlocal cells
        values = [[cell['value'] for cell in row] for row in grid['rows']]
        cells += sum(map(len, values))
        if cells > MAX_CELLS * 2:
            raise _Limited('cell_budget')
        return hashlib.sha256(json.dumps(values, ensure_ascii=False,
                             separators=(',', ':')).encode()).hexdigest()

    grids = before.get('_grids', [])
    for grid in grids:
        if any(cell.get('formula') and (cell['type'] == 'unknown' or not cell['value'])
               for row in grid['rows'] for cell in row):
            raise _Limited('cached_formula_values')
        expected[signature(grid)] += 1
    for path in paths:
        budget.check()
        value = _csv_summary(path, budget)
        observed[signature(value['_grids'][0])] += 1
    if any(_identity(path) != identity for path, identity in zip(paths, identities)):
        raise _Limited('result_changed_during_analysis')
    if grids and len(grids) == len(paths) and expected == observed:
        checked.add('all_sheet_values_unordered')
        unknown.add('sheet_names_and_identity')
    else:
        findings.append({'code': 'MULTI_SHEET_VALUES_NOT_MATCHED', 'category': 'uncertain',
            'feature': 'cell_values', 'message': 'Не удалось сопоставить все исходные листы и текстовые результаты.',
            'before': len(grids), 'after': len(paths), 'certainty': 'requires_review'})
        unknown.add('all_sheet_outputs_mapping')
    return {'kind': 'csv', 'documentClass': 'spreadsheet', 'complete': True,
            'counts': {'sheetFiles': len(paths)}, '_text': '', '_grids': [], '_identities': identities}


def analyze_pair(source, paths, target, *, cancel=None, source_evidence=None):
    """Compare one retained source/target pair without replacing either file."""
    source, paths = Path(source), [Path(path) for path in paths]
    target = str(target).lower().lstrip('.')
    before = source_evidence or collect_source_summary(source, cancel=cancel)
    report = {'schemaVersion': 1, 'status': 'not_checked', 'qualityCompared': False,
              'scope': 'verified_properties_only', 'findings': [], 'checkedFeatures': [],
              'unknown': list(UNCHECKED), 'recipeCandidates': [], 'before': _public(before), 'after': {},
              'outputScope': {'totalFiles': len(paths), 'primaryFiles': 0, 'supplementalFiles': len(paths),
                              'inspectedPrimaryFiles': 0, 'inspectedSupplementalFiles': 0}}
    findings, checked, unknown = report['findings'], set(), set(UNCHECKED)
    budget = _Budget(cancel)
    budget.check()
    try:
        if tuple(before.get('_identity', ())) != _identity(source):
            raise _Limited('source_changed_or_unavailable')
        if len(paths) > 10000:
            raise _Limited('output_file_budget')
        if target in RASTER:
            canonical = ALIASES.get(target, target)
            candidates = [p for p in paths if ALIASES.get(p.suffix.lower().lstrip('.'), p.suffix.lower().lstrip('.')) == canonical]
        else:
            candidates = [p for p in paths if p.suffix.lower().lstrip('.') == target]
        report['outputScope'].update(primaryFiles=len(candidates), supplementalFiles=len(paths) - len(candidates))
        if len(paths) != len(candidates):
            unknown.add('supplemental_outputs_not_inspected')
        multi_csv = len(candidates) > 1 and target in ('csv', 'tsv', 'tab')
        if len(candidates) != 1 and not multi_csv:
            raise _Limited('primary_output_mapping_not_determined')
        primary = candidates[0]
        after = (_multi_csv_summary(candidates, before, budget, findings, checked, unknown)
                 if multi_csv else collect_source_summary(primary, cancel=cancel))
        report['outputScope']['inspectedPrimaryFiles'] = len(candidates)
        report['after'] = _public(after)
        required_models = {'docx': ('ooxml', 'text'), 'xlsx': ('ooxml', 'spreadsheet'),
                           'pptx': ('ooxml', 'presentation'), 'odt': ('odf', 'text'),
                           'ods': ('odf', 'spreadsheet'), 'odp': ('odf', 'presentation'),
                           'odg': ('odf', 'drawing'), 'pdf': ('pdf', 'rendered')}
        if after.get('complete') and target in required_models and (
                after.get('kind'), after.get('documentClass')) != required_models[target]:
            raise GuardError('OUTPUT_FORMAT_MISMATCH', 'Структура результата не соответствует выбранному формату.', stage='comparison')
        unknown.update(before.get('unknown', ()))
        unknown.update(after.get('unknown', ()))
        if not before.get('complete') or not after.get('complete'):
            if not before.get('complete'):
                unknown.add('incomplete_source_inspection')
            if not after.get('complete'):
                unknown.add('incomplete_result_inspection')
            findings.append({'code': 'ANALYSIS_INCOMPLETE', 'category': 'not_checked',
                'message': 'Часть свойств не удалось надёжно проверить.',
                'certainty': 'unknown', 'reason': before.get('reason') or after.get('reason')})
        if not multi_csv and primary.suffix.lower() == source.suffix.lower() and primary.stat().st_size == source.stat().st_size:
            # Stream comparison is exact and bounded; do not retain whole files.
            identical = True
            with source.open('rb') as left, primary.open('rb') as right:
                while True:
                    budget.check()
                    a, b = left.read(65536), right.read(65536)
                    budget.check(len(a) + len(b))
                    if a != b:
                        identical = False
                        break
                    if not a:
                        break
            if identical:
                checked.add('identical_file_bytes')
        media_a, media_b = before.get('media'), after.get('media')
        if target in RASTER:
            # Source probing tolerates recoverable inputs. Output container
            # verification must not accept an incomplete SOF/BMP header alone.
            verified = inspect_media(primary, cancel=cancel, validate_container=True)
            if verified is None or verified['format'] != ALIASES.get(target, target):
                raise GuardError('MEDIA_OUTPUT_INVALID', 'Неверный контейнер результата.', stage='comparison')
        if media_a and media_b:
            _compare_media_sequences(media_a, media_b, target, findings, checked, unknown)
            if display_dimensions(media_a) != display_dimensions(media_b):
                findings.append({'code': 'IMAGE_DIMENSIONS_CHANGED', 'category': 'confirmed_loss', 'feature': 'image_dimensions',
                    'message': 'Размеры изображения изменились.', 'before': list(display_dimensions(media_a)),
                    'after': list(display_dimensions(media_b)), 'certainty': 'confirmed_subset'})
            else:
                checked.add('image_dimensions')
            if media_a.get('alpha') and not media_b.get('alpha'):
                expected = ALIASES.get(target, target) == 'jpeg'
                bmp_unknown = False
                if media_b['format'] == 'bmp':
                    with primary.open('rb') as stream:
                        head = stream.read(70)
                    depth = int.from_bytes(head[28:30], 'little') if len(head) >= 30 else None
                    # The generic BMP inspector does not resolve 32-bit alpha
                    # masks. Its default False is not proof of absent alpha.
                    bmp_unknown = depth == 32
                if bmp_unknown:
                    unknown.add('bmp_alpha_channel')
                else:
                    findings.append({'code': 'FORMAT_ALPHA_LIMITATION' if expected else 'IMAGE_ALPHA_LOST',
                        'category': 'format_limitation' if expected else 'confirmed_loss', 'feature': 'alpha_channel',
                        'message': 'JPEG не сохраняет канал прозрачности.' if expected else 'В результате нет канала прозрачности исходника.',
                        'before': True, 'after': False, 'certainty': 'format_contract' if expected else 'confirmed_subset'})
        if target in ('csv', 'tsv', 'tab', 'txt') and before.get('hasStyles'):
            findings.append({'code': 'FORMAT_STYLES_LIMITATION', 'category': 'format_limitation', 'feature': 'styles',
                'message': 'Текстовый формат не сохраняет оформление документа.', 'certainty': 'format_contract'})
        if before.get('counts', {}).get('formulas') and target in ('csv', 'tsv', 'tab', 'txt', 'pdf'):
            findings.append({'code': 'FORMAT_EDITABLE_FORMULAS_LIMITATION', 'category': 'format_limitation', 'feature': 'formulas',
                'message': 'В результате нет редактируемой модели формул исходной книги.', 'certainty': 'format_contract'})
        if before.get('complete') and after.get('complete'):
            if before.get('documentClass') == after.get('documentClass'):
                for field in ('tables', 'slides', 'sheets', 'images'):
                    left_count, right_count = before.get('counts', {}).get(field), after.get('counts', {}).get(field)
                    if left_count is not None and right_count is not None and left_count != right_count:
                        findings.append({'code': 'STRUCTURE_' + field.upper() + '_COUNT_CHANGED', 'category': 'uncertain',
                            'feature': field, 'message': 'Число структурных объектов изменилось; только по количеству нельзя доказать потерю содержимого.',
                            'before': left_count, 'after': right_count, 'certainty': 'requires_review'})
            if before.get('documentClass') == after.get('documentClass') and before.get('_grids') and after.get('_grids'):
                _compare_grids(before, after, findings, checked, unknown)
                _compare_formulas(before, after, findings, checked, unknown)
            text_a, text_b = _normal(before.get('_text', '')), _normal(after.get('_text', ''))
            text_complete = before.get('_text_complete', True) and after.get('_text_complete', True)
            if text_a and text_b and text_complete and before.get('documentClass') != 'spreadsheet':
                if text_a == text_b:
                    checked.add('normalized_text')
                else:
                    findings.append({'code': 'TEXT_REPRESENTATION_DIFFERS', 'category': 'uncertain', 'feature': 'text',
                        'message': 'Извлечённый текст отличается; это ещё не доказывает исчезновение текста в документе.',
                        'certainty': 'requires_review'})
            elif text_a and (not text_b or not text_complete) and after.get('kind') == 'pdf':
                findings.append({'code': 'PDF_TEXT_NOT_EXTRACTED', 'category': 'uncertain', 'feature': 'text',
                    'message': 'Текст PDF не извлечён. Это не доказательство потери видимого текста.', 'certainty': 'requires_review'})
            elif text_a and not text_b and after.get('kind') == 'text':
                findings.append({'code': 'TEXT_OUTPUT_EMPTY', 'category': 'confirmed_loss', 'feature': 'text',
                    'message': 'В исходнике извлечён текст, а текстовый результат пуст.',
                    'before': len(text_a), 'after': 0, 'certainty': 'confirmed_subset'})
            if after.get('kind') == 'pdf' and before.get('counts', {}).get('images', 0) > 0 and not media_a:
                findings.append({'code': 'PDF_IMAGE_PROPERTIES_UNVERIFIED', 'category': 'uncertain', 'feature': 'images',
                    'message': 'Встроенные изображения требуют проверки качества PDF; число объектов не доказывает сохранность пикселей.',
                    'certainty': 'requires_review'})
                report['recipeCandidates'].append('pdf_preserve_images')
            if target in ('csv', 'tsv', 'tab') and any(f['code'] in ('CELL_VALUES_CHANGED', 'CELL_REPRESENTATION_CHANGED') for f in findings):
                report['recipeCandidates'].append('csv_preserve_values')
        if (before.get('documentClass') == 'spreadsheet' and target in ('csv', 'tsv', 'tab')
                and len(before.get('_grids', [])) > 1 and 'all_sheet_values_unordered' not in checked):
            unknown.add('all_sheet_outputs_mapping')
        if 'identical_file_bytes' in checked:
            # Identical files have no conversion changes, even if a subset
            # inspector is unavailable. Unknown visual properties stay explicit.
            findings[:] = [f for f in findings if f['category'] == 'not_checked']
        budget.check()
        if tuple(before.get('_identity', ())) != _identity(source):
            raise _Limited('source_changed_during_analysis')
        if not multi_csv and after.get('_identity') and tuple(after['_identity']) != _identity(primary):
            raise _Limited('result_changed_during_analysis')
        if multi_csv and any(tuple(identity) != _identity(path)
                             for path, identity in zip(candidates, after.get('_identities', ()))):
            raise _Limited('result_changed_during_analysis')
        categories = {f['category'] for f in findings}
        report['status'] = ('confirmed_loss' if 'confirmed_loss' in categories else
                            'uncertain' if 'uncertain' in categories else
                            'format_limitation' if 'format_limitation' in categories else
                            'checked' if checked else 'not_checked')
    except GuardError as error:
        if error.code == 'CANCELLED':
            raise
        invalid = error.code in ('MEDIA_OUTPUT_INVALID', 'OUTPUT_FORMAT_MISMATCH')
        findings.append({'code': error.code, 'category': 'confirmed_loss' if invalid else 'not_checked',
            **({'feature': 'result_container'} if invalid else {}),
            'message': ('Структура результата не соответствует выбранному формату.' if error.code == 'OUTPUT_FORMAT_MISMATCH' else
                       'Контейнер результата повреждён.' if error.code == 'MEDIA_OUTPUT_INVALID' else 'Проверка результата не завершена.'),
            'certainty': 'confirmed_subset' if invalid else 'unknown'})
        report['status'] = 'confirmed_loss' if invalid else 'not_checked'
    except (_Limited, OSError, ValueError) as error:
        # A changed/deleted input or interrupted inspection invalidates the
        # comparison as a whole. Do not leave stale red claims or recipes.
        checked.clear()
        findings.clear()
        report['recipeCandidates'] = []
        report['status'] = 'not_checked'
        findings.append({'code': 'ANALYSIS_UNAVAILABLE', 'category': 'not_checked',
            'message': 'Анализ этого преобразования недоступен.', 'certainty': 'unknown',
            'reason': str(error) if isinstance(error, _Limited) else 'file_unavailable'})
    report['findings'] = findings[:MAX_FINDINGS]
    report['checkedFeatures'] = sorted(checked)
    report['evaluatedFeatures'] = sorted(checked | {f['feature'] for f in report['findings']
        if f.get('certainty') == 'confirmed_subset' and f.get('feature')})
    report['unknown'] = sorted(unknown)
    report['qualityCompared'] = bool(report['evaluatedFeatures'])
    return report


def preview_text(path, *, max_chars=100000, cancel=None):
    """A labelled text representation, never a fake visual document preview."""
    if type(max_chars) is not int or not 1 <= max_chars <= MAX_TEXT:
        raise ValueError('invalid_preview_budget')
    evidence = collect_source_summary(Path(path), cancel=cancel)
    text = evidence.get('_text', '')
    if not text and evidence.get('_grids'):
        parts = []
        size = 0
        for sheet in evidence['_grids']:
            for row in sheet['rows']:
                value = '\t'.join(cell['value'] for cell in row)
                parts.append(value)
                size += len(value)
                if size > max_chars:
                    break
            if size > max_chars:
                break
        text = '\n'.join(parts)
    if not evidence.get('complete') or not text:
        return {'kind': 'unavailable', 'text': '', 'truncated': False,
                'note': 'Текстовый предпросмотр этого файла недоступен.',
                'reason': evidence.get('reason', 'no_text_representation')}
    return {'kind': 'text', 'text': text[:max_chars], 'truncated': len(text) > max_chars or evidence.get('_text_truncated', False),
            'note': 'Текстовое представление: оформление и расположение объектов не показаны.'}
