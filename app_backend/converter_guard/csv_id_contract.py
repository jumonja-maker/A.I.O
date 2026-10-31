"""Bounded mandatory leading-zero ID checks, not full spreadsheet fidelity.

Only a content-confirmed installed Calc Text importer receives column hints.
LibreOffice produces the output. This module never rewrites original files.
Unsupported grammars and budgets remain explicit not_checked, not damage.
"""
from __future__ import annotations

from contextlib import closing
import csv
import json
from pathlib import Path
import re
import time
from xml.etree.ElementTree import ParseError
import zipfile

from .errors import GuardError

CALC_TEXT = 'Text - txt - csv (StarCalc)'
SPREADSHEET_TARGETS = frozenset(('xlsx', 'xlsm', 'xltx', 'xltm'))
DELIMITERS = {44: ',', 59: ';', 9: '\t', 124: '|'}
LEADING_ZERO_ID = re.compile(r'[+-]?0[0-9]+\Z')
MAX_COLUMNS = 256
MAX_TEXT_COLUMNS = 64
MAX_INPUT_BYTES = 32 * 1024 * 1024
MAX_OUTPUT_BYTES = 64 * 1024 * 1024
MAX_UNCOMPRESSED = 64 * 1024 * 1024
MAX_TRACKED_IDS = 50000
MAX_ROWS = 100000


class _Limited(ValueError):
    """Inspection is incomplete; it is not evidence of a damaged result."""


def _cancel(cancel):
    if cancel is not None and cancel.is_set():
        raise GuardError('CANCELLED', 'Проверка текстовых ID отменена.', stage='result')


def _stamp(path):
    path = Path(path)
    if path.is_symlink() or not path.is_file():
        raise _Limited('source_or_result_unavailable')
    info = path.stat()
    return (info.st_dev, info.st_ino, info.st_size, info.st_mtime_ns, info.st_ctime_ns)


class _Budget:
    def __init__(self, cancel, seconds):
        self.cancel = cancel
        self.deadline = time.monotonic() + seconds

    def check(self):
        _cancel(self.cancel)
        if time.monotonic() >= self.deadline:
            raise _Limited('csv_id_inspection_deadline')


def text_columns_in_sample(rows, width):
    """Return ordered one-based column/Text pairs from a bounded sample.

    A hint is not route admission or proof of a complete source scan. An
    inferred Text column remains textual, including its mixed numeric rows;
    columns without an observed leading-zero field retain native behavior.
    """
    if type(width) is not int or not 1 <= width <= MAX_COLUMNS:
        return ''
    columns = set()
    for row in rows:
        if len(row) != width:
            return ''
        for index, value in enumerate(row, 1):
            if LEADING_ZERO_ID.fullmatch(value):
                columns.add(index)
                if len(columns) > MAX_TEXT_COLUMNS:
                    return ''
    return '/'.join(str(column) + '/2' for column in sorted(columns))


def _import_dialect(import_filter, options):
    if import_filter != CALC_TEXT or not isinstance(options, str) or len(options) > 1024:
        return None
    parts = options.split(',')
    if len(parts) not in (4, 5) or parts[1:4] != ['34', '76', '1']:
        return None
    if parts[0] not in ('44', '59', '9', '124'):
        return None
    if len(parts) == 5:
        pairs = parts[4].split('/')
        if len(pairs) < 2 or len(pairs) % 2 or len(pairs) > 2 * MAX_TEXT_COLUMNS:
            return None
        last = 0
        for position, code in zip(pairs[::2], pairs[1::2]):
            # Decimal hints must not accidentally bypass the source-ID check.
            # Skip/date/unknown format codes still do not establish a mapping.
            if (re.fullmatch(r'[1-9][0-9]{0,2}', position) is None or
                    not last < int(position) <= MAX_COLUMNS or code not in ('2', '10')):
                return None
            last = int(position)
    return DELIMITERS[int(parts[0])]


def _source_ids(path, delimiter, budget):
    path = Path(path)
    budget.check()
    if path.stat().st_size > MAX_INPUT_BYTES:
        raise _Limited('csv_id_source_byte_budget')
    class Lines:
        def __init__(self, stream):
            self.stream, self.size = stream, 0
        def __iter__(self):
            return self
        def __next__(self):
            budget.check()
            value = self.stream.readline(256 * 1024 + 1)
            if not value:
                raise StopIteration
            self.size += len(value.encode('utf-8'))
            if len(value) > 256 * 1024 or self.size > MAX_INPUT_BYTES:
                raise _Limited('csv_id_source_byte_budget')
            if '\0' in value:
                raise _Limited('csv_id_source_binary')
            return value
    ids, width, rows = {}, None, 0
    with path.open('r', encoding='utf-8-sig', errors='strict', newline='') as stream:
        for record in csv.reader(Lines(stream), delimiter=delimiter, quotechar='"', strict=True):
            budget.check()
            rows += 1
            if rows > MAX_ROWS:
                raise _Limited('csv_id_source_row_budget')
            if width is None:
                width = len(record)
                if not 1 <= width <= MAX_COLUMNS:
                    raise _Limited('csv_id_source_column_budget')
            elif len(record) != width:
                raise _Limited('csv_id_source_row_mapping')
            for column, value in enumerate(record, 1):
                if LEADING_ZERO_ID.fullmatch(value):
                    if len(ids) >= MAX_TRACKED_IDS:
                        raise _Limited('csv_id_candidate_budget')
                    ids[(rows, column)] = value
    return ids, rows, width


def _verify_ooxml_budget(path, budget):
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if not infos or len(infos) > 10000 or len({i.filename for i in infos}) != len(infos):
            raise _Limited('csv_id_xlsx_package_budget')
        if sum(i.file_size for i in infos) > MAX_UNCOMPRESSED:
            raise _Limited('csv_id_xlsx_decompression_budget')
        used = 0
        for info in infos:
            budget.check()
            if info.flag_bits & 1 or info.file_size > MAX_UNCOMPRESSED:
                raise _Limited('csv_id_xlsx_encryption_or_budget')
            if info.filename.lower().endswith(('.xml', '.rels')):
                # Scan before the reader: native output remains untrusted.
                # Include null-normalization for UTF16/32 declarations and
                # a tail for declarations crossing chunk boundaries.
                tail = b''
                with archive.open(info) as stream:
                    while True:
                        budget.check()
                        block = stream.read(65536)
                        if not block:
                            break
                        used += len(block)
                        if used > MAX_UNCOMPRESSED:
                            raise _Limited('csv_id_xlsx_xml_budget')
                        data = tail + block.replace(b'\0', b'').upper()
                        if b'<!DOCTYPE' in data or b'<!ENTITY' in data:
                            raise _Limited('csv_id_xlsx_xml_entities')
                        tail = data[-32:]
    budget.check()


def check_csv_text_ids(paths, target, source, *, import_filter=None,
                       import_options=None, cancel=None, seconds=8.0):
    """Compare original leading-zero strings and actual Excel cell types.

    No cell text goes into error diagnostics. Number format 000 is not proof
    that an identifier survived as a string. No visual/fidelity claim is made.
    """
    target = str(target).lower().lstrip('.')
    delimiter = _import_dialect(import_filter, import_options)
    if target not in SPREADSHEET_TARGETS or delimiter is None:
        return None
    budget = _Budget(cancel, seconds)
    result = {'check': 'csv_text_ids', 'scope': 'source_leading_zero_id_cells_and_types',
              'qualityCompared': False, 'fullDocumentFidelity': False}
    try:
        budget.check()
        source = Path(source)
        candidates = [Path(p) for p in paths if Path(p).suffix.lower() == '.' + target]
        source_stamp = _stamp(source)
        if len(candidates) != 1:
            raise _Limited('csv_id_primary_output_mapping')
        output = candidates[0]
        output_stamp = _stamp(output)
        ids, rows, width = _source_ids(source, delimiter, budget)
        result.update(sourceRows=rows, sourceColumns=width, candidateIds=len(ids))
        if not ids:
            return None
        if output.stat().st_size > MAX_OUTPUT_BYTES:
            raise _Limited('csv_id_output_byte_budget')
        _verify_ooxml_budget(output, budget)
        from openpyxl import load_workbook
        # Own the underlying stream even if workbook construction raises;
        # close the worksheet generator on cancellation/budget exhaustion.
        differences, seen, checked = [], set(), 0
        with output.open('rb') as stream:
            workbook = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
            try:
                if len(workbook.worksheets) != 1:
                    raise _Limited('csv_id_output_sheet_mapping')
                sheet = workbook.worksheets[0]
                max_row = max(row for row, col in ids)
                max_col = max(col for row, col in ids)
                by_row = {}
                for (row, col), text in ids.items():
                    by_row.setdefault(row, []).append((col, text))
                sheet.reset_dimensions()
                with closing(sheet.iter_rows(min_row=1, max_row=max_row, max_col=max_col)) as iterator:
                    for row_number, cells in enumerate(iterator, 1):
                        budget.check()
                        for column, text in by_row.get(row_number, ()):
                            seen.add((row_number, column))
                            cell = cells[column - 1]
                            if cell.data_type != 's' or cell.value != text:
                                differences.append((row_number, column))
                            else:
                                checked += 1
                differences.extend(sorted(set(ids) - seen))
            finally:
                workbook.close()
        budget.check()
        if source_stamp != _stamp(source) or output_stamp != _stamp(output):
            raise GuardError('SOURCE_CHANGED_DURING_PREPARATION',
                             'Исходник или результат изменился во время проверки ID.', stage='result')
        result.update(idsCompared=checked, idsChanged=len(differences))
        if differences:
            details = {'check': 'csv_text_ids', 'candidateIds': len(ids),
                       'differentIds': len(differences),
                       'firstCoordinates': [[row, col] for row, col in differences[:12]]}
            raise GuardError('ENGINE_OUTPUT_CONTENT_LOSS',
                'Текстовые идентификаторы CSV потеряли исходную запись или тип после экспорта.',
                stage='result', diagnostic=json.dumps(details, separators=(',', ':')))
        return {**result, 'status': 'checked_subset'}
    except _Limited as error:
        _cancel(cancel)
        return {**result, 'status': 'not_checked', 'reason': str(error)[:100],
                'warning': 'CSV_TEXT_IDS_NOT_CHECKED'}
    except (OSError, UnicodeError, csv.Error, zipfile.BadZipFile, ImportError,
            ParseError, ValueError, IndexError, KeyError, RuntimeError) as error:
        _cancel(cancel)
        return {**result, 'status': 'not_checked',
                'reason': 'csv_id_inspection_' + type(error).__name__,
                'warning': 'CSV_TEXT_IDS_NOT_CHECKED'}
