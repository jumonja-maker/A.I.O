"""Safe, deliberately narrow merge for simple single-sheet XLSX files.

Never modify inputs. Reject advanced workbook features instead of silently losing
information. The output is written to a new streamed workbook and independently
re-opened before being exposed to the UI.
"""
from __future__ import annotations

from copy import copy
from io import BytesIO
from math import isfinite
from zipfile import ZipFile, BadZipFile
from pathlib import PurePath

import pickle
from pathlib import Path
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils.exceptions import InvalidFileException

MAX_INPUT_FILES = 100
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TOTAL_BYTES = 50 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 150 * 1024 * 1024
MAX_TOTAL_ROWS = 100_000
# A tiny XLSX can declare 100k rows x 16k columns; reject the logical
# rectangle before iter_rows allocates any of its empty cells.
MAX_GRID_CELLS = 1_000_000


class MergeError(ValueError):
    """A safe, user-displayable description of an unsupported input."""


def _check_zip(data: bytes, name: str) -> None:
    if len(data) > MAX_FILE_BYTES:
        raise MergeError(f'«{name}»: файл больше 20 МБ.')
    try:
        with ZipFile(BytesIO(data)) as archive:
            infos = archive.infolist()
            if not infos or len(infos) > 10_000:
                raise MergeError(f'«{name}»: недопустимая структура XLSX.')
            total = sum(x.file_size for x in infos)
            if total > MAX_UNCOMPRESSED_BYTES:
                raise MergeError(f'«{name}»: слишком большой распакованный XLSX.')
            if any(x.file_size > max(1, x.compress_size) * 300 for x in infos):
                raise MergeError(f'«{name}»: файл имеет подозрительно сильное сжатие.')
            if any('..' in PurePath(x.filename).parts or x.filename.startswith('/') for x in infos):
                raise MergeError(f'«{name}»: недопустимый путь внутри архива.')
            for part in ('[Content_Types].xml', 'xl/workbook.xml'):
                if part not in archive.namelist():
                    raise MergeError(f'«{name}»: отсутствует обязательная часть XLSX.')
            if any(x.startswith('xl/externalLinks/') or x.startswith('xl/vbaProject') for x in archive.namelist()):
                raise MergeError(f'«{name}»: внешние связи и макросы не поддерживаются.')
    except (BadZipFile, OSError) as exc:
        raise MergeError(f'«{name}»: файл повреждён или это не XLSX.') from exc


def _read_simple_table(data: bytes, name: str, row_sink=None):
    _check_zip(data, name)
    # Excel has its own OPC safety checks and NEVER imports converter routes.
    # This preserves the workbook merge's safety independently of Guard.
    from .xlsx_integrity import verify_excel_package
    try:
        verify_excel_package(data)
    except ValueError as exc:
        raise MergeError(f'«{name}»: {exc}') from exc
    try:
        wb = load_workbook(BytesIO(data), read_only=False, data_only=False, keep_links=False, rich_text=True)
    except (OSError, ValueError, TypeError, KeyError, InvalidFileException, BadZipFile) as exc:
        raise MergeError(f'«{name}»: невозможно открыть XLSX.') from exc
    try:
        # worksheets excludes chart-only tabs, which would otherwise be lost.
        if len(wb.sheetnames) != 1 or len(wb.worksheets) != 1:
            raise MergeError(f'«{name}»: требуется ровно один лист.')
        ws = wb.worksheets[0]
        if ws.merged_cells.ranges:
            raise MergeError(f'«{name}»: обнаружены объединённые ячейки.')
        if ws.tables or ws._charts or ws._images or ws.data_validations.count:
            raise MergeError(f'«{name}»: сложные элементы таблицы пока не поддерживаются.')
        if ws.auto_filter.ref or ws.protection.sheet or ws.conditional_formatting or ws.print_area:
            raise MergeError(f'«{name}»: фильтры, защита или правила оформления пока не поддерживаются.')
        if any(d.hidden for d in ws.row_dimensions.values()) or any(d.hidden for d in ws.column_dimensions.values()):
            raise MergeError(f'«{name}»: скрытые строки или столбцы не поддерживаются.')
        if wb._external_links or wb.defined_names:
            raise MergeError(f'«{name}»: внешние ссылки или именованные диапазоны не поддерживаются.')
        if ws.max_row > MAX_TOTAL_ROWS:
            raise MergeError(f'«{name}»: слишком много строк для первой версии.')
        if ws.max_row * ws.max_column > MAX_GRID_CELLS:
            raise MergeError(f'«{name}»: слишком большая таблица (более миллиона ячеек).')
        if not ws.max_row or not ws.max_column:
            raise MergeError(f'«{name}»: нет строки заголовков.')
        headers = []
        rows = []
        header_cells = []
        # Source style IDs belong to their own workbook; copying StyleArray
        # IDs into a new workbook silently changes colors/fonts. Capture the
        # actual objects once per distinct source style and re-register them
        # with the output workbook when copying a visible cell.
        from openpyxl.cell.rich_text import CellRichText
        style_cache = {}
        def visible_style(cell):
            style = cell._style
            if style is None or not any((style.fontId, style.fillId,
                                         style.borderId, style.alignmentId,
                                         style.protectionId, style.xfId)):
                return None
            key = cell.style_id
            if key not in style_cache:
                style_cache[key] = tuple(copy(part) for part in
                    (cell.font, cell.fill, cell.border, cell.alignment, cell.protection))
            return key
        for row_i, row in enumerate(ws.iter_rows()):
            content = []
            for cell in row:
                if isinstance(cell.value, CellRichText):
                    raise MergeError(f'«{name}»: разное оформление внутри ячейки (rich text) пока не поддерживается; объединение остановлено без потери стилей.')
                if cell.data_type in ('f', 'e'):
                    raise MergeError(f'«{name}»: формулы и значения-ошибки пока не поддерживаются.')
                if cell.hyperlink or cell.comment:
                    raise MergeError(f'«{name}»: ссылки и комментарии пока не поддерживаются.')
                # Retain visible styles and numeric display formats separately
                # while keeping references bounded by the source's style count.
                content.append((cell.value, cell.number_format, visible_style(cell)))
            if row_i == 0:
                header_cells = content
                headers = [v for v, _, _ in content]
            else:
                if row_sink is None:
                    rows.append(content)
                else:
                    row_sink(content)
        if not headers or any(not isinstance(h, str) or not h.strip() for h in headers):
            raise MergeError(f'«{name}»: заголовки должны быть непустым текстом.')
        if len(set(headers)) != len(headers):
            raise MergeError(f'«{name}»: названия столбцов повторяются.')
        # Column widths are layout guidance, not data; use the widest source
        # width when equally named columns have different settings.
        widths = {key: dim.width for key, dim in ws.column_dimensions.items()
                  if dim.width is not None}
        heights = {index: dim.height for index, dim in ws.row_dimensions.items()
                   if dim.height is not None}
        return (headers, rows, style_cache, header_cells, widths, heights,
                ws.freeze_panes, wb.loaded_theme,
                (ws.sheet_format.defaultRowHeight, ws.sheet_format.defaultColWidth, ws.sheet_format.baseColWidth))
    finally:
        wb.close()


def merge_xlsx(inputs: list[tuple[str, bytes]], progress=None) -> bytes:
    """Merge 2..100 XLSX byte sequences. Returns a validated, new XLSX.

    progress, if provided, receives only phase/current/total counters and
    never sees workbook contents or filenames.
    """
    # Private on-disk staging is bounded by input policy and removed on every
    # exit. Pickle is only written/read in this call, never loaded from uploads.
    with tempfile.TemporaryDirectory(prefix='aio-excel-') as directory:
        return _merge_xlsx(inputs, progress, Path(directory))


def _merge_xlsx(inputs, progress, directory):
    def report(phase, current=0, total=0):
        if progress is not None:
            progress(phase, current, total)
    if not 2 <= len(inputs) <= MAX_INPUT_FILES:
        raise MergeError('Выберите от 2 до 100 файлов XLSX.')
    if sum(len(data) for _, data in inputs) > MAX_TOTAL_BYTES:
        raise MergeError('Общий размер файлов превышает 50 МБ.')
    seen_names = set()
    expected = None
    stages = []
    count = 0
    header_style = None
    header_height = None
    freeze_panes = None
    max_widths = {}
    first_theme = None
    first_has_styles = False
    source_sheet_format = None
    for file_index, (name, data) in enumerate(inputs, start=1):
        report('read', file_index, len(inputs))
        if not isinstance(name, str) or not name.lower().endswith('.xlsx'):
            raise MergeError('Поддерживаются только файлы .xlsx (не .xls).')
        if not isinstance(data, bytes):
            raise MergeError('Неверный формат загруженных данных.')
        if name in seen_names:
            raise MergeError(f'Файл «{name}» выбран повторно.')
        seen_names.add(name)
        stage = directory / str(file_index)
        row_count = 0
        def sink(row):
            nonlocal row_count
            pickle.dump(row, staged, protocol=pickle.HIGHEST_PROTOCOL)
            row_count += 1
        with stage.open('xb') as staged:
            headers, rows, styles, first_row, widths, heights, source_freeze, theme, sheet_format = (
                _read_simple_table(data, name, row_sink=sink))
        if expected is None:
            expected = headers
            header_style = (first_row, styles)
            header_height = heights.get(1)
            freeze_panes = source_freeze
            first_theme = theme
            first_has_styles = bool(styles)
            source_sheet_format = sheet_format
        elif expected != headers:
            raise MergeError(f'«{name}»: названия или порядок столбцов отличаются.')
        if sheet_format != source_sheet_format:
            raise MergeError(f'«{name}»: разные размеры строк или столбцов по умолчанию; безопасное объединение такого оформления пока не поддерживается.')
        # Theme-indexed font/fill colors are not portable between unrelated
        # workbook color themes. Keep the first workbook's theme, and refuse
        # styled workbooks with a different theme rather than silently
        # reinterpreting the second workbook's color indices.
        if theme != first_theme and (styles or first_has_styles):
            raise MergeError(f'«{name}»: разные цветовые темы оформленных книг пока не поддерживаются.')
        for column, width in widths.items():
            max_widths[column] = max(width, max_widths.get(column, 0))
        if count + row_count + 1 > MAX_TOTAL_ROWS:
            raise MergeError('После объединения будет более 100 000 строк.')
        if (count + row_count + 1) * len(expected) > MAX_GRID_CELLS:
            raise MergeError('Объединённая таблица превысит предел миллиона ячеек.')
        stages.append((stage, row_count, styles, heights))
        count += row_count
    # The complete source workbook and its cell objects have been closed and
    # released. Only one staged row and the distinct style tables stay live.
    rows = styles = first_row = None
    out = Workbook(write_only=True)
    if first_theme:
        out.loaded_theme = first_theme
    ws = out.create_sheet('Объединённая таблица')
    ws.title = 'Объединённая таблица'
    default_height, default_width, base_col_width = source_sheet_format
    if default_height is not None:
        ws.sheet_format.defaultRowHeight = default_height
    if default_width is not None:
        ws.sheet_format.defaultColWidth = default_width
    if base_col_width is not None:
        ws.sheet_format.baseColWidth = base_col_width
    def apply_style(cell, metadata):
        if metadata is not None:
            cell.font, cell.fill, cell.border, cell.alignment, cell.protection = metadata
    if header_height is not None:
        ws.row_dimensions[1].height = header_height
    if freeze_panes is not None:
        ws.freeze_panes = freeze_panes
    for column, width in max_widths.items():
        ws.column_dimensions[column].width = width
    def output_cells(row, styles):
        cells = []
        for value, number_format, key in row:
            # An explicit empty inline cell preserves trailing column width
            # when a dimension-free streaming reader opens the workbook.
            cell = WriteOnlyCell(ws, value='' if value is None else value)
            if isinstance(value, str):
                cell.data_type = 's'  # Literal '=...' stays text.
            if isinstance(value, float):
                if not isfinite(value):
                    raise MergeError('В таблице обнаружено недопустимое число.')
                cell.value, cell.data_type = repr(value), 'n'
            if number_format and number_format != 'General':
                cell.number_format = number_format
            if key is not None:
                apply_style(cell, styles[key])
            cells.append(cell)
        return cells
    ws.append(output_cells(header_style[0], header_style[1]))
    report('write', 0, count)
    written = 0
    try:
        for stage, source_count, styles, heights in stages:
            with stage.open('rb') as staged:
                for source_row in range(2, source_count + 2):
                    row = pickle.load(staged)
                    if heights.get(source_row) is not None:
                        ws.row_dimensions[written + 2].height = heights[source_row]
                    ws.append(output_cells(row, styles))
                    written += 1
                    if written % 1024 == 0 or written == count:
                        report('write', written, count)
            stage.unlink()
    except BaseException:
        # Close streaming XML writers if serialization aborts mid-row.
        ws.close()
        out.close()
        raise
    report('save')
    stream = BytesIO()
    out.save(stream)
    out.close()
    blob = stream.getvalue()
    # Independently reopen the result and check the required logical structure.
    report('verify')
    check = load_workbook(BytesIO(blob), read_only=True, data_only=False)
    try:
        sheet = check.active
        # write_only deliberately omits worksheet dimensions; validate by
        # streaming all rows instead of trusting max_row from a header.
        actual_rows = sheet.values
        header = list(next(actual_rows))
        if sum(1 for _ in actual_rows) != count or header != expected:
            raise RuntimeError('Внутренняя ошибка проверки созданного файла.')
    finally:
        check.close()
    report('done', count, count)
    return blob
