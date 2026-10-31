"""V21 real HTTP lifecycle and independent Excel failure-cleanup regressions."""
from contextlib import contextmanager
import http.client
from io import BytesIO
import json
from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet._writer import ALL_TEMP_FILES
from app_backend.converter_guard import ConverterGuard, GuardLimits
from app_backend.converter_guard.errors import GuardError
from app_backend.converter_guard.http_api import PREFIX, _CancelableInput
from app_backend.excel import merge_xlsx
from test_converter_guard_http_v02 import start_test_server
from test_converter_guard_v02 import FakeAdapter


class GuardBoundaryLifecycle(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.guard = ConverterGuard(adapter=FakeAdapter(), spool_parent=self.folder.name,
            limits=GuardLimits(upload_idle_timeout=.15))
        self.server, self.thread, self.origin, self.token = start_test_server(self.guard)

    def tearDown(self):
        self.server.api.close(); self.guard.close()
        self.server.shutdown(); self.server.server_close(); self.thread.join(3)
        self.folder.cleanup()

    def request(self, path, body):
        connection = http.client.HTTPConnection('127.0.0.1', self.server.server_port, timeout=3)
        try:
            connection.request('POST', PREFIX + path, json.dumps(body).encode(),
                {'X-AIO-Token':self.token, 'Content-Type':'application/json'})
            response = connection.getresponse()
            return response.status, json.loads(response.read())
        finally:
            connection.close()

    def upload(self, job, *, path=None, size=4, first=b'x'):
        connection = http.client.HTTPConnection('127.0.0.1', self.server.server_port, timeout=.6)
        endpoint = '/input?job=' + job + '&file=f'
        if path is not None: endpoint = '/resource?job=' + job + '&file=f&path=' + path
        connection.putrequest('POST', PREFIX + endpoint)
        connection.putheader('X-AIO-Token', self.token)
        connection.putheader('Content-Type', 'application/octet-stream')
        connection.putheader('Content-Length', str(size))
        connection.endheaders(first)
        return connection

    def test_stalled_input_expires_idle_without_holding_slot_and_can_retry(self):
        _, state = self.request('/jobs', {'files':[{'id':'f','name':'one.txt','size':4}]})
        identity = state['jobId']
        client = self.upload(identity)
        try:
            response = client.getresponse()
            self.assertEqual(response.status, 408)
            self.assertEqual(json.loads(response.read())['code'], 'UPLOAD_IDLE_TIMEOUT')
        finally:
            client.close()
        job = self.guard._jobs[identity]
        self.assertEqual(job.io, 0)
        self.assertFalse(job.files[0].uploading)
        self.assertEqual(job.files[0].state, 'awaiting_upload')
        self.assertFalse(list(job.root.rglob('*.part')))
        completed = self.upload(identity, first=b'ABCD')
        try:
            response = completed.getresponse(); response.read()
            self.assertEqual(response.status, 200)
        finally: completed.close()
        self.assertEqual(job.files[0].source.read_bytes(), b'ABCD')

    def test_continuous_upload_can_exceed_idle_window_without_total_timeout(self):
        # A deterministic monotonic clock proves this is an idle deadline, not
        # a wall-clock file size/speed veto; it does not depend on runner load.
        clock = [0.0]
        class Connection:
            def recv(self, _size): return b'x'
        connection = Connection()
        def ready(*_args):
            clock[0] += .1
            return [connection], [], []
        with patch('app_backend.converter_guard.http_api.time.monotonic', side_effect=lambda:clock[0]), \
                patch('app_backend.converter_guard.http_api.select.select', side_effect=ready):
            reader = _CancelableInput(connection, threading.Event(), idle_timeout=.15)
            self.assertEqual(b''.join(reader.read(1) for _ in range(10)),b'x'*10)
        self.assertGreater(clock[0],.9)

    def test_stalled_html_resource_expires_and_releases_its_ownership(self):
        identity = self.guard.create([{'id':'f','name':'one.html','size':1}], prepare_only=True)['jobId']
        self.guard.upload(identity, 'f', BytesIO(b'x'), 1)
        self.guard.configure_resources(identity, 'f', [{'path':'image.png','size':4}])
        client = self.upload(identity, path='image.png')
        try:
            response = client.getresponse()
            self.assertEqual(response.status, 408)
            self.assertEqual(json.loads(response.read())['code'], 'UPLOAD_IDLE_TIMEOUT')
        finally:
            client.close()
        item = self.guard._jobs[identity].files[0]
        self.assertFalse(item.resources_uploading)
        self.assertFalse(item.resource_interrupts)
        self.assertFalse(item.resources['image.png']['ready'])
        self.assertFalse(list(item.source.parent.glob('.aio-resource-*')))
        self.assertEqual(self.guard._jobs[identity].io, 0)





def workbook_bytes(value):
    workbook = Workbook(); workbook.active.append(['code', 'value'])
    workbook.active.append(['00123', value]); result = BytesIO()
    workbook.save(result); workbook.close()
    return result.getvalue()


class ExcelFailureCleanup(unittest.TestCase):
    def setUp(self):
        self.inputs=[('a.xlsx',workbook_bytes(1)),('b.xlsx',workbook_bytes(2))]
        self.initial = set(ALL_TEMP_FILES)

    def tearDown(self):
        # Cleanup leaked baseline artifacts even on an intentionally red test.
        for name in tuple(ALL_TEMP_FILES):
            if name not in self.initial:
                Path(name).unlink(missing_ok=True); ALL_TEMP_FILES.remove(name)

    def assert_no_staged_xml(self):
        self.assertEqual(set(ALL_TEMP_FILES)-self.initial,set())

    def test_failed_excel_save_removes_private_streaming_xml(self):
        with patch.object(Workbook,'save',side_effect=OSError('disk full')):
            with self.assertRaisesRegex(OSError,'disk full'): merge_xlsx(self.inputs)
        self.assert_no_staged_xml()

    def test_failed_excel_progress_before_rows_cleans_header_writer(self):
        def progress(phase, *_):
            if phase == 'write': raise RuntimeError('observer failed')
        with self.assertRaisesRegex(RuntimeError,'observer failed'):
            merge_xlsx(self.inputs,progress=progress)
        self.assert_no_staged_xml()

    def test_failed_excel_progress_after_rows_cleans_output_writer(self):
        def progress(phase, current=0, total=0):
            if phase == 'write' and current: raise RuntimeError('observer failed')
        with self.assertRaisesRegex(RuntimeError,'observer failed'):
            merge_xlsx(self.inputs,progress=progress)
        self.assert_no_staged_xml()

    def test_failure_cleanup_never_deletes_another_workbooks_live_xml(self):
        foreign = Workbook(write_only=True); foreign_sheet = foreign.create_sheet('unrelated')
        foreign_sheet.append(['belongs to another operation'])
        owned = Path(foreign_sheet._writer.out)
        try:
            with patch.object(Workbook,'save',side_effect=OSError('disk full')):
                with self.assertRaises(OSError): merge_xlsx(self.inputs)
            self.assertTrue(owned.is_file())
            self.assertEqual(set(ALL_TEMP_FILES)-self.initial,{str(owned)})
        finally:
            foreign_sheet.close(); foreign_sheet._writer.cleanup(); foreign.close()

    def test_successful_excel_merge_remains_independent_and_preserves_values(self):
        with patch('app_backend.converter_guard.core.ConverterGuard.create',side_effect=AssertionError('Excel must remain independent')):
            result=merge_xlsx(self.inputs)
        with BytesIO(result) as buffer:
            workbook=load_workbook(buffer,read_only=True)
            try:
                self.assertEqual(list(workbook.active.values),[('code','value'),('00123',1),('00123',2)])
            finally: workbook.close()
        self.assert_no_staged_xml()


if __name__ == '__main__': unittest.main()
