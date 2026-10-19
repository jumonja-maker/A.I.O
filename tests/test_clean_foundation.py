"""Migration regression: UI, Excel and converter boundary remain independent."""
import io
import json
import threading
import unittest
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from openpyxl import Workbook, load_workbook
from app_backend.server import start_server, asset_root
from app_backend.converter_assets import read_designer_asset


def xlsx(value):
    book = Workbook()
    book.active.append(['Код', 'Значение'])
    book.active.append([value, 1])
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def multipart(files):
    boundary = 'aio-merge-boundary-42'
    sections = []
    for name, content in files:
        sections.append((
            '--' + boundary + '\r\n'
            'Content-Disposition: form-data; name="files"; filename="' + name + '"\r\n'
            'Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n'
            '\r\n').encode('ascii') + content + b'\r\n')
    return b''.join(sections) + ('--' + boundary + '--\r\n').encode('ascii'), boundary


class CleanFoundationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server, cls.token = start_server()
        cls.thread = threading.Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.base = 'http://127.0.0.1:' + str(cls.server.server_port)

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join(timeout=5)

    def fetch(self, path, token=None, data=None, headers=None):
        hdrs = headers or {}
        if token:
            hdrs['X-AIO-Token'] = token
        return urlopen(Request(self.base + path, data=data, headers=hdrs),
                       timeout=15)

    def test_shell_and_elva_load(self):
        html = self.fetch('/').read().decode('utf-8')
        self.assertIn('src/excel-tool.js', html)
        self.assertIn('src/converter-bridge.js', html)
        self.assertIn('elva-v13-happy-thumb.webp', html)
        self.assertIn('data-aio-token=', html)
        self.assertGreater(len(self.fetch('/assets/brand/elva-v13-happy-thumb.webp').read()), 2000)

    def test_designer_source_is_hash_checked_and_renderable(self):
        js = read_designer_asset(str(asset_root()), 'file-converter.js')
        css = read_designer_asset(str(asset_root()), 'file-converter.css')
        self.assertGreater(len(js), 25000)
        self.assertIn(b'registerTool', js)
        self.assertGreater(len(css), 9000)
        self.assertEqual(self.fetch('/addons/file-converter/file-converter.js').read(), js)

    def test_guard_boundary_is_authenticated_and_reports_coverage(self):
        with self.assertRaises(HTTPError) as ctx:
            self.fetch('/api/converter/status')
        self.assertEqual(ctx.exception.code, 403)
        result = json.loads(self.fetch('/api/converter/status', token=self.token).read())
        self.assertEqual(result['outputs'], ['pdf'])
        self.assertEqual(result['storage']['jobs'], 0)

    def test_authenticated_excel_merges_and_retains_literal_text(self):
        payload, boundary = multipart([
            ('a.xlsx', xlsx('00125')),
            ('b.xlsx', xlsx('00256')),
        ])
        response = self.fetch('/api/merge', token=self.token, data=payload,
            headers={'Content-Type':'multipart/form-data; boundary=' + boundary,
                     'X-AIO-Job':'a' * 32})
        self.assertEqual(response.status, 200)
        result = response.read()
        book = load_workbook(io.BytesIO(result))
        try:
            self.assertEqual(list(book.active.values),
                             [('Код', 'Значение'),('00125',1),('00256',1)])
        finally:
            book.close()
        state = json.loads(self.fetch('/api/merge-progress?job=' + 'a'*32,
            token=self.token).read())
        self.assertEqual(state['phase'],'done')

    def test_excel_requires_token_and_rejects_external_post(self):
        data, boundary = multipart([('a.xlsx', xlsx('a')),('b.xlsx',xlsx('b'))])
        with self.assertRaises(HTTPError) as err:
            self.fetch('/api/merge', data=data,
                headers={'Content-Type':'multipart/form-data; boundary='+boundary,
                         'X-AIO-Job':'a'*32})
        self.assertEqual(err.exception.code,403)
        with self.assertRaises(HTTPError) as err:
            self.fetch('/../../app_backend/excel.py')
        self.assertEqual(err.exception.code,404)


if __name__ == '__main__':
    unittest.main()
