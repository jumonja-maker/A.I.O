"""Deterministic XLSX engine checks. Run: python -m unittest discover -s tests."""
import io
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree
import unittest
from time import perf_counter
from datetime import date

from openpyxl import Workbook, load_workbook
from app_backend.excel import MergeError, merge_xlsx


def xlsx(headers=('Код', 'Цена'), rows=(), formula=False, bad_sheet=False, merged=False):
    w = Workbook()
    s = w.active
    s.append(headers)
    for row in rows:
        s.append(row)
    if formula:
        s['B2'] = '=SUM(1,2)'
    if merged:
        s.merge_cells('A3:B3')
    if bad_sheet:
        w.create_sheet('Второй')
    out = io.BytesIO()
    w.save(out)
    return out.getvalue()


def values(payload):
    w = load_workbook(io.BytesIO(payload), data_only=False)
    try:
        return list(w.active.values)
    finally:
        w.close()


class ExcelMergeTests(unittest.TestCase):
    def test_merge_rows_and_preserve_sources(self):
        a = xlsx(rows=[('00123', 22), ('0022', 35)])
        b = xlsx(rows=[('444', 9)])
        original = (a, b)
        out = merge_xlsx([('январь.xlsx', a), ('февраль.xlsx', b)])
        self.assertEqual(values(out), [('Код', 'Цена'), ('00123', 22), ('0022', 35), ('444', 9)])
        self.assertEqual(original, (a, b))

    def test_merge_preserves_basic_styles_dimensions_and_freeze_without_mixing_style_ids(self):
        from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
        def styled(value, *, width, color, height):
            book = Workbook()
            sheet = book.active
            sheet.append(("Код", "Цена"))
            sheet.append((value, 64.01))
            sheet["A1"].font = Font(bold=True, color="FF223344")
            sheet["A1"].fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")
            sheet["A1"].border = Border(bottom=Side(style="thin", color="FF112233"))
            sheet["A1"].alignment = Alignment(horizontal="center")
            sheet["A2"].font = Font(bold=True, color=color)
            sheet["A2"].alignment = Alignment(horizontal="right")
            sheet["B2"].number_format = "0.00"
            sheet.column_dimensions["A"].width = width
            sheet.row_dimensions[2].height = height
            sheet.freeze_panes = "A2"
            buffer = io.BytesIO()
            book.save(buffer)
            book.close()
            return buffer.getvalue()
        a = styled("00125", width=24, color="FFFF0000", height=25)
        b = styled("00777", width=32, color="FF0000FF", height=29)
        result = merge_xlsx([("a.xlsx", a), ("b.xlsx", b)])
        self.assertEqual(values(result), [
            ("Код", "Цена"), ("00125", 64.01), ("00777", 64.01)])
        book = load_workbook(io.BytesIO(result), data_only=False)
        try:
            sheet = book.active
            self.assertEqual(sheet["A1"].font.bold, True)
            self.assertEqual(sheet["A1"].fill.fgColor.rgb, "FFFFFF00")
            self.assertEqual(sheet["A1"].alignment.horizontal, "center")
            self.assertEqual(sheet["A1"].border.bottom.style, "thin")
            self.assertEqual(sheet["A2"].font.color.rgb, "FFFF0000")
            self.assertEqual(sheet["A3"].font.color.rgb, "FF0000FF")
            self.assertEqual(sheet["A3"].alignment.horizontal, "right")
            self.assertEqual(sheet["B2"].number_format, "0.00")
            self.assertEqual(sheet["B3"].number_format, "0.00")
            self.assertEqual(sheet.row_dimensions[2].height, 25)
            self.assertEqual(sheet.row_dimensions[3].height, 29)
            self.assertEqual(sheet.column_dimensions["A"].width, 32)
            self.assertEqual(sheet.freeze_panes, "A2")
        finally:
            book.close()
        self.assertNotEqual(a, result)
        self.assertNotEqual(b, result)

    def test_literal_formula_like_strings_do_not_become_formulas(self):
        def literal_source():
            book = Workbook()
            sheet = book.active
            sheet["A1"] = "=Заголовок"
            sheet["A1"].data_type = "s"
            sheet["B1"] = "Код"
            sheet["A2"] = "=1+1"
            sheet["A2"].data_type = "s"
            sheet["B2"] = "00125"
            payload = io.BytesIO()
            book.save(payload)
            book.close()
            return payload.getvalue()

        src = literal_source()
        result = merge_xlsx([("a.xlsx", src), ("b.xlsx", src)])
        book = load_workbook(io.BytesIO(result), data_only=False)
        try:
            self.assertEqual(book.active["A1"].value, "=Заголовок")
            self.assertEqual(book.active["A1"].data_type, "s")
            self.assertEqual(book.active["A2"].value, "=1+1")
            self.assertEqual(book.active["A2"].data_type, "s")
            self.assertEqual(book.active["A3"].data_type, "s")
            self.assertEqual(book.active["B2"].value, "00125")
        finally:
            book.close()

    def test_merge_rejects_encoded_external_xlsx_relationship(self):
        source = xlsx(rows=[("a", 1)])
        hostile = io.BytesIO()
        with (ZipFile(io.BytesIO(source)) as original,
              ZipFile(hostile, "w", ZIP_DEFLATED) as destination):
            for name in original.namelist():
                data = original.read(name)
                if name == "xl/_rels/workbook.xml.rels":
                    data = data.replace(
                        b"</Relationships>",
                        b'<Relationship Id="rId999" TargetMode="External" '
                        b'Target="h&#x74;tp://127.0.0.1:9/canary"/>'
                        b"</Relationships>")
                destination.writestr(name, data)
        with self.assertRaisesRegex(MergeError, 'Внешние'):
            merge_xlsx([("hostile.xlsx", hostile.getvalue()),
                        ("other.xlsx", source)])

    def test_sparse_billion_cell_rectangle_rejected_early(self):
        book = Workbook()
        book.active["A1"] = "Код"
        book.active["XFD100000"] = "marker"
        payload = io.BytesIO()
        book.save(payload)
        book.close()
        with self.assertRaisesRegex(MergeError, "миллиона ячеек"):
            merge_xlsx([("a.xlsx", payload.getvalue()),
                        ("b.xlsx", xlsx(headers=("Код",), rows=[]))])

    def test_aggregate_grid_budget_applies_before_output(self):
        from unittest.mock import patch
        source = xlsx(headers=("Код", "Цена"),
                      rows=[("a", 1), ("b", 2)])
        with patch("app_backend.excel.MAX_GRID_CELLS", 9):
            with self.assertRaisesRegex(MergeError, "миллиона ячеек"):
                merge_xlsx([("a.xlsx", source), ("b.xlsx", source)])

    def test_float_prices_stay_numeric_without_binary_decimal_tails(self):
        # User reproduced 64.01 => 64.01000000000001 in generated XML.
        prices = [64.01, 65.74, 70.93, 513.81, 524.19, 534.57,
                  0.1, 0.00000017, -89.96, 123456.789012345]
        a = xlsx(headers=('Артикул', 'Цена'),
                 rows=[(f'{i:05d}', n) for i, n in enumerate(prices, start=1)])
        b = xlsx(headers=('Артикул', 'Цена'), rows=[])
        out = merge_xlsx([('a.xlsx', a), ('b.xlsx', b)])
        with ZipFile(io.BytesIO(out)) as archive:
            tree = ElementTree.fromstring(archive.read('xl/worksheets/sheet1.xml'))
        cells = {cell.get('r'): cell for cell in tree.findall('.//{*}c')}
        for index, number in enumerate(prices, start=2):
            with self.subTest(number=number):
                price_cell = cells[f'B{index}']
                self.assertEqual(price_cell.get('t', 'n'), 'n')
                self.assertEqual(price_cell.find('{*}v').text, repr(number))
        book = load_workbook(io.BytesIO(out), read_only=True, data_only=True)
        try:
            self.assertEqual(
                [row[1] for row in list(book.active.values)[1:]], prices)
            self.assertTrue(all(isinstance(row[1], float)
                                for row in list(book.active.values)[1:]))
        finally:
            book.close()

    def test_progress_reports_actual_files_and_output_rows(self):
        stages = []
        a = xlsx(rows=[('a', 64.01)])
        b = xlsx(rows=[('b', 65.74)])
        merge_xlsx([('a.xlsx', a), ('b.xlsx', b)],
                   progress=lambda phase, current, total: stages.append(
                       (phase, current, total)))
        self.assertEqual(stages[0], ('read', 1, 2))
        self.assertEqual(stages[1], ('read', 2, 2))
        self.assertIn(('write', 0, 2), stages)
        self.assertIn(('write', 2, 2), stages)
        self.assertEqual([item[0] for item in stages[-3:]],
                         ['save', 'verify', 'done'])

    def test_ten_thousand_rows_with_eight_columns_fast_and_exact(self):
        # Regression: repeatedly reading Worksheet.max_row for each cell took
        # many minutes on a normal 10k-row merge. This must remain usable.
        headers = tuple(f'Поле {i}' for i in range(8))
        first_rows = [
            (f'{i:010d}', i, 'Москва', f'Товар {i}', i * 0.25,
             i % 17, 'Тест кириллицы', None)
            for i in range(1, 5001)
        ]
        second_rows = [
            (f'{i:010d}', i, 'Казань', f'Товар {i}', i * 0.25,
             i % 17, 'Тест кириллицы', None)
            for i in range(5001, 10001)
        ]
        a, b = xlsx(headers, first_rows), xlsx(headers, second_rows)
        started = perf_counter()
        merged = merge_xlsx([('Москва.xlsx', a), ('Казань.xlsx', b)])
        elapsed = perf_counter() - started
        self.assertLess(elapsed, 45.0,
                        f'10k-row merge exceeded 45 seconds: {elapsed:.1f}s')
        book = load_workbook(io.BytesIO(merged), read_only=True)
        try:
            result = book.active.iter_rows(values_only=True)
            self.assertEqual(next(result), headers)
            for expected, actual in zip(first_rows + second_rows, result):
                self.assertEqual(actual, expected)
            self.assertEqual(next(result, None), None)
        finally:
            book.close()
        print(f'EXCEL_PERFORMANCE_10000_ROWS: {elapsed:.2f}s; '
              '10000 rows x 8 cells exact')

    def test_empty_table_with_headers(self):
        a = xlsx(rows=[])
        b = xlsx(rows=[('Тест', 42)])
        self.assertEqual(values(merge_xlsx([('a.xlsx', a), ('b.xlsx', b)])), [('Код', 'Цена'), ('Тест', 42)])

    def test_date_and_number_format(self):
        def make():
            w = Workbook(); s = w.active
            s.append(['Дата', 'Номер'])
            s.append([date(2026, 10, 9), 123])
            s['B2'].number_format = '000000'
            buf = io.BytesIO(); w.save(buf)
            return buf.getvalue()
        out = merge_xlsx([('a.xlsx', make()), ('b.xlsx', xlsx(headers=('Дата', 'Номер'), rows=[]))])
        wb = load_workbook(io.BytesIO(out))
        self.assertEqual(wb.active['A2'].value.date(), date(2026, 10, 9))
        self.assertEqual(wb.active['B2'].number_format, '000000')
        wb.close()

    def test_structure_error(self):
        with self.assertRaisesRegex(MergeError, 'отличаются'):
            merge_xlsx([('a.xlsx', xlsx()), ('b.xlsx', xlsx(headers=('Цена', 'Код')))])

    def test_extra_chart_sheet_is_rejected_in_either_tab_order(self):
        from openpyxl.chart import BarChart, Reference

        for position in (0, 1):
            with self.subTest(chart_tab_position=position):
                book = Workbook()
                sheet = book.active
                sheet.append(['Код', 'Цена'])
                sheet.append(['001', 42])
                chart = BarChart()
                chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2),
                               titles_from_data=True)
                book.create_chartsheet('Диаграмма', position).add_chart(chart)
                stream = io.BytesIO()
                book.save(stream)
                book.close()
                original = stream.getvalue()
                stages = []
                with self.assertRaisesRegex(MergeError, 'ровно один лист'):
                    merge_xlsx([('с_диаграммой.xlsx', original),
                                ('b.xlsx', xlsx())],
                               progress=lambda phase, *args: stages.append(phase))
                self.assertNotIn('save', stages)
                self.assertEqual(stream.getvalue(), original)

    def test_reject_unsupported(self):
        cases = [
            ('формулы', xlsx(rows=[('a', 1)], formula=True)),
            ('листы', xlsx(bad_sheet=True)),
            ('объединённые', xlsx(rows=[('a', 2), ('b', 3)], merged=True)),
        ]
        for name, obj in cases:
            with self.subTest(name=name), self.assertRaises(MergeError):
                merge_xlsx([('a.xlsx', obj), ('b.xlsx', xlsx())])
        with self.assertRaisesRegex(MergeError, 'объединённые ячейки'):
            merge_xlsx([('объединено.xlsx', xlsx(rows=[('a', 2), ('b', 3)], merged=True)),
                        ('b.xlsx', xlsx())])
        with self.assertRaisesRegex(MergeError, 'только файлы .xlsx'):
            merge_xlsx([('a.xls', xlsx()), ('b.xlsx', xlsx())])
        with self.assertRaisesRegex(MergeError, 'повреждён'):
            merge_xlsx([('a.xlsx', b'not an xlsx'), ('b.xlsx', xlsx())])

    def test_limit_and_duplicate(self):
        with self.assertRaisesRegex(MergeError, 'от 2 до 100'):
            merge_xlsx([('one.xlsx', xlsx())])
        with self.assertRaisesRegex(MergeError, 'повторно'):
            merge_xlsx([('same.xlsx', xlsx()), ('same.xlsx', xlsx())])

    def test_no_filesystem_writes(self):
        self.assertIsInstance(merge_xlsx([('a.xlsx', xlsx()), ('b.xlsx', xlsx())]), bytes)


if __name__ == '__main__':
    unittest.main()
