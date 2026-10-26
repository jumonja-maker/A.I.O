"""Concrete before/after losses, limitations and deliberately unknowns."""
import io
from copy import copy
from datetime import datetime
import json
from pathlib import Path
import struct
import tempfile
import threading
import unittest
from unittest.mock import patch
import zipfile
import zlib

from app_backend.converter_guard.comparison_analysis import (
    analyze_pair, collect_source_summary, preview_text, MAX_TEXT)
from app_backend.converter_guard.errors import GuardError


def docx(path, text='Привет', cells=None):
    table = ''
    if cells is not None:
        table = '<w:tbl>' + ''.join('<w:tr>' + ''.join(
            '<w:tc><w:p><w:r><w:t>'+str(value)+'</w:t></w:r></w:p></w:tc>'
            for value in row) + '</w:tr>' for row in cells) + '</w:tbl>'
    xml = ('<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           '<w:body><w:p><w:r><w:t>'+text+'</w:t></w:r></w:p>'+table+'</w:body></w:document>')
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('word/document.xml', xml)
        archive.writestr('word/styles.xml', '<styles/>')


def png(path, *, width=2, height=2, alpha=True, frames=1):
    def chunk(name, data):
        return struct.pack('>I', len(data))+name+data+struct.pack('>I', zlib.crc32(name+data)&0xffffffff)
    data = b'\x89PNG\r\n\x1a\n'+chunk(b'IHDR', struct.pack('>IIBBBBB', width, height, 8, 6 if alpha else 2, 0, 0, 0))
    colour = 4 if alpha else 3
    compressed = zlib.compress(b''.join(b'\0'+bytes([0]*(width*colour)) for _ in range(height)))
    if frames > 1:
        data += chunk(b'acTL', struct.pack('>II', frames, 0))
    sequence = 0
    for frame in range(frames):
        if frames > 1:
            data += chunk(b'fcTL', struct.pack('>IIIIIHHBB', sequence, width, height, 0, 0, 1, 10, 0, 0))
            sequence += 1
        if frame == 0:
            data += chunk(b'IDAT', compressed)
        else:
            data += chunk(b'fdAT', struct.pack('>I', sequence)+compressed)
            sequence += 1
    path.write_bytes(data+chunk(b'IEND', b''))


def bmp(path):
    pixels = b'\0'*16
    dib = struct.pack('<IiiHHIIiiII', 40, 2, 2, 1, 24, 0, len(pixels), 0, 0, 0, 0)
    path.write_bytes(b'BM'+struct.pack('<IHHI', 54+len(pixels), 0, 0, 54)+dib+pixels)


def pptx(path, slides):
    p = 'http://schemas.openxmlformats.org/presentationml/2006/main'
    a = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    r = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('ppt/presentation.xml', '<p:presentation xmlns:p="'+p+'" xmlns:r="'+r+'">'
            '<p:sldIdLst>'+''.join('<p:sldId id="'+str(i+1)+'" r:id="r'+str(i)+'"/>' for i in range(len(slides)))+
            '</p:sldIdLst></p:presentation>')
        archive.writestr('ppt/_rels/presentation.xml.rels', '<Relationships>'+''.join(
            '<Relationship Id="r'+str(i)+'" Type="'+r+'/slide" Target="slides/slide'+str(i)+'.xml"/>'
            for i in range(len(slides)))+'</Relationships>')
        for i, text in enumerate(slides):
            archive.writestr('ppt/slides/slide'+str(i)+'.xml', '<p:sld xmlns:p="'+p+'" xmlns:a="'+a+'">'
                '<p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:t>'+text+'</a:t></a:r></a:p>'
                '</p:txBody></p:sp></p:spTree></p:cSld></p:sld>')


class ComparisonAnalysisV19(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def path(self, name):
        return self.root/name

    def test_intact_word_to_text_checks_only_normalized_text_and_explains_style_limit(self):
        source, output = self.path('source.docx'), self.path('result.txt')
        docx(source, 'Тест e\u0301')
        output.write_text('Тест é', encoding='utf-8')
        result = analyze_pair(source, [output], 'txt')
        self.assertEqual(result['status'], 'format_limitation')
        self.assertIn('normalized_text', result['checkedFeatures'])
        self.assertIn('visual_layout', result['unknown'])
        self.assertEqual(result['recipeCandidates'], [])
        self.assertNotIn(str(self.root), json.dumps(result, ensure_ascii=False))

    def test_missing_cells_are_confirmed_with_positions_not_a_count_only(self):
        source, output = self.path('source.docx'), self.path('result.docx')
        docx(source, cells=[['Name', 'Value'], ['ALPHA', '42']])
        docx(output, cells=[['Name'], ['ALPHA']])
        result = analyze_pair(source, [output], 'docx')
        self.assertEqual(result['status'], 'confirmed_loss')
        found = next(f for f in result['findings'] if f['code'] == 'CELL_VALUES_CHANGED')
        self.assertEqual(found['locations'][1], {'row': 2, 'column': 2, 'before': '42', 'after': ''})
        self.assertEqual(found['totalAffected'], 2)

    def test_csv_values_are_compared_without_losing_significant_spaces(self):
        source, output = self.path('source.csv'), self.path('result.csv')
        source.write_text('key,value\nSKU, a  b \n', encoding='utf-8')
        output.write_text('key,value\nSKU,a b\n', encoding='utf-8')
        result = analyze_pair(source, [output], 'csv')
        self.assertEqual(result['status'], 'confirmed_loss')
        self.assertIn('csv_preserve_values', result['recipeCandidates'])
        found = next(f for f in result['findings'] if f['code'] == 'CELL_VALUES_CHANGED')
        self.assertEqual(found['locations'][0]['before'], ' a  b ')

    def test_numeric_formatting_is_not_claimed_as_changed_value(self):
        source, output = self.path('source.csv'), self.path('result.csv')
        source.write_text('key,value\nA,001.00\n', encoding='utf-8')
        output.write_text('key,value\nA,1\n', encoding='utf-8')
        result = analyze_pair(source, [output], 'csv')
        self.assertEqual(result['status'], 'uncertain')
        self.assertFalse(any(f['code'] == 'CELL_VALUES_CHANGED' for f in result['findings']))
        self.assertIn('CELL_REPRESENTATION_CHANGED', [f['code'] for f in result['findings']])

    def test_text_identifier_leading_zeroes_remain_uncertain(self):
        source, output = self.path('source.csv'), self.path('result.csv')
        source.write_text('key,value\nA,00123\n', encoding='utf-8')
        output.write_text('key,value\nA,123\n', encoding='utf-8')
        result = analyze_pair(source, [output], 'csv')
        self.assertEqual(result['status'], 'uncertain')
        self.assertNotIn('cell_values_single_grid', result['checkedFeatures'])
        found = next(f for f in result['findings'] if f['code'] == 'CELL_REPRESENTATION_CHANGED')
        self.assertEqual(found['locations'][0]['before'], '00123')
        self.assertEqual(found['locations'][0]['after'], '123')
        self.assertNotIn('значение совпадает', found['message'])

    def test_localized_numbers_never_claim_equality_without_locale_and_precision_proof(self):
        from openpyxl import Workbook
        source, output = self.path('source.xlsx'), self.path('result.csv')
        book = Workbook(); sheet = book.active
        sheet.append(['key', 'value']); sheet.append(['A', 1.23456789012345])
        book.save(source); book.close()
        for display in ('1,23', '1,23456789012345', '€1,23', '1 234,56'):
            with self.subTest(display=display):
                output.write_text('key,value\nA,"'+display+'"\n', encoding='utf-8')
                result = analyze_pair(source, [output], 'csv')
                self.assertEqual(result['status'], 'uncertain')
                self.assertNotIn('cell_values_single_grid', result['checkedFeatures'])
                self.assertFalse(any(f['category'] == 'confirmed_loss' for f in result['findings']))
                found = next(f for f in result['findings'] if f['code'] == 'CELL_REPRESENTATION_CHANGED')
                self.assertEqual(found['locations'][0]['before'], '1.23456789012345')
                self.assertEqual(found['locations'][0]['after'], display)
                self.assertIn('точности требует проверки', found['message'])
                self.assertNotIn('значение совпадает', found['message'])
                self.assertIn('csv_preserve_values', result['recipeCandidates'])

    def test_different_numeric_value_is_confirmed(self):
        source, output = self.path('source.csv'), self.path('result.csv')
        source.write_text('key,value\nA,42\n', encoding='utf-8')
        output.write_text('key,value\nA,0\n', encoding='utf-8')
        self.assertEqual(analyze_pair(source, [output], 'csv')['status'], 'confirmed_loss')

    def test_long_cell_loss_uses_full_evidence_and_bounded_display(self):
        source, output = self.path('source.csv'), self.path('result.csv')
        source.write_text('key,value\nA,'+'x'*121+'SOURCE\n', encoding='utf-8')
        output.write_text('key,value\nA,'+'x'*121+'RESULT\n', encoding='utf-8')
        result = analyze_pair(source, [output], 'csv')
        self.assertEqual(result['status'], 'confirmed_loss')
        found = next(f for f in result['findings'] if f['code'] == 'CELL_VALUES_CHANGED')
        self.assertEqual(found['locations'][0]['before'], 'x'*120)
        self.assertTrue(found['locations'][0]['beforeTruncated'])
        self.assertTrue(found['locations'][0]['afterTruncated'])

    def test_confirmed_loss_counts_as_evaluated_but_uncertainty_is_not_a_proof(self):
        source, output = self.path('source.csv'), self.path('result.csv')
        source.write_text('A,B\n', encoding='utf-8')
        output.write_text('C,D\n', encoding='utf-8')
        result = analyze_pair(source, [output], 'csv')
        self.assertEqual(result['status'], 'confirmed_loss')
        self.assertEqual(result['checkedFeatures'], [])
        self.assertEqual(result['evaluatedFeatures'], ['cell_values'])
        self.assertTrue(result['qualityCompared'])
        source.write_text('key,value\nA,00123\n', encoding='utf-8')
        output.write_text('key,value\nA,123\n', encoding='utf-8')
        uncertain = analyze_pair(source, [output], 'csv')
        self.assertEqual(uncertain['status'], 'uncertain')
        self.assertEqual(uncertain['evaluatedFeatures'], [])
        self.assertFalse(uncertain['qualityCompared'])

    def test_added_result_cells_do_not_claim_full_grid_equality_or_source_loss(self):
        source, output = self.path('source.csv'), self.path('result.csv')
        source.write_text('key,value\nA,42\n', encoding='utf-8')
        for extra, expected in (('key,value\nA,42\nEXTRA,unexpected\n', [(3, 1), (3, 2)]),
                                ('key,value,extra\nA,42,unexpected\n', [(1, 3), (2, 3)])):
            with self.subTest(extra=extra):
                output.write_text(extra, encoding='utf-8')
                result = analyze_pair(source, [output], 'csv')
                self.assertEqual(result['status'], 'uncertain')
                self.assertEqual(result['checkedFeatures'], [])
                self.assertEqual(result['evaluatedFeatures'], [])
                self.assertFalse(result['qualityCompared'])
                found = next(f for f in result['findings'] if f['code'] == 'CELL_VALUES_ADDED')
                self.assertEqual([(entry['row'], entry['column']) for entry in found['locations']], expected)
                self.assertFalse(any(f['category'] == 'confirmed_loss' for f in result['findings']))

    def test_formula_added_outside_source_grid_also_prevents_value_equality(self):
        from openpyxl import Workbook
        source, output = self.path('source.xlsx'), self.path('result.xlsx')
        book = Workbook(); book.active.append(['key', 'value']); book.active.append(['A', 42]); book.save(source); book.close()
        book = Workbook(); book.active.append(['key', 'value']); book.active.append(['A', 42]); book.active['C3'] = '=1+2'; book.save(output); book.close()
        result = analyze_pair(source, [output], 'xlsx')
        self.assertEqual(result['status'], 'uncertain')
        self.assertNotIn('cell_values_single_grid', result['checkedFeatures'])
        found = next(f for f in result['findings'] if f['code'] == 'CELL_VALUES_ADDED')
        self.assertEqual(found['locations'], [{'row': 3, 'column': 3, 'before': '', 'after': '=1+2', 'afterKind': 'formula'}])

    def test_pages_and_frames_preserve_sequence_without_false_red(self):
        import app_backend.converter_guard.comparison_analysis as module
        source, output = self.path('source.png'), self.path('result.tiff')
        png(source); output.write_bytes(b'II*\0profile-boundary')
        def profile(format_, *, frames=1, pages=1):
            return {'format': format_, 'width': 2, 'height': 2, 'frames': frames, 'pages': pages,
                    'alpha': False, 'orientation': 1, 'durationsMs': [100]*frames, 'loopCount': 0}
        a, b = profile('png', frames=3), profile('tiff', pages=3)
        with patch.object(module, 'inspect_media', side_effect=lambda path, **kwargs: a if Path(path) == source else b):
            result = analyze_pair(source, [output], 'tiff')
        self.assertEqual(result['status'], 'format_limitation')
        self.assertIn('pages', result['checkedFeatures'])
        self.assertFalse(any(f['category'] == 'confirmed_loss' for f in result['findings']))
        with patch.object(module, 'inspect_media', side_effect=lambda path, **kwargs: b if Path(path) == output else a):
            reverse = analyze_pair(output, [source], 'png')
        self.assertEqual(reverse['status'], 'format_limitation')
        self.assertIn('frames', reverse['checkedFeatures'])
        self.assertIn('page_to_frame_timing', reverse['unknown'])
        self.assertFalse(any(f['code'] == 'ANIMATION_TIMING_CHANGED' for f in reverse['findings']))
        b['pages'] = 2
        with patch.object(module, 'inspect_media', side_effect=lambda path, **kwargs: a if Path(path) == source else b):
            missing = analyze_pair(source, [output], 'tiff')
        self.assertEqual(missing['status'], 'confirmed_loss')
        self.assertIn('MEDIA_PAGES_CHANGED', [f['code'] for f in missing['findings']])

    def test_animation_timing_loop_and_secondary_page_properties_are_checked(self):
        import app_backend.converter_guard.comparison_analysis as module
        source, output = self.path('source.png'), self.path('result.png')
        png(source, frames=2); png(output, frames=2, width=3)
        a = {'format': 'png', 'width': 2, 'height': 2, 'frames': 2, 'pages': 1,
             'alpha': False, 'orientation': 1, 'durationsMs': [100, 200], 'loopCount': 0, 'bitDepth': 16}
        b = dict(a, durationsMs=[100, 300], loopCount=1, bitDepth=8)
        with patch.object(module, 'inspect_media', side_effect=lambda path, **kwargs: a if Path(path) == source else b):
            result = analyze_pair(source, [output], 'png')
        self.assertTrue({'ANIMATION_TIMING_CHANGED', 'ANIMATION_LOOP_CHANGED', 'IMAGE_BIT_DEPTH_REDUCED'} <=
                        {f['code'] for f in result['findings']})
        a.update(format='tiff', frames=1, pages=2, pageDimensions=[[2, 2], [5, 4]], pageOrientations=[1, 1])
        b.update(format='tiff', frames=1, pages=2, pageDimensions=[[2, 2], [3, 4]], pageOrientations=[1, 1], bitDepth=16)
        first, second = self.path('source.tiff'), self.path('result.tiff')
        first.write_bytes(b'II*\0source'); second.write_bytes(b'II*\0result')
        with patch.object(module, 'inspect_media', side_effect=lambda path, **kwargs: a if Path(path) == first else b):
            result = analyze_pair(first, [second], 'tiff')
        found = next(f for f in result['findings'] if f['code'] == 'IMAGE_PAGE_DIMENSIONS_CHANGED')
        self.assertEqual(found['locations'], [{'page': 2, 'before': [5, 4], 'after': [3, 4]}])

    def test_row_reordering_is_not_falsely_reported_as_lost_values(self):
        source, output = self.path('source.csv'), self.path('result.csv')
        source.write_text('key,value\nA,42\nB,43\n', encoding='utf-8')
        output.write_text('key,value\nB,43\nA,42\n', encoding='utf-8')
        result = analyze_pair(source, [output], 'csv')
        self.assertEqual(result['status'], 'uncertain')
        self.assertIn('CELL_POSITIONS_CHANGED', [f['code'] for f in result['findings']])
        self.assertFalse(any(f['category'] == 'confirmed_loss' for f in result['findings']))

    def test_animation_loss_in_capable_png_and_static_bmp_limitation_are_separate(self):
        source, output, static = self.path('source.png'), self.path('out.png'), self.path('out.bmp')
        png(source, frames=2)
        png(output, frames=1)
        bmp(static)
        lost = analyze_pair(source, [output], 'png')
        expected = analyze_pair(source, [static], 'bmp')
        self.assertEqual(lost['status'], 'confirmed_loss')
        self.assertEqual(expected['status'], 'confirmed_loss')  # unexpected alpha loss, static animation separately expected
        self.assertIn('FORMAT_FRAMES_LIMITATION', [f['code'] for f in expected['findings']])
        self.assertNotIn('MEDIA_FRAMES_CHANGED', [f['code'] for f in expected['findings']])

    def test_canvas_change_is_a_concrete_property_change(self):
        source, output = self.path('source.png'), self.path('result.png')
        png(source, width=4)
        png(output, width=2)
        result = analyze_pair(source, [output], 'png')
        self.assertEqual(result['status'], 'confirmed_loss')
        finding = next(f for f in result['findings'] if f['code'] == 'IMAGE_DIMENSIONS_CHANGED')
        self.assertEqual(finding['before'], [4, 2])
        self.assertEqual(finding['after'], [2, 2])

    def test_unresolved_32bit_bmp_alpha_is_unknown_not_a_red_loss(self):
        source, output = self.path('source.png'), self.path('result.bmp')
        png(source); bmp(output)
        data = bytearray(output.read_bytes()); data[28:30] = (32).to_bytes(2, 'little'); output.write_bytes(data)
        result = analyze_pair(source, [output], 'bmp')
        self.assertIn('bmp_alpha_channel', result['unknown'])
        self.assertFalse(any(f['code'] == 'IMAGE_ALPHA_LOST' for f in result['findings']))

    def test_invalid_image_after_conversion_is_not_marked_checked(self):
        source, output = self.path('source.png'), self.path('result.bmp')
        png(source)
        bmp(output)
        output.write_bytes(output.read_bytes()[:26])
        result = analyze_pair(source, [output], 'bmp')
        self.assertEqual(result['status'], 'confirmed_loss')
        self.assertIn('MEDIA_OUTPUT_INVALID', [f['code'] for f in result['findings']])

    def test_invalid_container_is_evaluated_without_claiming_any_equality(self):
        source, invalid = self.path('source.txt'), self.path('invalid.bmp')
        source.write_text('Original', encoding='utf-8')
        invalid.write_bytes(b'BM'+bytes(24))
        result = analyze_pair(source, [invalid], 'bmp')
        self.assertEqual(result['status'], 'confirmed_loss')
        self.assertEqual(result['checkedFeatures'], [])
        self.assertEqual(result['evaluatedFeatures'], ['result_container'])
        self.assertTrue(result['qualityCompared'])
        self.assertIn('visual_layout', result['unknown'])
        found = next(f for f in result['findings'] if f['code'] == 'MEDIA_OUTPUT_INVALID')
        self.assertEqual(found['feature'], 'result_container')

    def test_pdf_extraction_difference_never_proves_lost_ink(self):
        source, output = self.path('source.docx'), self.path('result.pdf')
        docx(source, 'PDF text that the extractor misses')
        output.write_bytes(b'%PDF-fixture')
        from app_backend.converter_guard import comparison_analysis as analysis
        real = analysis.collect_source_summary
        def summary(path, **kwargs):
            if Path(path) == output:
                return {'kind': 'pdf', 'documentClass': 'rendered', 'complete': True,
                        'counts': {'pages': 1}, '_text': '', '_grids': []}
            return real(path, **kwargs)
        with patch.object(analysis, 'collect_source_summary', side_effect=summary):
            result = analyze_pair(source, [output], 'pdf')
        self.assertEqual(result['status'], 'uncertain')
        self.assertIn('PDF_TEXT_NOT_EXTRACTED', [f['code'] for f in result['findings']])
        self.assertFalse(any(f['category'] == 'confirmed_loss' for f in result['findings']))

    def test_document_images_offer_explicit_pdf_recipe_without_claiming_a_proven_loss(self):
        source, output = self.path('source.docx'), self.path('result.pdf')
        docx(source, 'A')
        with zipfile.ZipFile(source, 'a') as archive:
            archive.writestr('word/media/image.png', b'opaque source image fixture')
        output.write_bytes(b'%PDF-fixture')
        from app_backend.converter_guard import comparison_analysis as analysis
        real = analysis.collect_source_summary
        def summary(path, **kwargs):
            if Path(path) == output:
                return {'kind': 'pdf', 'documentClass': 'rendered', 'complete': True,
                        'counts': {'pages': 1}, '_text': 'A', '_grids': []}
            return real(path, **kwargs)
        with patch.object(analysis, 'collect_source_summary', side_effect=summary):
            result = analyze_pair(source, [output], 'pdf')
        self.assertEqual(result['status'], 'uncertain')
        self.assertEqual(result['recipeCandidates'], ['pdf_preserve_images'])
        self.assertFalse(any(f['category'] == 'confirmed_loss' for f in result['findings']))

    def test_real_pdf_reader_keeps_text_extraction_scope_explicit(self):
        from pypdf import PdfWriter
        source, output = self.path('source.txt'), self.path('result.pdf')
        source.write_text('Visible source text', encoding='utf-8')
        writer = PdfWriter(); writer.add_blank_page(width=100, height=100)
        with output.open('wb') as stream:
            writer.write(stream)
        result = analyze_pair(source, [output], 'pdf')
        self.assertEqual(result['status'], 'uncertain')
        self.assertEqual(result['after']['counts']['pages'], 1)

    def test_byte_identical_unknown_files_have_exact_bounded_evidence(self):
        source, output = self.path('source.unknown'), self.path('result.unknown')
        source.write_bytes(b'opaque\0payload')
        output.write_bytes(source.read_bytes())
        result = analyze_pair(source, [output], 'unknown')
        self.assertEqual(result['status'], 'checked')
        self.assertIn('identical_file_bytes', result['checkedFeatures'])
        self.assertIn('semantic_meaning', result['unknown'])

    def test_unknown_formats_and_missing_sources_are_not_checked(self):
        source, output = self.path('source.xyz'), self.path('result.zzz')
        source.write_bytes(b'opaque'); output.write_bytes(b'other')
        self.assertEqual(analyze_pair(source, [output], 'zzz')['status'], 'not_checked')
        source.unlink()
        self.assertEqual(analyze_pair(source, [output], 'zzz')['status'], 'not_checked')

    def test_source_cache_reused_for_multiple_targets_but_changed_source_declines(self):
        source, output = self.path('source.docx'), self.path('result.txt')
        docx(source, 'A'); output.write_text('A', encoding='utf-8')
        evidence = collect_source_summary(source)
        with patch('app_backend.converter_guard.comparison_analysis._paragraph_text', side_effect=AssertionError('source was rescanned')):
            result = analyze_pair(source, [output], 'txt', source_evidence=evidence)
        self.assertIn('normalized_text', result['checkedFeatures'])
        docx(source, 'Changed original')
        changed = analyze_pair(source, [output], 'txt', source_evidence=evidence)
        self.assertEqual(changed['status'], 'not_checked')
        self.assertEqual(changed['findings'][0]['reason'], 'source_changed_or_unavailable')

    def test_serialized_source_summary_and_mid_analysis_mutation_are_handled(self):
        source, output = self.path('source.csv'), self.path('result.csv')
        source.write_text('key,value\nA,42\n', encoding='utf-8')
        output.write_text('key,value\nA,0\n', encoding='utf-8')
        cached = json.loads(json.dumps(collect_source_summary(source)))
        self.assertEqual(analyze_pair(source, [output], 'csv', source_evidence=cached)['status'], 'confirmed_loss')
        from app_backend.converter_guard import comparison_analysis as module
        real = module._compare_grids
        def mutate(*args):
            real(*args)
            source.write_text('key,value\nA,123\n', encoding='utf-8')
        with patch.object(module, '_compare_grids', side_effect=mutate):
            result = analyze_pair(source, [output], 'csv', source_evidence=cached)
        self.assertEqual(result['status'], 'not_checked')
        self.assertEqual(result['recipeCandidates'], [])
        self.assertFalse(result['qualityCompared'])
        self.assertFalse(any(f['category'] == 'confirmed_loss' for f in result['findings']))

    def test_single_sheet_xlsx_to_csv_preserves_values_and_reports_real_format_limits(self):
        from openpyxl import Workbook
        source, output = self.path('source.xlsx'), self.path('result.csv')
        book = Workbook(); sheet = book.active
        sheet.append(['key', 'value']); sheet.append(['A', 42]); font = copy(sheet['A1'].font); font.bold = True; sheet['A1'].font = font
        book.save(source); book.close()
        output.write_text('key,value\nA,42\n', encoding='utf-8')
        result = analyze_pair(source, [output], 'csv')
        self.assertIn('cell_values_single_grid', result['checkedFeatures'])
        self.assertEqual(result['status'], 'format_limitation')
        self.assertIn('FORMAT_STYLES_LIMITATION', [f['code'] for f in result['findings']])

    def test_multiple_sheet_mapping_is_unknown_instead_of_invented_cell_losses(self):
        from openpyxl import Workbook
        source, output = self.path('source.xlsx'), self.path('result.csv')
        book = Workbook(); book.active.append(['A', 'B']); book.create_sheet('Second').append([42, 43]); book.save(source); book.close()
        output.write_text('A,B\n', encoding='utf-8')
        result = analyze_pair(source, [output], 'csv')
        self.assertFalse(any(f['category'] == 'confirmed_loss' for f in result['findings']))
        self.assertIn('all_sheet_outputs_mapping', result['unknown'])

    def test_localized_date_and_formatted_number_are_uncertain_not_confirmed_loss(self):
        from openpyxl import Workbook
        source, output = self.path('source.xlsx'), self.path('result.csv')
        book = Workbook(); sheet = book.active
        sheet.append(['date', 'amount']); sheet.append([datetime(2026, 10, 25), 42]); sheet['B2'].number_format = '"€"0.00'
        book.save(source); book.close()
        output.write_text('date,amount\n25.10.2026,€42.00\n', encoding='utf-8')
        result = analyze_pair(source, [output], 'csv')
        self.assertEqual(result['status'], 'uncertain')
        self.assertFalse(any(f['category'] == 'confirmed_loss' for f in result['findings']))
        self.assertNotIn('cell_values_single_grid', result['checkedFeatures'])
        found = next(f for f in result['findings'] if f['code'] == 'CELL_REPRESENTATION_CHANGED')
        self.assertNotIn('Числовое значение', found['message'])

    def test_odf_explicit_repeated_spaces_are_real_cell_data(self):
        source, output = self.path('source.ods'), self.path('result.csv')
        xml = ('<o:document-content xmlns:o="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
               'xmlns:t="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
               'xmlns:p="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
               '<o:body><o:spreadsheet><t:table><t:table-row>'
               '<t:table-cell><p:p>key</p:p></t:table-cell><t:table-cell><p:p>value</p:p></t:table-cell>'
               '</t:table-row><t:table-row><t:table-cell><p:p>A</p:p></t:table-cell>'
               '<t:table-cell><p:p>left<p:s p:c="2"/>right</p:p></t:table-cell>'
               '</t:table-row></t:table></o:spreadsheet></o:body></o:document-content>')
        with zipfile.ZipFile(source, 'w') as archive:
            archive.writestr('mimetype', 'application/vnd.oasis.opendocument.spreadsheet')
            archive.writestr('content.xml', xml)
        output.write_text('key,value\nA,left  right\n', encoding='utf-8')
        result = analyze_pair(source, [output], 'csv')
        self.assertIn('cell_values_single_grid', result['checkedFeatures'])
        self.assertFalse(any(f['category'] == 'confirmed_loss' for f in result['findings']))

    def test_odf_annotations_do_not_pollute_visible_cell_values(self):
        source, output = self.path('source.fods'), self.path('result.csv')
        source.write_text('<o:document xmlns:o="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
            'xmlns:t="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
            'xmlns:p="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
            'o:mimetype="application/vnd.oasis.opendocument.spreadsheet"><o:body><o:spreadsheet>'
            '<t:table><t:table-row><t:table-cell><p:p>key</p:p></t:table-cell>'
            '<t:table-cell><p:p>value</p:p></t:table-cell></t:table-row><t:table-row>'
            '<t:table-cell><p:p>A</p:p></t:table-cell><t:table-cell><p:p>Visible'
            '<o:annotation><p:p>Private comment</p:p></o:annotation></p:p></t:table-cell>'
            '</t:table-row></t:table></o:spreadsheet></o:body></o:document>', encoding='utf-8')
        output.write_text('key,value\nA,Visible\n', encoding='utf-8')
        result = analyze_pair(source, [output], 'csv')
        self.assertIn('cell_values_single_grid', result['checkedFeatures'])
        self.assertNotIn('Private comment', preview_text(source)['text'])
        self.assertIn('notes_and_comments', result['unknown'])
        self.assertFalse(any(f['category'] == 'confirmed_loss' for f in result['findings']))

    def test_word_related_text_and_revisions_have_explicit_coverage(self):
        source, output = self.path('source.docx'), self.path('result.docx')
        docx(source, 'Body'); docx(output, 'Body')
        for path, note in ((source, 'Original note'), (output, 'Changed note')):
            with zipfile.ZipFile(path, 'a') as archive:
                archive.writestr('word/footnotes.xml', '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    '<w:footnote w:type="separator"><w:p><w:r><w:t>Separator</w:t></w:r></w:p></w:footnote>'
                    '<w:footnote><w:p><w:r><w:t>'+note+'</w:t></w:r></w:p></w:footnote></w:footnotes>')
                archive.writestr('word/header1.xml', '<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    '<w:p><w:del><w:r><w:delText>Deleted</w:delText></w:r></w:del><w:r><w:t>Header</w:t></w:r></w:p></w:hdr>')
        result = analyze_pair(source, [output], 'docx')
        self.assertEqual(result['status'], 'uncertain')
        self.assertIn('footnotes', result['before']['textCoverage'])
        self.assertIn('headers', result['before']['textCoverage'])
        self.assertNotIn('normalized_text', result['checkedFeatures'])
        preview = preview_text(source)['text']
        self.assertIn('Original note', preview)
        self.assertNotIn('Deleted', preview)
        self.assertNotIn('Separator', preview)
        self.assertIn('comments_and_revisions', result['unknown'])

    def test_presentation_reads_ordered_slides_and_does_not_infer_quality_from_count(self):
        source, output = self.path('source.pptx'), self.path('result.pptx')
        pptx(source, ['First slide', 'Second slide'])
        pptx(output, ['First slide Second slide'])
        result = analyze_pair(source, [output], 'pptx')
        self.assertIn('normalized_text', result['checkedFeatures'])
        self.assertEqual(result['status'], 'uncertain')
        self.assertIn('STRUCTURE_SLIDES_COUNT_CHANGED', [f['code'] for f in result['findings']])
        self.assertEqual(result['before']['counts']['slides'], 2)
        self.assertIn('presentation_notes', result['unknown'])
        self.assertIn('First slide', preview_text(source)['text'])

    def test_formula_structure_is_distinct_from_unavailable_calculated_values(self):
        from openpyxl import Workbook
        source, output = self.path('source.xlsx'), self.path('result.xlsx')
        book = Workbook(); book.active.append(['key', 'value']); book.active.append(['A', '=6*7']); book.save(source); book.close()
        book = Workbook(); book.active.append(['key', 'value']); book.active.append(['A', 42]); book.save(output); book.close()
        result = analyze_pair(source, [output], 'xlsx')
        self.assertEqual(result['status'], 'confirmed_loss')
        self.assertIn('EDITABLE_FORMULAS_MISSING', [f['code'] for f in result['findings']])
        self.assertIn('cached_formula_values', result['unknown'])
        # An output formula without a stored cache must not look like an empty
        # value and generate a false red loss against a static numeric input.
        reverse = analyze_pair(output, [source], 'xlsx')
        self.assertFalse(any(f['category'] == 'confirmed_loss' for f in reverse['findings']))
        self.assertNotIn('cell_values_single_grid', reverse['checkedFeatures'])
        self.assertIn('cached_formula_values', reverse['unknown'])

    def test_bom_text_encodings_are_explicit_and_legacy_encoding_is_not_guessed(self):
        source, output = self.path('source.txt'), self.path('result.txt')
        source.write_text('Привет UTF', encoding='utf-8')
        for encoding in ('utf-16', 'utf-32'):
            with self.subTest(encoding=encoding):
                output.write_text('Привет UTF', encoding=encoding)
                result = analyze_pair(source, [output], 'txt')
                self.assertIn('normalized_text', result['checkedFeatures'])
                self.assertEqual(preview_text(output)['text'], 'Привет UTF')
        output.write_bytes('Привет UTF'.encode('cp1251'))
        result = analyze_pair(source, [output], 'txt')
        self.assertEqual(result['status'], 'not_checked')
        self.assertEqual(result['after']['reason'], 'text_encoding_not_inspected')
        output.write_text('', encoding='utf-8')
        empty = analyze_pair(source, [output], 'txt')
        self.assertEqual(empty['status'], 'confirmed_loss')
        self.assertIn('TEXT_OUTPUT_EMPTY', [f['code'] for f in empty['findings']])

    def test_csv_preview_truncation_does_not_hide_complete_value_inspection(self):
        source = self.path('source.csv')
        text = 'key,value\n'+''.join('K'+str(i)+','+'v'*1000+'\n' for i in range(600))
        # Keep fixture bytes identical on Windows: write_text would expand LF
        # to CRLF while the expected character count still describes LF text.
        source.write_bytes(text.encode('utf-8'))
        summary = collect_source_summary(source)
        self.assertTrue(summary['complete'])
        self.assertEqual(summary['counts']['characters'], len(text))
        preview = preview_text(source, max_chars=MAX_TEXT)
        self.assertTrue(preview['truncated'])
        self.assertEqual(len(preview['text']), MAX_TEXT)

    def test_many_empty_repeated_rows_have_a_budget_not_an_allocation_explosion(self):
        source = self.path('source.fods')
        source.write_text('<o:document xmlns:o="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
            'xmlns:t="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
            'o:mimetype="application/vnd.oasis.opendocument.spreadsheet"><o:body><o:spreadsheet>'
            '<t:table><t:table-row t:number-rows-repeated="50000"/>'
            '<t:table-row t:number-rows-repeated="50000"/></t:table>'
            '</o:spreadsheet></o:body></o:document>', encoding='utf-8')
        result = collect_source_summary(source)
        self.assertFalse(result['complete'])
        self.assertEqual(result['reason'], 'row_budget')

    def test_declared_custom_shared_strings_and_worksheet_bytes_are_budgeted_before_openpyxl(self):
        source = self.path('source.xlsx')
        with zipfile.ZipFile(source, 'w') as archive:
            archive.writestr('[Content_Types].xml', '<Types><Override PartName="/xl/custom.bin" '
                'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
            archive.writestr('xl/workbook.xml', '<workbook/>')
            archive.writestr('xl/custom.bin', b' ' * 2048)
        with patch('app_backend.converter_guard.comparison_analysis.MAX_READ', 1024), patch('openpyxl.load_workbook') as load:
            result = collect_source_summary(source)
        self.assertFalse(result['complete'])
        self.assertEqual(result['reason'], 'xlsx_xml_budget')
        load.assert_not_called()

    def test_xlsx_embedded_images_are_counted_for_the_pdf_recipe(self):
        from openpyxl import Workbook
        source = self.path('source.xlsx'); book = Workbook(); book.active.append(['A', 'B']); book.save(source); book.close()
        with zipfile.ZipFile(source, 'a') as archive:
            archive.writestr('xl/media/image.png', b'bounded image metadata fixture')
        evidence = collect_source_summary(source)
        self.assertTrue(evidence['complete'])
        self.assertEqual(evidence['counts']['images'], 1)

    def test_all_csv_sheet_values_are_compared_by_multiplicity_without_order_guesses(self):
        from openpyxl import Workbook
        source, first, second = self.path('source.xlsx'), self.path('a.csv'), self.path('b.csv')
        book = Workbook(); book.active.append(['A', 'B']); book.create_sheet('Second').append([42, 43]); book.save(source); book.close()
        first.write_text('42,43\n', encoding='utf-8'); second.write_text('A,B\n', encoding='utf-8')
        result = analyze_pair(source, [first, second], 'csv')
        self.assertIn('all_sheet_values_unordered', result['checkedFeatures'])
        self.assertNotIn('all_sheet_outputs_mapping', result['unknown'])
        self.assertIn('sheet_names_and_identity', result['unknown'])
        first.write_text('A,B\n', encoding='utf-8')  # duplicated sheet cannot mask the missing numeric sheet
        changed = analyze_pair(source, [first, second], 'csv')
        self.assertEqual(changed['status'], 'uncertain')
        self.assertNotIn('all_sheet_values_unordered', changed['checkedFeatures'])

    def test_multi_csv_mutation_invalidates_the_whole_comparison(self):
        from openpyxl import Workbook
        import app_backend.converter_guard.comparison_analysis as module
        source, first, second = self.path('source.xlsx'), self.path('first.csv'), self.path('second.csv')
        book = Workbook(); book.active.append(['A', 'B']); book.create_sheet('Second').append(['C', 'D']); book.save(source); book.close()
        first.write_text('A,B\n', encoding='utf-8'); second.write_text('C,D\n', encoding='utf-8')
        real = module._csv_summary
        def mutate(path, budget):
            value = real(path, budget)
            if path == second:
                first.write_text('CHANGED,B\n', encoding='utf-8')
            return value
        with patch.object(module, '_csv_summary', side_effect=mutate):
            result = analyze_pair(source, [first, second], 'csv')
        self.assertEqual(result['status'], 'not_checked')
        self.assertEqual(result['checkedFeatures'], [])
        self.assertEqual(result['findings'][0]['reason'], 'result_changed_during_analysis')

    def test_supplemental_outputs_and_alias_ambiguity_are_never_hidden(self):
        source, output, resource = self.path('source.csv'), self.path('result.csv'), self.path('image.png')
        source.write_text('A,B\n', encoding='utf-8'); output.write_text('A,B\n', encoding='utf-8'); png(resource)
        result = analyze_pair(source, [output, resource], 'csv')
        self.assertIn('supplemental_outputs_not_inspected', result['unknown'])
        self.assertEqual(result['outputScope'], {'totalFiles': 2, 'primaryFiles': 1, 'supplementalFiles': 1,
                                                'inspectedPrimaryFiles': 1, 'inspectedSupplementalFiles': 0})
        # Both extensions denote JPEG. Exact suffix must not silently hide the
        # other candidate and certify just one arbitrarily chosen output.
        one, two = self.path('one.jpg'), self.path('two.jpeg')
        one.write_bytes(b'one'); two.write_bytes(b'two')
        ambiguous = analyze_pair(source, [one, two], 'jpg')
        self.assertEqual(ambiguous['status'], 'not_checked')
        self.assertEqual(ambiguous['outputScope']['primaryFiles'], 2)
        self.assertEqual(ambiguous['findings'][0]['reason'], 'primary_output_mapping_not_determined')

    def test_duplicate_csv_paths_cannot_prove_two_sheets_saved(self):
        from openpyxl import Workbook
        source, output = self.path('source.xlsx'), self.path('result.csv')
        book = Workbook(); book.active.append(['A', 'B']); book.create_sheet('Second').append(['A', 'B']); book.save(source); book.close()
        output.write_text('A,B\n', encoding='utf-8')
        result = analyze_pair(source, [output, output], 'csv')
        self.assertEqual(result['status'], 'not_checked')
        self.assertEqual(result['findings'][0]['reason'], 'duplicate_primary_outputs')

    def test_byte_comparison_respects_read_quota(self):
        source, output = self.path('source.unrecognized'), self.path('result.unrecognized')
        source.write_bytes(b'x'*1024); output.write_bytes(b'x'*1024)
        with patch('app_backend.converter_guard.comparison_analysis.MAX_READ', 512):
            result = analyze_pair(source, [output], 'unrecognized')
        self.assertEqual(result['status'], 'not_checked')
        self.assertNotIn('identical_file_bytes', result['checkedFeatures'])

    def test_wrong_document_model_and_truncated_package_are_not_green(self):
        from openpyxl import Workbook
        source, wrong, partial = self.path('source.docx'), self.path('result.docx'), self.path('partial.docx')
        docx(source, 'Source text')
        real_xlsx = self.path('table.xlsx'); book = Workbook(); book.active.append(['A', 'B']); book.save(real_xlsx); book.close()
        wrong.write_bytes(real_xlsx.read_bytes())
        result = analyze_pair(source, [wrong], 'docx')
        self.assertEqual(result['status'], 'confirmed_loss')
        self.assertIn('OUTPUT_FORMAT_MISMATCH', [f['code'] for f in result['findings']])
        self.assertIn('result_container', result['evaluatedFeatures'])
        self.assertNotIn('result_container', result['checkedFeatures'])
        partial.write_bytes(source.read_bytes()[:80])
        result = analyze_pair(source, [partial], 'docx')
        self.assertEqual(result['status'], 'not_checked')
        self.assertIn('incomplete_result_inspection', result['unknown'])
        self.assertFalse(result['qualityCompared'])

    def test_utf16_entity_declarations_are_rejected_and_preview_is_literal_text(self):
        source = self.path('source.docx')
        data = ('<?xml version="1.0" encoding="UTF-16"?>'
                '<!DOCTYPE root [<!ENTITY x "expanded">]><root>&x;</root>').encode('utf-16')
        with zipfile.ZipFile(source, 'w') as archive:
            archive.writestr('word/document.xml', data)
        evidence = collect_source_summary(source)
        self.assertFalse(evidence['complete'])
        self.assertEqual(evidence['reason'], 'xml_declarations_not_inspected')
        text = self.path('safe.txt'); text.write_text('<script>never executed</script> abc', encoding='utf-8')
        preview = preview_text(text, max_chars=12)
        self.assertEqual(preview['kind'], 'text')
        self.assertEqual(preview['text'], '<script>neve')
        self.assertTrue(preview['truncated'])

    def test_cancel_and_preview_limits_remain_explicit(self):
        source = self.path('source.txt'); source.write_text('x', encoding='utf-8')
        cancel = threading.Event(); cancel.set()
        with self.assertRaises(GuardError) as caught:
            analyze_pair(source, [source], 'txt', cancel=cancel)
        self.assertEqual(caught.exception.code, 'CANCELLED')
        for value in (0, MAX_TEXT+1, True):
            with self.assertRaises(ValueError):
                preview_text(source, max_chars=value)


if __name__ == '__main__':
    unittest.main()
