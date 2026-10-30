"""A01/A03: unsafe Windows filenames and unsupported Excel semantics."""
import io
from pathlib import Path, PureWindowsPath
import tempfile
import threading
import unittest
from urllib.parse import quote

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

from app_backend.converter_upload_stream import MultipartUpload
from app_backend.excel import MergeError, merge_xlsx


def wire(name):
    boundary = 'aio-audit-file-name'
    field = ('filename="' + name + '"' if name.isascii()
             else "filename*=UTF-8''" + quote(name))
    head = ('--' + boundary + '\r\n'
            'Content-Disposition: form-data; name="files"; ' + field + '\r\n'
            'Content-Type: application/octet-stream\r\n\r\n').encode('ascii')
    body = head + b'input-content\r\n--' + boundary.encode() + b'--\r\n'
    return body, 'multipart/form-data; boundary=' + boundary


def sample_xlsx(configure=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Sheet'
    sheet.append(['Code', 'Value'])
    sheet.append(['00125', 123])
    if configure is not None:
        configure(sheet)
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


class WindowsUploadPathSafety(unittest.TestCase):
    def _one(self, name):
        body, content_type = wire(name)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            files, metadata = MultipartUpload(io.BytesIO(body), len(body),
                content_type, threading.Event()).read(root)
            self.assertEqual(metadata, {})
            self.assertEqual(len(files), 1)
            item = files[0]
            self.assertEqual(item['name'], name)
            self.assertEqual(item['path'].read_bytes(), b'input-content')
            self.assertTrue(item['path'].resolve().is_relative_to(root.resolve()))
            self.assertTrue((PureWindowsPath('D:/aio-work') / item['path'].name)
                .is_relative_to(PureWindowsPath('D:/aio-work')))
            return item['path'].name, item['warnings']

    def test_drive_relative_ads_device_and_invalid_leaves_are_normalized(self):
        for name in ('C:escaped.txt', 'D:other.pdf', 'invoice.txt:payload',
                     'NUL.txt', 'CON.txt', 'COM1.docx', 'LPT1.xlsx',
                     'question?.pdf', 'invalid|name.odt', 'trailing.',
                     'bad<name>.pdf'):
            with self.subTest(name=name):
                leaf, warnings = self._one(name)
                self.assertTrue(leaf.startswith('input.'), (name, leaf))
                self.assertEqual(warnings, ['SOURCE_NAME_NORMALIZED'])

    # BytesHeaderParser strips terminal whitespace from quoted MIME names.
    # That already prevents a trailing-space leaf reaching the filesystem.
    def test_valid_names_preserve_original_and_bytes(self):
        for name in ('normal.docx', 'Файл.xlsx', 'пример.txt', 'a.txt'):
            with self.subTest(name=name):
                leaf, warnings = self._one(name)
                self.assertEqual(leaf, name)
                self.assertEqual(warnings, [])


class ExcelNoSilentLoss(unittest.TestCase):
    def _rejected(self, config, match):
        with self.assertRaisesRegex(MergeError, match):
            merge_xlsx([('one.xlsx', sample_xlsx(config)),
                        ('two.xlsx', sample_xlsx())])

    def test_sheet_scoped_name_not_silently_discarded(self):
        def config(sheet):
            sheet.defined_names.add(DefinedName(
                'LocalCode', attr_text="'Sheet'!$A$2"))
        self._rejected(config, 'локальные именованные диапазоны')

    def test_print_title_rows_and_columns_not_silently_discarded(self):
        self._rejected(lambda sheet: setattr(sheet, 'print_title_rows', '1:1'),
                       'повторяемые')
        self._rejected(lambda sheet: setattr(sheet, 'print_title_cols', 'A:A'),
                       'повторяемые')

    def test_header_and_footer_every_variant(self):
        for attr, section in (('oddHeader','center'), ('oddFooter','left'),
                              ('evenHeader','right'), ('evenFooter','right'),
                              ('firstHeader','left'), ('firstFooter','center')):
            def config(sheet):
                getattr(getattr(sheet, attr), section).text = 'AUDIT_HEADER'
            with self.subTest(attr=attr):
                self._rejected(config, 'колонтитулы')

    def test_page_setup_is_not_silently_discarded(self):
        self._rejected(lambda sheet: setattr(
            sheet.page_setup, 'orientation', 'landscape'), 'настройки страницы')

    def test_supported_plain_data_still_merges(self):
        result = merge_xlsx([('one.xlsx', sample_xlsx()),
                             ('two.xlsx', sample_xlsx())])
        workbook = load_workbook(io.BytesIO(result), read_only=True)
        try:
            self.assertEqual([row[0] for row in workbook.active.values],
                             ['Code', '00125', '00125'])
        finally:
            workbook.close()
