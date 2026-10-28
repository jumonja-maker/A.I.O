"""B29: source strings and cell types, not appearance or a second converter."""
import csv
import hashlib
import json
from pathlib import Path
import tempfile
import threading
from types import SimpleNamespace
import unittest
from unittest.mock import patch
import zipfile

from openpyxl import Workbook
from app_backend.converter_guard import csv_id_contract as ids
from app_backend.converter_guard.csv_import import csv_import_hint
from app_backend.converter_guard.legacy_import import validate_options
from app_backend.converter_guard.errors import GuardError

FILTER = ids.CALC_TEXT
CALC = 'com.sun.star.sheet.SpreadsheetDocument'
NATIVE = SimpleNamespace(valid=True, fingerprint='b29-fixture',
    filters={FILTER: (CALC, frozenset({'csv', 'tab', 'tsv'}), frozenset({'IMPORT'}))},
    preferred_filters={})
ROWS = [['ID', 'Name', 'Amount'], ['001', 'Строка ё', 42], ['002', 'Tail', -7]]


class CsvTextIdsB29(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.csv(ROWS)

    def tearDown(self):
        self.temp.cleanup()

    def csv(self, rows, delimiter=',', name='source.csv'):
        path = self.root / name
        with path.open('w', encoding='utf-8', newline='') as stream:
            csv.writer(stream, delimiter=delimiter).writerows(rows)
        return path

    def xlsx(self, rows=None, extension='xlsx', number_format=None, second_sheet=False):
        path = self.root / ('output.' + extension)
        book = Workbook()
        for row in ROWS if rows is None else rows:
            book.active.append(row)
        if number_format:
            for row in book.active.iter_rows(min_row=2, max_col=1):
                row[0].number_format = number_format
        if second_sheet:
            book.create_sheet('other')
        book.save(path)
        book.close()
        return path

    def check(self, path, **kw):
        return ids.check_csv_text_ids([path], path.suffix[1:], self.source,
            import_filter=FILTER, import_options='44,34,76,1,1/2', **kw)

    def rewrite_zip(self, path, edit):
        with zipfile.ZipFile(path) as archive:
            members = {n: archive.read(n) for n in archive.namelist()}
        with zipfile.ZipFile(path, 'w') as archive:
            for name, raw in members.items():
                archive.writestr(name, edit(name, raw))

    def loss(self, rows=None, **kw):
        with self.assertRaises(GuardError) as ctx:
            self.check(self.xlsx(rows, **kw))
        self.assertEqual(ctx.exception.code, 'ENGINE_OUTPUT_CONTENT_LOSS')
        self.assertNotIn(str(self.root), ctx.exception.diagnostic)
        self.assertNotIn('Строка', ctx.exception.diagnostic)
        return json.loads(ctx.exception.diagnostic)

    def test_four_delimiters_use_only_id_column(self):
        for delimiter in ',;\t|':
            with self.subTest(delimiter=delimiter):
                source = self.csv(ROWS, delimiter)
                hint = csv_import_hint(source, NATIVE)
                self.assertEqual(hint, {'filter': FILTER, 'options': f'{ord(delimiter)},34,76,1,1/2'})
                self.assertTrue(validate_options(FILTER, hint['options'], NATIVE))

    def test_no_id_preserves_old_import_hint(self):
        source = self.csv([['name', 'value'], ['item', '12'], ['other', '7']])
        self.assertEqual(csv_import_hint(source, NATIVE)['options'], '44,34,76,1')

    def test_signed_and_long_ids_do_not_turn_whole_sheet_into_text(self):
        rows = [['Name', 'ID', 'Value'], ['item', '+001', '42'], ['other', '-002', '1.5'],
                ['long', '0' + '1' * 50, '2']]
        self.assertEqual(csv_import_hint(self.csv(rows), NATIVE)['options'], '44,34,76,1,2/2')

    def test_multiple_text_columns_sorted(self):
        rows = [['a', 'b', 'c'], ['001', '42', '007']]
        self.assertEqual(csv_import_hint(self.csv(rows), NATIVE)['options'], '44,34,76,1,1/2/3/2')

    def test_bom_and_renamed_source_still_use_content(self):
        source = self.csv(ROWS, name='renamed.bin')
        source.write_bytes(b'\xef\xbb\xbf' + source.read_bytes())
        self.assertEqual(csv_import_hint(source, NATIVE)['options'], '44,34,76,1,1/2')

    def test_ids_after_initial_eight_rows_are_seen(self):
        rows = [['id', 'amount']] + [[str(i), '42'] for i in range(1, 30)] + [['00030', '43']]
        self.assertEqual(csv_import_hint(self.csv(rows), NATIVE)['options'], '44,34,76,1,1/2')

    def test_unknown_orcus_filter_does_not_receive_calc_column_options(self):
        native = SimpleNamespace(valid=True, filters={
            'Orcus CSV': (CALC, {'csv'}, {'IMPORT'})}, preferred_filters={})
        self.assertEqual(csv_import_hint(self.source, native),
                         {'filter': 'Orcus CSV', 'options': '44,34,76,1'})

    def test_finite_option_grammar_rejects_other_codes_and_duplicates(self):
        for options in ('44,34,76,1,1/9', '44,34,76,1,1/2/1/2', '44,34,76,1,257/2',
                        '44,34,76,1,0/2', '44,34,76,1,1/2,extra', '44,34,76,1,',
                        '44,34,76,1,1/2/3', '44,34,76,1,2/2/1/2'):
            with self.subTest(options=options):
                self.assertFalse(validate_options(FILTER, options, NATIVE))
                self.assertIsNone(ids._import_dialect(FILTER, options))

    def test_four_excel_targets_keep_textual_ids(self):
        source_hash = hashlib.sha256(self.source.read_bytes()).hexdigest()
        for ext in sorted(ids.SPREADSHEET_TARGETS):
            with self.subTest(target=ext):
                result = self.check(self.xlsx(extension=ext))
                self.assertEqual(result['status'], 'checked_subset')
                self.assertEqual(result['candidateIds'], 2)
                self.assertEqual(result['idsCompared'], 2)
                self.assertFalse(result['qualityCompared'])
                self.assertFalse(result['fullDocumentFidelity'])
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), source_hash)

    def test_numeric_ids_rejected(self):
        detail = self.loss([ROWS[0], [1, 'Строка ё', 42], [2, 'Tail', -7]])
        self.assertEqual(detail['differentIds'], 2)

    def test_number_format_000_not_original_text(self):
        self.loss([ROWS[0], [1, 'Строка ё', 42], [2, 'Tail', -7]], number_format='000')

    def test_formulas_do_not_prove_text_id(self):
        self.loss([ROWS[0], ['="001"', 'Строка ё', 42], ROWS[2]])

    def test_same_count_changed_id_is_rejected(self):
        self.loss([ROWS[0], ['009', 'Строка ё', 42], ROWS[2]])

    def test_coordinate_swap_is_rejected(self):
        self.loss([ROWS[0], ROWS[2], ROWS[1]])

    def test_last_missing_id_is_rejected(self):
        self.loss(ROWS[:2])

    def test_all_output_cells_missing_is_not_green(self):
        self.loss([])

    def test_stale_dimension_does_not_hide_actual_ids(self):
        path = self.xlsx()
        def edit(name, raw):
            import re
            return re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', raw) if name.startswith('xl/worksheets/') else raw
        self.rewrite_zip(path, edit)
        self.assertEqual(self.check(path)['idsCompared'], 2)

    def test_zero_budget_is_unknown_not_loss(self):
        result = self.check(self.xlsx(), seconds=0)
        self.assertEqual(result['status'], 'not_checked')
        self.assertEqual(result['warning'], 'CSV_TEXT_IDS_NOT_CHECKED')

    def test_cancel_propagates(self):
        event = threading.Event(); event.set()
        with self.assertRaises(GuardError) as ctx:
            self.check(self.xlsx(), cancel=event)
        self.assertEqual(ctx.exception.code, 'CANCELLED')

    def test_input_byte_budget_is_not_damage(self):
        with patch.object(ids, 'MAX_INPUT_BYTES', 4):
            result = self.check(self.xlsx())
        self.assertEqual(result['status'], 'not_checked')

    def test_unknown_encoding_is_not_certified(self):
        self.source.write_bytes(b'id,label\n001,\xff\n')
        self.assertEqual(self.check(self.xlsx())['status'], 'not_checked')

    def test_ragged_row_mapping_is_unknown(self):
        self.source.write_text('id,x\n001,a\n002,b,extra\n', encoding='utf-8')
        self.assertEqual(self.check(self.xlsx())['status'], 'not_checked')

    def test_other_targets_and_filters_are_not_constrained(self):
        self.assertIsNone(ids.check_csv_text_ids([], 'pdf', self.source, import_filter=FILTER,
                                               import_options='44,34,76,1'))
        self.assertIsNone(ids.check_csv_text_ids([], 'xlsx', self.source, import_filter='DIF',
                                               import_options='76'))

    def test_no_ids_does_not_claim_any_quality(self):
        self.csv([['id', 'name'], ['1', 'item']])
        self.assertIsNone(self.check(self.xlsx()))

    def test_multiple_sheets_are_not_falsely_mapped(self):
        self.assertEqual(self.check(self.xlsx(second_sheet=True))['status'], 'not_checked')

    def test_duplicate_primary_outputs_unknown(self):
        path = self.xlsx()
        result = ids.check_csv_text_ids([path, path], 'xlsx', self.source,
                                       import_filter=FILTER, import_options='44,34,76,1')
        self.assertEqual(result['status'], 'not_checked')

    def test_bad_zip_not_green(self):
        path = self.root / 'bad.xlsx'; path.write_bytes(b'not a workbook')
        self.assertEqual(self.check(path)['status'], 'not_checked')

    def test_dtd_never_expanded(self):
        path = self.xlsx()
        self.rewrite_zip(path, lambda name, raw: b'<!DOCTYPE worksheet [<!ENTITY x "secret">]>' + raw
                         if name.startswith('xl/worksheets/') else raw)
        result = self.check(path)
        self.assertEqual(result['status'], 'not_checked')
        self.assertEqual(result['reason'], 'csv_id_xlsx_xml_entities')

    def test_source_identity_change_is_not_a_saved_verdict(self):
        path = self.xlsx()
        original = ids._stamp
        source_calls = 0
        def changed(p):
            nonlocal source_calls
            value = original(p)
            if Path(p) == self.source:
                source_calls += 1
                if source_calls > 1:
                    return (*value[:-1], value[-1] + 1)
            return value
        with patch.object(ids, '_stamp', side_effect=changed):
            with self.assertRaises(GuardError) as ctx:
                self.check(path)
        self.assertEqual(ctx.exception.code, 'SOURCE_CHANGED_DURING_PREPARATION')

    def test_cancel_mid_sheet_closes_handles(self):
        from openpyxl.worksheet._read_only import ReadOnlyWorksheet
        path = self.xlsx()
        event = threading.Event()
        original = ReadOnlyWorksheet._cells_by_row
        closed = []
        def interrupted(sheet, *args, **kwargs):
            iterator = original(sheet, *args, **kwargs)
            try:
                for number, row in enumerate(iterator, 1):
                    if number == 2:
                        event.set()
                    yield row
            finally:
                iterator.close()
                closed.append(True)
        with patch.object(ReadOnlyWorksheet, '_cells_by_row', interrupted):
            with self.assertRaises(GuardError) as ctx:
                self.check(path, cancel=event)
        self.assertEqual(ctx.exception.code, 'CANCELLED')
        self.assertTrue(closed)
        moved = path.with_name('moved.xlsx')
        path.rename(moved); moved.unlink()

    def test_late_id_not_in_hint_still_cannot_silently_pass(self):
        rows = [['id', 'amount']] + [[str(i), 42] for i in range(1, 260)] + [['000260', 42]]
        self.csv(rows)
        self.assertEqual(csv_import_hint(self.source, NATIVE)['options'], '44,34,76,1')
        rows[-1][0] = 260
        with self.assertRaises(GuardError) as ctx:
            ids.check_csv_text_ids([self.xlsx(rows)], 'xlsx', self.source,
                                  import_filter=FILTER, import_options='44,34,76,1')
        self.assertEqual(ctx.exception.code, 'ENGINE_OUTPUT_CONTENT_LOSS')

    def test_adapter_barrier_works_without_optional_quality(self):
        from app_backend.converter_guard.adapter import LibreOfficeAdapter
        exe = self.root / 'soffice.exe'; exe.touch()
        adapter = LibreOfficeAdapter(exe)
        adapter._registry = lambda: SimpleNamespace(fingerprint='b29-fixture', html_route_reason=lambda *_:None)
        adapter._local.import_filter = FILTER
        adapter._local.import_options = '44,34,76,1,1/2'
        broken = self.xlsx([ROWS[0], [1, 'Строка ё', 42], [2, 'Tail', -7]])
        with patch.object(adapter, '_convert_native', return_value=[broken]), \
                patch('app_backend.converter_guard.native_metafile.native_metafile_target', return_value=None), \
                patch('app_backend.converter_guard.adapter.check_native_output', return_value=None):
            with self.assertRaises(GuardError) as ctx:
                adapter.convert(self.source, 'xlsx', self.root / 'work', cancel=threading.Event(),
                                timeout=None, on_started=lambda: None)
        self.assertEqual(ctx.exception.code, 'ENGINE_OUTPUT_CONTENT_LOSS')


if __name__ == '__main__':
    unittest.main()
