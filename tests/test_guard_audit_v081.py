"""Regression tests for the 21 Oct 2026 F01-F12 Guard audit.

Run locally with: python -m unittest discover -s tests -p test_guard_audit_v081.py -v
This file never modifies application data; every fixture is temporary.
"""
import io
import json
import os
from pathlib import Path
from types import SimpleNamespace
import tempfile
import unittest
from unittest.mock import patch

from app_backend.converter_guard import ConverterGuard
from app_backend.converter_guard.core import _missing_html_resources
from app_backend.converter_guard import preflight
from app_backend.excel import merge_xlsx, MergeError
from test_converter_guard_v02 import FakeAdapter, wait_finished


class AuditGuardV081(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.guards = []

    def tearDown(self):
        for guard in self.guards:
            guard.close()
        self.tmp.cleanup()

    def guard(self):
        result = ConverterGuard(adapter=FakeAdapter(), recovery_parent=self.root / 'recovery')
        self.guards.append(result)
        return result

    def create_finished(self, names):
        guard = self.guard()
        files = [{'id':str(i),'name':name,'size':1} for i,name in enumerate(names)]
        identity = guard.create(files, ['pdf'])['jobId']
        for row in files:
            guard.upload(identity,row['id'],io.BytesIO(b'x'),1)
        guard.start(identity)
        state = wait_finished(guard,identity)
        self.assertEqual(state['succeeded'],len(names))
        return guard,identity,state

    def test_f02_corrupt_mid_ledger_never_releases_unseen_outputs(self):
        guard,job,state = self.create_finished(['a.docx','b.docx','c.docx'])
        folder=guard._jobs[job].root
        self.assertEqual(len(state['result']['outputs']),3)
        guard.close()
        ledger=folder/'results.jsonl'
        lines=ledger.read_text(encoding='utf-8').splitlines()
        for index,line in enumerate(lines):
            row=json.loads(line)
            if row['record'].get('kind')=='task_begin' and row['record'].get('index')==1:
                row['checksum']='0'*64
                lines[index]=json.dumps(row)
                break
        else:
            self.fail('Expected the second committed task')
        ledger.write_text('\n'.join(lines)+'\n',encoding='utf-8')
        resumed=self.guard()
        warnings=resumed.recoverable()['warnings']
        self.assertTrue(any(w['code']=='RECOVERY_LEDGER_INCOMPLETE' for w in warnings))
        snapshot=resumed.status(job)
        self.assertEqual(len(snapshot['result']['outputs']),3)
        self.assertTrue(snapshot['salvaged'])
        destination=self.root/'saved';destination.mkdir()
        copied=resumed.save_folder(snapshot['result']['ticket'],destination)
        self.assertTrue(copied['saved'])
        self.assertFalse(copied['savedAll'])
        self.assertEqual(len(list(destination.rglob('*.pdf'))),3)
        resumed.release(job)
        self.assertTrue(folder.exists(),'Never delete unindexed original outputs')
        self.assertEqual(len(list(folder.glob('work/*/output/*.pdf'))),3)

    def test_f03_interrupted_preparation_is_reported_not_silent(self):
        guard=self.guard()
        job=guard.create([{'id':'f','name':'staged.docx','size':5}],
                         prepare_only=True)['jobId']
        guard.upload(job,'f',io.BytesIO(b'hello'),5)
        folder=guard._jobs[job].root
        self.assertTrue((folder/'job.json').exists())
        # Mimic process death: release OS lock without calling clean close().
        guard._jobs[job].store.close()
        guard._jobs.clear()
        guard._closed=True
        guard._executor.shutdown(wait=True)
        resumed=self.guard()
        warnings=resumed.recoverable()['warnings']
        self.assertTrue(any(w['code']=='RECOVERY_UNFINISHED_PREPARATION' for w in warnings))
        self.assertTrue(folder.exists())

    def test_f05_cancel_during_packaging_sets_bundle_cancel(self):
        guard=self.guard()
        job=guard.create([{'id':'f','name':'x.docx','size':1}],['pdf'])['jobId']
        guard._jobs[job].state='packaging'
        guard.cancel(job)
        self.assertTrue(guard._jobs[job].bundle_cancel.is_set())

    def test_f10_mathml_content_is_not_an_extension_guess(self):
        from app_backend.converter_guard.preflight_native import _mathml_root
        mathml=self.root/'renamed.bin'
        mathml.write_text('<math xmlns="http://www.w3.org/1998/Math/MathML"><mfrac><mn>1</mn><mi>x</mi></mfrac></math>',encoding='utf-8')
        self.assertTrue(_mathml_root(mathml))
        mathml.write_text('<html><div>Not MathML</div></html>',encoding='utf-8')
        self.assertFalse(_mathml_root(mathml))

    def test_f10_formula_proof_requires_content_and_native_service(self):
        from app_backend.converter_guard.routes import Registry
        native=Registry.__new__(Registry)
        service='com.sun.star.formula.FormulaProperties'
        native.valid=True
        native.fingerprint='fixture'
        native.import_services=frozenset({service})
        native.exports={'pdf':{service}}
        native.uncertain_extensions=set()
        proof={'state':'detected','basis':'native_preflight_mathml_content',
               'nativeKind':4,'mathmlRoot':True,
               'registryFingerprint':'fixture','documentServices':[service]}
        self.assertEqual(native.observed_decision('pdf',proof),'registered')
        self.assertEqual(native.observed_decision('pdf',{**proof,'mathmlRoot':False}),'undetermined')
        self.assertEqual(native.observed_decision('pdf',{**proof,'nativeKind':5}),'undetermined')

    def test_f09_html_with_missing_local_asset_not_silently_approved(self):
        html=self.root/'index.html'
        html.write_text('<html><img src="relative.png"></html>',encoding='utf-8')
        self.assertTrue(_missing_html_resources(html))
        html.write_text('<html><img src="data:image/png;base64,abc"><img src="https://example.invalid/p.png"></html>',encoding='utf-8')
        self.assertFalse(_missing_html_resources(html))
        html.write_text('<html><img srcset="data:image/png;base64,aabb 1x, https://example.invalid/x.png 2x"></html>',encoding='utf-8')
        self.assertFalse(_missing_html_resources(html))
        html.write_text('<html><img srcset="data:image/png;base64,aabb 1x, local.png 2x"></html>',encoding='utf-8')
        self.assertTrue(_missing_html_resources(html))
        html.write_text('<html><a href="chapter2.html">next</a></html>',encoding='utf-8')
        self.assertFalse(_missing_html_resources(html))

    def test_f01_windows_fallback_uses_completed_worker_json(self):
        program=self.root/'program';program.mkdir()
        executable=program/'soffice.com';executable.write_bytes(b'office')
        (program/'python.exe').write_bytes(b'python')
        source=self.root/'document.docx';source.write_bytes(b'doc')
        service='com.sun.star.text.TextDocument'
        native=SimpleNamespace(valid=True,imports={'docx':{service}},
            filters={'testimport':(service,frozenset({'docx'}),frozenset({'IMPORT'}))},
            fingerprint='qa123')
        fake_os=SimpleNamespace(name='nt',environ=os.environ)
        def run_with(document):
            class StubTree:
                def __init__(self,command,*,cwd,env):
                    self.work=cwd
                    self.process=SimpleNamespace(poll=lambda:0,returncode=0)
                def __enter__(self):
                    (self.work/'result.json').write_text(json.dumps(document),encoding='utf-8')
                    return self
                def __exit__(self,*_):
                    return False
            with patch.object(preflight,'os',fake_os),patch.object(preflight,'ProcessTree',StubTree):
                return preflight.inspect_source(executable,source,self.root/('work'+str(len(list(self.root.glob('work*'))))),
                    cancel=SimpleNamespace(is_set=lambda:False,wait=lambda *_:False),timeout=1,
                    prepare_profile=lambda path:path.mkdir(),native=native)
        first=run_with({'state':'detected','modelLoaded':True,'documentServices':[service],
                        'filterName':'testimport','typeName':'writer'})
        self.assertEqual(first['basis'],'native_preflight_uno_document')
        self.assertEqual(first['selectedImportFilter'],'testimport')
        second=run_with({'state':'undetermined','code':'PREFLIGHT_IMPORT_FAILED','typeName':'unknown'})
        self.assertEqual(second['code'],'PREFLIGHT_IMPORT_FAILED')
        self.assertEqual(second['diagnosticType'],'unknown')


class AuditExcelV081(unittest.TestCase):
    @staticmethod
    def make_book(*, rich=False, row_height=34, col_width=28):
        from openpyxl import Workbook
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont
        wb=Workbook();ws=wb.active
        ws.sheet_format.defaultRowHeight=row_height
        ws.sheet_format.defaultColWidth=col_width
        ws.append(['Column'])
        if rich:
            ws.append([CellRichText(TextBlock(InlineFont(b=True),'Bold'),TextBlock(InlineFont(i=True),'Italic'))])
        else:
            ws.append(['Plain'])
        stream=io.BytesIO();wb.save(stream);wb.close()
        return stream.getvalue()

    def test_f11_preserve_default_geometry(self):
        from openpyxl import load_workbook
        data=self.make_book()
        merged=merge_xlsx([('one.xlsx',data),('two.xlsx',data)])
        wb=load_workbook(io.BytesIO(merged))
        self.assertEqual(wb.active.sheet_format.defaultRowHeight,34)
        self.assertEqual(wb.active.sheet_format.defaultColWidth,28)
        wb.close()

    def test_f11_reject_unsupported_rich_text_instead_of_losing_it(self):
        with self.assertRaises(MergeError) as error:
            merge_xlsx([('one.xlsx',self.make_book(rich=True)),
                        ('two.xlsx',self.make_book())])
        self.assertIn('rich text',str(error.exception))

    def test_f11_reject_unequal_default_geometry(self):
        with self.assertRaises(MergeError):
            merge_xlsx([('one.xlsx',self.make_book()),('two.xlsx',self.make_book(col_width=40))])


if __name__ == '__main__':
    unittest.main()
