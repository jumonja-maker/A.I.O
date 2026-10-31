"""Audit 101 regression contracts. Disposable files; no native-quality claims."""
from concurrent.futures import ThreadPoolExecutor
import hashlib
import io
import json
import os
from pathlib import Path
import secrets
import shutil
import tempfile
import threading
import unittest
from unittest.mock import patch
import zipfile

from app_backend.atomic_metadata import atomic_json
from app_backend.converter_guard import ConverterGuard, GuardError
from app_backend.converter_guard.result_store import ResultStore
from app_backend.converter_transport import ConverterTransport, Result
from app_backend.legacy_result_journal import LegacyWorkspace, recover_workspaces, file_hash
from app_backend.libreoffice_engine import LibreOfficeError
from app_backend.legacy_storage_budget import measured_bytes, LegacyRunControl
from test_converter_guard_v02 import FakeAdapter, wait_finished


class Disposable(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name).resolve()

    def transport(self, **options):
        value = ConverterTransport(recovery_parent=self.root/'legacy', **options)
        self.addCleanup(value.close)
        return value

    def published(self):
        workspace = LegacyWorkspace.create(self.root/'legacy')
        path = Path(workspace.name)/'native'/'0'/'0'/'output'/'test.pdf'
        path.parent.mkdir(parents=True)
        path.write_bytes(b'published')
        identity = secrets.token_hex(16)
        result = Result(workspace, path, path.name, 'application/pdf', request_id=identity,
                        size=path.stat().st_size, target_report={'succeeded':1,'failed':0}, format_ids=('pdf',))
        workspace.record(result)
        workspace.close()
        return Path(workspace.name), path, identity

    def orphan(self):
        root, path, identity = self.published()
        (root/'legacy-result.json').write_text('{}')
        return root, path, identity

    def ready(self, formats=('pdf', 'html')):
        guard = ConverterGuard(adapter=FakeAdapter(), recovery_parent=self.root/'guard')
        self.addCleanup(guard.close)
        identity = guard.create([{'id':'f','name':'source.docx','size':4}], list(formats))['jobId']
        guard.upload(identity, 'f', io.BytesIO(b'keep'), 4)
        guard.start(identity, archive_mode='on_demand')
        wait_finished(guard, identity)
        return guard, guard._jobs[identity]

    def bundle(self):
        guard, job = self.ready()
        guard.rebuild_bundle(job.id)
        wait_finished(guard, job.id)
        destination = self.root/'external.zip'
        with guard.open_result(job.bundle.ticket) as (source, _, _), destination.open('xb') as target:
            shutil.copyfileobj(source, target)
        return guard, job, destination

    def link(self, target, link):
        # Hard-link fallback still proves an existing alias is never truncated.
        try:
            link.symlink_to(target)
        except OSError:
            os.link(target, link)


class MetadataTests(Disposable):
    def test_existing_receipt_is_not_overwritten(self):
        target = self.root/'status.json'; target.write_bytes(b'private-old')
        with self.assertRaises(FileExistsError): atomic_json(target, {'new':1}, overwrite=False)
        self.assertEqual(target.read_bytes(), b'private-old')
        self.assertFalse(list(self.root.glob('.aio-json-*')))

    def test_predictable_new_alias_is_not_followed(self):
        other = self.root/'other'; other.write_bytes(b'private-old')
        target = self.root/'status.json'
        self.link(other, target.with_suffix('.new'))
        atomic_json(target, {'ok':True})
        self.assertEqual(other.read_bytes(), b'private-old')
        self.assertEqual(json.loads(target.read_text()), {'ok':True})

    def test_existing_final_link_is_not_followed_or_replaced(self):
        other = self.root/'other'; other.write_bytes(b'private-old')
        target = self.root/'status.json'; self.link(other, target)
        with self.assertRaises(FileExistsError): atomic_json(target, {'new':1}, overwrite=False)
        self.assertEqual(other.read_bytes(), b'private-old')
        self.assertEqual(target.read_bytes(), b'private-old')

    def test_fsync_failure_preserves_original_and_cleans_own_stage(self):
        target = self.root/'status.json'; target.write_bytes(b'old')
        with patch('app_backend.atomic_metadata.os.fsync', side_effect=OSError(28, 'full')):
            with self.assertRaises(OSError): atomic_json(target, {'new':1})
        self.assertEqual(target.read_bytes(), b'old')
        self.assertFalse(list(self.root.glob('.aio-json-*')))

    def test_fixed_random_collision_does_not_remove_unowned_stage(self):
        temporary = self.root/('.aio-json-'+'a'*32+'.new'); temporary.write_bytes(b'other writer')
        with patch('app_backend.atomic_metadata.secrets.token_hex', return_value='a'*32):
            with self.assertRaises(FileExistsError): atomic_json(self.root/'status.json', {'x':1})
        self.assertEqual(temporary.read_bytes(), b'other writer')

    def test_concurrent_publish_has_exactly_one_winner(self):
        target = self.root/'status.json'
        def attempt(number):
            try:
                atomic_json(target, {'writer':number}, overwrite=False)
                return number
            except FileExistsError:
                return None
        with ThreadPoolExecutor(max_workers=8) as pool:
            winners = [i for i in pool.map(attempt, range(32)) if i is not None]
        self.assertEqual(len(winners), 1)
        self.assertEqual(json.loads(target.read_text())['writer'], winners[0])
        self.assertFalse(list(self.root.glob('.aio-json-*')))

    def test_external_zip_uses_distinct_receipt_and_preserves_neighbor(self):
        guard, job, destination = self.bundle()
        other = destination.with_name(destination.name+'.AIO-save-status.json')
        other.write_bytes(b'user-owned')
        receipt = guard.verify_saved_download(job.bundle.ticket, destination)
        self.assertEqual(other.read_bytes(), b'user-owned')
        self.assertNotEqual(receipt['receiptFile'], other.name)
        self.assertEqual(json.loads((self.root/receipt['receiptFile']).read_text()), receipt)
        self.assertTrue(guard.confirm_saved(job.id, job.bundle.ticket, save_id=receipt['saveId'], compact=True)['savedAll'])

    def test_external_zip_predictable_stage_alias_stays_intact(self):
        guard, job, destination = self.bundle()
        other = self.root/'unrelated'; other.write_bytes(b'user-owned')
        receipt_path = destination.with_name(destination.name+'.AIO-save-status.json')
        self.link(other, receipt_path.with_suffix('.new'))
        self.assertTrue(guard.verify_saved_download(job.bundle.ticket, destination)['readbackVerified'])
        self.assertEqual(other.read_bytes(), b'user-owned')

    def test_two_verifications_preserve_first_immutable_receipt(self):
        guard, job, destination = self.bundle()
        first = guard.verify_saved_download(job.bundle.ticket, destination)
        raw = (self.root/first['receiptFile']).read_bytes()
        second = guard.verify_saved_download(job.bundle.ticket, destination)
        self.assertNotEqual(first['saveId'], second['saveId'])
        self.assertNotEqual(first['receiptFile'], second['receiptFile'])
        self.assertEqual((self.root/first['receiptFile']).read_bytes(), raw)
        with self.assertRaises(GuardError): guard.confirm_saved(job.id, job.bundle.ticket, save_id=first['saveId'])


class AvailabilityTests(Disposable):
    def test_guard_missing_file_is_unavailable_not_saved(self):
        guard, job = self.ready(('pdf',))
        artifact = job.artifacts[0]; artifact.path.unlink()
        state = guard.status(job.id)
        self.assertEqual(state['state'], 'failed')
        self.assertEqual(state['conversionState'], 'ready')
        self.assertEqual(state['resultAvailability'], 'unavailable')
        self.assertEqual(state['unavailableOutputs'], 1)
        self.assertEqual(state['unsavedDocuments'], 0)
        self.assertIsNone(state['result'])
        self.assertFalse(artifact.saved)
        self.assertEqual(guard.report(job.id)['succeeded'], 1)
        row = guard.recoverable(scan=True)['jobs'][0]
        self.assertEqual(row['unsavedOutputs'], 0)
        self.assertEqual(row['unavailableOutputs'], 1)

    def test_guard_one_missing_preserves_other_downloads(self):
        guard, job = self.ready()
        job.artifacts[0].path.unlink()
        state = guard.status(job.id)
        self.assertEqual(state['state'], 'partial')
        self.assertEqual(state['resultAvailability'], 'partial_unavailable')
        self.assertEqual(len(state['result']['outputs']), 2)
        with guard.open_result(state['result']['ticket']) as (stream, _, _):
            self.assertTrue(stream.read())

    def test_guard_same_size_replacement_is_unavailable(self):
        guard, job = self.ready(('pdf',))
        path = job.artifacts[0].path
        replacement = path.with_suffix('.replacement'); replacement.write_bytes(b'evil')
        os.replace(replacement, path)
        self.assertEqual(guard.status(job.id)['unavailableOutputs'], 1)
        with self.assertRaises(GuardError):
            with guard.open_result(job.artifacts[0].ticket): pass
        self.assertFalse(job.artifacts[0].saved)

    def test_guard_delta_and_inventory_poll_do_no_path_probes(self):
        guard, job = self.ready(('pdf',))
        revision = guard.status(job.id)['revision']
        with patch('app_backend.file_identity.file_identity', side_effect=AssertionError('poll did I/O')):
            self.assertTrue(guard.status(job.id, since_revision=revision, recover=False)['unchanged'])
            guard.recoverable(scan=False)

    def test_missing_before_save_marks_availability_and_retains_other_bytes(self):
        guard, job = self.ready()
        artifact = job.artifacts[0]; artifact.path.unlink()
        with self.assertRaises(GuardError): guard.save_folder(artifact.ticket, self.root)
        self.assertEqual(artifact.availability_error, 'RESULT_UNAVAILABLE')
        self.assertFalse(any(a.saved for a in job.artifacts))
        self.assertTrue(all(a.path.exists() for a in job.artifacts[1:]))
        self.assertFalse(job.saving)

    def test_legacy_manual_rescan_marks_missing_not_ready(self):
        root, path, identity = self.published()
        transport = self.transport()
        path.unlink(); transport.scan_recovery()
        row = transport.inventory()['items'][0]
        self.assertEqual(row['state'], 'unavailable')
        self.assertFalse(row['available'])
        self.assertEqual(transport.inventory()['saveableTotal'], 0)
        self.assertEqual(transport.recover(identity)['state'], 'unavailable')
        self.assertTrue(root.exists())
        self.assertTrue(transport.decline_recovery(request_id=identity)['declined'])

    def test_legacy_selected_read_detects_same_size_replacement_without_rescan(self):
        root, path, identity = self.published()
        transport = self.transport()
        replacement = root/'replacement'; replacement.write_bytes(b'x'*path.stat().st_size)
        os.replace(replacement, path)
        self.assertEqual(transport.recover(identity)['state'], 'unavailable')
        with self.assertRaises(LibreOfficeError):
            with transport.lease(transport.results_by_request[identity]): pass
        self.assertEqual(transport.storage_stats()['leasedResults'], 0)

    def test_legacy_inventory_poll_has_no_filesystem_read(self):
        self.published(); transport = self.transport()
        with patch('app_backend.converter_transport.owned_file_identity', side_effect=AssertionError('poll I/O')):
            self.assertEqual(transport.inventory()['total'], 1)


class LegacyOwnershipTests(Disposable):
    def test_completed_scan_cannot_publish_after_close(self):
        root, _, _ = self.published()
        transport = self.transport(lazy_recovery=True)
        acquired, resume = threading.Event(), threading.Event()
        original = recover_workspaces
        def paused(*args, **kwargs):
            result = original(*args, **kwargs)
            acquired.set(); self.assertTrue(resume.wait(5))
            return result
        outcomes = []
        with patch('app_backend.converter_transport.recover_workspaces', paused):
            worker = threading.Thread(target=lambda: outcomes.append(transport.scan_recovery()))
            worker.start(); self.assertTrue(acquired.wait(5))
            transport.close(); resume.set(); worker.join(5)
        self.assertFalse(worker.is_alive())
        self.assertEqual(outcomes[0]['state'], 'cancelled')
        self.assertFalse(transport.results)
        owner = LegacyWorkspace(root); owner.close()

    def test_root_limit_applies_before_salvage_writes(self):
        root, path, _ = self.orphan(); original = path.read_bytes()
        transport = self.transport(max_owned_roots=1)
        with patch('app_backend.converter_transport.tempfile.TemporaryDirectory') as create:
            with self.assertRaises(LibreOfficeError) as caught: transport.salvage_workspace(root.name)
        self.assertEqual(caught.exception.code, 'LEGACY_STORAGE_LIMIT')
        create.assert_not_called()
        self.assertEqual(transport.storage_stats()['ownedRoots'], 1)
        self.assertEqual(path.read_bytes(), original)

    def test_repeated_salvage_cannot_accumulate_duplicate_copies(self):
        root, path, _ = self.orphan()
        transport = self.transport(max_owned_roots=3)
        first = transport.salvage_workspace(root.name)
        with self.assertRaises(LibreOfficeError) as caught: transport.salvage_workspace(root.name)
        self.assertEqual(caught.exception.code, 'LEGACY_SALVAGE_EXISTS')
        self.assertEqual(transport.storage_stats()['ownedRoots'], 2)
        self.assertEqual(len(transport.results), 1)
        self.assertFalse(first['conversionSuccess'])
        self.assertTrue(path.exists())

    def test_salvage_slot_is_reserved_while_native_copy_is_in_flight(self):
        root, _, _ = self.orphan(); transport = self.transport(max_owned_roots=2)
        from app_backend import converter_transport as module
        original = module.salvage_unverified_workspace
        def inspect(*args, **kwargs):
            self.assertEqual(transport.storage_stats()['ownedRoots'], 2)
            with self.assertRaises(LibreOfficeError) as caught:
                transport.begin(secrets.token_hex(16), None)
            self.assertEqual(caught.exception.code, 'LEGACY_STORAGE_LIMIT')
            return original(*args, **kwargs)
        with patch.object(module, 'salvage_unverified_workspace', inspect): transport.salvage_workspace(root.name)
        self.assertFalse(transport._salvage_roots)

    def test_inflight_salvage_bytes_are_counted(self):
        root, _, _ = self.orphan()
        transport = self.transport(max_owned_bytes=16*1024*1024)
        with tempfile.TemporaryDirectory(dir=self.root) as folder:
            class Owner: name = folder
            (Path(folder)/'work').write_bytes(b'x'*2000)
            transport._salvage_roots[root.name] = Owner()
            transport.check_storage()
            self.assertGreaterEqual(transport.used_bytes, 2000)
            transport._salvage_roots.clear()

    def test_salvage_failure_releases_slot_and_keeps_source(self):
        root, path, _ = self.orphan(); transport = self.transport()
        before = transport.storage_stats()['ownedRoots']
        with patch('app_backend.converter_transport.salvage_unverified_workspace', side_effect=OSError(28, 'full')):
            with self.assertRaises(OSError): transport.salvage_workspace(root.name)
        self.assertEqual(transport.storage_stats()['ownedRoots'], before)
        self.assertFalse(transport.results)
        self.assertFalse(transport._salvage_roots)
        self.assertTrue(path.exists())

    def test_bulk_decline_removes_original_and_derivative(self):
        root, path, _ = self.orphan(); transport = self.transport()
        saved = transport.salvage_workspace(root.name)
        derivative = transport.results[saved['result']['ticket']].path
        outcome = transport.decline_all_recovery()
        self.assertEqual(outcome['notDeclined'], 0)
        self.assertFalse(root.exists())
        self.assertFalse(derivative.exists())
        self.assertEqual(transport.inventory()['total'], 0)
        self.assertTrue(outcome['historyUntouched'])

    def test_individual_decline_hides_leased_derivative_until_read_finishes(self):
        root, _, _ = self.orphan(); transport = self.transport()
        ticket = transport.salvage_workspace(root.name)['result']['ticket']
        with transport.lease(ticket) as result:
            outcome = transport.decline_recovery(workspace_id=root.name)
            self.assertTrue(outcome['cleanupPending'])
            self.assertTrue(result.path.exists())
            self.assertEqual(transport.inventory()['total'], 0)
        self.assertNotIn(ticket, transport.results)
        self.assertFalse(result.path.exists())

    def test_bulk_decline_reports_unreadable_storage(self):
        transport = self.transport(lazy_recovery=True)
        with patch('app_backend.converter_transport.recover_workspaces', side_effect=PermissionError()):
            with self.assertRaises(LibreOfficeError) as caught: transport.decline_all_recovery()
        self.assertEqual(caught.exception.code, 'LEGACY_STORAGE_UNAVAILABLE')
        self.assertFalse(transport.recovery_available)

    def test_bulk_decline_does_not_touch_current_request(self):
        root, _, _ = self.orphan(); transport = self.transport()
        current = transport.begin(secrets.token_hex(16), None)
        (current.root/'input').write_bytes(b'current')
        transport.decline_all_recovery()
        self.assertEqual((current.root/'input').read_bytes(), b'current')
        transport.finish(current)

    def test_cancelled_hash_never_opens_payload(self):
        cancel = threading.Event(); cancel.set()
        with patch.object(Path, 'open') as opening:
            with self.assertRaises(InterruptedError): file_hash(self.root/'missing', cancel=cancel)
        opening.assert_not_called()

    def test_salvage_readback_has_no_redundant_testzip_pass(self):
        root, _, _ = self.orphan(); transport = self.transport()
        with patch.object(zipfile.ZipFile, 'testzip', side_effect=AssertionError('redundant read')):
            result = transport.salvage_workspace(root.name)
        with zipfile.ZipFile(transport.results[result['result']['ticket']].path) as archive:
            report = json.loads(archive.read('AIO-salvage-report.json'))
            for entry in report['files']:
                self.assertEqual(hashlib.sha256(archive.read(entry['path'])).hexdigest(), entry['sha256'])

    def test_shutdown_during_salvage_cannot_publish(self):
        root, path, _ = self.orphan(); transport = self.transport()
        from app_backend import converter_transport as module
        original = module.salvage_unverified_workspace
        def finish_then_close(*args, **kwargs):
            value = original(*args, **kwargs); transport.close(); return value
        with patch.object(module, 'salvage_unverified_workspace', finish_then_close):
            with self.assertRaises(LibreOfficeError): transport.salvage_workspace(root.name)
        self.assertFalse(transport.results)
        self.assertFalse(transport._salvage_roots)
        self.assertTrue(path.exists())


class JournalSchemaTests(Disposable):
    def test_missing_path_and_wrong_schema_do_not_hide_other_valid_roots(self):
        bad_root, _, _ = self.published()
        good_root, _, good_id = self.published()
        raw = json.loads((bad_root/'legacy-result.json').read_text())
        raw.pop('checksum'); raw.pop('path'); raw['checksum'] = ResultStore._digest(raw)
        (bad_root/'legacy-result.json').write_text(json.dumps(raw))
        recovered, issues = recover_workspaces(self.root/'legacy')
        try:
            self.assertEqual([v['requestId'] for _, v in recovered], [good_id])
            self.assertEqual(issues[0]['code'], 'LEGACY_RECOVERY_UNVERIFIED')
            self.assertTrue(bad_root.exists())
        finally:
            for workspace, _ in recovered: workspace.close()

    def test_non_object_journal_is_isolated(self):
        root, _, _ = self.published()
        for value in ([], None, 'text', 123, {'path':None}):
            with self.subTest(value=value):
                (root/'legacy-result.json').write_text(json.dumps(value))
                recovered, issues = recover_workspaces(self.root/'legacy')
                self.assertEqual(recovered, [])
                self.assertEqual(issues[0]['code'], 'LEGACY_RECOVERY_UNVERIFIED')


class QuotaTests(Disposable):
    def test_measurement_counts_all_files_and_sees_later_writes(self):
        expected = 0
        for number in range(100):
            path = self.root/str(number % 3)/str(number)
            path.parent.mkdir(exist_ok=True); path.write_bytes(b'x'*number); expected += number
        self.assertEqual(measured_bytes([self.root]), expected)
        (self.root/'new').write_bytes(b'changed')
        self.assertEqual(measured_bytes([self.root]), expected + 7)

    def test_entry_and_depth_bounds_preserved(self):
        path = self.root/'a'/'b'/'c'; path.mkdir(parents=True)
        (path/'file').write_bytes(b'x')
        with patch('app_backend.legacy_storage_budget.MAX_SCAN_ENTRIES', 2):
            with self.assertRaises(OSError): measured_bytes([self.root])
        with patch('app_backend.legacy_storage_budget.MAX_SCAN_DEPTH', 1):
            with self.assertRaises(OSError): measured_bytes([self.root])

    def test_poll_cooldown_starts_after_expensive_measurement(self):
        checks = []
        control = LegacyRunControl(threading.Event(), lambda: checks.append(True))
        with patch('app_backend.legacy_storage_budget.time.monotonic', side_effect=[0.0, 0.4, 0.6]):
            self.assertFalse(control.is_set()); self.assertFalse(control.is_set())
        self.assertEqual(len(checks), 1)
        self.assertEqual(control.next_scan, .9)


class ExcelMetadataTests(Disposable):
    @staticmethod
    def book(change=lambda ws:None):
        from openpyxl import Workbook
        workbook = Workbook(); sheet = workbook.active
        sheet.append(['Name','Code','Flag']); sheet.append(['00123', 12.5, True])
        change(sheet)
        stream = io.BytesIO(); workbook.save(stream); workbook.close()
        return stream.getvalue()

    def test_range_width_expands_to_every_column(self):
        from openpyxl import load_workbook
        from app_backend.excel import merge_xlsx
        def change(sheet):
            sheet.column_dimensions['B'].min=2; sheet.column_dimensions['B'].max=3
            sheet.column_dimensions['B'].width=35
        data = self.book(change); output = merge_xlsx([('one.xlsx',data),('two.xlsx',data)])
        workbook=load_workbook(io.BytesIO(output)); self.addCleanup(workbook.close)
        self.assertEqual(workbook.active.column_dimensions['B'].width, 35)
        self.assertEqual(workbook.active.column_dimensions['C'].width, 35)
        self.assertEqual(list(workbook.active.values)[1:], [('00123',12.5,True)]*2)

    def test_unsupported_layout_is_rejected_before_silent_loss(self):
        from app_backend.excel import merge_xlsx, MergeError
        changes = [lambda s:setattr(s.row_dimensions[2],'outlineLevel',2),
                   lambda s:setattr(s.page_margins,'left',1.75),
                   lambda s:setattr(s.sheet_view,'rightToLeft',True),
                   lambda s:setattr(s.sheet_view,'showGridLines',False),
                   lambda s:setattr(s.column_dimensions['B'],'collapsed',True)]
        for number, change in enumerate(changes):
            with self.subTest(number=number):
                data=self.book(change); before=hashlib.sha256(data).digest()
                with self.assertRaises(MergeError): merge_xlsx([('one.xlsx',data),('two.xlsx',self.book())])
                self.assertEqual(hashlib.sha256(data).digest(), before)

    def test_overlapping_width_ranges_are_not_silently_reinterpreted(self):
        from app_backend.excel import merge_xlsx, MergeError
        def change(sheet):
            sheet.column_dimensions['B'].min=2; sheet.column_dimensions['B'].max=3
            sheet.column_dimensions['B'].width=35; sheet.column_dimensions['C'].width=12
        with self.assertRaises(MergeError): merge_xlsx([('one.xlsx',self.book(change)),('two.xlsx',self.book())])


class ExplicitAvailabilityIntegration(Disposable):
    def test_current_lookup_refresh_reconciles_without_loading_old_jobs(self):
        guard, job = self.ready(('pdf',))
        job.artifacts[0].path.unlink()
        with patch.object(guard, '_load_recovery', side_effect=AssertionError('Old jobs must not auto-load')):
            state = guard.status(job.id, recover=False)
        self.assertEqual(state['state'], 'failed')
        self.assertEqual(state['conversionState'], 'ready')
        self.assertEqual(state['unavailableOutputs'], 1)
        self.assertEqual(state['unsavedDocuments'], 0)
        self.assertIsNone(state['result'])

    def test_explicit_manual_scan_reconciles_even_when_history_already_loaded(self):
        guard, job = self.ready(('pdf',))
        guard._recovery_loaded = True
        job.artifacts[0].path.unlink()
        with patch.object(guard, '_recover_results', side_effect=AssertionError('Already loaded, no old-root rescan')):
            guard.start_recovery_scan()
            guard._recovery_future.result(timeout=5)
        snapshot = guard.recoverable(scan=False)
        self.assertEqual(snapshot['recoveryScan']['state'], 'completed')
        row = next(item for item in snapshot['jobs'] if item['jobId'] == job.id)
        self.assertEqual(row['unsavedOutputs'], 0)
        self.assertEqual(row['unavailableOutputs'], 1)
        self.assertEqual(row['state'], 'failed')

    def test_current_lookup_delta_poll_never_reads_file_metadata(self):
        guard, job = self.ready(('pdf',))
        revision = job.revision
        job.artifacts[0].path.unlink()
        with patch('app_backend.file_identity.file_identity', side_effect=AssertionError('No disk poll')):
            state = guard.status(job.id, recover=False, since_revision=revision, delta=True)
        self.assertTrue(state['unchanged'])
        # The next explicit refresh, not the poll, reconciles the missing output.
        self.assertEqual(guard.status(job.id, recover=False)['unavailableOutputs'], 1)


if __name__ == '__main__': unittest.main()
