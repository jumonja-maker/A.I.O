"""Cross-session comparison journeys over real HTTP and isolated analyzers.

Only native document export is a deterministic fixture. Reports, processes,
leases, cancellation, recovery journals, HTTP authentication and saved bytes
are real. Filesystem-failure injection is explicitly limited to its test.
"""
from concurrent.futures import ThreadPoolExecutor
import csv
import hashlib
import io
import json
from pathlib import Path
import sqlite3
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
from urllib.error import HTTPError
from urllib.parse import urlencode
from urllib.request import Request, urlopen
import zipfile

from openpyxl import Workbook, load_workbook
from pypdf import PdfWriter

from app_backend.converter_guard import ConverterGuard, GuardError
from test_comparison_service_v19 import ControlledNativeAdapter
from test_converter_guard_http_v02 import start_test_server
from test_converter_guard_v02 import wait_finished


def workbook_bytes(*, sheets=2, image=True):
    book = Workbook()
    for index in range(sheets):
        sheet = book.active if index == 0 else book.create_sheet()
        sheet.title = ('First', 'Second')[index] if index < 2 else 'Sheet' + str(index)
        sheet.append(['code', 'amount'])
        sheet.append(['00123' if index == 0 else 'SECOND_SHEET', 42 + index])
        sheet['B2'].number_format = '0.00'
    stream = io.BytesIO(); book.save(stream); book.close()
    if image:
        # Add a real OOXML drawing relationship without introducing Pillow as
        # a QA dependency. The first sheet displays the embedded 1px image.
        edited = io.BytesIO()
        with zipfile.ZipFile(stream) as source, zipfile.ZipFile(edited, 'w', zipfile.ZIP_DEFLATED) as archive:
            for item in source.infolist():
                raw = source.read(item.filename)
                if item.filename == 'xl/worksheets/sheet1.xml':
                    raw = raw.replace(b'</worksheet>',
                        b'<drawing xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rIdImage"/></worksheet>')
                elif item.filename == '[Content_Types].xml':
                    raw = raw.replace(b'</Types>',
                        b'<Default Extension="png" ContentType="image/png"/>'
                        b'<Override PartName="/xl/drawings/drawing1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/></Types>')
                archive.writestr(item.filename, raw)
            archive.writestr('xl/worksheets/_rels/sheet1.xml.rels',
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rIdImage" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing" Target="/xl/drawings/drawing1.xml"/></Relationships>')
            archive.writestr('xl/drawings/drawing1.xml',
                '<wsDr xmlns="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
                'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
                'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                '<oneCellAnchor><from><col>3</col><colOff>0</colOff><row>1</row><rowOff>0</rowOff></from>'
                '<ext cx="9525" cy="9525"/><pic><nvPicPr><cNvPr id="1" name="Probe"/><cNvPicPr/></nvPicPr>'
                '<blipFill><a:blip r:embed="rId1"/><a:stretch><a:fillRect/></a:stretch></blipFill>'
                '<spPr><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></spPr></pic><clientData/></oneCellAnchor></wsDr>')
            archive.writestr('xl/drawings/_rels/drawing1.xml.rels',
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="/xl/media/probe.png"/></Relationships>')
            archive.writestr('xl/media/probe.png', bytes.fromhex(
                '89504e470d0a1a0a0000000d4948445200000001000000010802000000907753de'
                '0000000c49444154789c63606060000000040001f61738550000000049454e44ae426082'))
        stream = edited
    return stream.getvalue()


def saved_payload_hashes(destination, job_id, expected_count):
    """Independently read saved payloads and validate all three save records."""
    names = ('AIO-conversion-report.json', 'AIO-results-manifest.json', 'AIO-save-status.json')
    metadata = {}
    for name in names:
        matches = list(Path(destination).rglob(name))
        assert len(matches) == 1, (name, matches)
        metadata[name] = matches[0]
    report = json.loads(metadata[names[0]].read_text(encoding='utf-8'))
    tasks = [target['taskId'] for row in report['files'] for target in row['targets']]
    assert len(tasks) == expected_count and all(task.startswith(job_id + ':') for task in tasks)
    manifest = json.loads(metadata[names[1]].read_text(encoding='utf-8'))
    status = json.loads(metadata[names[2]].read_text(encoding='utf-8'))
    assert manifest['jobId'] == status['jobId'] == job_id
    assert manifest['readbackVerified'] is True and status['state'] == 'completed'
    assert status['outputsWritten'] == expected_count and len(manifest['outputs']) == expected_count
    folder = metadata[names[1]].parent
    payloads = {}
    paths = set(metadata.values())
    for row in manifest['outputs']:
        relative = Path(row['path'])
        assert not relative.is_absolute() and '..' not in relative.parts
        path = folder / relative
        assert path.resolve().is_relative_to(folder.resolve()) and path.is_file()
        with path.open('rb') as stream:
            digest = hashlib.file_digest(stream, 'sha256').hexdigest()
        assert digest == row['sha256'] and path.stat().st_size == row['size']
        assert path.name not in payloads, 'Synthetic journey payload names must remain distinct'
        payloads[path.name] = digest; paths.add(path)
    assert set(Path(destination).rglob('*')) >= paths
    assert {path for path in Path(destination).rglob('*') if path.is_file()} == paths
    return payloads


class JourneyAdapter(ControlledNativeAdapter):
    """Native boundary only: CSV gets corrected; PDF remains unverified."""
    def __init__(self):
        super().__init__()
        self.local = threading.local()
        self.fail_pdf_retry = False

    def prepare_retry_recipe(self, recipe):
        self.local.recipe = recipe

    def export_info(self):
        recipe = getattr(self.local, 'recipe', None)
        return {'automaticRecipe': recipe, 'automaticParameters': {'fixture': True, 'recipe': recipe}}

    def convert(self, source, target, workspace, *, cancel, timeout, on_started,
                output_limit, on_document_service=None):
        if cancel.is_set():
            raise GuardError('CANCELLED', 'Controlled native boundary cancelled')
        recipe = getattr(self.local, 'recipe', None)
        if self.fail_pdf_retry and recipe == 'pdf_preserve_images':
            raise GuardError('NATIVE_EXPORT_FAILED', 'Controlled PDF retry rejection')
        on_started()
        if on_document_service:
            on_document_service({'state': 'detected', 'documentServices': ['com.sun.star.sheet.SpreadsheetDocument']})
        self.calls.append((source.name, target, recipe))
        folder = workspace / 'output'; folder.mkdir(parents=True)
        if target == 'csv' and source.suffix.lower() == '.xlsx':
            book = load_workbook(source, read_only=True, data_only=True)
            paths = []
            try:
                for index, sheet in enumerate(book.worksheets):
                    path = folder / (source.stem + '-' + sheet.title + '.csv')
                    with path.open('w', newline='', encoding='utf-8') as stream:
                        writer = csv.writer(stream)
                        for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                            values = list(row)
                            if recipe is None and index == 0 and row_index == 1:
                                values[1] = 0
                            writer.writerow(values)
                    paths.append(path)
                return paths
            finally:
                book.close()
        destination = folder / (source.stem + '.' + target)
        if target == 'pdf':
            writer = PdfWriter(); writer.add_blank_page(width=400, height=300)
            with destination.open('wb') as stream:
                writer.write(stream)
            writer.close()
        else:
            destination.write_bytes(source.read_bytes())
        return [destination]


class ComparisonJourneys(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='aio-comparison-journeys-')
        self.root = Path(self.temp.name)
        self.recovery = self.root / 'recovery'
        self._start_host()

    def _start_host(self):
        self.adapter = JourneyAdapter()
        self.guard = ConverterGuard(adapter=self.adapter, recovery_parent=self.recovery)
        self.server, self.thread, self.origin, self.token = start_test_server(self.guard)

    def _stop_host(self):
        self.server.api.close()
        self.guard.close()
        self.server.shutdown(); self.server.server_close(); self.thread.join(3)

    def tearDown(self):
        self._stop_host()
        self.temp.cleanup()

    def request(self, path, body=None, *, token=True, raw=False):
        headers = {'X-AIO-Token': self.token if token is True else str(token)} if token else {}
        if body is not None:
            headers['Content-Type'] = 'application/json'
            body = body if raw else json.dumps(body).encode()
        return urlopen(Request(self.origin + '/api/converter-guard/' + path, data=body, headers=headers), timeout=20)

    def call(self, path, body=None, **kwargs):
        with self.request(path, body, **kwargs) as response:
            return json.loads(response.read())

    def error(self, status, path, body=None, **kwargs):
        with self.assertRaises(HTTPError) as caught:
            self.request(path, body, **kwargs)
        self.assertEqual(caught.exception.code, status)
        data = caught.exception.read()
        self.assertNotIn(str(self.root).encode(), data)
        return json.loads(data)

    def converted(self, count=1, *, formats=None, workbook=False, sheets=2):
        data = workbook_bytes(sheets=sheets) if workbook else b'Preserved original paragraph.\n'
        suffix = '.xlsx' if workbook else '.txt'
        rows = [{'id': str(index), 'name': 'source-' + str(index) + suffix, 'size': len(data)}
                for index in range(count)]
        identity = self.guard.create(rows, formats or (['csv'] if workbook else ['txt']), retain_sources=True)['jobId']
        for row in rows:
            self.guard.upload(identity, row['id'], io.BytesIO(data), len(data))
        self.guard.start(identity, archive_mode='on_demand')
        return wait_finished(self.guard, identity, seconds=30)

    def wait(self, identity, predicate=None, seconds=30):
        end = time.monotonic() + seconds
        packet = None
        while time.monotonic() < end:
            packet = self.call('compare?' + urlencode({'job': identity}))
            if (predicate(packet) if predicate else packet['state'] in ('completed', 'cancelled', 'failed', 'interrupted')):
                return packet
            time.sleep(.025)
        self.fail('Comparison journey deadline: ' + repr(packet))

    def analyze(self, identity):
        self.call('compare/start', {'jobId': identity})
        result = self.wait(identity)
        self.assertEqual(result['state'], 'completed', result)
        return result

    def retry(self, identity, analysis, tasks=None):
        tasks = tasks or [item['taskId'] for item in analysis['items'] if item.get('recipeCandidates')]
        self.call('compare/retry', {'jobId': identity, 'analysisId': analysis['analysisId'], 'taskIds': tasks})
        return self.wait(identity, lambda packet: packet.get('retry', {}).get('state') in ('completed', 'failed', 'cancelled'))

    def test_completed_report_and_source_survive_real_host_recreation_without_new_conversion(self):
        state = self.converted(workbook=True)
        identity = state['jobId']; before = self.analyze(identity)
        source = self.guard._jobs[identity].files[0].source
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        self._stop_host(); self._start_host()
        restored = self.call('compare?' + urlencode({'job': identity}))
        self.assertEqual(restored['state'], 'completed')
        self.assertEqual(restored['analysisId'], before['analysisId'])
        self.assertEqual(restored['items'], before['items'])
        self.assertTrue(restored['recovered'])
        frozen = self.call('compare/start', {'jobId': identity})
        self.assertEqual(frozen['analysisId'], before['analysisId'])
        self.assertEqual(self.adapter.calls, [], 'Opening a recovered report must not export or reanalyze documents')
        with self.guard.comparison_inputs(identity) as records:
            self.assertEqual(hashlib.sha256(records[0]['source'].read_bytes()).hexdigest(), digest)
        self.assertTrue(self.guard.status(identity)['comparisonAvailable'])

    def test_multi_output_preview_selects_second_sheet_and_rejects_foreign_index(self):
        state = self.converted(workbook=True); identity = state['jobId']
        analysis = self.analyze(identity); task = analysis['items'][0]['taskId']
        detail = self.call('compare/item?' + urlencode({'job': identity, 'task': task}))
        previews = detail['outputPreviews']
        self.assertEqual(len(previews), 2)
        self.assertEqual([row['outputIndex'] for row in previews], [0, 1])
        second = self.call('compare/preview?' + urlencode({'job': identity, 'task': task, 'side': 'result', 'outputIndex': 1}))
        self.assertEqual(second['kind'], 'text'); self.assertIn('SECOND_SHEET', second['text'])
        first = self.call('compare/preview?' + urlencode({'job': identity, 'task': task, 'side': 'result', 'outputIndex': 0}))
        self.assertNotIn('SECOND_SHEET', first['text'])
        for invalid in (-1, 999, '../../private', '1.5'):
            self.error(400, 'compare/preview?' + urlencode({'job': identity, 'task': task, 'side': 'result', 'outputIndex': invalid}))

    def test_mixed_recipes_keep_each_child_verdict_and_all_prior_retry_cycles_after_recovery(self):
        state = self.converted(workbook=True, sheets=1, formats=['csv', 'pdf']); identity = state['jobId']
        analysis = self.analyze(identity)
        self.assertEqual({recipe for row in analysis['items'] for recipe in row['recipeCandidates']},
                         {'csv_preserve_values', 'pdf_preserve_images'})
        first = self.retry(identity, analysis)
        self.assertEqual(first['retry']['state'], 'completed', first['retry'])
        children = first['retry']['children']
        self.assertEqual(len(children), 2)
        verdicts = {child['recipe']: child['verification']['assessment'] for child in children}
        self.assertEqual(verdicts['csv_preserve_values'], 'checked_subset_improved')
        self.assertNotEqual(verdicts['pdf_preserve_images'], 'checked_subset_improved')
        self.assertNotEqual(first['retry']['verification']['assessment'], 'checked_subset_improved')
        original_run = first['retry']['runId']; old_jobs = set(first['retry']['jobIds'])
        per_task = []
        for offset in (0, 1):
            page = self.call('compare/retry-history?' + urlencode(
                {'job': identity, 'runId': original_run, 'offset': offset, 'limit': 1}))
            self.assertEqual(page['runId'], original_run)
            self.assertEqual(page['total'], 2); self.assertEqual(len(page['items']), 1)
            per_task.extend(page['items'])
        self.assertEqual({row['parentTaskId'] for row in per_task}, {row['taskId'] for row in analysis['items']})
        self.assertEqual({row['format']: row['assessment'] for row in per_task},
                         {'csv': 'checked_subset_improved', 'pdf': verdicts['pdf_preserve_images']})
        second = self.retry(identity, first)
        self.assertNotEqual(second['retry']['runId'], original_run)
        self.assertEqual(second['historyTotal'], 2)
        history = self.call('compare/retry-history?' + urlencode({'job': identity, 'offset': 0, 'limit': 1}))
        self.assertEqual(history['total'], 2); self.assertEqual(len(history['items']), 1)
        self._stop_host(); self._start_host()
        restored = self.call('compare?' + urlencode({'job': identity}))
        self.assertEqual(restored['historyTotal'], 2)
        original = next(run for run in restored['retryHistory'] if run['runId'] == original_run)
        self.assertEqual(set(original['jobIds']), old_jobs)
        self.assertEqual(len(original['children']), 2)
        self.assertEqual({child['recipe']: child['verification']['assessment'] for child in original['children']}, verdicts)
        released = next(iter(old_jobs)); self.guard.release(released, discard=True)
        refreshed = self.call('compare/retry-history?' + urlencode({'job': identity}))
        old = next(run for run in refreshed['items'] if run['runId'] == original_run)
        self.assertFalse(next(child for child in old['children'] if child['jobId'] == released)['available'])
        self.assertTrue(self.guard.status(identity)['result']['outputs'])

    def test_later_recipe_failure_keeps_the_earlier_candidate_and_honest_per_task_verdicts(self):
        state = self.converted(workbook=True, sheets=1, formats=['csv', 'pdf']); identity = state['jobId']
        analysis = self.analyze(identity)
        self.adapter.fail_pdf_retry = True
        final = self.retry(identity, analysis)
        csv_child = next(child for child in final['retry']['children'] if child['recipe'] == 'csv_preserve_values')
        self.assertTrue(csv_child['available'])
        self.assertEqual(csv_child['verification']['assessment'], 'checked_subset_improved')
        good = self.guard.status(csv_child['jobId'])
        self.assertEqual(good['succeeded'], 1)
        self.assertTrue(good['result']['outputs'])
        pdf_child = next(child for child in final['retry']['children'] if child['recipe'] == 'pdf_preserve_images')
        self.assertNotEqual(pdf_child['verification']['assessment'], 'checked_subset_improved')
        self.assertEqual(self.guard.status(pdf_child['jobId'])['failed'], 1)
        self.assertNotEqual(final['retry']['verification']['assessment'], 'checked_subset_improved')
        self.assertEqual(self.guard.status(identity)['result']['outputs'], state['result']['outputs'])
        destination = self.root / 'earlier-good-candidate'; destination.mkdir()
        self.assertTrue(self.guard.save_folder(good['result']['ticket'], destination, verify=True)['savedAll'])

    def test_partial_cancel_resume_keeps_report_identity_and_restart_rejects_old_selection(self):
        state = self.converted(24); identity = state['jobId']
        self.call('compare/start', {'jobId': identity})
        progress = self.wait(identity, lambda packet: 0 < packet['processed'] < packet['total'])
        self.call('compare/control', {'jobId': identity, 'action': 'cancel'})
        cancelled = self.wait(identity)
        self.assertEqual(cancelled['state'], 'cancelled'); self.assertLess(cancelled['processed'], cancelled['total'])
        partial_ids = {row['taskId'] for row in cancelled['items']}
        frozen = self.call('compare/start', {'jobId': identity})
        self.assertEqual(frozen['state'], 'cancelled'); self.assertEqual(frozen['analysisId'], progress['analysisId'])
        self.call('compare/control', {'jobId': identity, 'action': 'resume'})
        resumed = self.wait(identity)
        self.assertEqual(resumed['state'], 'completed'); self.assertEqual(resumed['analysisId'], progress['analysisId'])
        self.assertEqual(resumed['processed'], 24); self.assertEqual(sum(resumed['summary'].values()), 24)
        self.assertEqual(len({row['taskId'] for row in resumed['items']}), 24)
        self.assertTrue(partial_ids.issubset({row['taskId'] for row in resumed['items']}))
        restarted = self.call('compare/start', {'jobId': identity, 'restart': True})
        self.assertNotEqual(restarted['analysisId'], resumed['analysisId'])
        self.wait(identity)
        self.error(409, 'compare/retry', {'jobId': identity, 'analysisId': resumed['analysisId'],
                                        'taskIds': [identity + ':0']})

    def test_save_during_real_analysis_preserves_saved_bytes_and_pending_reports(self):
        state = self.converted(20); identity = state['jobId']
        self.call('compare/start', {'jobId': identity})
        active = self.wait(identity, lambda packet: 0 < packet['processed'] < packet['total'])
        self.assertGreater(self.guard._jobs[identity].io, 0)
        output_hashes = {artifact.path.name: hashlib.sha256(artifact.path.read_bytes()).hexdigest()
                         for artifact in self.guard._jobs[identity].artifacts}
        destination = self.root / 'saved'; destination.mkdir()
        saved = self.guard.save_folder(state['result']['ticket'], destination, verify=True)
        self.assertTrue(saved['savedAll'])
        final = self.wait(identity)
        self.assertEqual(final['state'], 'completed'); self.assertEqual(final['analysisId'], active['analysisId'])
        actual = saved_payload_hashes(destination, identity, 20)
        self.assertEqual(actual, output_hashes)
        self.assertEqual(self.guard._jobs[identity].io, 0)

    def test_discard_during_real_analysis_cannot_erase_the_leased_files_early(self):
        state = self.converted(20); identity = state['jobId']
        self.call('compare/start', {'jobId': identity})
        self.wait(identity, lambda packet: 0 < packet['processed'] < packet['total'])
        job = self.guard._jobs[identity]
        root = job.root
        originals = [item.source for item in job.files]
        self.call('release', {'jobId': identity, 'discard': True})
        with self.guard._lock:
            if job.io:
                self.assertTrue(all(path.is_file() for path in originals), 'A live analysis lease protects every original')
        try:
            self.call('compare/control', {'jobId': identity, 'action': 'cancel'})
        except HTTPError as error:
            # Automatic release observation may have already stopped/reaped
            # the worker and removed the whole job before this request.
            self.assertEqual(error.code, 404)
            error.close()
        until = time.monotonic() + 12
        while root.exists() and time.monotonic() < until:
            time.sleep(.03)
        self.assertFalse(root.exists(), 'Cancelled analysis must release the deferred discard lease')
        self.error(404, 'compare?' + urlencode({'job': identity}))

    def test_two_jobs_do_not_cross_results_and_old_credentials_or_malformed_commands_are_rejected(self):
        first = self.converted(12); second = self.converted(workbook=True)
        a, b = first['jobId'], second['jobId']
        self.call('compare/start', {'jobId': a})
        self.error(409, 'compare/start', {'jobId': b})
        self.call('compare/control', {'jobId': a, 'action': 'cancel'}); self.wait(a)
        b_report = self.analyze(b)
        self.assertTrue(all(row['taskId'].startswith(b + ':') for row in b_report['items']))
        self.error(404, 'compare/item?' + urlencode({'job': b, 'task': a + ':0'}))
        for credential in (False, 'stale-session-token'):
            self.error(403, 'compare?' + urlencode({'job': b}), token=credential)
        for raw in (b'{', b'[]', b'{"jobId":null}', b'{"jobId":"'+b.encode()+b'","restart":"yes"}'):
            self.error(400, 'compare/start', raw, raw=True)
        self.error(404, 'compare?' + urlencode({'job': 'f' * 32}))
        self.assertEqual(self.guard.status(a)['succeeded'], 12)
        self.assertEqual(self.guard.status(b)['succeeded'], 1)

    def test_missing_report_database_after_recovery_never_becomes_an_empty_success(self):
        state = self.converted(); identity = state['jobId']; self.analyze(identity)
        database = self.server.api.comparison._get(identity).db
        self._stop_host(); Path(database).unlink(); self._start_host()
        try:
            report = self.call('compare?' + urlencode({'job': identity}))
        except HTTPError as error:
            self.assertEqual(error.code, 404)
        else:
            self.assertFalse(report.get('reportAvailable', False))
            self.assertNotEqual(report['state'], 'completed')
            self.assertFalse(report.get('qualityCompared', False))
        self.assertTrue(self.guard.status(identity)['result']['outputs'])

    def test_concurrent_retry_cancel_and_other_job_release_keep_task_ownership(self):
        original = self.converted(workbook=True, sheets=1); parent = original['jobId']
        analysis = self.analyze(parent)
        task = analysis['items'][0]['taskId']
        running = self.converted(18); active = running['jobId']
        discarded = self.converted()['jobId']
        self.call('compare/start', {'jobId': active})
        self.wait(active, lambda packet: packet['state'] == 'running')
        operations = (
            ('compare/retry', {'jobId': parent, 'analysisId': analysis['analysisId'], 'taskIds': [task]}),
            ('compare/control', {'jobId': active, 'action': 'cancel'}),
            ('release', {'jobId': discarded, 'discard': True}),
        )
        with ThreadPoolExecutor(max_workers=3) as executor:
            futures = [executor.submit(self.call, endpoint, body) for endpoint, body in operations]
            for future in futures:
                future.result(timeout=20)
        self.assertEqual(self.wait(active)['state'], 'cancelled')
        final = self.wait(parent, lambda packet: packet.get('retry', {}).get('state') in ('completed', 'failed', 'cancelled'))
        self.assertEqual(final['retry']['state'], 'completed', final['retry'])
        child = self.guard.report(final['retry']['jobId'])
        self.assertEqual(child['comparisonParentJobId'], parent)
        self.assertEqual(child['comparisonParentTasks'], [task])
        self.assertEqual(self.guard.status(parent)['result']['outputs'], original['result']['outputs'])
        self.assertEqual(self.guard.status(active)['succeeded'], 18)
        self.error(404, 'jobs?' + urlencode({'job': discarded}))

    def test_report_storage_open_failure_is_explicit_and_does_not_destroy_conversion_outputs(self):
        state = self.converted(); identity = state['jobId']; before = self.analyze(identity)
        # Only the report database-open boundary is injected. No analyzer,
        # process, result or native-save behavior is replaced by this fixture.
        with patch('app_backend.converter_guard.comparison_service.sqlite3.connect',
                   side_effect=sqlite3.OperationalError('Controlled report storage access failure')):
            self.error(404, 'compare?' + urlencode({'job': identity}))
        restored = self.call('compare?' + urlencode({'job': identity}))
        self.assertEqual(restored['analysisId'], before['analysisId'])
        self.assertEqual(restored['items'], before['items'])
        self.assertEqual(self.guard.status(identity)['result']['outputs'], state['result']['outputs'])
        self.assertEqual(len(self.adapter.calls), 1)


if __name__ == '__main__':
    unittest.main(verbosity=2)
