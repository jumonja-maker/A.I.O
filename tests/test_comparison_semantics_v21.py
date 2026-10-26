from pathlib import Path
import re
import tempfile
import unittest
from unittest.mock import patch
import zipfile
import xml.etree.ElementTree as ET

from app_backend.converter_guard.worksheet_geometry import (
    GeometryUnavailable, worksheet_geometry, tree_events)
from openpyxl import Workbook
from pypdf import PdfWriter
from app_backend.converter_guard.comparison_analysis import analyze_pair, collect_source_summary
from app_backend.converter_guard.quality import inspect_document

OFFICE = 'urn:oasis:names:tc:opendocument:xmlns:office:1.0'
DRAW = 'urn:oasis:names:tc:opendocument:xmlns:drawing:1.0'
TEXT = 'urn:oasis:names:tc:opendocument:xmlns:text:1.0'


def workbook(path, rows, dimension=None, cell_coordinate=None):
    book = Workbook()
    for row in rows: book.active.append(row)
    book.save(path)
    if dimension is not None or cell_coordinate is not None:
        with zipfile.ZipFile(path) as archive:
            files = {name: archive.read(name) for name in archive.namelist()}
        xml = files['xl/worksheets/sheet1.xml']
        if dimension is not None:
            xml = re.sub(rb'<dimension ref="[^"]+"', ('<dimension ref="'+dimension+'"').encode(), xml)
        if cell_coordinate is not None:
            xml = xml.replace(b'<c r="B2"', ('<c r="'+cell_coordinate+'"').encode())
        files['xl/worksheets/sheet1.xml'] = xml
        with zipfile.ZipFile(path, 'w') as archive:
            for name, value in files.items(): archive.writestr(name, value)


def strict_workbook(path, *, whole_package=False):
    """Controlled Strict XML namespace fixture, preserving real cell contents."""
    workbook(path, [['one', 'TWO'], ['third', 'FOUR']])
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    for name, value in parts.items():
        if name == 'xl/worksheets/sheet1.xml' or whole_package and name.endswith(('.xml', '.rels')):
            parts[name] = value.replace(
                b'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
                b'http://purl.oclc.org/ooxml/spreadsheetml/main')
            if whole_package:
                parts[name] = parts[name].replace(
                    b'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
                    b'http://purl.oclc.org/ooxml/officeDocument/relationships')
    with zipfile.ZipFile(path, 'w') as archive:
        for name, value in parts.items(): archive.writestr(name, value)


def flat(path, kind='presentation', body='presentation', text='same', pages=1):
    content = ''.join('<draw:page><text:p>'+text+'</text:p></draw:page>' for _ in range(pages))
    path.write_text('<office:document xmlns:office="'+OFFICE+'" xmlns:draw="'+DRAW+'" xmlns:text="'+TEXT+
        '" office:mimetype="application/vnd.oasis.opendocument.'+kind+'"><office:body><office:'+body+'>'+content+
        '</office:'+body+'></office:body></office:document>', encoding='utf-8')


class ComparisonSemanticsV21(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_strict_worksheet_source_cannot_certify_missing_values(self):
        source, result = self.root/'strict.xlsx', self.root/'empty.xlsx'
        strict_workbook(source); workbook(result, [])
        report = analyze_pair(source, [result], 'xlsx')
        self.assertEqual(report['status'], 'not_checked')
        self.assertFalse(report['before']['complete'])
        self.assertEqual(report['before']['reason'], 'xlsx_strict_values_not_inspected')
        self.assertNotIn('cell_values_single_grid', report['checkedFeatures'])
        self.assertEqual(report['recipeCandidates'], [])

    def test_strict_worksheet_result_cannot_hide_added_values(self):
        source, result = self.root/'empty.xlsx', self.root/'strict.xlsx'
        workbook(source, []); strict_workbook(result)
        report = analyze_pair(source, [result], 'xlsx')
        self.assertEqual(report['status'], 'not_checked')
        self.assertFalse(report['after']['complete'])
        self.assertEqual(report['after']['reason'], 'xlsx_strict_values_not_inspected')
        self.assertNotIn('cell_values_single_grid', report['checkedFeatures'])

    def test_quality_cannot_hash_strict_cells_as_an_empty_sheet(self):
        source = self.root/'strict.xlsx'; strict_workbook(source)
        report = inspect_document(source)
        self.assertEqual(report['status'], 'not_checked')
        self.assertFalse(report['complete'])
        self.assertEqual(report['reason'], 'xlsx_strict_values_not_inspected')
        self.assertFalse(report.get('cellValues', {}).get('valuesDigest'))

    def test_full_strict_workbook_unrecognized_by_reader_is_incomplete(self):
        source = self.root/'strict.xlsx'; strict_workbook(source, whole_package=True)
        original = source.read_bytes()
        summary, quality = collect_source_summary(source), inspect_document(source)
        self.assertFalse(summary['complete'])
        self.assertEqual(summary['reason'], 'xlsx_strict_values_not_inspected')
        self.assertFalse(quality['complete'])
        self.assertEqual(quality['status'], 'not_checked')
        self.assertEqual(quality['reason'], 'xlsx_worksheets_not_inspected')
        self.assertFalse(quality.get('cellValues', {}).get('valuesDigest'))
        self.assertEqual(source.read_bytes(), original)

    def test_understated_xlsx_dimension_cannot_hide_three_lost_values(self):
        a,b=self.root/'a.xlsx',self.root/'b.xlsx'
        workbook(a,[['one','TWO'],['third','FOUR']],dimension='A1');workbook(b,[['one']])
        report=analyze_pair(a,[b],'xlsx')
        self.assertEqual(report['status'],'confirmed_loss')
        self.assertNotIn('cell_values_single_grid',report['checkedFeatures'])
        loss=next(f for f in report['findings'] if f['code']=='CELL_VALUES_CHANGED')
        self.assertEqual(loss['totalAffected'],3)

    def test_quality_digest_reads_real_cells_outside_declared_dimension(self):
        a,b=self.root/'a.xlsx',self.root/'b.xlsx'
        workbook(a,[['one','TWO'],['third','FOUR']],dimension='A1');workbook(b,[['one']])
        source,result=inspect_document(a),inspect_document(b)
        self.assertTrue(source['complete']);self.assertTrue(result['complete'])
        self.assertNotEqual(source['cellValues']['valuesDigest'],result['cellValues']['valuesDigest'])
        self.assertEqual((source['cellValues']['rows'],source['cellValues']['columns']),(2,2))

    def test_overstated_metadata_does_not_force_large_empty_allocation(self):
        a,b=self.root/'a.xlsx',self.root/'b.xlsx'
        workbook(a,[['one','TWO']],dimension='A1:XFD1048576');workbook(b,[['one','TWO']])
        report=analyze_pair(a,[b],'xlsx')
        self.assertIn('cell_values_single_grid',report['checkedFeatures'])
        self.assertTrue(inspect_document(a)['complete'])

    def test_actual_sparse_coordinate_is_bounded_before_openpyxl_iteration(self):
        a=self.root/'huge.xlsx'
        workbook(a,[['one','TWO'],['third','FOUR']],dimension='A1',cell_coordinate='XFD1048576')
        report=collect_source_summary(a)
        self.assertFalse(report['complete'])
        self.assertIn(report['reason'],('xlsx_cell_geometry_budget','xlsx_cell_coordinate_inconsistent'))
        quality=inspect_document(a)
        self.assertFalse(quality['complete'])

    def test_flat_presentation_model_and_slide_counts_are_preserved(self):
        a=self.root/'a.fodp';flat(a,kind='presentation-flat-xml',pages=2)
        summary=collect_source_summary(a)
        self.assertTrue(summary['complete']);self.assertEqual(summary['documentClass'],'presentation')
        self.assertEqual(summary['counts']['slides'],2)

    def test_flat_wrong_body_does_not_certify_a_model(self):
        a=self.root/'a.fodp';flat(a,kind='presentation',body='text')
        summary=collect_source_summary(a)
        self.assertFalse(summary['complete'])
        self.assertEqual(summary['reason'],'odf_mimetype_body_mismatch')

    def test_pdf_page_count_change_is_explicit_even_if_no_text_extracted(self):
        a,b=self.root/'a.pdf',self.root/'b.pdf'
        for path,pages in ((a,2),(b,1)):
            writer=PdfWriter()
            for _ in range(pages):writer.add_blank_page(width=200,height=300)
            with path.open('wb') as stream:writer.write(stream)
        report=analyze_pair(a,[b],'pdf')
        self.assertEqual(report['status'],'uncertain')
        difference=next(f for f in report['findings'] if f['code']=='STRUCTURE_PAGES_COUNT_CHANGED')
        self.assertEqual((difference['before'],difference['after']),(2,1))
        self.assertEqual(difference['category'],'uncertain')

    def test_understated_result_dimension_cannot_hide_injected_cells(self):
        source, result = self.root/'source.xlsx', self.root/'result.xlsx'
        workbook(source, [['one']])
        workbook(result, [['one', 'TWO'], ['third', 'FOUR']], dimension='A1')
        report = analyze_pair(source, [result], 'xlsx')
        self.assertEqual(report['status'], 'uncertain')
        self.assertNotIn('cell_values_single_grid', report['checkedFeatures'])
        addition = next(f for f in report['findings'] if f['code'] == 'CELL_VALUES_ADDED')
        self.assertEqual(addition['totalAffected'], 3)

    def test_geometry_rejects_ambiguous_cell_and_row_references(self):
        prefix = '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
        suffix = '</sheetData></worksheet>'
        for rows in (
            '<row r="0"><c r="A1"/></row>',
            '<row r="-1"><c r="A1"/></row>',
            '<row r="1"><c r="A0"/></row>',
            '<row r="1"><c r=""/></row>',
            '<row r="1"><c r="A1"/><c r="A1"/></row>',
            '<row r="1"><c r="B1"/><c r="A1"/></row>',
            '<row r="1"><c r="A2"/></row>',
            '<row r="1"><c r="A1"/></row><row r="1"><c r="A1"/></row>',
        ):
            with self.subTest(rows=rows), self.assertRaises(GeometryUnavailable):
                worksheet_geometry(tree_events(ET.fromstring(prefix+rows+suffix)), check=lambda: None)
        rows = '<row><c/><c/></row><row><c/><c/></row>'
        self.assertEqual(worksheet_geometry(tree_events(ET.fromstring(prefix+rows+suffix)),
                                           check=lambda: None), (2, 2))

    def test_real_sparse_extent_rejected_before_readonly_iteration(self):
        from openpyxl.worksheet._read_only import ReadOnlyWorksheet
        source = self.root/'sparse.xlsx'
        workbook(source, [['one'], ['third','FOUR']], dimension='A1')
        with zipfile.ZipFile(source) as archive:
            files = {name: archive.read(name) for name in archive.namelist()}
        files['xl/worksheets/sheet1.xml'] = files['xl/worksheets/sheet1.xml'].replace(
            b'<row r="2">', b'<row r="50000">').replace(b'r="A2"', b'r="A50000"').replace(
            b'r="B2"', b'r="Z50000"')
        with zipfile.ZipFile(source, 'w') as archive:
            for name, value in files.items(): archive.writestr(name, value)
        with patch.object(ReadOnlyWorksheet, 'iter_rows', side_effect=AssertionError('iteration must not start')) as reader:
            self.assertEqual(collect_source_summary(source)['reason'], 'xlsx_cell_geometry_budget')
            self.assertEqual(inspect_document(source)['reason'], 'xlsx_cell_geometry_budget')
            reader.assert_not_called()

    def test_sheet_streams_close_even_while_failure_traceback_is_retained(self):
        from openpyxl.worksheet._read_only import ReadOnlyWorksheet
        from app_backend.converter_guard import comparison_analysis as analysis, quality
        source = self.root/'stream.xlsx'; workbook(source, [['one','TWO']])
        streams = []
        original = ReadOnlyWorksheet._get_source
        def capture(sheet):
            stream = original(sheet); streams.append(stream); return stream
        for module, failing in ((analysis, '_value'), (quality, 'json.dumps')):
            streams.clear()
            with patch.object(ReadOnlyWorksheet, '_get_source', capture):
                if module is analysis:
                    manager = patch.object(analysis, failing, side_effect=RuntimeError('injected cell inspection failure'))
                else:
                    manager = patch.object(quality.json, 'dumps', side_effect=RuntimeError('injected cell digest failure'))
                with manager:
                    try:
                        if module is analysis:
                            with zipfile.ZipFile(source) as archive:
                                analysis._xlsx_summary(source, archive, analysis._Budget(None))
                        else:
                            quality._xlsx_values(source, quality._Budget(None, 5, 1024*1024))
                    except RuntimeError as error:
                        self.assertIn('injected cell', str(error))
                        self.assertTrue(streams)
                        self.assertTrue(all(stream.closed for stream in streams))
                    else:
                        self.fail('injected failure did not reach inspection')

    def test_geometry_fix_keeps_formula_cache_uncertainty_and_date_types(self):
        from datetime import datetime
        source, result = self.root/'formula.xlsx', self.root/'result.xlsx'
        book = Workbook(); sheet = book.active
        sheet.append(['label','value']); sheet.append(['formula','=1+2'])
        sheet.append(['date',datetime(2026,10,26)]); book.save(source)
        evidence = collect_source_summary(source)
        self.assertTrue(evidence['complete'])
        self.assertEqual(evidence['_grids'][0]['rows'][1][1]['formula'], '=1+2')
        self.assertEqual(evidence['_grids'][0]['rows'][2][1]['type'], 'date')
        workbook(result,[['label','value'],['formula',3],['date','2026-10-26']])
        report=analyze_pair(source,[result],'xlsx')
        self.assertIn('cached_formula_values', report['unknown'])
        self.assertNotIn('cell_values_single_grid', report['checkedFeatures'])

    def test_odf_presentation_template_is_same_family_with_real_slides(self):
        source = self.root/'source.otp'
        xml = '<office:document-content xmlns:office="'+OFFICE+'" xmlns:draw="'+DRAW+'" xmlns:text="'+TEXT+'"><office:body><office:presentation><draw:page><text:p>slide</text:p></draw:page></office:presentation></office:body></office:document-content>'
        with zipfile.ZipFile(source, 'w') as archive:
            archive.writestr('mimetype', 'application/vnd.oasis.opendocument.presentation-template')
            archive.writestr('content.xml', xml)
        evidence=collect_source_summary(source)
        self.assertTrue(evidence['complete'])
        self.assertEqual(evidence['documentClass'],'presentation')
        self.assertEqual(evidence['counts']['slides'],1)


    def test_shared_strings_outside_dimension_are_checked_without_type_guessing(self):
        source, result = self.root/'source.xlsx', self.root/'result.xlsx'
        workbook(source, [['one','SECRET']], dimension='A1')
        workbook(result, [['one']])
        with zipfile.ZipFile(source) as archive:
            files = {name: archive.read(name) for name in archive.namelist()}
        namespace = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
        root = ET.fromstring(files['xl/worksheets/sheet1.xml'])
        for index, cell in enumerate(root.iter('{'+namespace+'}c')):
            for child in list(cell): cell.remove(child)
            cell.set('t','s'); ET.SubElement(cell,'{'+namespace+'}v').text=str(index)
        files['xl/worksheets/sheet1.xml'] = ET.tostring(root)
        files['xl/sharedStrings.xml'] = ('<sst xmlns="'+namespace+'" count="2" uniqueCount="2"><si><t>one</t></si><si><t>SECRET</t></si></sst>').encode()
        files['[Content_Types].xml'] = files['[Content_Types].xml'].replace(b'</Types>',
            b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
        with zipfile.ZipFile(source,'w') as archive:
            for name,value in files.items(): archive.writestr(name,value)
        evidence=collect_source_summary(source)
        self.assertTrue(evidence['complete'])
        self.assertEqual(evidence['_grids'][0]['rows'][0][1]['value'],'SECRET')
        self.assertEqual(analyze_pair(source,[result],'xlsx')['status'],'confirmed_loss')
        self.assertNotEqual(inspect_document(source)['cellValues']['valuesDigest'],
                            inspect_document(result)['cellValues']['valuesDigest'])

    def test_template_and_macro_enabled_extensions_keep_document_model_contract(self):
        source = self.root/'source.xlsx'; workbook(source,[['one','two']])
        for extension in ('docm','dotx','dotm','potx','ppsm','ott','otp'):
            with self.subTest(extension=extension):
                result=self.root/('result.'+extension)
                result.write_bytes(source.read_bytes())
                report=analyze_pair(source,[result],extension)
                self.assertEqual(report['status'],'confirmed_loss')
                self.assertIn('OUTPUT_FORMAT_MISMATCH',[f['code'] for f in report['findings']])
                self.assertNotIn('cell_values_single_grid',report['checkedFeatures'])


    def test_orphan_word_header_cannot_prove_text_loss_from_empty_body(self):
        source,result=self.root/'orphan.docx',self.root/'empty.txt'
        namespace='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
        with zipfile.ZipFile(source,'w') as archive:
            archive.writestr('word/document.xml','<w:document xmlns:w="'+namespace+'"><w:body><w:p/></w:body></w:document>')
            archive.writestr('word/header1.xml','<w:hdr xmlns:w="'+namespace+'"><w:p><w:r><w:t>Detached old header</w:t></w:r></w:p></w:hdr>')
        result.write_text('')
        report=analyze_pair(source,[result],'txt')
        self.assertEqual(report['status'],'uncertain')
        self.assertNotIn('TEXT_OUTPUT_EMPTY',[f['code'] for f in report['findings']])
        self.assertIn('WORD_AUXILIARY_TEXT_NOT_CONFIRMED',[f['code'] for f in report['findings']])

    def test_word_body_loss_remains_confirmed_even_with_auxiliary_parts(self):
        source,result=self.root/'body.docx',self.root/'empty.txt'
        namespace='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
        with zipfile.ZipFile(source,'w') as archive:
            archive.writestr('word/document.xml','<w:document xmlns:w="'+namespace+'"><w:body><w:p><w:r><w:t>Body content</w:t></w:r></w:p></w:body></w:document>')
            archive.writestr('word/header1.xml','<w:hdr xmlns:w="'+namespace+'"><w:p><w:r><w:t>Detached old header</w:t></w:r></w:p></w:hdr>')
        result.write_text('')
        report=analyze_pair(source,[result],'txt')
        self.assertEqual(report['status'],'confirmed_loss')
        self.assertIn('TEXT_OUTPUT_EMPTY',[f['code'] for f in report['findings']])


if __name__=='__main__':unittest.main()
