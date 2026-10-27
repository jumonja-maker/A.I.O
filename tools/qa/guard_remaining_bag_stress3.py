"""Separate targeted stress3 corpus and acceptance, never the user's FULLTEST.

Run only in a new directory. Child generators use synthetic controls and the
installed native stack. This driver requires the reviewed feature PRs together.
"""
import argparse
import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import threading
import zipfile
from xml.etree import ElementTree as ET

ROOT=Path(__file__).resolve().parents[2]
sys.path[:0]=[str(ROOT),str(ROOT/'tests'),str(ROOT/'tools/qa')]

# Expected decisions are fixed before any child starts. A successful exit code
# alone is insufficient: exact result counts and each expected verdict matter.
STAGES=(
    ('dif-regression','guard_b28_dif_native_acceptance.py',{'cases':48}),
    ('text-id-regression','guard_b29_csv_native_acceptance.py',{'cases':108}),
    ('all-sheets','guard_b30_native_acceptance.py',{'cases':9}),
    ('dbf-jpeg','guard_native_routes_b13_b31.py',{'dbf':8,'jpeg':120,'negative':6}),
    ('tiff','guard_b32_tiff_native_acceptance.py',{'cases':24,'negative':6}),
    ('html','guard_b33_html_native_acceptance.py',{'positive':18,'excluded':30,'negative':3}),
    ('session','guard_session_native_acceptance.py',{'checks':3}),
    ('viewer-100','guard_before_after_b19_b25_acceptance.py',{'checks':5}),
    ('viewer-150','guard_before_after_b19_b25_acceptance.py',{'checks':5}),
)

# These are parser/identity decisions, not claims about a native conversion.
# Native corrupt IFD/DBF, timeout/cancel and restart are separate child stages.
SAFETY_DECISIONS={
    'truncated-zip':'not_checked','truncated-xml':'not_checked',
    'dtd-entity':'not_checked','missing-html-resource':'not_checked',
    'dynamic-hidden-html':'not_checked','text-mismatch':'uncertain',
    'empty-result':'confirmed_loss','conflicting-mime':'invalid',
    'renamed-writer':'recognized','inspection-budget':'not_checked',
    'changed-source-preview':'unavailable','changed-result-preview':'unavailable',
}


def digest(path):
    with Path(path).open('rb') as stream:return hashlib.file_digest(stream,'sha256').hexdigest()


def write(path,data):path.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')


def validate_stage(name,data,expected):
    if not isinstance(data,dict) or data.get('status')!='PASS':raise ValueError(name+': child evidence is not PASS')
    for key,count in expected.items():
        rows=data.get(key)
        if not isinstance(rows,list) or len(rows)!=count:raise ValueError(name+': expected '+str(count)+' '+key)
        allowed=('SAFE_REJECTION',) if key=='negative' else ('SAFE_UNAVAILABLE',) if key=='excluded' else ('PASS',)
        if any(not isinstance(row,dict) or row.get('status') not in allowed for row in rows):
            raise ValueError(name+': unexpected '+key+' verdict')
    return {'name':name,'status':'PASS','expectedCounts':expected}


def referenced_sources(folder,data):
    """Bind reported synthetic input hashes to actual regular corpus bytes."""
    references=[]
    mapping=data.get('sourceSHA256',{})
    if isinstance(mapping,dict):
        for name,sha in mapping.items():references.append((folder/name,sha,'native_input_control'))
    for key in ('positive','excluded','negative'):
        for row in data.get(key,[]):
            if isinstance(row,dict) and isinstance(row.get('source'),str):
                expected=row.get('sourceSha256') or row.get('sha256')
                if expected:references.append((folder/row['source'],expected,'expected_'+key))
    if data.get('sourceSha256') and (folder/'native-4k.png').is_file():
        references.append((folder/'native-4k.png',data['sourceSha256'],'qt_4k_alpha_control'))
    # Exact existing DIF/CSV controls retain their checked input hash in each
    # case. Index matching files too; never invent an original corpus path.
    hashes={row.get('check',{}).get('sourceSHA256') for row in data.get('cases',[]) if isinstance(row,dict)}
    hashes.update(row.get('independent',{}).get('originalSha256') for row in data.get('cases',[]) if isinstance(row,dict))
    hashes.discard(None)
    if hashes:
        for path in folder.rglob('*'):
            if path.is_file() and not path.is_symlink() and path.stat().st_size<4*1024*1024:
                sha=digest(path)
                if sha in hashes:references.append((path,sha,'exact_existing_regression_input'))
    rows=[];seen=set();root=folder.resolve()
    for path,expected,role in references:
        resolved=path.resolve(strict=True)
        if path.is_symlink() or not resolved.is_relative_to(root) or not path.is_file():raise ValueError('unsafe corpus input path')
        relative=resolved.relative_to(root).as_posix()
        if relative in seen:continue
        seen.add(relative)
        actual=digest(path)
        if actual!=expected:raise ValueError('corpus source changed: '+relative)
        rows.append({'file':relative,'bytes':path.stat().st_size,'sha256':actual,'role':role})
    return rows


def safety_negative_corpus(folder):
    """New bytes through real inert parsers and the isolated preview worker.

    No importer mocks, network requests, external entity expansion or reduced
    production limits. Inspection-budget exercises the public bounded API.
    """
    from app_backend.converter_guard.comparison_analysis import analyze_pair
    from app_backend.converter_guard.source_contract import inspect_source_contract
    folder.mkdir();records=[]
    office='urn:oasis:names:tc:opendocument:xmlns:office:1.0'
    text='urn:oasis:names:tc:opendocument:xmlns:text:1.0'
    writer=(f'<office:document xmlns:office="{office}" xmlns:text="{text}" '
        'office:mimetype="application/vnd.oasis.opendocument.text"><office:body>'
        '<office:text><text:p>ALPHA 001 OMEGA</text:p></office:text></office:body></office:document>')
    result=folder/'result.txt';result.write_text('ALPHA 001 OMEGA',encoding='utf-8')
    sentinel=folder/'entity-private.txt';sentinel.write_text('ENTITY_MUST_NEVER_BE_EXPANDED',encoding='utf-8')
    controls={
        'truncated-zip':('broken.docx',b'PK\x03\x04'+b'x'*64),
        'truncated-xml':('broken.fodt',f'<office:document xmlns:office="{office}"'.encode()),
        'dtd-entity':('entity.fodt',('<!DOCTYPE office:document [<!ENTITY private SYSTEM "'+sentinel.resolve().as_uri()+'">]>'+writer.replace('ALPHA 001 OMEGA','&private;')).encode()),
        'missing-html-resource':('missing.html',b'<html><body><p>ALPHA 001 OMEGA</p><img src="absent.png"></body></html>'),
        'dynamic-hidden-html':('dynamic.html',b'<html><body><p hidden>ALPHA</p><script>throw Error("MUST_NOT_EXECUTE")</script><p>OMEGA</p></body></html>'),
        'text-mismatch':('mismatch.txt',b'OTHER 123 VALUE'),
        'empty-result':('empty-source.txt',b'ALPHA 001 OMEGA'),
        'conflicting-mime':('conflicting.fods',writer.replace('opendocument.text','opendocument.spreadsheet').encode()),
        'renamed-writer':('writer-renamed.xlsx',writer.encode()),
        'inspection-budget':('bounded.fodt',writer.encode()),
    }
    for name,(filename,raw) in controls.items():
        source=folder/filename;source.write_bytes(raw);before=digest(source)
        if name in ('conflicting-mime','renamed-writer','inspection-budget'):
            outcome=inspect_source_contract(source,**({'budget':8} if name=='inspection-budget' else {}))
            actual=outcome['state']
        else:
            output=result
            if name=='empty-result':output=folder/'empty.txt';output.write_bytes(b'')
            outcome=analyze_pair(source,[output],'txt');actual=outcome['status']
            if actual in ('not_checked','uncertain') and outcome['checkedFeatures']:
                raise ValueError(name+': unsupported control acquired quality proof')
        if actual!=SAFETY_DECISIONS[name]:raise ValueError(name+': unexpected decision '+actual)
        if digest(source)!=before:raise ValueError(name+': parser changed original')
        if 'ENTITY_MUST_NEVER_BE_EXPANDED' in json.dumps(outcome):raise ValueError('entity expansion leaked')
        records.append({'name':name,'file':filename,'sha256':before,'status':'PASS',
            'expectedDecision':SAFETY_DECISIONS[name],'actualDecision':actual,
            'scope':'bounded_source_declaration' if 'state' in outcome else 'inert_pair_analysis',
            'evidence':outcome})
    for side in ('source','result'):
        name='changed-'+side+'-preview';case=folder/name;case.mkdir()
        source=case/'data.txt';source.write_bytes(b'NEW BYTES');before=digest(source)
        request=case/'request.json';write(request,{'source':str(source.resolve()),'mode':'preview',
            'previewSide':side,'expectedSha256':hashlib.sha256(b'OLD BYTES').hexdigest()})
        run=subprocess.run([sys.executable,'-m','app_backend.converter_guard.comparison_worker',str(request.resolve())],
            cwd=ROOT,capture_output=True,timeout=30)
        if run.returncode:raise ValueError(name+': preview worker failed '+run.stderr.decode(errors='replace')[-1000:])
        outcome=json.loads((case/'preview.json').read_text(encoding='utf-8'))
        if outcome.get('kind')!='unavailable' or outcome.get('code')!=side.upper()+'_CHANGED' or (case/'preview.png').exists():
            raise ValueError(name+': changed bytes were published as the old report')
        if digest(source)!=before:raise ValueError(name+': preview changed original')
        records.append({'name':name,'file':source.relative_to(folder).as_posix(),'sha256':before,
            'status':'PASS','expectedDecision':'unavailable','actualDecision':outcome['kind'],
            'scope':'preview_identity_boundary','evidence':outcome})
    if set(row['name'] for row in records)!=set(SAFETY_DECISIONS):raise ValueError('incomplete safety corpus')
    write(folder/'manifest.json',{'status':'PASS','controls':records})
    return records


def native_positive_corpus(folder,executable):
    """Native OOXML positives plus a second native PDF import; not macro proof."""
    from app_backend.converter_guard import LibreOfficeAdapter
    from app_backend.converter_guard.source_contract import inspect_source_contract
    from guard_v16_bugbag_acceptance import writer_seed,presentation_seed
    from guard_v08_native_acceptance import fixtures
    from openpyxl import Workbook
    folder.mkdir();seeds=fixtures(folder/'seeds');adapter=LibreOfficeAdapter(executable);cancel=threading.Event();records=[]
    with adapter.batch(folder/'office'):
        for level in range(1,9):
            for kind,target in (('writer','docx'),('calc','xlsx'),('impress','pptx')):
                if kind=='writer':
                    text=writer_seed().replace('AIO v16 explicit four-column control','AIO STRESS3 L'+str(level))
                    extra=''.join('<text:p>Уровень %d строка %d ID %04d</text:p>'%(level,index,index) for index in range(level*16))
                    seed=folder/f'writer-{level}.fodt';seed.write_text(text.replace('</office:text>',extra+'</office:text>'),encoding='utf-8')
                elif kind=='impress':
                    root=ET.fromstring(presentation_seed().replace('AIO_V16_PRESENTATION','AIO_STRESS3_L'+str(level)))
                    container=root.find('{urn:oasis:names:tc:opendocument:xmlns:office:1.0}body')[0];page=container[0]
                    for index in range(1,level):
                        duplicate=ET.fromstring(ET.tostring(page));duplicate.set('{urn:oasis:names:tc:opendocument:xmlns:drawing:1.0}name','Page'+str(index));container.append(duplicate)
                    seed=folder/f'impress-{level}.fodp';seed.write_bytes(ET.tostring(root,encoding='utf-8',xml_declaration=True))
                else:
                    seed=folder/f'calc-{level}.xlsx';book=Workbook();book.remove(book.active)
                    for index in range(1+(level-1)//2):
                        sheet=book.create_sheet('Лист '+str(index+1));sheet.append(['ID','Long ID','Value','Formula A','Formula B','Text'])
                        for row in range(2,302+level*5):sheet.append([f'{row:05d}','12345678901234567890',row/4,f'=C{row}*3',f'=SUM(C{row},D{row})','Строка '+str(row)])
                    book.save(seed);book.close()
                proof=adapter.preflight(seed,folder/f'inspect-seed-{kind}-{level}',cancel=cancel,timeout=90)
                if proof.get('state')!='detected':raise ValueError('positive native seed failed '+kind+str(level))
                adapter.prepare_source(proof)
                paths=adapter.convert(seed,target,folder/f'native-{kind}-{level}',cancel=cancel,timeout=120,on_started=lambda:None)
                if adapter.execution_mode()!='native_uno':raise ValueError('positive OOXML was not created by native UNO')
                output=next(p for p in paths if p.suffix=='.'+target);before=digest(output)
                if kind=='calc':
                    from openpyxl import load_workbook
                    cached=load_workbook(output,read_only=True,data_only=True,keep_links=False)
                    try:
                        for sheet in cached:
                            for index,values in enumerate(sheet.values):
                                if not index:continue
                                if (values[0]!=f'{index+1:05d}' or values[1]!='12345678901234567890'
                                        or values[3]!=values[2]*3 or values[4]!=values[2]*4):
                                    raise ValueError('native positive formula/ID values failed')
                    finally:cached.close()
                contract=inspect_source_contract(output)
                if contract.get('state')!='recognized':raise ValueError('native OOXML package contract failed')
                with zipfile.ZipFile(output) as archive:
                    if archive.testzip() is not None:raise ValueError('native positive ZIP CRC failed')
                    parts=len(archive.namelist())
                second=adapter.preflight(output,folder/f'inspect-positive-{kind}-{level}',cancel=cancel,timeout=90)
                if second.get('state')!='detected' or second.get('basis')!='native_preflight_uno_document':raise ValueError('positive requires actual native model import')
                adapter.prepare_source(second)
                pdf=next(p for p in adapter.convert(output,'pdf',folder/f'pdf-{kind}-{level}',cancel=cancel,timeout=120,on_started=lambda:None) if p.suffix=='.pdf')
                pdf_proof=adapter._preflight_native(pdf,folder/f'pdf-reimport-{kind}-{level}',cancel=cancel,timeout=90)
                if pdf_proof.get('state')!='detected' or pdf_proof.get('basis')!='native_preflight_uno_document':raise ValueError('second native PDF import failed')
                if digest(output)!=before:raise ValueError('native positive source changed')
                records.append({'level':level,'kind':kind,'target':target,'status':'PASS',
                    'file':output.relative_to(folder).as_posix(),'sha256':before,'bytes':output.stat().st_size,
                    'packageParts':parts,'sourceContract':contract,'nativeImportFilter':second.get('selectedImportFilter'),
                    'pdfSha256':digest(pdf),'independentNativePdfModel':pdf_proof.get('documentServices'),
                    'macros':'not_present_not_tested'})
    if len(records)!=24:raise ValueError('incomplete native positive corpus')
    write(folder/'manifest.json',{'schema':1,'status':'PASS','scope':'native_generated_OOXML_positive_controls_only','files':records})
    return records


def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--output',type=Path,required=True)
    p.add_argument('--executable',type=Path);p.add_argument('--dry-run',action='store_true');args=p.parse_args()
    # Never remove/overwrite an earlier QA corpus, even after failure.
    args.output.mkdir(parents=True,exist_ok=False)
    plan={'schema':1,'scope':'remaining_bag_targeted_stress3_not_FULLTEST','releaseBuild':False,'historicalInputs':'not_available',
          'stages':[{'name':name,'script':script,'expected':expected} for name,script,expected in STAGES],
          'parserAndIdentityControls':SAFETY_DECISIONS,
          'coverageExceptions':[{'area':'Base ORP','status':'not_tested','reason':'no genuine Base ReportDefinition fixture'},
              {'area':'installed EXE/manual user acceptance','status':'not_tested','reason':'no build or FULLTEST authorization'},
              {'area':'macro-bearing positives and original 680 corpus','status':'not_tested','reason':'fixtures not supplied'}]}
    write(args.output/'expected-decisions.json',plan)
    if args.dry_run:
        write(args.output/'evidence.json',{'status':'NOT_RUN','reason':'dry_run','expectedDecisions':plan});return 2
    if not args.executable or not args.executable.is_file():raise ValueError('installed executable required')
    report={'schema':1,'status':'RUNNING','scope':plan['scope'],'sourceSHA':os.environ.get('AIO_BUILD_REVISION'),
            'stages':[],'corpus':[],'coverageExceptions':plan['coverageExceptions'],
            'falseRejections':None,'silentSuccesses':None,'expectedUnavailable':None}
    try:
        report['parserAndIdentityControls']=safety_negative_corpus(args.output/'safety-controls')
        positives=native_positive_corpus(args.output/'native-positives',args.executable)
        report['nativePositiveControls']=positives
        for name,script,expected in STAGES:
            path=ROOT/'tools/qa'/script
            if not path.is_file():raise ValueError('required reviewed feature missing: '+script)
            folder=args.output/name
            env=dict(os.environ)
            env['QT_SCALE_FACTOR']='1.5' if name=='viewer-150' else '1'
            result=subprocess.run([sys.executable,str(path),'--output',str(folder),'--executable',str(args.executable)],env=env,cwd=ROOT,timeout=1800)
            evidence=folder/'evidence.json'
            if result.returncode or not evidence.is_file():raise ValueError(name+': child failed, see retained evidence')
            data=json.loads(evidence.read_text(encoding='utf-8'));outcome=validate_stage(name,data,expected)
            outcome.update(scriptSha256=digest(path),evidenceSha256=digest(evidence))
            report['stages'].append(outcome)
            report['corpus'].append({'stage':name,'generator':script,'generatorSha256':digest(path),
                'expectedDecisions':expected,'files':referenced_sources(folder,data)})
        report.update(status='PASS',falseRejections=0,silentSuccesses=0,expectedUnavailable=30,
                      countingScope='only fixed controls in expected-decisions.json; not historical corpus rates')
    except Exception as error:report.update(status='FAIL',reason=str(error));raise
    finally:
        write(args.output/'corpus-manifest.json',{'schema':1,'nativePositiveControls':report.get('nativePositiveControls',[]),
            'parserAndIdentityControls':report.get('parserAndIdentityControls',[]),'stages':report['corpus']})
        write(args.output/'evidence.json',report)
        print('STRESS3_SUMMARY='+json.dumps({k:v for k,v in report.items() if k not in ('corpus','nativePositiveControls')},ensure_ascii=True),flush=True)
    return 0

if __name__=='__main__':raise SystemExit(main())
