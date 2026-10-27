"""Exact stress2 DIF corpus, mandatory native B28 controls. No user FULLTEST."""
import argparse
import hashlib
import json
from pathlib import Path
import sys
import threading
import time

ROOT = Path(__file__).resolve().parents[2]
sys.path[:0] = [str(ROOT), str(ROOT / 'tests')]
from b28_dif_fixtures import original_dif, SHA256
from app_backend.converter_guard import LibreOfficeAdapter
from app_backend.converter_guard.dif_table_contract import check_dif_table


def digest(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def independent_control(path, target, level):
    """Independent high-value checks, not just a rerun of the product contract."""
    rows = 7 + 9 * level
    if target in ('xlsx','xlsm','xltx'):
        from openpyxl import load_workbook
        with path.open('rb') as stream:
            book = load_workbook(stream, read_only=True, data_only=True)
            try:
                assert len(book.worksheets) == 1
                sheet = book.worksheets[0];sheet.reset_dimensions()
                data = list(sheet.iter_rows())
                assert sum(c.value is not None for row in data for c in row) == rows * 6
                assert data[1][0].value == '001' and data[1][0].data_type == 's'
                assert data[-1][0].value == f'{rows-1:03d}'
                assert data[-1][1].value == f'Строка {rows-1} / QA L{level}'
                assert float(data[-1][2].value) == float(f'{(rows-1)*1.43:.2f}')
            finally:
                book.close()
    else:
        from html.parser import HTMLParser
        class Table(HTMLParser):
            def __init__(self):super().__init__();self.inside=False;self.values=[];self.value='';self.nesting=0
            def handle_starttag(self,tag,attrs):
                if tag in ('td','th'):self.inside=True;self.value=''
            def handle_data(self,value):
                if self.inside:self.value+=value
            def handle_endtag(self,tag):
                if tag in ('td','th') and self.inside:self.values.append(self.value.strip());self.inside=False
        table=Table();table.feed(path.read_text(encoding='utf-8-sig'));table.close()
        assert len([v for v in table.values if v]) == rows*6
        assert table.values[6] == '001'
        assert table.values[-6] == f'{rows-1:03d}'
        assert table.values[-5] == f'Строка {rows-1} / QA L{level}'
    return {'method':'openpyxl-actual-cells' if target in ('xlsx','xlsm','xltx') else 'independent-html-text',
            'nonEmptyCells':rows*6,'headTailAndStringIds':True}


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--output',type=Path,required=True)
    p.add_argument('--executable',type=Path,required=True)
    p.add_argument('--baseline',action='store_true')
    args=p.parse_args();args.output.mkdir(parents=True,exist_ok=False)
    report={'schema':1,'status':'RUNNING','scope':'B28 exact-source targeted native QA, not FULLTEST',
            'baseline':args.baseline,'cases':[],'fullQualityRequested':False,'sourceBytesUnchanged':False}
    sources={}
    for level in (range(1,9) if not args.baseline else (1,8)):
        path=args.output/f'{level:02d}_LEVEL_{level}.dif';path.write_bytes(original_dif(level));sources[level]=path
    cancel=threading.Event();adapter=LibreOfficeAdapter(args.executable)
    proofs={}
    def one(level,target,mode):
        source=sources[level];work=args.output/f'{mode}-{level}-{target}'
        tick=time.monotonic()
        try:
            paths=adapter.convert(source,target,work,cancel=cancel,timeout=45,on_started=lambda:None)
            primary=[x for x in paths if x.suffix.lower()=='.'+target]
            assert len(primary)==1
            product=adapter.export_info()
            check=check_dif_table(primary,target,source,cancel=cancel)
            control=independent_control(primary[0],target,level)
            assert check['status']=='checked_subset' and check['matchedCells']==(7+9*level)*6,check
            if not args.baseline:
                emitted=[r for r in product.get('resultChecks',()) if r.get('check')=='dif_table_content']
                assert emitted and emitted[-1]['status']=='checked_subset','mandatory guard missing'
            case={'level':level,'target':target,'mode':adapter.execution_mode(),'status':'PASS',
                  'sha256':digest(primary[0]),'bytes':primary[0].stat().st_size,'check':check,
                  'independent':control,'preparation':product.get('sourcePreparation',{}).get('kind'),
                  'executionReason':product.get('executionReason')}
        except Exception as error:
            case={'level':level,'target':target,'mode':adapter.execution_mode(),'status':'FAIL',
                  'code':getattr(error,'code',type(error).__name__),'reason':str(error)[:1024]}
            if not args.baseline:
                report['cases'].append(case);raise
        case['elapsedMs']=round((time.monotonic()-tick)*1000)
        report['cases'].append(case);print('B28_NATIVE_CASE='+json.dumps(case,ensure_ascii=True),flush=True)
    try:
        with adapter.batch(args.output/'office'):
            for level,source in sources.items():
                proof=adapter.preflight(source,args.output/f'inspect-{level}',cancel=cancel,timeout=45)
                assert proof.get('state')=='detected',(level,proof)
                proofs[level]=proof;adapter.prepare_source(proof)
                for target in ('html','xhtml','xlsm','xltx') + (() if args.baseline else ('xlsx',)):
                    one(level,target,'batch')
        if not args.baseline:
            for level in (1,8):
                adapter.prepare_source(proofs[level])
                for target in ('html','xhtml','xlsm','xltx'):one(level,target,'cli')
        assert all(digest(path)==SHA256[level-1] for level,path in sources.items())
        report.update(status='OBSERVED' if args.baseline else 'PASS',sourceBytesUnchanged=True)
    except Exception as error:
        report.update(status='FAIL',error=type(error).__name__,reason=str(error)[:2048]);raise
    finally:
        report['passed']=sum(c['status']=='PASS' for c in report['cases'])
        report['failed']=sum(c['status']=='FAIL' for c in report['cases'])
        (args.output/'evidence.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
        print('B28_NATIVE_SUMMARY='+json.dumps({k:v for k,v in report.items() if k!='cases'},ensure_ascii=True),flush=True)


if __name__=='__main__':main()
