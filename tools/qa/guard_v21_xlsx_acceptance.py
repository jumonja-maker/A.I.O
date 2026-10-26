"""Offline native gate: inspection limits must not veto real XLSX conversion."""
import argparse
import hashlib
import json
from pathlib import Path
import re
import sys
import threading
import time
import zipfile

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from openpyxl import Workbook, load_workbook
from app_backend.converter_guard import LibreOfficeAdapter
from app_backend.converter_guard.comparison_analysis import collect_source_summary
from app_backend.converter_guard.quality import inspect_document

EXPECTED = [('one', 'TWO'), ('third', 'FOUR')]


def digest(path):
    with path.open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def fixture(path, kind):
    book = Workbook()
    for row in EXPECTED: book.active.append(row)
    book.save(path); book.close()
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    for name, value in parts.items():
        if kind == 'strict' and name.endswith(('.xml', '.rels')):
            parts[name] = value.replace(
                b'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
                b'http://purl.oclc.org/ooxml/spreadsheetml/main').replace(
                b'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
                b'http://purl.oclc.org/ooxml/officeDocument/relationships')
        elif kind == 'stale_dimension' and name == 'xl/worksheets/sheet1.xml':
            parts[name] = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', value)
    with zipfile.ZipFile(path, 'w') as archive:
        for name, value in parts.items(): archive.writestr(name, value)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--executable', type=Path)
    parser.add_argument('--require-mode', choices=['native_uno', 'native_cli'])
    args = parser.parse_args(); args.output.mkdir(parents=True, exist_ok=False)
    adapter = LibreOfficeAdapter(args.executable); cancel = threading.Event(); cases = []
    if not adapter.available: raise RuntimeError('Required LibreOffice unavailable')
    with adapter.batch(args.output/'office'):
        for kind in ('strict', 'stale_dimension'):
            source = args.output/(kind+'.xlsx'); fixture(source, kind); before = digest(source)
            summary, quality = collect_source_summary(source), inspect_document(source)
            if kind == 'strict':
                assert not summary['complete'] and not quality['complete'], (summary, quality)
            else:
                assert summary['complete'] and quality['complete'], (summary, quality)
                assert quality['cellValues']['rows'] == quality['cellValues']['columns'] == 2
            tick = time.monotonic()
            proof = adapter.preflight(source, args.output/('inspect-'+kind), cancel=cancel, timeout=60)
            assert proof.get('state') == 'detected', (kind, proof)
            assert adapter.route_decision('xlsx', proof) == 'registered', (kind, proof)
            adapter.prepare_source(proof); adapter.prepare_retry_recipe(None)
            paths = adapter.convert(source, 'xlsx', args.output/('export-'+kind), cancel=cancel,
                                    timeout=90, on_started=lambda: None)
            mode = adapter.execution_mode()
            if args.require_mode: assert mode == args.require_mode, mode
            result = next(path for path in paths if path.suffix.lower() == '.xlsx')
            book = load_workbook(result, read_only=True, data_only=False, keep_links=False)
            try:
                assert len(book.worksheets) == 1
                values = list(book.worksheets[0].iter_rows(values_only=True))
                assert values == EXPECTED, (kind, values)
            finally: book.close()
            assert digest(source) == before, 'Source bytes changed'
            cases.append({'case': kind, 'status': 'PASS', 'mode': mode, 'verifiedCells': 4,
                'inspectionComplete': summary['complete'], 'qualityComplete': quality['complete'],
                'sourceSHA256': before, 'resultSHA256': digest(result), 'sourceUnchanged': True,
                'elapsedMs': round((time.monotonic()-tick)*1000)})
            print('AIO_V21_XLSX_CASE='+json.dumps(cases[-1]), flush=True)
    result = {'status': 'PASS', 'scope': '2 synthetic native XLSX value/route controls', 'cases': cases}
    (args.output/'evidence.json').write_text(json.dumps(result, indent=2), encoding='utf-8')
    print('AIO_V21_XLSX='+json.dumps(result), flush=True)


if __name__ == '__main__': main()
