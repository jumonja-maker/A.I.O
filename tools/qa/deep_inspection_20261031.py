"""Second independent audit pass. Diagnostic completion is NOT product acceptance.
No product files are patched, no personal files are read, no releases are built.
"""
from __future__ import annotations
import ast
import copy
from contextlib import contextmanager
import datetime
import hashlib
import io
import json
import os
import platform
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import types
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT));sys.path.insert(0,str(ROOT/'tests'))
BASE='97fabe65b64076dc8063e0813865ac70b42bfb14'
FIRST='d11b52ac9354182b6cc6f9daac8f4967f0a9dead'
original=subprocess.check_output(['git','show',FIRST+':tools/qa/deep_inspection_20261031.py'],cwd=ROOT).decode()
old=types.ModuleType('first_audit');old.__file__=str(Path(__file__))
exec(compile(original,'verified-first-audit','exec'),old.__dict__)
emit,probe=old.emit,old.probe

def immutable_product():
    changed=subprocess.check_output(['git','diff','--name-only',BASE,'HEAD'],cwd=ROOT).decode().splitlines()
    permitted={'tools/qa/deep_inspection_20261031.py','.github/workflows/aio-deep-inspection-20261031.yml'}
    assert not set(changed)-permitted, 'Unexpected product changes'
    emit('BASELINE',{'baseline':BASE,'auditCommit':subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT).decode().strip(),
                    'platform':platform.platform(),'python':sys.version,'productChanged':False})

@contextmanager
def guarded_fixture(formats=('pdf','html'), bundle=False):
    from app_backend.converter_guard import ConverterGuard
    from test_converter_guard_v02 import FakeAdapter,wait_finished
    with tempfile.TemporaryDirectory() as temp:
        base=Path(temp);owned=base/'owned';owned.mkdir()
        guard=ConverterGuard(adapter=FakeAdapter(),spool_parent=owned)
        try:
            identity=guard.create([{'id':'f','name':'source.docx','size':4}],list(formats))['jobId']
            guard.upload(identity,'f',io.BytesIO(b'keep'),4)
            guard.start(identity,archive_mode='on_demand');wait_finished(guard,identity)
            job=guard._jobs[identity]
            if bundle:
                guard.rebuild_bundle(identity);wait_finished(guard,identity)
                assert job.bundle is not None
            yield base,guard,job
        finally:guard.close()

def zip_receipt_collision():
    with guarded_fixture(bundle=True) as (base,guard,job):
        destination=base/'chosen.zip';shutil.copyfile(job.bundle.path,destination)
        receipt=destination.with_name(destination.name+'.AIO-save-status.json')
        sentinel=b'PREEXISTING USER NOTES MUST NOT BE REPLACED'
        receipt.write_bytes(sentinel)
        result=guard.verify_saved_download(job.bundle.ticket,destination)
        return {'issue_reproduced':receipt.read_bytes()!=sentinel,'verificationSucceeded':result['readbackVerified'],
                'preexistingAdjacentFileOverwritten':receipt.read_bytes()!=sentinel,
                'originalResultsPresent':all(a.path.exists() for a in job.artifacts)}

def zip_receipt_temporary_link():
    with guarded_fixture(bundle=True) as (base,guard,job):
        destination=base/'chosen.zip';shutil.copyfile(job.bundle.path,destination)
        receipt=destination.with_name(destination.name+'.AIO-save-status.json')
        unrelated=base/'unrelated.txt';sentinel=b'UNRELATED SYNTHETIC FILE';unrelated.write_bytes(sentinel)
        temporary=receipt.with_suffix('.new')
        try:temporary.symlink_to(unrelated)
        except OSError as error:
            return {'issue_reproduced':None,'not_tested_reason':'OS refused symlink fixture: '+type(error).__name__}
        error=None
        try:guard.verify_saved_download(job.bundle.ticket,destination)
        except Exception as e:error=type(e).__name__
        return {'issue_reproduced':unrelated.read_bytes()!=sentinel,'unrelatedSyntheticFileChanged':unrelated.read_bytes()!=sentinel,
                'exception':error,'scope':'only pre-created links inside disposable audit directory'}

def zip_corruption_fail_closed():
    from app_backend.converter_guard import GuardError
    with guarded_fixture(bundle=True) as (base,guard,job):
        destination=base/'damaged.zip';destination.write_bytes(b'not a zip')
        code=None
        try:guard.verify_saved_download(job.bundle.ticket,destination)
        except GuardError as e:code=e.code
        return {'issue_reproduced':not bool(code),'control_pass':code=='SAVE_VERIFY_FAILED' and all(not a.saved and a.path.exists() for a in job.artifacts),
                'code':code,'saveReceiptPresent':job.bundle.save_receipt is not None}

def folder_cancel_boundaries():
    from app_backend.converter_guard import GuardError
    cases=[]
    for phase in ('before_copy','after_first_output','before_readback'):
        with guarded_fixture() as (base,guard,job):
            destination=base/'user';destination.mkdir();cancel=threading.Event()
            if phase=='before_copy':cancel.set()
            def progress(current,total):
                if phase=='after_first_output' and current==1:cancel.set()
            def verifying():
                if phase=='before_readback':cancel.set()
            code=None
            try:guard.save_folder(job.artifacts[0].ticket,destination,cancel=cancel,on_progress=progress,on_verifying=verifying)
            except GuardError as e:code=e.code
            safe=all(a.path.exists() and not a.saved for a in job.artifacts)
            cases.append({'phase':phase,'code':code,'sourceRetainedUnsaved':safe,'savingReleased':not job.saving,'io':job.io})
    return {'issue_reproduced':any(not c['sourceRetainedUnsaved'] or not c['code'] or c['io'] for c in cases),
            'control_pass':all(c['sourceRetainedUnsaved'] and c['code']=='SAVE_CANCELLED' and c['io']==0 for c in cases),'cases':cases}

def folder_collision_control():
    with guarded_fixture() as (base,guard,job):
        destination=base/'user';destination.mkdir();existing=destination/('AIO_Results_'+job.id[:8]);existing.mkdir()
        sentinel=existing/'keep.txt';sentinel.write_bytes(b'keep user data')
        outcome=guard.save_folder(job.artifacts[0].ticket,destination)
        return {'issue_reproduced':sentinel.read_bytes()!=b'keep user data','control_pass':outcome['saved'] and outcome['name'].endswith('_1'),
                'oldFileIntact':sentinel.read_bytes()==b'keep user data','distinctDestination':outcome['name']}

def folder_metadata_failure():
    from app_backend.converter_guard import result_store,GuardError
    with guarded_fixture() as (base,guard,job):
        dest=base/'user';dest.mkdir();original=result_store.atomic_json
        def failing(path,value):
            if Path(path).name=='AIO-save-status.json' and value.get('state')=='completed':
                raise OSError('controlled metadata disk full')
            return original(path,value)
        error=None
        with patch.object(result_store,'atomic_json',side_effect=failing):
            try:guard.save_folder(job.artifacts[0].ticket,dest)
            except GuardError as e:error=e.code
        retained=all(a.path.exists() and not a.saved for a in job.artifacts)
        return {'issue_reproduced':not retained or not error,'control_pass':retained and error=='SAVE_STORAGE_FAILED',
                'code':error,'allOriginalsUnsaved':retained,'leaseCount':job.io}

def excel_feature_matrix():
    from openpyxl import Workbook,load_workbook
    from app_backend.excel import merge_xlsx,MergeError
    cases=[]
    def attributes(ws):
        return {'Bwidth':ws.column_dimensions['B'].width,'Cwidth':ws.column_dimensions['C'].width,
                'Bmin':ws.column_dimensions['B'].min,'Bmax':ws.column_dimensions['B'].max,
                'row2Outline':ws.row_dimensions[2].outlineLevel,
                'leftMargin':ws.page_margins.left,'rightToLeft':ws.sheet_view.rightToLeft,
                'gridLines':ws.sheet_view.showGridLines}
    for case in ('ordinary_values','grouped_column_width','row_outline','page_margin','right_to_left','gridlines_hidden'):
        wb=Workbook();ws=wb.active;ws.append(['A','B','C']);ws.append(['value',123,True])
        if case=='grouped_column_width':
            ws.column_dimensions.group('B','C',outline_level=0,hidden=False);ws.column_dimensions['B'].width=35
        elif case=='row_outline':ws.row_dimensions[2].outlineLevel=2
        elif case=='page_margin':ws.page_margins.left=1.75
        elif case=='right_to_left':ws.sheet_view.rightToLeft=True
        elif case=='gridlines_hidden':ws.sheet_view.showGridLines=False
        data=io.BytesIO();wb.save(data);wb.close()
        reloaded=load_workbook(io.BytesIO(data.getvalue()));before=attributes(reloaded.active);reloaded.close()
        try:result=merge_xlsx([('one.xlsx',data.getvalue()),('two.xlsx',data.getvalue())])
        except MergeError as e:
            cases.append({'case':case,'rejected':True,'reason':str(e)});continue
        merged=load_workbook(io.BytesIO(result));after=attributes(merged.active)
        cases.append({'case':case,'rejected':False,'before':before,'after':after,'rows':merged.active.max_row,
                      'dataRow':list(next(merged.active.iter_rows(min_row=2,max_row=2,values_only=True)))})
        merged.close()
    loss=[]
    for c in cases:
        if c.get('rejected'):continue
        key={'grouped_column_width':'Bmax','row_outline':'row2Outline','page_margin':'leftMargin',
             'right_to_left':'rightToLeft','gridlines_hidden':'gridLines'}.get(c['case'])
        if key and c['before'][key]!=c['after'][key]:loss.append(c['case'])
    return {'issue_reproduced':bool(loss),'silentlyChangedFeatures':loss,'cases':cases,
            'scope':'simple-table merger promises reject unsupported features; visual metadata only, not numeric data loss'}

def excel_value_fidelity():
    from openpyxl import Workbook,load_workbook
    from app_backend.excel import merge_xlsx
    wb=Workbook();ws=wb.active;ws.append(['text','number','boolean','date'])
    for i in range(1,80):ws.append(['000'+str(i),i/7, bool(i%2),datetime.datetime(2026,1,1)+datetime.timedelta(days=i)])
    data=io.BytesIO();wb.save(data);wb.close()
    source=load_workbook(io.BytesIO(data.getvalue()));expected=list(source.active.values)[1:];source.close()
    merged=load_workbook(io.BytesIO(merge_xlsx([('a.xlsx',data.getvalue()),('b.xlsx',data.getvalue())])))
    observed=list(merged.active.values)[1:];merged.close()
    return {'issue_reproduced':observed!=expected+expected,'control_pass':observed==expected+expected,
            'rowsChecked':len(observed),'types':['text-leading-zero','float','boolean','datetime']}

def guard_snapshot_cost():
    from app_backend.converter_guard import ConverterGuard
    from test_converter_guard_v02 import FakeAdapter
    results=[]
    for count in (100,1000,5000):
        with tempfile.TemporaryDirectory() as temp:
            g=ConverterGuard(adapter=FakeAdapter(),spool_parent=temp)
            try:
                row=[{'id':'f'+str(i),'name':'input'+str(i)+'.txt','size':1} for i in range(count)]
                created=g.create(row,['pdf','txt','html']);job=g._jobs[created['jobId']]
                timings=[]
                for _ in range(3):
                    tick=time.perf_counter()
                    for __ in range(100):snapshot=g._snapshot(job,only_files=['f0'],result_delta=True)
                    timings.append((time.perf_counter()-tick)*10)
                tick=time.perf_counter();report=g._report(job);report_seconds=time.perf_counter()-tick
                results.append({'files':count,'tasks':count*3,'selectedFilesReturned':len(snapshot['files']),
                                'deltaMeanMs':timings,'fullReportMs':report_seconds*1000,
                                'fullReportBytes':len(json.dumps(report).encode())})
            finally:g.close()
    return {'issue_reproduced':False,'measurement':results,
            'scope':'CPU snapshot/serialization only; no native export or user-disk benchmark'}

def legacy_unavailable_decline_all():
    import app_backend.converter_transport as mod
    with tempfile.TemporaryDirectory() as temp:
        t=mod.ConverterTransport(recovery_parent=Path(temp)/'legacy',lazy_recovery=True)
        try:
            with patch.object(mod,'recover_workspaces',side_effect=PermissionError('audit denied')):t.inventory()
            outcome=t.decline_all_recovery()
            return {'issue_reproduced':not t.recovery_available and outcome.get('notDeclined')==0,
                    'storageAvailable':t.recovery_available,'bulkOutcome':outcome,
                    'scope':'unreadable storage cannot be enumerated, so all-old-files refusal is not established'}
        finally:t.close()

def atomic_metadata_failed_write():
    from app_backend.converter_guard.result_store import atomic_json
    with tempfile.TemporaryDirectory() as temp:
        root=Path(temp);path=root/'meta.json';path.write_bytes(b'original-valid-metadata')
        caught=None
        with patch('app_backend.converter_guard.result_store.os.fsync',side_effect=OSError('audit no space')):
            try:atomic_json(path,{'next':'value'})
            except OSError as e:caught=type(e).__name__
        return {'issue_reproduced':False,'control_pass':bool(caught) and path.read_bytes()==b'original-valid-metadata',
                'oldMetadataPreserved':path.read_bytes()==b'original-valid-metadata','temporaryLeft':path.with_suffix('.new').exists()}

def legacy_ui_cancel_keeps_salvage():
    import app_backend.legacy_result_journal as mod
    source=(ROOT/'src/legacy-recovery.js').read_text()
    return {'issue_reproduced':False,'static_review_only':True,
            'releaseCalls':source.count("api('release'"),'salvageHasRelease':"api('release'" in source[source.index('async function salvage'):source.index('async function refresh')],
            'salvageChecksCancellation': 'cancel' in __import__('inspect').signature(mod.salvage_unverified_workspace).parameters,
            'note':'static evidence only; standalone salvage ZIP lifecycle requires a bounded owner and cancellation'}

if __name__=='__main__':
    immutable_product()
    for name,op in [('salvage_bypasses_root_limit',old.salvage_slot_limit),('scan_publishes_after_shutdown',old.scan_shutdown_race),
       ('malformed_journal_path',old.missing_journal_path),('manual_rescan_leaves_missing_file_ready',old.stale_inventory),
       ('quota_whole_store_cost',old.quota_scan_cost),('decline_all_keeps_salvage_copy',old.bulk_decline_temporary_salvage),
       ('zip_receipt_overwrites_adjacent_file',zip_receipt_collision),('zip_receipt_temporary_follows_link',zip_receipt_temporary_link),
       ('corrupt_zip_rejected_without_release',zip_corruption_fail_closed),('folder_save_cancel_boundaries',folder_cancel_boundaries),
       ('folder_save_no_overwrite_control',folder_collision_control),('folder_final_metadata_failure',folder_metadata_failure),
       ('excel_feature_matrix',excel_feature_matrix),('excel_value_fidelity_control',excel_value_fidelity),
       ('guard_delta_snapshot_and_report_cost',guard_snapshot_cost),('unavailable_storage_bulk_decline_claim',legacy_unavailable_decline_all),
       ('atomic_json_fsync_failure_control',atomic_metadata_failed_write),('legacy_salvage_lifecycle_static',legacy_ui_cancel_keeps_salvage)]:
        probe(name,op)
    emit('SUMMARY',{'diagnosticCompleted':True,'releaseReady':False,'baseline':BASE,'productChanged':False,
         'probes':len(old.RESULTS),'issueReproductions':sum(r.get('issue_reproduced') is True for r in old.RESULTS),
         'positiveControls':sum(r.get('control_pass') is True for r in old.RESULTS),
         'harnessErrors':sum(r['execution']=='harness_error' for r in old.RESULTS)})
