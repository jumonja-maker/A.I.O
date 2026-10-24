"""Offline targeted acceptance for table routes, EPUB and raster preservation.

Run explicitly on the reviewed source with an installed LibreOffice executable.
Generates owned fixtures; never downloads a corpus or changes original files.
Unsupported SVM/ORP exports are reported from the installed native registry,
separately from actual conversion tests. This is not FULLTEST or GUI approval.
"""
import argparse
import base64
import csv
from decimal import Decimal
import hashlib
from html.parser import HTMLParser
import importlib.util
import io
import json
import os
from pathlib import Path
import platform
import posixpath
import re
import shutil
import struct
import subprocess
import sys
import threading
import time
from urllib.parse import unquote, unquote_to_bytes, urlsplit
import zipfile
import zlib
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from app_backend.converter_guard import ConverterGuard, GuardError, GuardLimits, LibreOfficeAdapter
from app_backend.converter_guard.media_fidelity import inspect_media
from app_backend.converter_guard.routes import registry
from guard_v16_bugbag_acceptance import MARKERS, native_export, writer_seed, check_markers, digest

TABLE_ROWS = [['NAME', 'AMOUNT', 'FLAG'], ['ALPHA', Decimal('123.5'), 'YES'],
              ['БЕТА', Decimal('-2.25'), 'NO'], ['AIO_SENTINEL', Decimal('99'), 'YES']]
HIDDEN_CSS = r'(?:display\s*:\s*none|visibility\s*:\s*hidden|opacity\s*:\s*0(?:\.0+)?\s*(?:!important\s*)?(?:;|}|$))'


def _media_fixture_helpers():
    # These are real decodable fixture bytes and standard-library generators,
    # shared with source regressions; production has no dependency on tests.
    spec = importlib.util.spec_from_file_location('aio_v16_media_fixtures', ROOT / 'tests/test_guard_media_fidelity_v16.py')
    module = importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
    return module


def machine_facts(executable):
    def git(*args):
        try:
            result = subprocess.run(['git', *args], cwd=ROOT, capture_output=True, text=True, timeout=15)
            return result.stdout.strip() if result.returncode == 0 else None
        except (OSError, subprocess.TimeoutExpired):
            return None
    version = subprocess.run([str(executable), '--version'], capture_output=True, text=True, timeout=20)
    working_tree = git('status', '--porcelain')
    row = {'platform': platform.platform(), 'osName': os.name, 'machine': platform.machine(),
        'python': platform.python_version(), 'sourceCommit': git('rev-parse', 'HEAD'),
        'sourceTree': git('rev-parse', 'HEAD^{tree}'), 'engineVersion': version.stdout.strip(),
        'sourceWorkingTreeDirty': bool(working_tree) if working_tree is not None else None,
        'runnerName': os.environ.get('RUNNER_NAME'), 'runnerOS': os.environ.get('RUNNER_OS'),
        'windowsNativeTested': os.name == 'nt', 'windowsGuiTested': False, 'fulltest510': 'NOT_RUN'}
    try:
        from PySide6.QtCore import qVersion
        from PySide6.QtGui import QImageReader, QImageWriter
        row['qt'] = {'version': qVersion(),
            'readFormats': [bytes(v).decode('ascii') for v in QImageReader.supportedImageFormats()],
            'writeFormats': [bytes(v).decode('ascii') for v in QImageWriter.supportedImageFormats()]}
    except ImportError:
        row['qt'] = {'status': 'unavailable'}
    return row


def dbf_descriptor(path):
    data = Path(path).read_bytes()
    if len(data) < 33:
        raise ValueError('Truncated DBF')
    records, header, width = struct.unpack_from('<IHH', data, 4)
    if header < 33 or (header - 33) % 32 or header + records * width > len(data) or data[header - 1] != 13:
        raise ValueError('Invalid DBF record/header sizes')
    fields = []
    for offset in range(32, header - 1, 32):
        name = data[offset:offset + 11].split(b'\0', 1)[0].decode('ascii')
        fields.append((name.split(',')[0], chr(data[offset + 11]), data[offset + 16], data[offset + 17]))
    return data, records, header, width, fields


def dbf_records(path):
    data, records, header, width, fields = dbf_descriptor(path)
    rows = [[f[0] for f in fields]]
    for number in range(records):
        record = data[header + number * width:header + (number + 1) * width]
        if record[:1] == b'*':
            continue
        position, row = 1, []
        for name, kind, length, decimals in fields:
            raw = record[position:position + length]; position += length
            text = raw.decode('utf-8').strip()
            if kind == 'M':
                pointer = int(text or '0')
                if pointer:
                    memo_path = Path(path).with_suffix('.dbt')
                    if not memo_path.is_file(): memo_path = memo_path.with_suffix('.DBT')
                    memo = memo_path.read_bytes(); start = pointer * 512
                    end = memo.find(b'\x1a\x1a', start)
                    if start >= len(memo) or end < start: raise ValueError('Missing or truncated DBF memo block')
                    text = memo[start:end].decode('utf-8')
                else: text = ''
            row.append(Decimal(text) if kind in ('N', 'F') and text else text)
        rows.append(row)
    return rows


def dbf_import_rows(path):
    """Calc exposes native DBF field metadata in its first-row labels."""
    data, records, header, width, fields = dbf_descriptor(path)
    rows = dbf_records(path)
    rows[0] = [name + ',' + kind + (',' + str(length) if kind in ('C', 'N', 'F') else '') +
        (',' + str(decimals) if kind in ('N', 'F') else '') for name, kind, length, decimals in fields]
    return rows


def dif_records(path):
    lines = Path(path).read_text(encoding='utf-8-sig').splitlines()
    position = lines.index('DATA') + 3
    rows, current = [], None
    while position + 1 < len(lines):
        type_value, value = lines[position], lines[position + 1]; position += 2
        kind, numeric = type_value.split(',', 1)
        if kind == '-1':
            if value == 'BOT':
                if current is not None: rows.append(current)
                current = []
            elif value == 'EOD':
                if current is not None: rows.append(current)
                return rows
            else:
                raise ValueError('Unexpected DIF control')
        elif kind == '1':
            if current is None: raise ValueError('DIF cell before BOT')
            current.append(next(csv.reader([value], strict=True))[0])
        elif kind == '0':
            if current is None: raise ValueError('DIF number before BOT')
            current.append(Decimal(numeric))
        else:
            raise ValueError('Unsupported DIF cell type')
    raise ValueError('Missing DIF EOD')


def xlsx_records(path):
    from openpyxl import load_workbook
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        return [[list(row) for row in sheet.iter_rows(values_only=True)] for sheet in workbook.worksheets]
    finally:
        workbook.close()


def normalize_rows(rows):
    result = []
    for row in rows:
        normalized = []
        for value in row:
            normalized.append(Decimal(str(value)) if isinstance(value, (int, float, Decimal)) else str(value or '').strip())
        result.append(normalized)
    return result


def table_check(paths, target, expected):
    if target == 'xlsx':
        actual = [rows for p in paths for rows in xlsx_records(p)]
    else:
        parser = dbf_records if target == 'dbf' else dif_records
        actual = [parser(p) for p in paths]
    normalized_actual = [normalize_rows(rows) for rows in actual]
    normalized_expected = [normalize_rows(rows) for rows in expected]
    # Fan-out filenames and native sheet order are not content identity.
    remaining = list(normalized_actual)
    for sheet in normalized_expected:
        if sheet not in remaining:
            raise AssertionError('Missing or changed table cells in ' + target)
        remaining.remove(sheet)
    if remaining:
        raise AssertionError('Unexpected extra table output in ' + target)
    return {'status': 'pass', 'sheets': len(actual), 'rowsPerSheet': [len(r) for r in actual],
            'columnsPerSheet': [len(r[0]) if r else 0 for r in actual],
            'method': 'independent_DBF_DIF_records_or_openpyxl_cells'}


def dif_charset_check(paths):
    """Read the saved declaration independently and bind it to exact DIF bytes."""
    primary = [p for p in paths if p.suffix.lower() == '.dif']
    sidecars = [p for p in paths if p.name.endswith('.aio-encoding.json')]
    if not primary or len(sidecars) != 1:
        raise AssertionError('Missing or ambiguous DIF UTF-8 declaration')
    metadata = json.loads(sidecars[0].read_text(encoding='utf-8'))
    if (metadata.get('format') != 'DIF' or metadata.get('encoding') != 'UTF-8'
            or metadata.get('bom') is not False or metadata.get('dataBytesChanged') is not False
            or metadata.get('libreOfficeImport') != {'FilterName': 'DIF', 'FilterOptions': '76'}):
        raise AssertionError('Incorrect DIF charset or native import declaration')
    records = metadata.get('files')
    if not isinstance(records, list) or len(records) != len(primary):
        raise AssertionError('DIF declaration does not cover every sheet output')
    wanted = {p.name: {'name': p.name, 'bytes': p.stat().st_size, 'sha256': digest(p)} for p in primary}
    if len(wanted) != len(primary) or any(not isinstance(r, dict) for r in records):
        raise AssertionError('Invalid DIF file declaration')
    actual = {r.get('name'): r for r in records}
    if len(actual) != len(records) or actual != wanted:
        raise AssertionError('DIF declaration bytes or hashes do not match saved outputs')
    for path in primary:
        raw = path.read_bytes()
        if not raw.startswith((b'TABLE\n', b'TABLE\r')):
            raise AssertionError('DIF native TABLE header changed or BOM was inserted')
        raw.decode('utf-8', errors='strict')
    return {'encoding': 'UTF-8', 'encodingDeclaration': 'pass', 'encodingDeclaredFiles': len(primary),
            'encodingNativeFilterOptions': '76'}


def mixed_dbf_check(paths, expected, source_values):
    primary = [p for p in paths if p.suffix.lower() == '.dbf']
    check = table_check(primary, 'dbf', [expected])
    dbf = primary[0]
    if not any(kind == 'M' for name, kind, length, decimals in dbf_descriptor(dbf)[4]):
        raise AssertionError('Mixed column was not represented by a DBF memo')
    mapping = dbf.with_name(dbf.name + '.aio-columns.jsonl')
    records = [json.loads(line) for line in mapping.read_text(encoding='utf-8').splitlines()]
    if records[0].get('memoColumns') != [0] or records[0].get('numericColumnsRemainTyped') is not True:
        raise AssertionError('Mixed DBF column declaration missing')
    cells = {}
    for record in records[1:]:
        if 'row' not in record: continue  # Additional schema/sheet metadata.
        row, column = record['row'], record['column']
        height, width = record.get('rowRepeat', 1), record.get('columnRepeat', 1)
        if any(type(n) is not int for n in (row, column, height, width)) or min(row, column) < 0 or min(height, width) < 1:
            raise AssertionError('Invalid mixed DBF mapping coordinates')
        for wanted in range(1, len(source_values) + 1):
            if row <= wanted < row + height and column <= 0 < column + width:
                if (wanted, 0) in cells: raise AssertionError('Overlapping mixed DBF source mappings')
                cells[wanted, 0] = record
    reconstructed = 0
    for row, value in enumerate(source_values, 1):
        cell = cells.get((row, 0), {})
        raw, text, kind = cell.get('sourceValue'), cell.get('exportedText'), cell.get('sourceType')
        if isinstance(value, str):
            # Older CLI snapshots omit office:string-value. Their explicit
            # string type and exact displayed text can reconstruct that scalar.
            if kind != 'string' or text != value or raw not in (None, value):
                raise AssertionError('Mixed DBF source string value/type changed')
            reconstructed += raw is None
        else:
            try:
                numeric = kind in ('number', 'float', 'currency', 'percentage') and raw is not None and text is not None
                if not numeric or Decimal(str(raw)) != Decimal(str(value)) or Decimal(str(text)) != Decimal(str(value)):
                    raise AssertionError('Mixed DBF source number value/type changed')
            except (ValueError, ArithmeticError) as error:
                raise AssertionError('Mixed DBF source number value/type missing') from error
    return {**check, 'memoSidecar': 'pass', 'sourceTypeMapping': 'pass',
            'reconstructedStringValues': reconstructed,
            'longUnicodeBytes': len(str(source_values[-1]).encode('utf-8'))}


class _DocumentHTML(HTMLParser):
    def __init__(self):
        super().__init__(); self.images, self.text, self.tables, self.stack, self.open_tables = [], [], [], [], []
        self.saw_body = False; self.stylesheets = []; self.css = []
    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'link' and 'stylesheet' in attrs.get('rel', '').lower().split():
            self.stylesheets.append(attrs.get('href', ''))
        hidden = (bool(self.stack and self.stack[-1][1]) or tag in ('head', 'script', 'style', 'template') or
            'hidden' in attrs or attrs.get('aria-hidden') == 'true' or
            bool(re.search(HIDDEN_CSS, attrs.get('style', ''), re.I)))
        if tag == 'body': self.saw_body = True
        active = not hidden and (tag == 'body' or any(name == 'body' for name, _ in self.stack))
        if tag not in ('img', 'br', 'hr', 'meta', 'link', 'input', 'source', 'wbr', 'area', 'base', 'embed', 'param', 'col'):
            self.stack.append((tag, hidden))
        if not active: return
        if tag == 'img':
            if any(re.fullmatch(r'0+(?:\.0+)?(?:px|%|cm|mm|in|pt)?', attrs.get(key, '').strip()) for key in ('width', 'height')):
                raise AssertionError('Document image has zero visible dimensions')
            self.images.append(attrs.get('src', ''))
        elif tag == 'table':
            table = {'rows': [], 'row': None, 'cell': None}; self.open_tables.append(table); self.tables.append(table['rows'])
        elif self.open_tables:
            table = self.open_tables[-1]
            if tag == 'tr': table['row'] = []; table['rows'].append(table['row'])
            elif tag in ('td', 'th') and table['row'] is not None:
                if int(attrs.get('rowspan', '1')) != 1 or int(attrs.get('colspan', '1')) != 1:
                    raise AssertionError('Control table cells were merged')
                table['cell'] = []; table['row'].append(table['cell'])
    def handle_endtag(self, tag):
        hidden = bool(self.stack and self.stack[-1][1])
        if not hidden and self.open_tables:
            if tag == 'table': self.open_tables.pop()
            elif tag in ('td', 'th'): self.open_tables[-1]['cell'] = None
            elif tag == 'tr': self.open_tables[-1]['row'] = None
        for position in range(len(self.stack) - 1, -1, -1):
            if self.stack[position][0] == tag:
                del self.stack[position:]; break
    def handle_data(self, data):
        if any(tag == 'style' for tag, _ in self.stack): self.css.append(data)
        if not any(tag == 'body' for tag, _ in self.stack) or (self.stack and self.stack[-1][1]): return
        self.text.append(data)
        if self.open_tables and self.open_tables[-1]['cell'] is not None:
            self.open_tables[-1]['cell'].append(data)


def _control_table(tables):
    expected = [[MARKERS[row * 4 + column] for column in range(4)] for row in range(12)]
    actual = [[[' '.join(''.join(cell).split()) for cell in row] for row in table] for table in tables]
    if actual.count(expected) != 1:
        raise AssertionError('Missing or changed 12-row four-column control table cells')
    return 48


def _visible_styles(parsed, *, archive=None, base=None):
    styles = list(parsed.css)
    styles.extend(_image_payload(ref, archive=archive, base=base).decode('utf-8-sig') for ref in parsed.stylesheets)
    # This generated control has no conditional/hidden styling. Fail closed on
    # such rules rather than treating a DOM cell hidden by CSS as visible proof.
    if any(re.search(HIDDEN_CSS, css, re.I) for css in styles):
        raise AssertionError('Hidden CSS content cannot prove the visible control')


def _package_member(base, reference):
    uri = urlsplit(reference)
    if uri.scheme or uri.netloc or not uri.path: raise AssertionError('Nonlocal or missing document image/content reference')
    path = posixpath.normpath(posixpath.join(posixpath.dirname(base), unquote(uri.path)))
    if path.startswith('../') or path.startswith('/'): raise AssertionError('Document reference escapes package')
    return path


def _image_payload(reference, *, archive=None, base=None):
    if reference.startswith('data:'):
        header, separator, encoded = reference.partition(',')
        if not separator or not header.lower().startswith('data:image/'):
            raise AssertionError('Invalid embedded document image')
        return base64.b64decode(encoded, validate=True) if header.lower().endswith(';base64') else unquote_to_bytes(encoded)
    if archive is not None: return archive.read(_package_member(base, reference))
    uri = urlsplit(reference)
    if uri.scheme or uri.netloc or not uri.path: raise AssertionError('Nonlocal or missing HTML image')
    path = (base.parent / unquote(uri.path)).resolve()
    if not path.is_relative_to(base.parent.resolve()): raise AssertionError('HTML image escapes saved folder')
    return path.read_bytes()


def _document_image(data):
    from PySide6.QtGui import QImage
    magic = (data.startswith((b'\x89PNG\r\n\x1a\n', b'\xff\xd8\xff', b'GIF87a', b'GIF89a', b'BM', b'II*\0', b'MM\0*')) or
             data.startswith(b'RIFF') and data[8:12] == b'WEBP')
    image = QImage.fromData(data) if magic else QImage()
    if image.isNull(): raise AssertionError('Referenced document image cannot be decoded')
    # Every illustrated control contains this exact generated PNG. Checking its
    # decoded pixels rejects a substituted picture as well as a missing one.
    if image.width() != 3 or image.height() != 5 or any(image.pixelColor(x, y).getRgb() != (30, 80, 120, 100)
            for y in range(image.height()) for x in range(image.width())):
        raise AssertionError('Referenced document control image pixels changed')
    return 1


def _odf_body(root):
    office = 'urn:oasis:names:tc:opendocument:xmlns:office:1.0'
    body = root.find('{%s}body/{%s}text' % (office, office))
    if body is None: raise AssertionError('Missing Writer document body')
    return body


def _visible_xml(body, *, word=False):
    word_ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    blocked = {'hidden-text', 'hidden-paragraph', 'tracked-changes', 'annotation', 'scripts'}
    def visit(node):
        if node.tag.rsplit('}', 1)[-1] in blocked: return None
        if any(key.rsplit('}', 1)[-1] in ('display', 'visibility') and value in ('none', 'hidden', 'collapse') for key, value in node.attrib.items()): return None
        if word and (node.tag in (word_ns + 'del', word_ns + 'moveFrom', word_ns + 'instrText') or
                     node.tag == word_ns + 'r' and node.find(word_ns + 'rPr/' + word_ns + 'vanish') is not None): return None
        result = ET.Element(node.tag, node.attrib); result.text, result.tail = node.text, node.tail
        for child in node:
            visible = visit(child)
            if visible is not None: result.append(visible)
        return result
    return visit(body)


def _xml_tables(body, *, word=False):
    namespace = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main' if word else 'urn:oasis:names:tc:opendocument:xmlns:table:1.0'
    table, row, cell = ('tbl', 'tr', 'tc') if word else ('table', 'table-row', 'table-cell')
    tables = []
    for element in body.iter('{%s}%s' % (namespace, table)):
        rows = []
        for line in element.iter('{%s}%s' % (namespace, row)):
            cells = []
            for item in line:
                if item.tag != '{%s}%s' % (namespace, cell): continue
                repeat = int(item.get('{%s}number-columns-repeated' % namespace, '1')) if not word else 1
                if repeat != 1: raise AssertionError('Repeated control table cell changed')
                cells.append(list(item.itertext()))
            rows.append(cells)
        tables.append(rows)
    return tables


def painted_pdf_images(page, pdf):
    """Count painted Image XObjects, not unused/shared page resources."""
    from pypdf.generic import ContentStream
    def visit(content, resources, stack):
        if content is None: return 0
        if len(stack) > 32: raise ValueError('PDF Form nesting too deep for QA control')
        total = 0
        for operands, operator in content.operations:
            if operator == b'INLINE IMAGE':
                total += 1; continue
            if operator != b'Do' or not operands: continue
            objects = resources.get('/XObject', {})
            if hasattr(objects, 'get_object'): objects = objects.get_object()
            item = objects[operands[0]].get_object()
            if item.get('/Subtype') == '/Image': total += 1
            elif item.get('/Subtype') == '/Form':
                if id(item) in stack: raise ValueError('PDF Form cycle')
                child_resources = item.get('/Resources', resources)
                if hasattr(child_resources, 'get_object'): child_resources = child_resources.get_object()
                total += visit(ContentStream(item, pdf), child_resources, stack + [id(item)])
        return total
    resources = page.get('/Resources', {})
    if hasattr(resources, 'get_object'): resources = resources.get_object()
    return visit(page.get_contents(), resources, [])


def document_check(paths, target, *, image_required):
    primary = next(p for p in paths if p.suffix.lower() in {'.' + target, '.html' if target == 'htm' else '.' + target})
    cells = None
    if target == 'epub':
        with zipfile.ZipFile(primary) as archive:
            if archive.testzip() is not None: raise AssertionError('EPUB ZIP CRC')
            if archive.read('mimetype') != b'application/epub+zip': raise AssertionError('EPUB MIME')
            container = ET.fromstring(archive.read('META-INF/container.xml'))
            roots = container.findall('{urn:oasis:names:tc:opendocument:xmlns:container}rootfiles/'
                '{urn:oasis:names:tc:opendocument:xmlns:container}rootfile')
            packages = [r.get('full-path') for r in roots if r.get('media-type') == 'application/oebps-package+xml']
            if len(packages) != 1: raise AssertionError('EPUB package declaration')
            package_path = _package_member('', packages[0])
            package = ET.fromstring(archive.read(package_path)); opf = '{http://www.idpf.org/2007/opf}'
            manifest = {}
            for item in package.findall(opf + 'manifest/' + opf + 'item'):
                if not item.get('id') or item.get('id') in manifest: raise AssertionError('EPUB duplicate or missing manifest identity')
                manifest[item.get('id')] = item
            spine = package.find(opf + 'spine')
            if spine is None or not len(spine): raise AssertionError('EPUB reading order')
            texts, tables, images, seen = [], [], 0, set()
            for entry in spine:
                identity = entry.get('idref'); item = manifest.get(identity)
                if item is None or item.get('media-type') != 'application/xhtml+xml' or identity in seen:
                    raise AssertionError('Missing, duplicate or non-XHTML EPUB spine item')
                seen.add(identity); section = _package_member(package_path, item.get('href', ''))
                payload = archive.read(section); ET.fromstring(payload)  # XHTML must be well formed.
                parsed = _DocumentHTML(); parsed.feed(payload.decode('utf-8-sig')); parsed.close()
                if not parsed.saw_body: raise AssertionError('EPUB spine section has no body')
                _visible_styles(parsed, archive=archive, base=section)
                texts.extend(parsed.text); tables.extend(parsed.tables)
                images += sum(_document_image(_image_payload(ref, archive=archive, base=section)) for ref in parsed.images)
            text = ''.join(texts); cells = _control_table(tables)
    elif target in ('odt', 'docm', 'dotx', 'docx'):
        with zipfile.ZipFile(primary) as archive:
            if archive.testzip() is not None: raise AssertionError('Office ZIP CRC')
            member = 'content.xml' if target == 'odt' else 'word/document.xml'
            root = ET.fromstring(archive.read(member))
            body = _odf_body(root) if target == 'odt' else root.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}body')
            if body is None: raise AssertionError('Missing document body')
            body = _visible_xml(body, word=target != 'odt')
            text = ''.join(body.itertext()); cells = _control_table(_xml_tables(body, word=target != 'odt')); images = 0
            if target == 'odt':
                draw = '{urn:oasis:names:tc:opendocument:xmlns:drawing:1.0}'
                for frame in body.iter(draw + 'frame'):
                    for image in frame.findall(draw + 'image'):
                        svg = '{urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0}'
                        if any(not re.match(r'^(?:[1-9]\d*(?:\.\d+)?|0?\.\d*[1-9]\d*)(?:cm|mm|in|pt|px)$', frame.get(svg + key, '')) for key in ('width', 'height')):
                            raise AssertionError('Document image frame has no positive visible dimensions')
                        reference = image.get('{http://www.w3.org/1999/xlink}href')
                        binary = image.find('{urn:oasis:names:tc:opendocument:xmlns:office:1.0}binary-data')
                        data = _image_payload(reference, archive=archive, base=member) if reference else base64.b64decode(''.join(binary.itertext()), validate=True) if binary is not None else b''
                        images += _document_image(data)
            if target in ('docm', 'dotx', 'docx'):
                content_types = ET.fromstring(archive.read('[Content_Types].xml'))
                main = next(e.get('ContentType') for e in content_types if e.get('PartName') == '/word/document.xml')
                expected = {'docm': 'application/vnd.ms-word.document.macroEnabled.main+xml',
                    'dotx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml',
                    'docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml'}[target]
                if main != expected: raise AssertionError('OOXML main Content-Type does not match target')
    elif target == 'pdf':
        from pypdf import PdfReader
        pdf = PdfReader(primary); text = '\n'.join(p.extract_text() or '' for p in pdf.pages)
        images = sum(painted_pdf_images(p, pdf) for p in pdf.pages)
    elif target in ('html', 'htm'):
        parser = _DocumentHTML(); parser.feed(primary.read_text(encoding='utf-8-sig')); parser.close()
        if not parser.saw_body: raise AssertionError('Missing HTML document body')
        _visible_styles(parser, base=primary)
        text = ''.join(parser.text); cells = _control_table(parser.tables)
        images = sum(_document_image(_image_payload(ref, base=primary)) for ref in parser.images)
    elif target == 'xml':
        xml = ET.parse(primary).getroot()
        if xml.get('{urn:oasis:names:tc:opendocument:xmlns:office:1.0}mimetype') != 'application/vnd.oasis.opendocument.text':
            raise AssertionError('Generic XML does not preserve the native Writer dialect')
        body = _visible_xml(_odf_body(xml)); text = ''.join(body.itertext()); images = 0
        cells = _control_table(_xml_tables(body))
    else:
        text = primary.read_text(encoding='utf-8-sig'); images = 0
    missing = [m for m in MARKERS if m not in text]
    if missing: raise AssertionError('Missing document/table markers: ' + str(len(missing)))
    if image_required and images < 1: raise AssertionError('Document image lost')
    return {'status': 'pass', 'markers': len(MARKERS), 'images': images,
            'tableCells': cells, 'method': 'independent_visible_body_table_cells_and_referenced_image_pixels'}


def first_frame_hash(path):
    from PySide6.QtGui import QImage, QImageReader
    reader = QImageReader(str(path)); image = reader.read()
    if image.isNull():
        if Path(path).suffix.lower() in ('.tif', '.tiff'):
            try:
                return uncompressed_tiff_frame_hash(path)
            except (ValueError, KeyError, struct.error, IndexError):
                pass
        return {'status': 'not_run', 'reason': reader.errorString()}
    image = image.convertToFormat(QImage.Format.Format_RGBA8888)
    raw = bytes(image.constBits())
    pixels = b''.join(raw[y * image.bytesPerLine():y * image.bytesPerLine() + image.width() * 4]
                      for y in range(image.height()))
    return {'status': 'pass', 'rgbaSha256': hashlib.sha256(pixels).hexdigest(),
            'width': image.width(), 'height': image.height(), 'alpha': image.hasAlphaChannel(),
            'decoder': 'QImageReader'}


def decoded_frame_hashes(path):
    """Decode every generated full-canvas animation frame independently."""
    from PySide6.QtGui import QImage, QImageReader
    path = Path(path)
    info = inspect_media(path)
    if info['format'] == 'png' and info['frames'] > 1:
        data = path.read_bytes()
        position, ihdr, ancillary, frame = 8, None, [], []
        compressed, hashes = [], []
        def png_chunk(kind, payload):
            return struct.pack('>I', len(payload)) + kind + payload + struct.pack('>I', zlib.crc32(kind + payload) & 0xffffffff)
        def finish():
            if not compressed: return
            control = struct.unpack('>IIIIIHHBB', frame[0])
            width, height, x, y = control[1:5]
            if (x, y, width, height) != (0, 0, info['width'], info['height']):
                raise ValueError('QA APNG decoder requires generated full-canvas controls')
            standalone = b'\x89PNG\r\n\x1a\n' + png_chunk(b'IHDR', ihdr)
            standalone += b''.join(ancillary) + png_chunk(b'IDAT', b''.join(compressed)) + png_chunk(b'IEND', b'')
            image = QImage.fromData(standalone)
            if image.isNull(): raise ValueError('APNG frame pixel decode failed')
            image = image.convertToFormat(QImage.Format.Format_RGBA8888)
            raw = bytes(image.constBits())
            pixels = b''.join(raw[y * image.bytesPerLine():y * image.bytesPerLine() + image.width() * 4] for y in range(image.height()))
            hashes.append(hashlib.sha256(pixels).hexdigest())
        while position + 12 <= len(data):
            length = struct.unpack_from('>I', data, position)[0]
            kind, payload = data[position + 4:position + 8], data[position + 8:position + 8 + length]
            if kind == b'IHDR': ihdr = payload
            elif kind in (b'PLTE', b'tRNS'): ancillary.append(png_chunk(kind, payload))
            elif kind == b'fcTL':
                finish(); compressed.clear(); frame[:] = [payload]
            elif kind == b'IDAT': compressed.append(payload)
            elif kind == b'fdAT': compressed.append(payload[4:])
            elif kind == b'IEND': finish(); break
            position += 12 + length
        if len(hashes) != info['frames']: raise ValueError('Incomplete APNG frame decode')
        return hashes
    if info['frames'] == 1:
        proof = first_frame_hash(path)
        return [proof['rgbaSha256']] if proof['status'] == 'pass' else None
    reader = QImageReader(str(path)); hashes = []
    for _ in range(info['frames']):
        image = reader.read()
        if image.isNull(): raise ValueError('Animation pixel decode failed: ' + reader.errorString())
        image = image.convertToFormat(QImage.Format.Format_RGBA8888)
        raw = bytes(image.constBits())
        pixels = b''.join(raw[y * image.bytesPerLine():y * image.bytesPerLine() + image.width() * 4] for y in range(image.height()))
        hashes.append(hashlib.sha256(pixels).hexdigest())
    return hashes


def uncompressed_tiff_frame_hash(path, page_index=0):
    """Independent pixel reader for small uncompressed classic/BigTIFF controls."""
    data = Path(path).read_bytes()
    order = '<' if data[:2] == b'II' else '>' if data[:2] == b'MM' else None
    if order is None:
        raise ValueError('Unsupported TIFF header')
    version = struct.unpack_from(order + 'H', data, 2)[0]
    if version == 42:
        position = struct.unpack_from(order + 'I', data, 4)[0]
        count = struct.unpack_from(order + 'H', data, position)[0]
        count_size, entry_size, inline_size, amount_code, offset_code = 2, 12, 4, 'I', 'I'
    elif version == 43 and data[4:8] == struct.pack(order + 'HH', 8, 0):
        position = struct.unpack_from(order + 'Q', data, 8)[0]
        count = struct.unpack_from(order + 'Q', data, position)[0]
        count_size, entry_size, inline_size, amount_code, offset_code = 8, 20, 8, 'Q', 'Q'
    else:
        raise ValueError('Unsupported TIFF header')
    for _ in range(page_index):
        position = struct.unpack_from(order + offset_code, data,
                                      position + count_size + count * entry_size)[0]
        if not position: raise ValueError('TIFF page index outside file')
        count = struct.unpack_from(order + ('H' if version == 42 else 'Q'), data, position)[0]
    tags = {}
    for index in range(count):
        offset = position + count_size + index * entry_size
        tag, kind = struct.unpack_from(order + 'HH', data, offset)
        amount = struct.unpack_from(order + amount_code, data, offset + 4)[0]
        if kind not in (1, 3, 4, 16): continue
        code, size = {1: ('B', 1), 3: ('H', 2), 4: ('I', 4), 16: ('Q', 8)}[kind]
        value = offset + entry_size - inline_size
        payload = value if amount * size <= inline_size else struct.unpack_from(order + offset_code, data, value)[0]
        tags[tag] = struct.unpack_from(order + code * amount, data, payload)
    width, height = tags[256][0], tags[257][0]
    channels = tags[277][0]
    if (tags.get(259, (1,))[0] != 1 or tags.get(262, (2,))[0] != 2 or
            tags.get(284, (1,))[0] != 1 or channels not in (3, 4) or any(v != 8 for v in tags[258])):
        raise ValueError('Unsupported TIFF pixel layout')
    pixels = b''.join(data[o:o + length] for o, length in zip(tags[273], tags[279]))
    if len(pixels) != width * height * channels:
        raise ValueError('Truncated TIFF strips')
    if channels == 3:
        pixels = b''.join(pixels[p:p + 3] + b'\xff' for p in range(0, len(pixels), 3))
    elif tags.get(338, (2,))[0] != 2:
        raise ValueError('Associated TIFF alpha not supported by QA control')
    return {'status': 'pass', 'rgbaSha256': hashlib.sha256(pixels).hexdigest(),
            'width': width, 'height': height, 'alpha': channels == 4,
            'decoder': 'independent_uncompressed_TIFF'}


def media_check(source, paths, target):
    primary = next(p for p in paths if p.suffix.lower() == '.' + target)
    before, after = inspect_media(source), inspect_media(primary)
    keys = ('width', 'height', 'frames', 'pages', 'alpha', 'durationsMs') + (('loopCount',) if before['frames'] > 1 else ())
    for key in keys:
        if before.get(key) != after.get(key): raise AssertionError('Changed media property: ' + key)
    identity = source.suffix.lower() == primary.suffix.lower()
    if identity and digest(source) != digest(primary): raise AssertionError('Identity bytes changed')
    decoded_source, decoded_result = first_frame_hash(source), first_frame_hash(primary)
    source_frames, result_frames = decoded_frame_hashes(source), decoded_frame_hashes(primary)
    if source_frames is not None and result_frames is not None and source_frames != result_frames:
        raise AssertionError('Decoded animation frame pixels changed')
    if decoded_source['status'] == decoded_result['status'] == 'pass':
        if decoded_source['rgbaSha256'] != decoded_result['rgbaSha256']:
            raise AssertionError('Decoded first-frame pixels changed')
    elif not identity:
        raise AssertionError('Cross-format output decoder unavailable')
    return {'status': 'pass', 'source': before, 'result': after,
            'byteIdentity': identity, 'decodedSource': decoded_source, 'decodedResult': decoded_result,
            'decodedFrameCount': len(result_frames) if result_frames is not None else None,
            'method': 'whole_file_identity_or_properties_and_decoded_RGBA'}


def svm_check(paths, work, executable, seconds, *, expected_pages):
    primary = [p for p in paths if p.suffix.lower() == '.svm']
    if not primary or len(primary) != expected_pages:
        raise AssertionError('SVM output page count changed')
    checks = []
    from PySide6.QtGui import QImage, QImageReader
    for number, path in enumerate(primary):
        with path.open('rb') as stream:
            if stream.read(6) != b'VCLMTF': raise AssertionError('SVM native container signature')
        exports = native_export(path, 'png', work / str(number), executable, seconds, 'draw_png_Export')
        image_path = next(p for p in exports if p.suffix.lower() == '.png')
        image = QImageReader(str(image_path)).read()
        if image.isNull(): raise AssertionError('Native SVM re-import did not produce a decodable image')
        # The generated Draw control has a coloured rectangle and text. A plain
        # white or transparent canvas cannot satisfy this independent render.
        sample = image.scaled(64, 64)
        coloured = sum(1 for y in range(sample.height()) for x in range(sample.width())
            if sample.pixelColor(x, y).alpha() and
            min(sample.pixelColor(x, y).red(), sample.pixelColor(x, y).green(), sample.pixelColor(x, y).blue()) < 220)
        if coloured < 3: raise AssertionError('Native SVM control rendered an empty canvas')
        rgba = image.convertToFormat(QImage.Format.Format_RGBA8888)
        checks.append({'size': path.stat().st_size, 'nativeReimportWidth': image.width(),
                       'nativeReimportHeight': image.height(), 'nonwhiteSamples': coloured,
                       'rgbaSha256': hashlib.sha256(bytes(rgba.constBits())).hexdigest()})
    if len({row['rgbaSha256'] for row in checks}) != expected_pages:
        raise AssertionError('Distinct SVM fixture pages rendered as duplicate images')
    return {'status': 'pass', 'pages': len(primary), 'controls': checks,
            'method': 'SVM_signature_native_reimport_and_independent_QImage_decode'}


def draw_html_check(paths, target, source):
    from urllib.parse import unquote, urlsplit
    from PySide6.QtGui import QImage
    primary = next(p for p in paths if p.suffix.lower() == '.' + target)
    tree = ET.parse(primary).getroot()
    objects = [e for e in tree.iter() if e.tag.rsplit('}', 1)[-1] == 'object']
    markers = source['pageMarkers']
    if len(objects) != len(markers): raise AssertionError('Drawing HTML page count changed')
    pages = []
    def safe_nodes(root):
        for node in root.iter():
            if node.tag.rsplit('}', 1)[-1].lower() in ('script', 'foreignObject'.lower()):
                raise AssertionError('Active drawing output')
            for key, value in node.attrib.items():
                local = key.rsplit('}', 1)[-1].lower()
                if local.startswith('on'): raise AssertionError('Drawing event handler')
                if local == 'style' and any(x in value.lower() for x in ('@import', 'javascript:', 'http:', 'https:', 'file:')):
                    raise AssertionError('External or active drawing CSS')
                if local in ('href', 'src', 'data') and value and not value.startswith(('#', 'data:')):
                    parsed = urlsplit(value)
                    if parsed.scheme or parsed.netloc: raise AssertionError('External drawing resource')
                    dependency = (primary.parent / unquote(parsed.path)).resolve(strict=True)
                    if not dependency.is_relative_to(primary.parent.resolve()): raise AssertionError('Escaped drawing dependency')
            if node.tag.rsplit('}', 1)[-1].lower() == 'style':
                css = ''.join(node.itertext()).lower()
                if '@import' in css or 'javascript:' in css or 'http:' in css or 'https:' in css:
                    raise AssertionError('External or active drawing CSS')
    safe_nodes(tree)
    for index, obj in enumerate(objects):
        parsed = urlsplit(obj.get('data', ''))
        if parsed.scheme or parsed.netloc: raise AssertionError('External SVG page')
        page = (primary.parent / unquote(parsed.path)).resolve(strict=True)
        if not page.is_relative_to(primary.parent.resolve()) or page.suffix.lower() != '.svg':
            raise AssertionError('Invalid SVG page dependency')
        svg = ET.parse(page).getroot(); safe_nodes(svg)
        text = ''.join(svg.itertext())
        if markers[index] not in text or any(m in text for m in markers if m != markers[index]):
            raise AssertionError('Drawing pages flattened, reordered or missing their text')
        images = [e for e in svg.iter() if e.tag.rsplit('}', 1)[-1] == 'image']
        matched = 0
        for image in images:
            href = image.get('href') or image.get('{http://www.w3.org/1999/xlink}href', '')
            if not href.startswith('data:image/png;base64,'): raise AssertionError('Drawing PNG not embedded')
            payload = base64.b64decode(''.join(href.split(',', 1)[1].split()), validate=True)
            if hashlib.sha256(payload).hexdigest() != source['imageSha256']: raise AssertionError('Drawing PNG bytes changed')
            decoded = QImage.fromData(payload)
            if decoded.isNull() or decoded.width() != 1 or decoded.height() != 1:
                raise AssertionError('Drawing embedded PNG pixel decode failed')
            matched += 1
        if not matched: raise AssertionError('Drawing page image missing')
        if not any(e.tag.rsplit('}', 1)[-1] in ('path', 'rect', 'polygon', 'polyline') for e in svg.iter()):
            raise AssertionError('Drawing vector geometry missing')
        pages.append({'marker': markers[index], 'svgSha256': digest(page), 'embeddedOriginalPngs': matched,
                      'embeddedPngSha256': source['imageSha256']})
    return {'status': 'pass', 'pages': len(pages), 'pageChecks': pages,
            'method': 'independent_SVG_page_order_text_original_PNG_decode_and_local_dependency_checks'}


def pdf_pages_check(paths):
    primary = next(p for p in paths if p.suffix.lower() == '.tiff')
    facts = inspect_media(primary)
    if facts['pages'] != 3 or facts.get('pageDimensions') != [[300, 400]] * 3:
        raise AssertionError('PDF page count or 300 DPI page dimensions changed')
    hashes = [uncompressed_tiff_frame_hash(primary, i)['rgbaSha256'] for i in range(3)]
    expected = [hashlib.sha256(bytes(rgba) * 300 * 400).hexdigest()
                for rgba in [(255, 0, 0, 255), (0, 255, 0, 255), (0, 0, 255, 255)]]
    if hashes != expected: raise AssertionError('Rendered PDF page colours/pixels changed')
    return {'status': 'pass', 'pages': 3, 'pageDimensions': facts['pageDimensions'],
            'pageRgbaSha256': hashes, 'method': 'independent_BigTIFF_all_page_pixel_hashes_300_DPI'}


def _raster_sources(fixtures, helper):
    sources = {}
    raster = [('png', helper.png(width=3, height=5), ['png', 'webp', 'tiff']),
        ('gif', helper.gif(frames=8, delay=15), ['gif', 'apng', 'webp']),
        ('apng', helper.png(width=3, height=5, frames=8, delay=150), ['apng', 'webp']),
        ('webp', helper.WEBP_ANIMATED, ['webp', 'apng']),
        ('tiff', helper.tiff(pages=3, alpha=True, width=3, height=5), ['tiff'])]
    for ext, data, targets in raster:
        path = fixtures / ('animated.' + ext if ext in ('gif', 'apng', 'webp') else 'picture.' + ext)
        path.write_bytes(data); inspect_media(path)
        sources[ext] = {'path': path, 'targets': targets, 'kind': 'media'}
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    pdf = PdfWriter()
    for colour in ('1 0 0', '0 1 0', '0 0 1'):
        page = pdf.add_blank_page(width=72, height=96)
        page[NameObject('/Resources')] = DictionaryObject()
        content = DecodedStreamObject(); content.set_data(('q ' + colour + ' rg 0 0 72 96 re f Q').encode('ascii'))
        page[NameObject('/Contents')] = pdf._add_object(content)
    path = fixtures / 'three-pages.pdf'
    with path.open('wb') as output: pdf.write(output)
    sources['pdf_pages'] = {'path': path, 'targets': ['tiff'], 'kind': 'pdf_pages'}
    return sources


def prepare(root, executable, seconds, suite='all'):
    from openpyxl import Workbook
    helper = _media_fixture_helpers()
    fixtures = root / 'fixtures'; fixtures.mkdir()
    if suite == 'media':
        return _raster_sources(fixtures, helper)
    def save_workbook(path, rows_by_sheet):
        book = Workbook(); book.remove(book.active)
        for index, rows in enumerate(rows_by_sheet):
            sheet = book.create_sheet('Sheet%d' % (index + 1))
            for row in rows:
                sheet.append([float(v) if isinstance(v, Decimal) else v for v in row])
        book.save(path)
    table = fixtures / 'single.xlsx'; save_workbook(table, [TABLE_ROWS])
    second = [list(row) for row in TABLE_ROWS]; second[1][0] = 'SECOND_ALPHA'; second[3][0] = 'SECOND_SENTINEL'
    multi = fixtures / 'multiple.xlsx'; save_workbook(multi, [TABLE_ROWS, second])
    sources = {'table': {'path': table, 'targets': ['dbf', 'dif'], 'kind': 'table', 'expected': [TABLE_ROWS]},
        'multisheet': {'path': multi, 'targets': ['dbf', 'dif'], 'kind': 'table', 'expected': [TABLE_ROWS, second]}}
    long_text = 'Я🙂' * 200
    mixed_rows = [['MIXED', 'COUNT'], [7, 1], ['TEXT_VALUE', 2], [long_text, 3]]
    mixed = fixtures / 'mixed.xlsx'; save_workbook(mixed, [mixed_rows])
    sources['mixed'] = {'path': mixed, 'targets': ['dbf'], 'kind': 'mixed_dbf',
        'expected': [['MIXED', 'COUNT'], ['7', 1], ['TEXT_VALUE', 2], [long_text, 3]],
        'sourceValues': [7, 'TEXT_VALUE', long_text]}
    for ext, fil in [('dbf', 'dBase:76'), ('dif', 'DIF:76')]:
        paths = native_export(table, ext, root / ('make-' + ext), executable, seconds, fil)
        path = next(p for p in paths if p.suffix.lower() == '.' + ext)
        table_check([path], ext, [TABLE_ROWS])
        sources[ext] = {'path': path, 'targets': ['xlsx', 'dif' if ext == 'dbf' else 'dbf'],
                        'kind': 'table', 'expected': [dbf_import_rows(path) if ext == 'dbf' else TABLE_ROWS]}
    if suite != 'routes': sources.update(_raster_sources(fixtures, helper))
    image = base64.b64encode(helper.png(width=3, height=5)).decode('ascii')
    xml = writer_seed().replace('xmlns:table=',
        'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
        'xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" xmlns:table=')
    xml = xml.replace('</office:text>', '<text:p><draw:frame text:anchor-type="as-char" '
        'svg:width="1cm" svg:height="1cm"><draw:image draw:mime-type="image/png">'
        '<office:binary-data>' + image + '</office:binary-data></draw:image></draw:frame></text:p></office:text>')
    writer = fixtures / 'illustrated.fodt'; writer.write_text(xml, encoding='utf-8')
    odts = native_export(writer, 'odt', root / 'make-odt', executable, seconds, 'writer8')
    odt = next(p for p in odts if p.suffix.lower() == '.odt')
    document_check([odt], 'odt', image_required=True)
    sources['writer'] = {'path': odt, 'targets': ['txt', 'epub'], 'kind': 'document', 'imageRequired': True}
    epubs = native_export(odt, 'epub', root / 'make-epub', executable, seconds, 'EPUB')
    epub = next(p for p in epubs if p.suffix.lower() == '.epub')
    document_check([epub], 'epub', image_required=True)
    sources['epub'] = {'path': epub, 'targets': ['odt', 'pdf', 'html'], 'kind': 'document', 'imageRequired': True}
    html = fixtures / 'web-writer.html'
    html.write_text('<!doctype html><html><head><meta charset="utf-8"></head><body>'
        '<p>HTML Writer native routes Проверка</p><table>' + ''.join('<tr>' + ''.join(
            '<td>' + marker + '</td>' for marker in MARKERS[row * 4:row * 4 + 4]) + '</tr>'
            for row in range(12)) + '</table></body></html>', encoding='utf-8')
    sources['html_writer'] = {'path': html, 'targets': ['docm', 'dotx', 'xml'],
                              'kind': 'document', 'imageRequired': False}
    spec = importlib.util.spec_from_file_location('aio_v18_drawing_fixtures', ROOT / 'tests/test_guard_native_draw_html_v18.py')
    drawing = importlib.util.module_from_spec(spec); spec.loader.exec_module(drawing)
    for key, presentation in [('drawing', False)] + ([('slides', True)] if os.name == 'nt' else []):
        flat = drawing.make_flat_drawing(fixtures / (key + ('.fodp' if presentation else '.fodg')),
                                         presentation=presentation)
        from app_backend.converter_guard.adapter import _RunControl, prepare_profile
        from app_backend.libreoffice_engine import convert_path
        work = root / ('make-' + key)
        prepare_profile(work / 'profile')
        target, exporter = ('odp', 'impress8') if presentation else ('odg', 'draw8')
        importer = 'OpenDocument Presentation Flat XML' if presentation else 'OpenDocument Drawing Flat XML'
        native_paths = convert_path(flat, target, work, executable=executable,
            cancel=_RunControl(threading.Event(), work, seconds, None),
            convert_spec=target + ':' + exporter, import_filter=importer)
        path = next(p for p in native_paths if p.suffix.lower() == '.' + target)
        sources[key] = {'path': path, 'targets': ['html', 'htm', 'xhtml'] +
            (['svm'] if key == 'drawing' and os.name == 'nt' else []), 'kind': 'draw_html',
            'pageMarkers': list(drawing.MARKERS), 'imageSha256': hashlib.sha256(drawing.PNG).hexdigest()}
    return sources


def run(root, sources, executable, seconds):
    guard = ConverterGuard(adapter=LibreOfficeAdapter(executable),
        limits=GuardLimits(preflight_timeout=seconds, file_timeout=seconds), recovery_parent=root / 'owned')
    row = {'status': 'fail', 'requestedPairs': sum(len(s['targets']) for s in sources.values()),
           'unplannedPairs': [], 'checks': []}
    before = {key: digest(s['path']) for key, s in sources.items()}
    try:
        metadata = [{'id': key, 'name': s['path'].name, 'size': s['path'].stat().st_size} for key, s in sources.items()]
        job = guard.create(metadata, prepare_only=True)['jobId']
        for item in metadata:
            with sources[item['id']]['path'].open('rb') as stream:
                guard.upload(job, item['id'], stream, item['size'])
        inspection = guard.preflight(job); row['preflight'] = inspection['files']
        formats = {f['id']: f for f in inspection['formats']}
        plan = []
        for key, source in sources.items():
            proof = next(f['preflight'] for f in inspection['files'] if f['id'] == key)
            admitted = []
            for target in source['targets']:
                fmt = formats.get(target, {})
                if key in fmt.get('selectableFileIds', []): admitted.append(target)
                else: row['unplannedPairs'].append({'source': key, 'target': target,
                    'reason': proof.get('code') or fmt.get('unavailableReasons', {}).get(key) or fmt.get('reasonCode')})
            if admitted: plan.append({'fileId': key, 'formats': admitted})
        if not plan: return {**row, 'reason': 'no_admitted_pairs'}
        guard.start(job, file_ids=[p['fileId'] for p in plan], plan=plan,
                    only_verified=True, archive_mode='on_demand')
        until = time.monotonic() + seconds * sum(len(p['formats']) for p in plan) + 120
        last_completed = -1
        while True:
            state = guard.status(job)
            if state['completed'] != last_completed:
                last_completed = state['completed']; print(json.dumps({'completed': last_completed, 'total': state['total']}), flush=True)
            if state['finished']: break
            if time.monotonic() > until: guard.cancel(job); raise RuntimeError('QA supervisor deadline')
            time.sleep(.05)
        row['report'] = guard.report(job)
        if state.get('result'):
            row['save'] = guard.save_folder(state['result']['ticket'], root, verify=True)
            saved = root / row['save']['name']
            manifest = json.loads((saved / 'AIO-results-manifest.json').read_text(encoding='utf-8'))
            grouped = {}
            for item in manifest['outputs']:
                grouped.setdefault((item['fileId'], item['format']), []).append(saved / item['path'])
            for (key, target), paths in grouped.items():
                # Exclude sidecars from table and raster primary counts.
                primary = [p for p in paths if p.suffix.lower() == '.' + target]
                source = sources[key]
                try:
                    if source['kind'] == 'table':
                        check = table_check(primary, target, source['expected'])
                        if target == 'dif': check.update(dif_charset_check(paths))
                    elif source['kind'] == 'mixed_dbf': check = mixed_dbf_check(paths, source['expected'], source['sourceValues'])
                    elif source['kind'] == 'media': check = media_check(source['path'], primary, target)
                    elif target == 'svm': check = svm_check(paths, root / ('svm-control-' + key), executable, seconds,
                        expected_pages=len(source['pageMarkers']))
                    elif source['kind'] == 'draw_html': check = draw_html_check(paths, target, source)
                    elif source['kind'] == 'pdf_pages': check = pdf_pages_check(paths)
                    else: check = document_check(paths, target,
                        image_required=source.get('imageRequired', False) and target != 'txt')
                    row['checks'].append({'source': key, 'target': target, **check})
                except Exception as error:
                    row['checks'].append({'source': key, 'target': target, 'status': 'fail',
                                          'reason': type(error).__name__ + ':' + str(error)})
        row['originalsUnchanged'] = all(digest(s['path']) == before[k] for k, s in sources.items())
        expected = row['requestedPairs']
        row['status'] = ('pass' if not row['unplannedPairs'] and state['failed'] == state['cancelled'] == 0
            and state['succeeded'] == expected and len(row['checks']) == expected
            and all(c['status'] == 'pass' for c in row['checks']) and row['originalsUnchanged'] else 'fail')
        return row
    finally:
        guard.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--executable', type=Path, required=True)
    parser.add_argument('--seconds', type=float, default=45)
    parser.add_argument('--suite', choices=('all', 'routes', 'media'), default='all')
    args = parser.parse_args()
    if args.seconds <= 0 or not args.executable.is_file(): parser.error('Existing native executable and positive timeout required')
    args.executable = args.executable.resolve()
    native = registry(args.executable)
    if not native.valid: parser.error('Use actual LibreOffice program executable with native registry')
    os.environ.setdefault('QT_QPA_PLATFORM', 'offscreen')
    from PySide6.QtGui import QGuiApplication
    application = QGuiApplication.instance() or QGuiApplication([])
    args.output.mkdir(parents=True, exist_ok=False)
    started = time.monotonic()
    evidence = {'schema': 1, 'status': 'fail', 'scope': 'targeted_native_routes_and_media', 'suite': args.suite,
        'machine': machine_facts(args.executable), 'qualityCertified': False,
        'nativeExportCoverage': {target: {'registeredServices': sorted(native.exports.get(target, ())),
            'status': 'available_requires_compatible_fixture' if native.exports.get(target) else 'unavailable',
            'reason': None if native.exports.get(target) else 'NATIVE_EXPORT_FILTER_ABSENT'} for target in ('txt', 'svm', 'orp')}}
    evidence['nativeExportCoverage']['svm'].update(
        status='native_graphic_route_test_required' if os.name == 'nt' else 'NOT_RUN',
        graphicExport='requires_live_UNO_GraphicExportFilter_image/x-svm_proof',
        acceptance='required_Draw_to_SVM' if os.name == 'nt' else 'NOT_RUN_Linux_LOK_has_no_graphic_encoder_proof')
    evidence['nativeExportCoverage']['presentationHtml'] = {
        'status': 'required_two_slide_HTML_HTM_XHTML' if os.name == 'nt' else 'NOT_RUN',
        'reason': None if os.name == 'nt' else 'requires_observed_UNO_per_page_SVG_graphic_export_proof'}
    try:
        sources = prepare(args.output, args.executable, args.seconds, args.suite)
        evidence['fixtures'] = {key: {'name': s['path'].name, 'sha256': digest(s['path']),
            'kind': s['kind'], 'targets': s['targets']} for key, s in sources.items()}
        evidence['acceptance'] = run(args.output, sources, args.executable, args.seconds)
        evidence['status'] = evidence['acceptance']['status']
        for target in ('txt', 'svm'):
            checked = [c for c in evidence['acceptance']['checks'] if c['target'] == target]
            if checked:
                evidence['nativeExportCoverage'][target].update(
                    status='pass' if all(c['status'] == 'pass' for c in checked) else 'fail',
                    actualContentChecks=len(checked), reason=None)
    except Exception as error:
        evidence['error'] = {'type': type(error).__name__, 'code': getattr(error, 'code', None), 'message': str(error)}
    finally:
        evidence['elapsedMs'] = round((time.monotonic() - started) * 1000)
        (args.output / 'evidence.json').write_text(json.dumps(evidence, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
        print(json.dumps({'status': evidence['status'], 'elapsedMs': evidence['elapsedMs'],
                          'error': evidence.get('error')}), flush=True)
    return int(evidence['status'] != 'pass')


if __name__ == '__main__':
    sys.exit(main())
