"""Guarded audit 101 edits. Plain resulting source is committed, not patched at runtime."""
import ast
import hashlib
from pathlib import Path
import subprocess
from deep_audit_methods import ASSETS

BASE = '97fabe65b64076dc8063e0813865ac70b42bfb14'
changes = {}
originals = {}
PREVIOUS = {}


def source(path):
    value = subprocess.check_output(['git', 'show', BASE + ':' + path]).decode('utf-8-sig')
    originals[path] = value
    return value


def replace(value, before, after, count=1):
    if value.count(before) != count:
        raise RuntimeError('Source anchor mismatch: ' + repr(before[:160]))
    return value.replace(before, after)


def method(value, cls, name, replacement):
    tree = ast.parse(value)
    container = tree if cls is None else next(n for n in tree.body if isinstance(n, ast.ClassDef) and n.name == cls)
    node = next(n for n in container.body if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef)) and n.name == name)
    start = min([node.lineno] + [n.lineno for n in node.decorator_list]) - 1
    lines = value.splitlines(keepends=True)
    return ''.join(lines[:start]) + replacement.rstrip() + '\n' + ''.join(lines[node.end_lineno:])

p = 'app_backend/converter_guard/result_store.py'
s = source(p)
s = method(s, None, 'atomic_json', 'from ..atomic_metadata import atomic_json')
changes[p] = s

p = 'app_backend/legacy_result_journal.py'
s = source(p)
s = method(s, None, 'file_hash', '''def _check_cancel(cancel):
    if cancel is not None and cancel.is_set():
        raise InterruptedError('Legacy recovery stopped')


def file_hash(path, *, cancel=None):
    _check_cancel(cancel)
    digest = hashlib.sha256()
    with path.open('rb') as stream:
        while block := stream.read(1024 * 1024):
            _check_cancel(cancel)
            digest.update(block)
    _check_cancel(cancel)
    return digest.hexdigest()''')
s = replace(s, '    def load(self):\n        root = Path(self.name)',
               '    def load(self, *, cancel=None):\n        _check_cancel(cancel)\n        root = Path(self.name)')
s = replace(s, "        checksum = value.pop('checksum', None)",
               "        if not isinstance(value, dict):\n            raise ValueError('Invalid legacy journal object')\n        checksum = value.pop('checksum', None)")
s = replace(s, "                value.get('kind') != 'aio_legacy_result' or",
               "                value.get('kind') != 'aio_legacy_result' or\n                not isinstance(value.get('path'), str) or")
s = replace(s, "file_hash(output) != value['sha256']", "file_hash(output, cancel=cancel) != value['sha256']")
s = method(s, None, 'recover_workspaces', ASSETS['recover_workspaces.txt'])
s = method(s, None, 'salvage_unverified_workspace', ASSETS['salvage_unverified_workspace.txt'])
changes[p] = s

p = 'app_backend/converter_transport.py'
s = source(p)
s = replace(s, 'from .legacy_storage_budget import LegacyRunControl, measured_bytes, storage_error',
               'from .legacy_storage_budget import LegacyRunControl, measured_bytes, storage_error\nfrom .file_identity import owned_file_identity')
s = replace(s, '    salvaged: bool = False\n', '''    salvaged: bool = False
    source_workspace_id: str | None = None
    availability_error: str | None = None
    file_identity: tuple | None = field(default=None, init=False)

    def __post_init__(self):
        self.file_identity = owned_file_identity(self.path, self.temporary.name)
''')
s = replace(s, '        self._recovery_scan_lock = threading.Lock()', '''        self._recovery_scan_lock = threading.Lock()
        self._admission_lock = threading.Lock()
        self._shutdown = threading.Event()
        self._salvage_roots = {}''')
s = method(s, 'ConverterTransport', 'scan_recovery', ASSETS['transport_scan.txt'])
s = method(s, 'ConverterTransport', 'salvage_workspace', ASSETS['transport_salvage.txt'])
s = method(s, 'ConverterTransport', '_owned_root_count', ASSETS['transport_root_count.txt'])
old = '''            roots = ([self.recovery_parent] + [result.temporary.name
                      for result in self.results.values() if result.salvaged]
                     if self.recovery_parent is not None else
                     list({item.temporary.name for item in self.active.values()} |
                          {item.temporary.name for item in self.results.values()} |
                          {item.name for item in self.cleanup_pending}))'''
s = replace(s, old, '''            owned = {item.temporary.name for item in self.active.values()} | {
                item.temporary.name for item in self.results.values()} | {
                item.name for item in self.cleanup_pending} | {
                item.name for item in self._salvage_roots.values() if item is not None}
            roots = list(owned)
            if self.recovery_parent is not None:
                parent = self.recovery_parent.absolute()
                roots = [parent] + [name for name in owned
                                    if not Path(name).absolute().is_relative_to(parent)]''')
s = replace(s, '    def begin(self, identity, connection, *, expected_upload_bytes=0):', '''    def begin(self, identity, connection, *, expected_upload_bytes=0):
        # Serialize admission/reservation, not status, cancellation or native work.
        with self._admission_lock:
            return self._begin_owned(identity, connection, expected_upload_bytes=expected_upload_bytes)

    def _begin_owned(self, identity, connection, *, expected_upload_bytes=0):''')
start = s.index('        disk_roots = 0\n', s.index('    def _begin_owned'))
end = s.index('        with self.lock:\n', start)
s = s[:start] + '        disk_roots = self._disk_roots()\n' + s[end:]
s = replace(s, 'max(self._owned_root_count(), disk_roots)', 'self._owned_root_count(disk_roots)')
s = replace(s, "                if request_id is not None:\n                    deleted = self.release(ticket)",
'''                with self.lock:
                    derivatives = [key for key, result in self.results.items()
                                   if result.source_workspace_id == Path(workspace.name).name]
                copies_deleted = all([self.release(key) for key in derivatives])
                if request_id is not None:
                    deleted = self.release(ticket)''')
s = replace(s, "return {'declined': True, 'cleanupPending': not deleted,",
               "return {'declined': True, 'cleanupPending': not (deleted and copies_deleted),")
a = s.index('    def decline_all_recovery(')
b = s.index('    def salvage_workspace(', a)
part = s[a:b]
part = replace(part, '        if not self._recovery_loaded:\n            self.scan_recovery()',
'''        self.scan_recovery()
        with self.lock:
            if self.closed:
                raise LibreOfficeError('Программа закрывается.', 'CANCELLED')
            if not self.recovery_available:
                raise LibreOfficeError('Не удалось прочитать хранилище. Общий отказ не подтверждён.',
                                       'LEGACY_STORAGE_UNAVAILABLE')''')
part = replace(part, "        return {'declined': completed, 'notDeclined': len(errors),", '''        # Also retire derivatives whose source disappeared externally.
        with self.lock:
            derivatives = [key for key, result in self.results.items()
                           if result.salvaged and result.source_workspace_id and not result.released]
        for key in derivatives:
            pending += not self.release(key)
        return {'declined': completed, 'derivedDiscarded': len(derivatives), 'notDeclined': len(errors),''')
s = s[:a] + part + s[b:]
s = replace(s, "items = [{'state': 'salvaged_unverified' if result.salvaged else 'ready',", 
'''items = [{'state': 'unavailable' if result.availability_error else
                                'salvaged_unverified' if result.salvaged else 'ready',
                      'available': result.availability_error is None,
                      'availabilityCode': result.availability_error,''')
s = replace(s, "'recoverUrl': '/api/converter/recover?request=' + result.request_id}",
               "'recoverUrl': None if result.availability_error else '/api/converter/recover?request=' + result.request_id}")
s = replace(s, "'total': len(available), 'offset': offset, 'limit': limit,", 
''' 'total': len(available), 'saveableTotal': sum(not r.availability_error for r in available),
                    'unavailableResults': sum(bool(r.availability_error) for r in available),
                    'salvagedCount': sum(r.salvaged for r in available),
                    'offset': offset, 'limit': limit,'''.lstrip())
a = s.index('    def recover(self, identity):')
b = s.index('    def cancel(',a)
part=s[a:b]
part=replace(part, '        with self.lock:\n            ticket = self.results_by_request.get(identity)',
'''        with self.lock:
            selected = self.results_by_request.get(identity)
        self._refresh_result_availability({selected})
        with self.lock:
            ticket = self.results_by_request.get(identity)''')
part=replace(part, "            if result is not None and not result.released:\n                return {'state': 'ready',", 
'''            if result is not None and not result.released:
                if result.availability_error:
                    return {'state': 'unavailable', 'requestId': identity,
                            'code': result.availability_error}
                return {'state': 'ready',''')
s=s[:a]+part+s[b:]
s=replace(s, '        try:\n            yield result\n        finally:', '''        try:
            self._refresh_result_availability({ticket})
            if result.availability_error:
                raise LibreOfficeError('Файл результата отсутствует или изменён. Обновите список.',
                                       'RESULT_NOT_FOUND')
            yield result
        finally:''')
s=replace(s, '            self.closed = True\n', '            self.closed = True\n            self._shutdown.set()\n')
changes[p]=s

p='app_backend/legacy_storage_budget.py'
s=source(p)
s=method(s,None,'measured_bytes',ASSETS['measured_bytes.txt'])
s=replace(s, '            self.next_scan = now + .5\n            self.check()',
               '            self.check()\n            self.next_scan = time.monotonic() + .5')
changes[p]=s

p='app_backend/converter_guard/core.py'
s=source(p)
s=replace(s, '    save_receipt: dict | None = None\n', '''    save_receipt: dict | None = None
    availability_error: str | None = None
    file_identity: tuple | None = field(default=None, init=False)

    def __post_init__(self):
        from ..file_identity import file_identity
        self.file_identity = file_identity(self.path)
''')
s=replace(s, "return {'supportingFile': self.companion, 'fileId': self.file_id,", 
'''return {'available': self.availability_error is None, 'availabilityCode': self.availability_error,
                'supportingFile': self.companion, 'fileId': self.file_id,''')
s=replace(s, "                'ticket': self.ticket,", "                'ticket': None if self.availability_error else self.ticket,")
s=replace(s, "'downloadUrl': '/api/converter-guard/result?ticket=' + self.ticket,", 
               "'downloadUrl': None if self.availability_error else '/api/converter-guard/result?ticket=' + self.ticket,")
s=method(s,'ConverterGuard','verify_saved_download',ASSETS['verify_saved_download.txt'])
s=method(s,'ConverterGuard','_owned_artifact_path',ASSETS['owned_artifact_path.txt'])
pos=s.index('    def status(self, job_id,')
s=s[:pos]+ASSETS['guard_reconcile.txt'].rstrip()+'\n\n'+s[pos:]
a=s.index('    def status(self, job_id,');b=s.index('    def retry_plan(',a)
part=s[a:b]
part=replace(part, '        self._prune(defer_cleanup=not recover)\n',
'''        self._prune(defer_cleanup=not recover)
        if recover and since_revision is None:
            self._reconcile_results(job_id)
''')
s=s[:a]+part+s[b:]
a=s.index('    def recoverable(');b=s.index('    def _cached_proof(',a)
part=s[a:b]
part=replace(part, '        if scan:\n            self._load_recovery()', '''        if scan:
            self._load_recovery()
            with self._lock:
                identities = [job.id for job in self._jobs.values() if not job.released]
            for identity in identities:
                try:
                    self._reconcile_results(identity)
                except GuardError as error:
                    if error.code not in ('JOB_NOT_FOUND', 'GUARD_CLOSED'):
                        raise''')
part=replace(part, "'state': j.state, 'running': j.running,", "'state': self._availability_state(j), 'conversionState': j.state, 'running': j.running,")
part=replace(part, "'outputs': len(j.artifacts), 'bytes': sum(a.size for a in j.artifacts),", 
''' 'outputs': sum(not a.availability_error for a in j.artifacts),
                              'knownOutputs': len(j.artifacts),
                              'unavailableOutputs': sum(bool(a.availability_error) for a in j.artifacts),
                              'bytes': sum(a.size for a in j.artifacts if not a.availability_error),'''.lstrip())
part=part.replace('sum(not a.saved', 'sum(not a.saved and not a.availability_error')
part=replace(part, 'sum(a.size for a in j.artifacts if not a.saved)',
                  'sum(a.size for a in j.artifacts if not a.saved and not a.availability_error)')
s=s[:a]+part+s[b:]
a=s.index('    def _snapshot(');b=s.index('    def _reconcile_results(',a)
part=s[a:b]
part=replace(part, '        result = copy.deepcopy(job.result)', '''        available = [a for a in job.artifacts if not a.availability_error]
        result = copy.deepcopy(job.result)
        if result is not None:
            anchor = self._tickets.get(result.get('ticket'))
            if anchor is not None and anchor.availability_error:
                result = ({**available[0].metadata(), 'jobId': job.id, 'folderOnly': True}
                          if available else None)''')
part=replace(part, "result['outputs'] = [a.metadata() for a in artifacts]", "result['outputs'] = [a.metadata() for a in artifacts if not a.availability_error]")
part=replace(part, "'state': job.state, 'finished': finished,", 
''' 'state': self._availability_state(job), 'conversionState': job.state, 'finished': finished,
                'unavailableOutputs': sum(bool(a.availability_error) for a in job.artifacts),'''.lstrip())
part=part.replace('sum(not a.saved', 'sum(not a.saved and not a.availability_error')
part=replace(part, "'resultAvailability': ('recovered_unverified' if job.salvage else", 
''' 'resultAvailability': ('unavailable' if job.artifacts and not available and
                                        (job.bundle is None or job.bundle.availability_error) else
                                       'partial_unavailable' if any(a.availability_error for a in job.artifacts) else
                                       'recovered_unverified' if job.salvage else'''.lstrip())
part=replace(part, "'partialResult': ({**job.artifacts[0].metadata(), 'jobId': job.id,\n                    'outputsCount': len(job.artifacts), 'folderOnly': True}\n                    if job.artifacts else None)",
                  "'partialResult': ({**available[0].metadata(), 'jobId': job.id,\n                    'outputsCount': len(available), 'folderOnly': True}\n                    if available else None)")
s=s[:a]+part+s[b:]
s=replace(s, "'saved':a.saved, 'readbackVerified':a.readback_verified,", 
               "'saved':a.saved, 'available':not a.availability_error,\n                                       'availabilityCode':a.availability_error, 'readbackVerified':a.readback_verified,")
changes[p]=s

p='app_backend/excel.py'
s=source(p)
s=replace(s, 'from openpyxl.utils.exceptions import InvalidFileException', 
'''from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.worksheet.page import PageMargins''')
s=replace(s, '        if ws.max_row > MAX_TOTAL_ROWS:', '''        if any(d.outlineLevel or d.collapsed for d in
               (*ws.row_dimensions.values(), *ws.column_dimensions.values())):
            raise MergeError(f'«{name}»: группировка строк или столбцов пока не поддерживается.')
        if ws.page_margins != PageMargins():
            raise MergeError(f'«{name}»: нестандартные поля печати пока не поддерживаются.')
        if ws.sheet_view.rightToLeft is True or ws.sheet_view.showGridLines is False:
            raise MergeError(f'«{name}»: направление справа налево или скрытая сетка пока не поддерживаются.')
        if ws.max_row > MAX_TOTAL_ROWS:''')
s=replace(s, '''        widths = {key: dim.width for key, dim in ws.column_dimensions.items()
                  if dim.width is not None}''', '''        widths = {}
        for key, dim in ws.column_dimensions.items():
            if dim.width is None:
                continue
            first = dim.min or column_index_from_string(key)
            last = dim.max or first
            if not 1 <= first <= last <= 16384:
                raise MergeError(f'«{name}»: недопустимый диапазон ширины столбцов.')
            for index in range(first, last + 1):
                column = get_column_letter(index)
                if column in widths and widths[column] != dim.width:
                    raise MergeError(f'«{name}»: пересекающиеся настройки ширины столбцов.')
                widths[column] = dim.width''')
changes[p]=s

p='tests/test_m02_m30_recovery_save.py'
s=source(p)
s=replace(s, "        guard.verify_saved_download(job.bundle.ticket, path)\n        path.write_bytes(b'broken zip')", 
'''        first_receipt = guard.verify_saved_download(job.bundle.ticket, path)
        first_path = path.parent / first_receipt['receiptFile']
        first_bytes = first_path.read_bytes()
        path.write_bytes(b'broken zip')''')
s=replace(s, "        status = json.loads(path.with_name(path.name+'.AIO-save-status.json').read_text())\n        self.assertEqual(status['state'], 'failed')", 
'''        self.assertEqual(first_path.read_bytes(), first_bytes)
        attempts = [json.loads(p.read_text()) for p in path.parent.glob(path.name+'.AIO-save-status-*.json')]
        status = next(item for item in attempts if item['state'] == 'failed')
        self.assertNotEqual(status['saveId'], first_receipt['saveId'])
        self.assertEqual(status['state'], 'failed')''')
changes[p]=s

p='src/legacy-recovery.js'
s=source(p)
s=replace(s, 'button.disabled=pending||loading;\n      get', "button.disabled=pending||loading||button.dataset.unavailable==='true';\n      get")
s=replace(s, "(item.salvagedUnverified===true?' · НЕПРОВЕРЕННЫЕ файлы':item.partial===true?' · есть ошибки преобразования':'');", 
               "(item.available===false?' · ФАЙЛ НЕДОСТУПЕН':item.salvagedUnverified===true?' · НЕПРОВЕРЕННЫЕ файлы':item.partial===true?' · есть ошибки преобразования':'');")
s=replace(s, "            button.setAttribute('data-l-action',discard?'discard':'save');", 
'''            button.setAttribute('data-l-action',discard?'discard':'save');
            if(item.available===false){button.dataset.unavailable='true';button.disabled=true;}''')
s=replace(s, "get('[data-l-decline-all]').hidden=!(packet.restoredCount>0||", "get('[data-l-decline-all]').hidden=!(packet.restoredCount>0||packet.salvagedCount>0||")
s=replace(s, "'Найдено результатов: '+total+'. Целостность восстановленных файлов проверена; качество конвертации не подтверждено.'", 
               "'Найдено записей: '+total+'. Недоступных файлов: '+(packet.unavailableResults||0)+'. Качество конвертации не подтверждено.'")
changes[p]=s

for path, value in changes.items():
    if path.endswith('.py'):
        ast.parse(value, filename=path)
    current = Path(path).read_text(encoding='utf-8-sig')
    digest = hashlib.sha256(current.encode()).hexdigest()
    if current not in (originals[path], value) and digest not in PREVIOUS.get(path, ()):
        raise RuntimeError('Refusing an unexpected concurrent edit: ' + path)
for path, value in changes.items():
    Path(path).write_text(value, encoding='utf-8', newline='\n')
    print('SOURCE_FIXED', path, hashlib.sha256(value.encode()).hexdigest())
