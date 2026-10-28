"""Native Windows profile, all-descendant drain and concurrent-session QA."""
import argparse
import ctypes
import hashlib
import json
import os
from pathlib import Path
import sys
import threading
import time

ROOT=Path(__file__).resolve().parents[2];sys.path[:0]=[str(ROOT),str(ROOT/'tests')]
from openpyxl import Workbook,load_workbook
from app_backend.converter_guard import ConverterGuard,GuardLimits,LibreOfficeAdapter


def sha(path):
    with path.open('rb') as stream:return hashlib.file_digest(stream,'sha256').hexdigest()


def fixture(path):
    book=Workbook();book.remove(book.active)
    for index in range(4):
        sheet=book.create_sheet('Лист '+str(index+1));sheet.append(['ID','Long ID','Value','Double','Next','Note','Sheet','Tail'])
        for row in range(1,301):sheet.append([f'{row:03d}','12345678901234567890',row+index*1000,
            f'=C{row+1}*2',f'=D{row+1}+1','строка\nдва' if row==1 else 'Кириллица',index,f'END-{index}-{row}'])
    book.save(path);book.close()


def verify(path):
    book=load_workbook(path,read_only=True,data_only=True,keep_links=False)
    try:
        assert len(book.sheetnames)==4
        for index,sheet in enumerate(book):
            rows=list(sheet.values);assert len(rows)==301 and all(len(r)==8 for r in rows)
            assert rows[0]==('ID','Long ID','Value','Double','Next','Note','Sheet','Tail')
            for row,v in enumerate(rows[1:],1):
                assert isinstance(v[0],str) and v[0]==f'{row:03d}'
                assert isinstance(v[1],str) and v[1]=='12345678901234567890'
                value=row+index*1000;assert v[2:5]==(value,value*2,value*2+1)
                assert v[5]==('строка\nдва' if row==1 else 'Кириллица') and v[6]==index and v[7]==f'END-{index}-{row}'
    finally:book.close()


def own_process_handles(tree):
    from ctypes import wintypes as w
    class Pids(ctypes.Structure):
        _fields_=[('assigned',ctypes.c_uint32),('count',ctypes.c_uint32),('ids',ctypes.c_size_t*128)]
    pids=Pids()
    assert tree._kernel.QueryInformationJobObject(tree._job,3,ctypes.byref(pids),ctypes.sizeof(pids),None)
    assert 2<=pids.count<=128,'Expected real UNO worker plus native office descendants'
    kernel=ctypes.WinDLL('kernel32',use_last_error=True)
    kernel.OpenProcess.argtypes=[w.DWORD,w.BOOL,w.DWORD];kernel.OpenProcess.restype=w.HANDLE
    kernel.IsProcessInJob.argtypes=[w.HANDLE,w.HANDLE,ctypes.POINTER(w.BOOL)];kernel.IsProcessInJob.restype=w.BOOL
    kernel.WaitForSingleObject.argtypes=[w.HANDLE,w.DWORD];kernel.WaitForSingleObject.restype=w.DWORD
    kernel.CloseHandle.argtypes=[w.HANDLE];kernel.CloseHandle.restype=w.BOOL
    handles=[]
    try:
        for pid in pids.ids[:pids.count]:
            handle=kernel.OpenProcess(0x100000|0x1000,False,pid);assert handle
            handles.append((int(pid),handle));owned=w.BOOL()
            assert kernel.IsProcessInJob(handle,tree._job,ctypes.byref(owned)) and owned.value
            assert kernel.WaitForSingleObject(handle,0)==258
        return kernel,handles
    except BaseException:
        for _,handle in handles:kernel.CloseHandle(handle)
        raise


def main():
    p=argparse.ArgumentParser();p.add_argument('--output',type=Path,required=True);p.add_argument('--executable',type=Path,required=True);args=p.parse_args()
    assert os.name=='nt','Actual Windows native acceptance required'
    args.output.mkdir(parents=True,exist_ok=False);source=args.output/'source.xlsx';fixture(source);before=sha(source)
    report={'schema':1,'status':'RUNNING','scope':'actual native Windows sessions, not FULLTEST or installed EXE',
            'sourceSHA':os.environ.get('AIO_BUILD_REVISION'),'sourceSha256':before,'sourceSHA256':{source.name:before},'checks':[]}
    handles=[];kernel=None
    try:
        cancel=threading.Event();adapter=LibreOfficeAdapter(args.executable);neighbor=LibreOfficeAdapter(args.executable)
        base=args.output/'deep';work=base/('w'*max(1,175-len(str(base))-1))
        with neighbor.batch(args.output/'neighbor') as other:
            other_proof=neighbor.preflight(source,args.output/'neighbor-inspect',cancel=cancel,timeout=90)
            assert other_proof.get('state')=='detected' and not other.disabled
            other_pid=other.tree.process.pid
            with adapter.batch(work) as office:
                proof=adapter.preflight(source,args.output/'deep-inspect',cancel=cancel,timeout=90)
                assert proof.get('state')=='detected' and not office.disabled,proof
                assert proof.get('nativeImportMs') is not None,'Fallback-only preflight is not a warm UNO proof'
                assert proof.get('documentServices')==['com.sun.star.sheet.SpreadsheetDocument']
                assert office._profile_owner is not None
                profile=office.profile;marker=office._profile_owner.marker;tree=office.tree
                assert not profile.is_relative_to(work)
                assert profile.is_dir() and marker.is_file()
                adapter.prepare_source(proof)
                output=next(path for path in adapter.convert(source,'xlsx',args.output/'result',cancel=cancel,timeout=180,on_started=lambda:None) if path.suffix=='.xlsx')
                assert adapter.execution_mode()=='native_uno';verify(output)
                kernel,handles=own_process_handles(tree)
                report['checks'].append({'name':'deep-workspace-warm-native-roundtrip','status':'PASS',
                    'workspaceLength':len(str(work)),'profileLength':len(str(profile)),'fallbackUsed':office.disabled,
                    'nativeImportFilter':proof.get('selectedImportFilter'),'exportFilter':adapter.export_info().get('filter'),
                    'executionMode':adapter.execution_mode(),'sheets':4,'independentlyCheckedCells':4*301*8,
                    'outputSha256':sha(output),'qualityCompared':False,'fullVisualFidelity':'not_tested'})
            exited=[kernel.WaitForSingleObject(handle,0)==0 for _,handle in handles]
            assert all(exited) and tree.job_drain_verified
            assert not profile.exists() and not marker.exists()
            assert other.tree.process.poll() is None and other.tree.process.pid==other_pid
            again=neighbor.preflight(source,args.output/'neighbor-still-live',cancel=cancel,timeout=90)
            assert again.get('state')=='detected' and not other.disabled
            report['checks'].append({'name':'all-owned-descendants-and-private-profile-released','status':'PASS',
                'ownedPids':[pid for pid,_ in handles],'independentHandleExitStates':exited,'jobDrainVerified':True,
                'profileRemoved':True,'ownerMarkerRemoved':True})
            report['checks'].append({'name':'concurrent-native-session-survives','status':'PASS','survivingWorkerPid':other_pid,
                'reimportFilter':again.get('selectedImportFilter'),'otherSourceBytesUnchanged':sha(source)==before})
        # Change only our external fixture during an observed real native
        # import. The job must use its uploaded owned snapshot, not reopen the
        # now-corrupt external original. No importer/engine is replaced.
        external=args.output/'external-changing.xlsx';external.write_bytes(source.read_bytes());changed=threading.Event()
        class ObservedAdapter(LibreOfficeAdapter):
            def convert(self,*values,**options):
                callback=options.get('on_stage')
                def stage(value):
                    if value=='import' and not changed.is_set():
                        external.write_bytes(b'QA deliberate external modification during native import');changed.set()
                    if callback:callback(value)
                options['on_stage']=stage
                return super().convert(*values,**options)
        guard=ConverterGuard(adapter=ObservedAdapter(args.executable),recovery_parent=args.output/'snapshot-owned',limits=GuardLimits(file_timeout=180))
        try:
            job=guard.create([{'id':'race','name':external.name,'size':external.stat().st_size}],prepare_only=True,retain_sources=True)['jobId']
            with external.open('rb') as stream:guard.upload(job,'race',stream,external.stat().st_size)
            guard.preflight(job,timeout_seconds=120)
            guard.start(job,formats=['xlsx'],file_ids=['race'],only_verified=True,archive_mode='on_demand')
            until=time.monotonic()+180
            while not guard.status(job)['finished']:
                if time.monotonic()>until:raise TimeoutError('Native snapshot conversion did not finish')
                time.sleep(.02)
            state=guard.status(job);assert changed.is_set(),'Actual native import stage was not observed'
            assert state['succeeded']==1 and state['failed']==0,guard.report(job)
            assert sha(guard._jobs[job].files[0].source)==before
            selected=next(row for row in state['result']['outputs'] if row['name'].endswith('.xlsx'))
            saved=args.output/'independently-saved-snapshot.xlsx'
            with guard.open_result(selected['ticket']) as (stream,_,__):saved.write_bytes(stream.read())
            verify(saved);guard.verify_saved_download(selected['ticket'],saved);guard.confirm_saved(job,selected['ticket'])
            saved_sha=sha(saved);external_sha=sha(external);assert external_sha!=before
        finally:guard.close()
        assert sha(saved)==saved_sha and sha(external)==external_sha
        report['sourceSHA256'][external.name]=external_sha
        report['checks'].append({'name':'external-source-mutation-during-native-import','status':'PASS',
            'expectedDecision':'ready_from_uploaded_snapshot','observedNativePhase':'import','externalModifiedBy':'QA_fixture',
            'externalShaBefore':before,'externalShaAfter':external_sha,'ownedSnapshotSha256':before,
            'savedOutputSha256':saved_sha,'independentReadback':True,'externalAndSavedFilesRetained':True})
        assert sha(source)==before;report['status']='PASS'
    except Exception as error:
        report.update(status='FAIL',error=type(error).__name__,reason=str(error));raise
    finally:
        for _,handle in handles:kernel.CloseHandle(handle)
        (args.output/'evidence.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
        print('NATIVE_PROFILE_SUMMARY='+json.dumps(report,ensure_ascii=True),flush=True)


if __name__=='__main__':main()
