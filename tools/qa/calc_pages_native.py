"""Actual AIO Calc -> PNG/JPEG print-layout acceptance, not historical FULLTEST.

Independent coloured row bands detect repeated pages and clipped lower content;
counting files or trusting SourceRange alone cannot establish those properties.
"""
import hashlib
import json
from pathlib import Path
import threading
import time


_BANDS = ('red', 'blue', 'gold')
_BAND_COLOURS = ('FF4040', '4040FF', 'FFC840')


def _digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def _marker_counts(pixels):
    """Coarse, JPEG-tolerant test-fixture oracle; never used by the product."""
    counts = dict.fromkeys((*_BANDS, 'green', 'magenta', 'cyan'), 0)
    for r, g, b in pixels:
        if r > 180 and g < 125 and b < 125:
            counts['red'] += 1
        elif b > 180 and r < 125 and g < 125:
            counts['blue'] += 1
        elif r > 180 and g > 140 and b < 125:
            counts['gold'] += 1
        elif g > 180 and r < 125 and b < 125:
            counts['green'] += 1
        elif r > 180 and b > 180 and g < 125:
            counts['magenta'] += 1
        elif g > 180 and b > 180 and r < 125:
            counts['cyan'] += 1
    return counts


def _verify_markers(counts, expected):
    """Require visible source bands; unexpected hidden/outside regions fail."""
    expected = set(expected)
    missing = sorted(key for key in expected if counts.get(key, 0) < 10)
    extra = sorted(key for key, value in counts.items() if key not in expected and value >= 10)
    if missing or extra:
        raise AssertionError({'missingPageBands': missing, 'unexpectedPageBands': extra,
                              'observedMarkerPixels': counts})


def _fixture(path, mode):
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill
    from openpyxl.worksheet.pagebreak import Break
    from openpyxl.worksheet.properties import PageSetupProperties
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Portrait'
    second = workbook.create_sheet('Landscape')
    hidden = workbook.create_sheet('Hidden control')
    hidden.sheet_state = 'hidden'
    # Visibly distinctive controls, excluded by explicit print settings.
    for row in range(1, 11):
        for col in (1, 2):
            cell = hidden.cell(row, col, 'HIDDEN_CONTROL')
            cell.fill = PatternFill('solid', fgColor='FF40FF')
    for current, count in ((sheet, 60), (second, 10)):
        current.column_dimensions['A'].width = 24
        current.column_dimensions['B'].width = 15
        for row in range(1, count + 1):
            current.cell(row, 1, '%s_ROW_%03d' % (current.title, row))
            current.cell(row, 2, row)
            current.row_dimensions[row].height = 16
            colour = _BAND_COLOURS[(row - 1) // 20] if current is sheet else '40FF40'
            for cell in current[row][:2]:
                cell.fill = PatternFill('solid', fgColor=colour)
        current.print_area = 'A1:B%d' % count
        for row in range(200, 210):
            for col in (1, 2):
                cell = current.cell(row, col, 'OUTSIDE_PRINT_AREA')
                cell.fill = PatternFill('solid', fgColor='40FFFF')
        current.page_setup.paperSize = current.PAPERSIZE_A4
        current.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=False)
        current.page_setup.scale = 100
    second.page_setup.orientation = 'landscape'
    if mode == 'breaks':
        sheet.row_breaks.append(Break(id=20))
        sheet.row_breaks.append(Break(id=40))
    elif mode == 'scale':
        sheet.page_setup.scale = 50
    elif mode == 'fit':
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.scale = None
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
    else:
        raise ValueError(mode)
    workbook.save(path)
    workbook.close()


def run(executable, root):
    from PIL import Image
    from app_backend.converter_guard import LibreOfficeAdapter
    root = Path(root)
    root.mkdir(parents=True, exist_ok=False)
    cancel = threading.Event()
    adapter = LibreOfficeAdapter(executable)
    results = []
    started = time.monotonic()
    current_case = None
    report = {'status': 'FAIL', 'scope': 'synthetic Calc print-layout controls, not original FULLTEST',
              'cases': results}
    try:
        for mode, expected in (('breaks', 4), ('scale', 2), ('fit', 2)):
            source = root / ('calc-' + mode + '.xlsx')
            _fixture(source, mode)
            original = _digest(source)
            with adapter.batch(root / ('office-' + mode)):
                proof = adapter.preflight(source, root / ('probe-' + mode), cancel=cancel, timeout=90)
                assert proof.get('state') == 'detected', proof
                targets = ('png','jpg','jpeg','jfif','jif','jpe') if mode == 'breaks' else ('png','jpg')
                for target in targets:
                    current_case = {'mode': mode, 'target': target}
                    adapter.prepare_source(proof)
                    tick = time.monotonic()
                    paths = adapter.convert(source, target, root / (mode + '-' + target),
                        cancel=cancel, timeout=120, on_started=lambda: None)
                    assert adapter.execution_mode() == 'native_uno', adapter.execution_mode()
                    info = adapter.export_info()
                    receipt = next(check for check in info.get('resultChecks', [])
                                   if check.get('policyId') == 'native_all_calc_print_pages_v1')
                    assert receipt['expectedUnits'] == receipt['observedOutputs'] == expected, receipt
                    assert receipt['qualityCompared'] is False
                    assert 'NATIVE_CALC_RASTER_PAGES' in info.get('warnings', ()), info
                    manifest = json.loads(next(p for p in paths if p.name == 'AIO-pages.json').read_text(encoding='utf-8'))
                    assert manifest['pageCount'] == len(manifest['pages']) == expected, manifest
                    assert manifest['scope'] == 'loaded_calc_print_layout'
                    assert manifest['respectsPrintSettings'] is True
                    assert [p['page'] for p in manifest['pages']] == list(range(1, expected+1))
                    grouped = {0: set(), 1: set()}
                    hashes, marker_evidence = [], []
                    for position, row in enumerate(manifest['pages']):
                        path = next(p for p in paths if p.name == row['file'])
                        index = row['sheetIndex']
                        assert index in grouped, row
                        first_col, first_row, last_col, last_row = row['sourceRange']
                        assert first_col == 0 and last_col == 1, row
                        assert 0 <= first_row <= last_row < (60 if index == 0 else 10), row
                        grouped[index].update(range(first_row, last_row+1))
                        # The colour expectation comes from the synthetic input,
                        # not the output's self-reported sheet/range metadata.
                        expected_bands = ({_BANDS[position]} if mode == 'breaks' and position < 3 else
                                          set(_BANDS) if mode != 'breaks' and position == 0 else {'green'})
                        with Image.open(path) as image:
                            assert image.format == ('PNG' if target == 'png' else 'JPEG'), image.format
                            assert list(image.size) == row['pixelSize'], (image.size,row)
                            width, height = image.size
                            assert (width < height) if index == 0 else (width > height), row
                            rgb = image.convert('RGB')
                            hashes.append(hashlib.sha256(rgb.tobytes()).hexdigest())
                            rgb.thumbnail((180, 180))
                            counts = _marker_counts(rgb.getdata())
                            _verify_markers(counts, expected_bands)
                            marker_evidence.append({'page': position + 1, 'expectedBands': sorted(expected_bands),
                                                    'markerPixels': counts})
                    assert grouped[0] == set(range(60)) and grouped[1] == set(range(10)), grouped
                    if mode == 'breaks':
                        assert len(set(hashes)) == expected, (mode,target,'repeated page content')
                    assert _digest(source) == original, 'source changed'
                    result = {**current_case, 'status': 'PASS', 'pages': expected, 'sourceUnchanged': True,
                              'pixelMarkers': marker_evidence, 'ms': round((time.monotonic()-tick)*1000)}
                    results.append(result)
                    print('CALC_PAGE_CASE='+json.dumps(result, ensure_ascii=True), flush=True)
        report['status'] = 'PASS'
        return report
    except Exception as error:
        report.update(failedCase=current_case, error=type(error).__name__, reason=str(error),
                      code=getattr(error, 'code', None), diagnostic=getattr(error, 'diagnostic', None))
        raise
    finally:
        report['elapsedSeconds'] = round(time.monotonic()-started, 3)
        (root/'evidence.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        print('CALC_PAGES_NATIVE='+json.dumps(report, ensure_ascii=True), flush=True)
