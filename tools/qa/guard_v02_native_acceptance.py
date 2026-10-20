"""Opt-in synthetic Guard QA; no UI integration, user corpus or Windows host.

QA-only: python-docx, python-pptx, Pillow, openpyxl, pypdf, optional Poppler.
python tools/qa/guard_v02_native_acceptance.py --work /new/empty/directory
--large-fixtures may reuse large-one.docx/large-two.docx from prior synthetic QA.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
from pathlib import Path
import random
import re
import shutil
import subprocess
import sys
import threading
import time
import tracemalloc
from urllib.parse import unquote
from urllib.request import Request, urlopen
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[2]
sys.path[:0] = [str(ROOT), str(ROOT / 'tests')]
from app_backend.converter_guard import ConverterGuard, LibreOfficeAdapter
from app_backend.libreoffice_engine import convert_path
from test_converter_guard_http_v02 import start_test_server


def digest(path):
    with path.open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def main():
    from PIL import Image, ImageChops
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.shared import Inches
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches as SlideInches
    from pypdf import PdfReader

    parser = argparse.ArgumentParser()
    parser.add_argument('--work', required=True, type=Path)
    parser.add_argument('--large-fixtures', type=Path)
    args = parser.parse_args()
    work = args.work.resolve()
    work.mkdir(parents=True, exist_ok=False)
    fixtures = work / 'fixtures'
    fixtures.mkdir()
    small = Image.frombytes('RGB', (96, 64), random.Random(7410).randbytes(96 * 64 * 3))
    small.save(fixtures / 'image.bmp')
    small.save(fixtures / 'image.png')
    document = Document()
    document.add_heading('GUARD_BEGIN_20261019', 1)
    paragraph = document.add_paragraph('FILENAME ')
    field = OxmlElement('w:fldSimple')
    field.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instr', ' FILENAME ')
    run, text = OxmlElement('w:r'), OxmlElement('w:t')
    text.text = 'writer.docx'
    run.append(text); field.append(run); paragraph._p.append(field)
    table = document.add_table(rows=2, cols=2)
    for cell, value in zip([c for row in table.rows for c in row.cells],
                           ['GUARD_TABLE_A', '00125', 'GUARD_TABLE_B', '42']):
        cell.text = value
    document.add_picture(str(fixtures / 'image.png'), width=Inches(1))
    math = OxmlElement('m:oMath')
    run, text = OxmlElement('m:r'), OxmlElement('m:t')
    text.text = '1+2'
    run.append(text); math.append(run)
    document.add_paragraph()._p.append(math)
    document.add_page_break()
    document.add_paragraph('GUARD_END_20261019')
    document.save(fixtures / 'writer.docx')
    book = Workbook()
    for sheet, marker in [(book.active, 'GUARD_SHEET_A'), (book.create_sheet('Second'), 'GUARD_SHEET_B')]:
        sheet.append([marker, 'Value']); sheet.append(['00125', 42])
        sheet.column_dimensions['A'].width = 30
    book.save(fixtures / 'calc.xlsx'); book.close()
    slides = Presentation()
    for marker in ('GUARD_SLIDE_A', 'GUARD_SLIDE_B'):
        slide = slides.slides.add_slide(slides.slide_layouts[6])
        slide.shapes.add_textbox(SlideInches(1), SlideInches(1), SlideInches(7), SlideInches(2)).text = marker
    slides.save(fixtures / 'impress.pptx')
    (fixtures / 'text.unlisted').write_text('GUARD_UNLISTED_20261019\n', encoding='utf-8')
    with ZipFile(fixtures / 'bad.docx', 'w') as archive:
        archive.writestr('[Content_Types].xml', '<bad broken="')
        archive.writestr('word/document.xml', '<not-xml')
    if args.large_fixtures:
        for name in ('large-one.docx', 'large-two.docx'):
            shutil.copyfile(args.large_fixtures / name, fixtures / name)
        large_markers = ['LARGE_BEGIN_7411', 'LARGE_END_7411']
    else:
        large_image = Image.frombytes('RGB', (3600, 2600), random.Random(7411).randbytes(3600*2600*3))
        large_image.save(fixtures / 'large-image.png')
        large = Document()
        large.add_paragraph('LARGE_BEGIN_7411')
        large.add_picture(str(fixtures / 'large-image.png'), width=Inches(2))
        large.add_paragraph('LARGE_END_7411')
        large.save(fixtures / 'large-one.docx')
        shutil.copyfile(fixtures / 'large-one.docx', fixtures / 'large-two.docx')
        large_markers = ['LARGE_BEGIN_7411', 'LARGE_END_7411']
    big = [fixtures / 'large-one.docx', fixtures / 'large-two.docx']
    assert min(p.stat().st_size for p in big) > 20*1024**2
    assert sum(p.stat().st_size for p in big) > 50*1024**2
    originals = {p: digest(p) for p in fixtures.iterdir() if p.is_file()}
    expected_digests = {p.name: value for p, value in originals.items()}

    class ObservedAdapter(LibreOfficeAdapter):
        calls = 0
        cancel_request = None
        def convert(self, source, *args, **kwargs):
            assert digest(source) == expected_digests[source.name], 'Staged input bytes changed'
            self.calls += 1
            request = self.cancel_request
            if request is not None and request[1] == source.name:
                original = kwargs['on_started']
                def skip_started_file():
                    original()
                    self.cancel_request = None
                    guard.cancel_file(request[0], '0')
                kwargs['on_started'] = skip_started_file
            return super().convert(source, *args, **kwargs)

    adapter = ObservedAdapter()
    if not adapter.available:
        raise SystemExit('NOT_RUN: installed LibreOffice required')
    spool = work / 'spool'
    spool.mkdir()
    guard = ConverterGuard(adapter=adapter, spool_parent=spool)
    evidence = {'scope': 'Synthetic Linux/source Guard v0.2; no UI or Windows integration',
                'engine': subprocess.run([str(adapter.executable), '--version'], capture_output=True,
                    text=True, timeout=15, check=True).stdout.strip(), 'cases': [],
                'nativeInputBytesUnchangedAtEveryLaunch': True}

    def run_case(label, sources, formats, *, successes=None, measure=False, skip_first=False):
        rows = [{'id':str(i), 'name':p.name, 'size':p.stat().st_size} for i,p in enumerate(sources)]
        started = time.monotonic()
        if measure: tracemalloc.start()
        try:
            identity = guard.create(rows, list(formats))['jobId']
            if skip_first:
                adapter.cancel_request = (identity, sources[0].name)
            for i,source in enumerate(sources):
                with source.open('rb') as stream:
                    guard.upload(identity, str(i), stream, source.stat().st_size)
            guard.start(identity)
            deadline = time.monotonic() + 120
            while time.monotonic() < deadline:
                state = guard.status(identity)
                if state['finished']: break
                time.sleep(.025)
            else:
                raise AssertionError('Synthetic QA deadline reached')
            if measure:
                _, peak = tracemalloc.get_traced_memory()
        finally:
            if measure: tracemalloc.stop()
        tasks = [t for f in state['files'] for t in f['targets']]
        passed = sum(t['state'] == 'ready' for t in tasks)
        if successes is not None:
            assert passed == successes, state
        assert all(t['libreOfficeStarted'] for t in tasks), state
        row = {'label':label, 'formats':list(formats),
               'sources':[{'name':p.name,'bytes':p.stat().st_size} for p in sources],
               'seconds':round(time.monotonic()-started,3),
               'succeeded':passed, 'failed':len(tasks)-passed,
               'states':[{k:f[k] for k in ('fileId','name','state','targets')} for f in state['files']]}
        if measure:
            row['pythonTracedPeakBytes'] = peak
            assert peak < 8*1024**2, row
        saved = []
        for output in state['result']['outputs']:
            target = work / 'results' / label / output['archiveName']
            target.parent.mkdir(parents=True, exist_ok=True)
            key = output['downloadUrl'].split('ticket=')[1]
            with guard.open_result(key) as (stream,_,_), target.open('wb') as destination:
                shutil.copyfileobj(stream,destination,64*1024)
            saved.append(target)
        result_key = state['result']['downloadUrl'].split('ticket=')[1]
        if state['result']['name'].endswith('.zip'):
            target = work / 'results' / label / 'AIO_Results.zip'
            with guard.open_result(result_key) as (stream,_,_), target.open('wb') as destination:
                shutil.copyfileobj(stream,destination,64*1024)
            with ZipFile(target) as archive:
                for output,path in zip(state['result']['outputs'],saved):
                    with archive.open(output['archiveName']) as stream:
                        assert hashlib.file_digest(stream,'sha256').hexdigest() == digest(path)
        guard.confirm_saved(identity,result_key)
        guard.release(identity)
        assert guard.storage_stats()['jobs'] == 0 and guard.storage_stats()['reservedBytes'] == 0
        # Bearer tickets are transient and must not be committed in QA evidence.
        for file in row['states']:
            for task in file['targets']:
                for output in task['outputs']:
                    output.pop('downloadUrl',None)
        evidence['cases'].append(row)
        print(json.dumps({'case':label,'succeeded':passed,'failed':len(tasks)-passed,
                          'seconds':row['seconds']},ensure_ascii=False),flush=True)
        return saved,row

    try:
        pdfs,_ = run_case('bad-first-mixed-inputs',
            [fixtures/name for name in ('bad.docx','writer.docx','calc.xlsx','impress.pptx','image.bmp')],
            ['pdf'],successes=4)
        writer_pdf = next(p for p in pdfs if p.name == 'writer.pdf')
        for name,markers in [('writer',['GUARD_BEGIN_20261019','GUARD_END_20261019','GUARD_TABLE_A','writer.docx','1+2']),
                             ('calc',['GUARD_SHEET_A','GUARD_SHEET_B']),
                             ('impress',['GUARD_SLIDE_A','GUARD_SLIDE_B'])]:
            path = next(p for p in pdfs if p.name == name+'.pdf')
            text = ''.join(''.join((p.extract_text() or '').split()) for p in PdfReader(path).pages)
            assert all(marker in text for marker in markers), (name,text)
        run_case('bmp-multiple-targets',[fixtures/'image.bmp'],['png','pdf','tiff','svg','webp'],successes=5)
        outputs,_ = run_case('writer-native-targets',[fixtures/'writer.docx'],
                             ['docx','odt','html','txt','pdf'],successes=5)
        html = next(p for p in outputs if p.suffix == '.html')
        for src in re.findall(r'<img[^>]+src="([^"]+)"',html.read_text(),flags=re.I):
            if not src.startswith('data:'):
                assert (html.parent/unquote(src)).is_file(),src
        run_case('calc-native-targets',[fixtures/'calc.xlsx'],['xlsx','ods','csv','pdf'],successes=4)
        run_case('impress-native-targets',[fixtures/'impress.pptx'],['pptx','odp','pdf'],successes=3)
        run_case('unlisted-input-extension',[fixtures/'text.unlisted'],['docx','pdf'],successes=2)
        expected_digests[writer_pdf.name] = digest(writer_pdf)
        run_case('pdf-native-import',[writer_pdf],['png'],successes=1)
        large_pdfs,large_row = run_case('over-50-mib-batch',big,['pdf'],successes=2,measure=True)
        for path in large_pdfs:
            text=''.join(p.extract_text() or '' for p in PdfReader(path).pages)
            assert all(marker in text for marker in large_markers)
        _, skipped = run_case('skip-running-native-file',
            [fixtures/'writer.docx',fixtures/'calc.xlsx'],['pdf'],successes=1,skip_first=True)
        assert skipped['states'][0]['targets'][0]['code'] == 'CANCELLED'
        assert skipped['states'][1]['state'] == 'ready'

        direct = convert_path(fixtures/'writer.docx','pdf',work/'direct-reference',
                              executable=adapter.executable)[0]
        first,second = PdfReader(writer_pdf),PdfReader(direct)
        assert len(first.pages) == len(second.pages) == 2
        assert [p.extract_text() for p in first.pages] == [p.extract_text() for p in second.pages]
        def image_hashes(reader):
            return [[hashlib.sha256(image.data).hexdigest() for image in page.images] for page in reader.pages]
        assert image_hashes(first) == image_hashes(second)
        baseline = {'pages':2,'textEqual':True,'encodedImagesEqual':True,'renderedPixelComparison':'NOT_RUN'}
        renderer=shutil.which('pdftoppm')
        if renderer:
            for i,path in enumerate((writer_pdf,direct)):
                subprocess.run([renderer,'-r','72','-png',str(path),str(work/('render-'+str(i)))],
                               check=True,capture_output=True,timeout=30)
            for page in (1,2):
                with Image.open(work/f'render-0-{page}.png') as a,Image.open(work/f'render-1-{page}.png') as b:
                    assert a.size == b.size and ImageChops.difference(a,b).getbbox() is None
            baseline['renderedPixelComparison']='PASS: both pages, 72 DPI, identical pixels'
        evidence['directNativePDFBaseline'] = baseline

        server,thread,origin,token=start_test_server(guard)
        try:
            def http(path,body):
                data=json.dumps(body).encode('utf-8')
                with urlopen(Request(origin+'/api/converter-guard/'+path,data=data,
                        headers={'X-AIO-Token':token,'Content-Type':'application/json'}),timeout=15) as response:
                    return json.load(response)
            source=fixtures/'image.bmp'
            identity=http('jobs',{'files':[{'id':'bmp','name':source.name,'size':source.stat().st_size}],
                                  'formats':['pdf']})['jobId']
            with source.open('rb') as stream:
                with urlopen(Request(origin+'/api/converter-guard/input?job='+identity+'&file=bmp',
                        data=stream,headers={'X-AIO-Token':token,'Content-Type':'application/octet-stream',
                                             'Content-Length':str(source.stat().st_size)}),timeout=15) as response:
                    assert json.load(response)['ok']
            http('start',{'jobId':identity})
            deadline=time.monotonic()+30
            while time.monotonic()<deadline:
                with urlopen(Request(origin+'/api/converter-guard/jobs?job='+identity,
                        headers={'X-AIO-Token':token}),timeout=15) as response:
                    state=json.load(response)
                if state['finished']:break
                time.sleep(.025)
            assert state['state']=='ready',state
            with urlopen(origin+state['result']['downloadUrl'],timeout=15) as response:
                destination=work/'native-http-result.pdf'
                with destination.open('wb') as output:shutil.copyfileobj(response,output,64*1024)
            assert len(PdfReader(destination).pages)==1
            key=state['result']['downloadUrl'].split('ticket=')[1]
            assert http('saved',{'jobId':identity,'ticket':key})['files'][0]['state']=='saved'
            http('release',{'jobId':identity})
            evidence['nativeHTTP']='PASS: BMP → native PDF → HTTP stream → file → save acknowledgment → release'
        finally:
            server.api.close()
            server.shutdown();server.server_close();thread.join(3)
        assert all(digest(path)==expected for path,expected in originals.items())
        assert guard.storage_stats()=={'jobs':0,'reservedBytes':0,'tickets':0,'running':0,'queued':0}
        evidence.update(status='PASS',sourcesUnchanged=True,
            nativeAttempts=adapter.calls,largeBatchBytes=sum(p.stat().st_size for p in big),
            pythonTracedPeakLargeBatchBytes=large_row['pythonTracedPeakBytes'],
            notRun=['Windows','NIKITA','installed EXE','GUI integration','full user corpus'])
        (work/'evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2))
        print(json.dumps({'status':'PASS','cases':len(evidence['cases']),'nativeAttempts':adapter.calls,
                          'largeBatchBytes':evidence['largeBatchBytes'],
                          'pythonTracedPeakLargeBatchBytes':evidence['pythonTracedPeakLargeBatchBytes'],
                          'nativeHTTP':'PASS','baseline':baseline},ensure_ascii=False),flush=True)
    finally:
        guard.close()


if __name__=='__main__':
    main()
