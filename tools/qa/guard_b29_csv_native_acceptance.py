"""B29 native Windows acceptance on original CSV/TAB/TSV stress2 bytes.

No user FULLTEST or installation. Every workbook is read independently of
Guard's new contract. Old-import controls use the same native engine with
its original four-token CSV options, never the mandatory final-result API.
"""
import argparse
import copy
import hashlib
import json
from pathlib import Path
import sys
import threading
import time

ROOT = Path(__file__).resolve().parents[2]
sys.path[:0] = [str(ROOT), str(ROOT / 'tests')]
from b29_csv_fixtures import original_csv, SHA256
from app_backend.converter_guard import LibreOfficeAdapter
from app_backend.converter_guard.csv_id_contract import CALC_TEXT
from openpyxl import load_workbook

TARGETS = ('xlsx', 'xlsm', 'xltx', 'xltm')


def sha(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def read_cells(path):
    with Path(path).open('rb') as stream:
        book = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
        try:
            assert len(book.worksheets) == 1, 'one native table expected'
            sheet = book.worksheets[0]
            sheet.reset_dimensions()
            return [[(cell.value, cell.data_type, cell.number_format) for cell in row]
                    for row in sheet.iter_rows()]
        finally:
            book.close()


def verify(path, source, old_cells, level, extension):
    cells = read_cells(path)
    total = 7 + 9 * level
    assert len(cells) == total and all(len(row) == 6 for row in cells), 'table geometry changed'
    for row in range(1, total):
        assert cells[row][0][0] == f'{row:03d}' and cells[row][0][1] == 's', ('ID', row)
        assert cells[row][1][0] == f'Строка {row} / QA L{level}', ('label', row)
    # B29 must not convert all money/date/rate columns to text or change
    # their old native values/types/formats. The CSV source has no schema;
    # compare to the actual native baseline, not a guessed Python type.
    assert [row[1:] for row in cells] == [row[1:] for row in old_cells], 'non-ID native cells changed'
    assert sha(source) == SHA256[extension][level - 1], 'source bytes changed'
    return {'ids': total - 1, 'type': 'string', 'nonIdCellParity': (total * 5),
            'originalSha256': SHA256[extension][level - 1], 'outputSha256': sha(path)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--executable', type=Path, required=True)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=False)
    adapter = LibreOfficeAdapter(args.executable)
    if not adapter.available:
        raise RuntimeError('native LibreOffice is required')
    cancel = threading.Event()
    sources, proofs, controls = {}, {}, {}
    report = {'schema': 1, 'status': 'RUNNING', 'scope': 'B29 exact-source native routes, not FULLTEST',
              'sourceBytesUnchanged': False, 'cases': [], 'baseline': [], 'numericControl': None}
    for extension in ('csv', 'tab', 'tsv'):
        folder = args.output / extension
        folder.mkdir()
        for level in range(1, 9):
            source = folder / f'{level:02d}_LEVEL_{level}.{extension}'
            source.write_bytes(original_csv(level, extension))
            sources[extension, level] = source
    def one(extension, level, target, kind):
        source = sources[extension, level]
        tick = time.monotonic()
        try:
            adapter.prepare_source(proofs[extension, level])
            paths = adapter.convert(source, target, args.output / f'{kind}-{extension}-{level}-{target}',
                cancel=cancel, timeout=90, on_started=lambda: None)
            primary = [path for path in paths if path.suffix.lower() == '.' + target]
            assert len(primary) == 1
            evidence = verify(primary[0], source, controls[extension, level], level, extension)
            checks = [item for item in adapter.export_info().get('resultChecks', ())
                      if item.get('check') == 'csv_text_ids']
            assert checks and checks[-1]['status'] == 'checked_subset'
            assert checks[-1]['idsCompared'] == 6 + 9 * level, 'mandatory ID contract missing'
            mode = adapter.execution_mode()
            if kind == 'cli':
                assert mode == 'native_cli', 'direct CLI control was not CLI'
            case = {'extension': extension, 'level': level, 'target': target, 'kind': kind,
                    'mode': mode, 'status': 'PASS', 'check': checks[-1], 'independent': evidence}
        except Exception as error:
            case = {'extension': extension, 'level': level, 'target': target, 'kind': kind,
                    'status': 'FAIL', 'code': getattr(error, 'code', type(error).__name__),
                    'reason': str(error)[:1024]}
            report['cases'].append(case)
            print('B29_NATIVE_CASE=' + json.dumps(case, ensure_ascii=True), flush=True)
            raise
        case['elapsedMs'] = round((time.monotonic() - tick) * 1000)
        report['cases'].append(case)
        print('B29_NATIVE_CASE=' + json.dumps(case, ensure_ascii=True), flush=True)
    try:
        with adapter.batch(args.output / 'office'):
            for (extension, level), source in sources.items():
                proof = adapter.preflight(source, args.output / f'inspect-{extension}-{level}',
                                          cancel=cancel, timeout=60)
                assert proof.get('state') == 'detected', proof
                assert proof.get('selectedImportFilter') == CALC_TEXT, proof
                assert proof.get('selectedImportOptions') == ('44' if extension == 'csv' else '9') + ',34,76,1,1/2', proof
                proofs[extension, level] = proof
                old_proof = copy.deepcopy(proof)
                old_proof['selectedImportOptions'] = ','.join(proof['selectedImportOptions'].split(',')[:4])
                adapter.prepare_source(old_proof)
                outputs = adapter._convert_native(source, 'xlsx', args.output / f'old-{extension}-{level}',
                    cancel=cancel, timeout=90, on_started=lambda: None)
                reference = next(path for path in outputs if path.suffix.lower() == '.xlsx')
                old = read_cells(reference)
                assert len(old) == 7 + 9 * level
                assert old[1][0][0] == 1 and old[1][0][1] == 'n', 'old ID-loss control did not reproduce'
                controls[extension, level] = old
                report['baseline'].append({'extension': extension, 'level': level, 'status': 'OBSERVED',
                    'oldImportOptions': old_proof['selectedImportOptions'], 'firstId': 1,
                    'firstIdType': 'n', 'outputSha256': sha(reference)})
                for target in TARGETS:
                    assert adapter.route_decision(target, proof) == 'registered', (extension, target)
                    one(extension, level, target, 'batch')
            # Synthetic integer column proves the patch did not simply turn
            # every column into strings, independently of locale decimals.
            control = args.output / 'numeric-control.csv'
            control.write_bytes(b'ID,Amount\r\n001,42\r\n002,-7\r\n')
            proof = adapter.preflight(control, args.output / 'inspect-numeric', cancel=cancel, timeout=60)
            assert proof.get('state') == 'detected'
            adapter.prepare_source(proof)
            paths = adapter.convert(control, 'xlsx', args.output / 'numeric-result', cancel=cancel,
                                    timeout=90, on_started=lambda: None)
            cells = read_cells(next(p for p in paths if p.suffix == '.xlsx'))
            assert cells[1][0][:2] == ('001', 's') and cells[2][0][:2] == ('002', 's')
            assert cells[1][1][:2] == (42, 'n') and cells[2][1][:2] == (-7, 'n')
            report['numericControl'] = {'status': 'PASS', 'idsAsText': 2, 'unaffectedNumbers': 2}
        for extension in ('csv', 'tab', 'tsv'):
            for target in TARGETS:
                one(extension, 8, target, 'cli')
        assert all(sha(path) == SHA256[ext][level - 1] for (ext, level), path in sources.items())
        assert len(report['cases']) == 108 and len(report['baseline']) == 24
        report.update(status='PASS', sourceBytesUnchanged=True)
    except Exception as error:
        report.update(status='FAIL', error=type(error).__name__, reason=str(error)[:2048])
        raise
    finally:
        report['passed'] = sum(case['status'] == 'PASS' for case in report['cases'])
        report['failed'] = sum(case['status'] == 'FAIL' for case in report['cases'])
        (args.output / 'evidence.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        print('B29_NATIVE_SUMMARY=' + json.dumps({k: v for k, v in report.items()
              if k not in ('cases', 'baseline')}, ensure_ascii=True), flush=True)


if __name__ == '__main__':
    main()
