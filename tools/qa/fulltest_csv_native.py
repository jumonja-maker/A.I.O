"""M24 native controls: original corpus bytes, typed values and protected IDs.

Only the acceptance harness appends locale tokens for adversarial locale
controls; production keeps per-column hints and does not change user locale.
This is 36 focused exports, not the historical 5723-task FULLTEST.
"""
import csv
from decimal import Decimal
import hashlib
import io
import json
from pathlib import Path
import sys
import threading
import time
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from app_backend.converter_guard import LibreOfficeAdapter

NS = {'o': 'urn:oasis:names:tc:opendocument:xmlns:office:1.0',
      't': 'urn:oasis:names:tc:opendocument:xmlns:table:1.0',
      'text': 'urn:oasis:names:tc:opendocument:xmlns:text:1.0'}


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def materialize_controls(output):
    fixture = json.loads((ROOT / 'tests/fixtures/fulltest_csv_2levels.json').read_text(encoding='utf-8'))
    if fixture['schema'] != 2 or fixture['level1DataRows'] != 7 or fixture['level2DataRows'] != 40:
        raise AssertionError('Unexpected original corpus fixture schema')
    lines = fixture['csvLevel2'].splitlines(keepends=True)
    if len(lines) != 41:
        raise AssertionError('Historical CSV row count changed')
    controls = []
    for level, count in ((1, 7), (2, 40)):
        for extension in ('csv', 'tab', 'tsv'):
            name = f'0{level}_LEVEL_{level}.{extension}'
            text = ''.join(lines[:count + 1])
            delimiter = ',' if extension == 'csv' else '\t'
            if delimiter == '\t':
                text = text.replace(',', '\t')
            raw = text.encode('utf-8')
            if hashlib.sha256(raw).hexdigest() != fixture['sha256'][name]:
                raise AssertionError('Original corpus SHA256 mismatch: ' + name)
            source = output / name
            with source.open('xb') as stream:
                stream.write(raw)
            controls.append((source, list(csv.reader(io.StringIO(text), delimiter=delimiter)), fixture['sha256'][name]))
    return controls


def xlsx_cells(path, row_count):
    from openpyxl import load_workbook
    with Path(path).open('rb') as stream:
        book = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
        try:
            if len(book.worksheets) != 1:
                raise AssertionError('Unexpected XLSX sheet count')
            rows = [[(cell.data_type, cell.value) for cell in row]
                    for row in book.worksheets[0].iter_rows(min_row=1, max_row=row_count, max_col=5)]
        finally:
            book.close()
    return rows


def fods_cells(path, row_count):
    root = ET.parse(path).getroot()
    tables = root.findall('o:body/o:spreadsheet/t:table', NS)
    if len(tables) != 1:
        raise AssertionError('Unexpected FODS sheet count')
    rows = []
    for row in tables[0].findall('t:table-row', NS):
        values = []
        for cell in row.findall('t:table-cell', NS):
            kind = cell.get('{' + NS['o'] + '}value-type')
            value = cell.get('{' + NS['o'] + '}value') if kind == 'float' else '\n'.join(
                ''.join(p.itertext()) for p in cell.findall('text:p', NS))
            repeated = int(cell.get('{' + NS['t'] + '}number-columns-repeated', '1'))
            values.extend([(kind, value)] * min(repeated, max(0, 5 - len(values))))
            if len(values) >= 5:
                break
        repeated = int(row.get('{' + NS['t'] + '}number-rows-repeated', '1'))
        rows.extend([values] * min(repeated, max(0, row_count - len(rows))))
        if len(rows) >= row_count:
            break
    return rows


def verify_cells(path, target, expected):
    observed = xlsx_cells(path, len(expected)) if target == 'xlsx' else fods_cells(path, len(expected))
    if len(observed) != len(expected):
        raise AssertionError('Source/result row count differs')
    string_type, numeric_type = ('s', 'n') if target == 'xlsx' else ('string', 'float')
    checked_numbers = 0
    for row_number, (before, after) in enumerate(zip(expected, observed), 1):
        if len(after) != 5:
            raise AssertionError('Source/result column count differs')
        for column, text in enumerate(before):
            kind, value = after[column]
            if row_number == 1 or column < 2:
                if kind != string_type or value != text:
                    raise AssertionError(f'String or leading-zero ID changed at {row_number}:{column+1}')
            else:
                if kind != numeric_type or isinstance(value, bool) or Decimal(str(value)) != Decimal(text):
                    raise AssertionError(f'Number or numeric type changed at {row_number}:{column+1}')
                checked_numbers += 1
    return {'numericCells': checked_numbers, 'idCells': len(expected) - 1}


def run(executable, output):
    output = Path(output)
    output.mkdir(parents=True, exist_ok=False)
    cases = []
    report = {'status': 'FAIL', 'scope': 'six original CSV/TAB/TSV -> XLSX/FODS; three import locales; not full historical corpus'}
    start = time.monotonic()
    try:
        controls = materialize_controls(output)
        cancel = threading.Event()
        adapter = LibreOfficeAdapter(Path(executable))
        counter = 0
        with adapter.batch(output / 'office'):
            for source, expected, original_sha in controls:
                counter += 1
                proven = adapter.preflight(source, output / f'probe-{counter}', cancel=cancel, timeout=90)
                if proven.get('state') != 'detected':
                    raise AssertionError('Native CSV model not detected: ' + source.name)
                options = proven.get('selectedImportOptions')
                separator = '44' if source.suffix == '.csv' else '9'
                if options != separator + ',34,76,1,1/2/4/10/5/10':
                    raise AssertionError('Required numeric/Text column hints absent: ' + str(options))
                for locale_name, locale_id in (('production_default', None), ('de-DE', '1031'), ('ru-RU', '1049')):
                    for target in ('xlsx', 'fods'):
                        counter += 1
                        adapter.prepare_source(proven)
                        adapter.prepare_retry_recipe(None)
                        if adapter._local.import_options != options:
                            raise AssertionError('Validated production hints were discarded before native import')
                        if locale_id is not None:
                            # Test-only adversarial environment, AFTER the real
                            # production validation above. Never broaden accepted
                            # production options just to pass a locale control.
                            # This changes one disposable import, not any profile.
                            adapter._local.import_options = options + ',' + locale_id
                        tick = time.monotonic()
                        paths = adapter.convert(source, target, output / f'task-{counter}',
                            cancel=cancel, timeout=120, on_started=lambda: None)
                        primary = [path for path in paths if path.suffix == '.' + target]
                        if len(primary) != 1:
                            raise AssertionError('Ambiguous native output')
                        counts = verify_cells(primary[0], target, expected)
                        if locale_id is None and target == 'xlsx':
                            checks = adapter.export_info().get('resultChecks', [])
                            if not any(item.get('check') == 'csv_text_ids' and item.get('status') == 'checked_subset'
                                       for item in checks):
                                raise AssertionError('Mandatory production CSV ID check did not execute')
                        if digest(source) != original_sha:
                            raise AssertionError('Original corpus input changed')
                        cases.append({'source': source.name, 'sourceSha256': original_sha,
                                      'target': target, 'localeControl': locale_name,
                                      'status': 'PASS', **counts,
                                      'ms': round((time.monotonic() - tick) * 1000),
                                      'mode': adapter.execution_mode()})
        if len(cases) != 36:
            raise AssertionError('Incomplete native CSV acceptance')
        report.update(status='PASS', sourceUnchanged=True, cases=cases,
                      nativeExports=36, originalInputs=6, targets=['xlsx', 'fods'],
                      globalLocaleChanged=False)
    except Exception as error:
        report.update(cases=cases, error=type(error).__name__, reason=str(error),
                      code=getattr(error, 'code', None), diagnostic=getattr(error, 'diagnostic', None))
        raise
    finally:
        report['elapsedSeconds'] = round(time.monotonic() - start, 3)
        (output / 'evidence.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
        print('FULLTEST_CSV_NATIVE=' + json.dumps(report, ensure_ascii=True), flush=True)
    return report


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument('--executable', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    run(args.executable, args.output)
