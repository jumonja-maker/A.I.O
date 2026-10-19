"""Opt-in synthetic LibreOffice qualification, never a user-document corpus.

QA-only dependencies: python-docx, python-pptx, Pillow, pypdf, openpyxl.
These are NOT imported by the application or added as runtime dependencies.
Usage: python tools/qa/guard_native_smoke.py --work /path/to/new/empty/directory
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
from pathlib import Path
import random
import subprocess
import sys
import time
import unicodedata
import zipfile

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from app_backend.converter_guard import ConverterGuard, LibreOfficeAdapter
from app_backend.converter_guard.core import GuardLimits
from app_backend.converter_guard.libreoffice import find_libreoffice, prepare_profile
from app_backend.converter_guard.process import ProcessTree


def digest(path):
    with path.open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def direct_export(executable, source, directory, extension, filter_name, profile):
    directory.mkdir(parents=True)
    prepare_profile(profile)
    command = [str(executable), '-env:UserInstallation=' + profile.resolve().as_uri(),
        '--headless', '--nologo', '--nodefault', '--norestore',
        '--convert-to', extension + ':' + filter_name,
        '--outdir', str(directory.resolve()), str(source.resolve())]
    started = time.monotonic()
    with ProcessTree(command, cwd=directory) as process:
        while process.process.poll() is None:
            if time.monotonic() - started > 60:
                raise AssertionError('Direct fixture export timed out')
            time.sleep(.05)
        assert process.process.returncode == 0, process.process.returncode
    result = directory / (source.stem + '.' + extension)
    assert result.is_file(), str(result)
    return result


def docx_fixture(path, image):
    from lxml import etree as ET  # Preserve the package's original default namespaces.
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.shared import Inches
    document = Document()
    document.add_heading('GUARD_BEGIN_7402', level=1)
    document.add_paragraph('Проверка исходного документа. Кириллица, таблица, формула и сноска.')
    paragraph = document.add_paragraph('MATH_GUARD_7402 ')
    math = OxmlElement('m:oMath')
    run, text = OxmlElement('m:r'), OxmlElement('m:t')
    text.text = 'x=7'
    run.append(text); math.append(run); paragraph._p.append(math)
    reference = OxmlElement('w:footnoteReference')
    reference.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}id', '1')
    footnote_run = OxmlElement('w:r'); footnote_run.append(reference); paragraph._p.append(footnote_run)
    table = document.add_table(rows=2, cols=2)
    for cell, value in zip([c for row in table.rows for c in row.cells], ['CELL_A_7402','00125','CELL_B_7402','42']):
        cell.text = value
    document.add_picture(str(image), width=Inches(1.5))
    document.add_page_break()
    document.add_paragraph('GUARD_END_7402')
    document.save(path)
    with zipfile.ZipFile(path) as archive:
        entries = {item.filename: archive.read(item) for item in archive.infolist()}
    rel_ns = 'http://schemas.openxmlformats.org/package/2006/relationships'
    rels = ET.fromstring(entries['word/_rels/document.xml.rels'])
    ET.SubElement(rels, '{' + rel_ns + '}Relationship', {
        'Id': 'rIdGuardFootnotes', 'Type': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes',
        'Target': 'footnotes.xml'})
    entries['word/_rels/document.xml.rels'] = ET.tostring(rels, encoding='utf-8', xml_declaration=True)
    types = ET.fromstring(entries['[Content_Types].xml'])
    ET.SubElement(types, '{http://schemas.openxmlformats.org/package/2006/content-types}Override', {
        'PartName': '/word/footnotes.xml',
        'ContentType': 'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml'})
    entries['[Content_Types].xml'] = ET.tostring(types, encoding='utf-8', xml_declaration=True)
    entries['word/footnotes.xml'] = b'''<?xml version="1.0" encoding="UTF-8"?>
<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
 <w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>
 <w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>
 <w:footnote w:id="1"><w:p><w:r><w:t>FOOTNOTE_GUARD_7402</w:t></w:r></w:p></w:footnote>
</w:footnotes>'''
    with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items(): archive.writestr(name, data)


def make_fixtures(root, executable):
    from PIL import Image
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches
    from docx import Document
    from docx.shared import Inches as DocxInches
    root.mkdir()
    image = Image.frombytes('RGB', (128,128), random.Random(7402).randbytes(128*128*3))
    image_path = root / 'source-image.png'; image.save(image_path)
    writer = root / 'writer.docx'; docx_fixture(writer, image_path)
    calc = root / 'calc.xlsx'
    workbook = Workbook()
    workbook.active.title = 'First'
    for sheet, marker in [(workbook.active,'SHEET_A_7402'), (workbook.create_sheet('Second'),'SHEET_B_7402')]:
        sheet.append([marker, 'Value']); sheet.append(['00125', '=1+2'])
        sheet.column_dimensions['A'].width = 28
    workbook.save(calc); workbook.close()
    impress = root / 'impress.pptx'; presentation = Presentation()
    for marker in ('SLIDE_A_7402','SLIDE_B_7402'):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1),Inches(1), Inches(7), Inches(2)).text = marker
    presentation.save(impress)
    inputs = {'docx': writer, 'xlsx': calc, 'pptx': impress}
    for source, extension, filter_name in [(writer,'odt','writer8'), (writer,'doc','MS Word 97'),
            (writer,'rtf','Rich Text Format'), (calc,'ods','calc8'), (calc,'xls','MS Excel 97'),
            (impress,'odp','impress8'), (impress,'ppt','MS PowerPoint 97')]:
        inputs[extension] = direct_export(executable, source, root / ('make-' + extension),
            extension, filter_name, root / ('profile-' + extension))
    # A real DOCX beyond the historical 20 MB boundary, with an embedded RGB PNG.
    large_image = Image.frombytes('RGB',(3200,2400),random.Random(7403).randbytes(3200*2400*3))
    large_image.save(root / 'large-image.png')
    large = Document(); large.add_paragraph('LARGE_BEGIN_7403')
    large.add_picture(str(root / 'large-image.png'),width=DocxInches(2))
    large.add_paragraph('LARGE_END_7403'); large.save(root / 'large.docx')
    inputs['large_docx'] = root / 'large.docx'
    assert inputs['large_docx'].stat().st_size > 20 * 1024 * 1024
    return inputs, image, large_image


def pdf_facts(path, expected, image=None):
    from pypdf import PdfReader
    reader = PdfReader(path)
    text = unicodedata.normalize('NFKC', '\n'.join(page.extract_text() or '' for page in reader.pages))
    compact = ''.join(text.split())
    markers = {marker: ''.join(marker.split()) in compact for marker in expected}
    facts = {'pages':len(reader.pages), 'markers':markers, 'size':path.stat().st_size}
    assert all(markers.values()), (path.name, markers, text)
    if image is not None:
        pixel_hash = hashlib.sha256(image.convert('RGB').tobytes()).hexdigest()
        images = [item.image for page in reader.pages for item in page.images]
        matches = [i for i in images if i.size == image.size and
                   hashlib.sha256(i.convert('RGB').tobytes()).hexdigest() == pixel_hash]
        facts.update({'embeddedImagePixelsExact':bool(matches),'imageSize':list(image.size),'sourcePixelSHA256':pixel_hash})
        assert matches, (path.name, 'Image pixels changed')
    return facts


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--work', required=True, type=Path)
    args = parser.parse_args()
    root = args.work.resolve(); root.mkdir(parents=True, exist_ok=False)
    executable = find_libreoffice()
    if not executable: raise SystemExit('NOT_RUN: LibreOffice not found')
    evidence = {'scope':'Synthetic Linux/source qualification; not Windows, installer or full visual corpus acceptance.'}
    evidence['libreOfficeVersion'] = subprocess.run([str(executable),'--version'], capture_output=True,
        text=True,timeout=10,check=True).stdout.strip()
    fixtures, image, large_image = make_fixtures(root/'fixtures', executable)
    before = {key: {'sha256':digest(path),'size':path.stat().st_size} for key,path in fixtures.items()}
    guard = ConverterGuard(adapter=LibreOfficeAdapter(executable), limits=GuardLimits(file_timeout=60), spool_parent=root)
    try:
        rows = [{'id':'empty_first','name':'empty.docx','size':0}] + [
            {'id':key,'name':path.name,'size':path.stat().st_size} for key,path in fixtures.items()]
        job = guard.create(rows)['jobId']
        guard.upload(job,'empty_first',io.BytesIO(b''),0)
        for key,path in fixtures.items():
            with path.open('rb') as stream: guard.upload(job,key,stream,path.stat().st_size)
        started = time.monotonic(); guard.start(job)
        while True:
            state = guard.status(job)
            if state['finished']: break
            if time.monotonic()-started > 700: raise AssertionError('Native batch did not finish')
            time.sleep(.1)
        evidence['elapsedConversionSeconds'] = round(time.monotonic()-started,3)
        evidence['files'] = state['files']
        assert state['files'][0]['code'] == 'INPUT_EMPTY', state['files'][0]
        assert all(f['state']=='ready' for f in state['files'][1:]), state['files']
        outputs = root/'outputs'; outputs.mkdir()
        facts = {}
        for output in state['result']['outputs']:
            key = output['fileId']; destination = outputs/(key+'.pdf')
            with guard.open_result(output['downloadUrl'].split('ticket=')[1]) as (stream,_,_):
                with destination.open('wb') as saved:
                    while block:=stream.read(65536): saved.write(block)
            family = 'calc' if key in ('xlsx','xls','ods') else 'impress' if key in ('pptx','ppt','odp') else 'writer'
            markers = (['LARGE_BEGIN_7403','LARGE_END_7403'] if key=='large_docx' else
                ['SHEET_A_7402','SHEET_B_7402','00125'] if family=='calc' else
                ['SLIDE_A_7402','SLIDE_B_7402'] if family=='impress' else ['GUARD_BEGIN_7402','GUARD_END_7402','CELL_A_7402'])
            if key=='docx': markers += ['FOOTNOTE_GUARD_7402','x=7']
            facts[key] = pdf_facts(destination,markers,large_image if key=='large_docx' else image if key=='docx' else None)
        evidence['pdfFacts'] = facts
        evidence['sourcesUnchanged'] = all(before[key]['sha256']==digest(path) for key,path in fixtures.items())
        assert evidence['sourcesUnchanged']
        evidence['inputs'] = before
        direct = direct_export(executable, fixtures['docx'], root/'direct-control', 'pdf',
            'writer_pdf_Export:{"UseLosslessCompression":{"type":"boolean","value":"true"},"ReduceImageResolution":{"type":"boolean","value":"false"}}',
            root/'direct-control-profile')
        evidence['directControl'] = pdf_facts(direct,['GUARD_BEGIN_7402','GUARD_END_7402','FOOTNOTE_GUARD_7402','x=7'],image)
        assert evidence['directControl']['pages']==facts['docx']['pages']
        evidence['trace'] = state['trace']
        guard.release(job)
    finally:
        guard.close()
        evidence['storageAfterClose'] = guard.storage_stats()
        (root/'evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    assert evidence['storageAfterClose']=={'jobs':0,'reservedBytes':0,'running':0}
    print(json.dumps({'status':'PASS','inputs':len(fixtures),'emptyFirstContinued':True,
        'seconds':evidence['elapsedConversionSeconds'],'sourcesUnchanged':True,
        'largeSourceBytes':before['large_docx']['size'],'evidence':str(root/'evidence.json')},ensure_ascii=False))


if __name__=='__main__': main()
