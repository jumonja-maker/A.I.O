"""Opt-in synthetic native LO acceptance. No user documents or Windows runner.

QA-only dependencies: python-docx, python-pptx, Pillow, openpyxl, pypdf.
Run: python tools/qa/native_converter_acceptance.py --work /new/empty/directory
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
from pathlib import Path
import random
import re
import shutil
import subprocess
import sys
import threading
import time
import tracemalloc
from urllib.parse import unquote
from urllib.request import Request, urlopen
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[2]
sys.path[:0] = [str(ROOT), str(ROOT / 'tests')]
from app_backend.libreoffice_engine import engine_binary, convert_path
from app_backend.server import start_server
from test_converter_native_stream import multipart


def digest(path):
    with path.open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def main():
    from PIL import Image
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.shared import Inches
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches as SlideInches
    from pypdf import PdfReader

    parser = argparse.ArgumentParser()
    parser.add_argument('--work', required=True, type=Path)
    args = parser.parse_args()
    work = args.work.resolve()
    work.mkdir(parents=True, exist_ok=False)
    fixtures = work / 'fixtures'
    fixtures.mkdir()
    executable = engine_binary()
    if executable is None:
        raise SystemExit('NOT_RUN: LibreOffice not found')
    evidence = {'scope': 'Synthetic Linux/source HTTP + native LO; not a user corpus',
                'engine': subprocess.run([str(executable), '--version'], capture_output=True,
                    text=True, timeout=15, check=True).stdout.strip(), 'cases': []}
    small = Image.frombytes('RGB', (96, 64), random.Random(7410).randbytes(96 * 64 * 3))
    for extension, kind in [('bmp', 'BMP'), ('png', 'PNG'), ('jpg', 'JPEG'),
                            ('gif', 'GIF'), ('tiff', 'TIFF'), ('webp', 'WEBP')]:
        small.save(fixtures / ('image.' + extension), format=kind)
    (fixtures / 'image.svg').write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="96" height="64">'
        '<rect width="96" height="64" fill="#0871df"/><text x="4" y="30">AIO</text></svg>')
    document = Document()
    document.add_heading('NATIVE_BEGIN_7410', 1)
    paragraph = document.add_paragraph('FILENAME_7410 ')
    field = OxmlElement('w:fldSimple')
    field.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instr', ' FILENAME ')
    run, text = OxmlElement('w:r'), OxmlElement('w:t')
    text.text = 'writer.docx'
    run.append(text); field.append(run); paragraph._p.append(field)
    table = document.add_table(rows=2, cols=2)
    for cell, value in zip([c for row in table.rows for c in row.cells],
                           ['CELL_A_7410', '00125', 'CELL_B_7410', '42']):
        cell.text = value
    document.add_picture(str(fixtures / 'image.png'), width=Inches(1))
    document.add_page_break()
    document.add_paragraph('NATIVE_END_7410')
    document.save(fixtures / 'writer.docx')
    book = Workbook()
    for sheet, marker in [(book.active, 'SHEET_A_7410'), (book.create_sheet('Second'), 'SHEET_B_7410')]:
        sheet.append([marker, 'Value']); sheet.append(['00125', 42])
        sheet.column_dimensions['A'].width = 30
    book.save(fixtures / 'calc.xlsx'); book.close()
    slides = Presentation()
    for marker in ('SLIDE_A_7410', 'SLIDE_B_7410'):
        slide = slides.slides.add_slide(slides.slide_layouts[6])
        slide.shapes.add_textbox(SlideInches(1), SlideInches(1), SlideInches(7), SlideInches(2)).text = marker
    slides.save(fixtures / 'impress.pptx')
    inputs = [fixtures / name for name in ('writer.docx', 'calc.xlsx', 'impress.pptx')]
    for source, extensions in [(inputs[0], ('doc', 'odt', 'rtf')),
                               (inputs[1], ('xls', 'ods')),
                               (inputs[2], ('ppt', 'odp'))]:
        for extension in extensions:
            output = convert_path(source, extension, work / ('make-' + extension), executable=executable)[0]
            target = fixtures / output.name
            shutil.copyfile(output, target)
            inputs.append(target)
    (fixtures / 'text.txt').write_text('NATIVE_TEXT_7410\nПрямой LibreOffice', encoding='utf-8')
    (fixtures / 'table.csv').write_text('NATIVE_CSV_7410,Value\n00125,42\n', encoding='utf-8')
    (fixtures / 'page.html').write_text('<!doctype html><html><body><p>NATIVE_HTML_7410</p></body></html>')
    (fixtures / 'text.unlisted').write_text('NATIVE_UNLISTED_7410\n', encoding='utf-8')
    inputs.extend(fixtures / name for name in ('text.txt', 'table.csv', 'page.html', 'text.unlisted',
        'image.bmp', 'image.png', 'image.jpg', 'image.gif', 'image.tiff', 'image.svg', 'image.webp'))
    large_image = Image.frombytes('RGB', (3600, 2600),
                                  random.Random(7411).randbytes(3600 * 2600 * 3))
    large_image.save(fixtures / 'large-image.png')
    large = Document()
    large.add_paragraph('LARGE_BEGIN_7411')
    large.add_picture(str(fixtures / 'large-image.png'), width=Inches(2))
    large.add_paragraph('LARGE_END_7411')
    large.save(fixtures / 'large-one.docx')
    shutil.copyfile(fixtures / 'large-one.docx', fixtures / 'large-two.docx')
    big = [fixtures / 'large-one.docx', fixtures / 'large-two.docx']
    assert all(p.stat().st_size > 20 * 1024 * 1024 for p in big)
    assert sum(p.stat().st_size for p in big) > 50 * 1024 * 1024
    originals = {p: digest(p) for p in inputs + big}
    server, token = start_server()
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    base = 'http://127.0.0.1:' + str(server.server_port)
    try:
        def convert(label, sources, formats=('pdf',)):
            started = time.monotonic()
            wire, headers = multipart([(p.name, p) for p in sources], formats=formats)
            try:
                with urlopen(Request(base + '/api/converter/convert', data=wire,
                        headers={'X-AIO-Token': token, **headers}), timeout=120) as response:
                    value = json.load(response)
            finally:
                wire.close()
            row = {'label': label, 'sources': [{'name': p.name, 'bytes': p.stat().st_size}
                                              for p in sources],
                   'formats': list(formats), 'report': value['result']['targetReport'],
                   'seconds': round(time.monotonic() - started, 3)}
            evidence['cases'].append(row)
            target = work / 'results' / label / value['result']['name']
            target.parent.mkdir(parents=True)
            with urlopen(base + value['result']['downloadUrl'], timeout=120) as response, target.open('wb') as output:
                while chunk := response.read(64 * 1024):
                    output.write(chunk)
            row['outputBytes'] = target.stat().st_size
            row['outputSHA256'] = digest(target)
            body = json.dumps({'ticket': value['result']['ticket']}).encode()
            with urlopen(Request(base + '/api/converter/release', data=body,
                    headers={'X-AIO-Token': token, 'Content-Type': 'application/json'})) as response:
                assert json.load(response)['ok']
            return target, row

        target, batch = convert('native-inputs', inputs)
        assert batch['report']['failed'] == 0, batch
        with ZipFile(target) as archive:
            archive.extractall(work / 'batch-pdfs')
        pdfs = sorted((work / 'batch-pdfs').rglob('*.pdf'))
        assert len(pdfs) == len(inputs)
        for name, markers in [('writer', ('NATIVE_BEGIN_7410', 'NATIVE_END_7410', 'CELL_A_7410', 'writer.docx')),
                              ('calc', ('SHEET_A_7410', 'SHEET_B_7410')),
                              ('impress', ('SLIDE_A_7410', 'SLIDE_B_7410'))]:
            source = next(p for p in pdfs if p.name == name + '.pdf' and '/00' in str(p))
            text = ''.join(''.join((page.extract_text() or '').split()) for page in PdfReader(source).pages)
            assert all(marker in text for marker in markers), (source, markers, text)
        tracemalloc.start()
        target, big_row = convert('over-50-mib-batch', big)
        _current, peak = tracemalloc.get_traced_memory()
        tracemalloc.stop()
        big_row['pythonTracedPeakBytes'] = peak
        assert big_row['report']['succeeded'] == 2 and not big_row['report']['failed']
        with ZipFile(target) as archive:
            for name in archive.namelist():
                if name.endswith('.pdf'):
                    text = ''.join(page.extract_text() or '' for page in PdfReader(io.BytesIO(archive.read(name))).pages)
                    assert 'LARGE_BEGIN_7411' in text and 'LARGE_END_7411' in text
        for extension in ('png', 'jpg', 'bmp', 'gif', 'tiff', 'svg', 'webp'):
            result, row = convert('bmp-to-' + extension, [fixtures / 'image.bmp'], (extension,))
            assert row['report']['succeeded'] == 1
            if extension != 'svg':
                with Image.open(result) as image:
                    row['imageSize'] = list(image.size)
        result, row = convert('html-with-companions', [fixtures / 'writer.docx'], ('html',))
        assert result.suffix == '.zip', 'HTML sidecar images must travel with the HTML file'
        with ZipFile(result) as archive:
            html_name = next(name for name in archive.namelist() if name.endswith('.html'))
            html = archive.read(html_name).decode('utf-8')
            for image in re.findall(r'<img[^>]+src="([^"]+)"', html, flags=re.I):
                if not image.startswith('data:'):
                    assert str(Path(html_name).parent / unquote(image)) in archive.namelist(), image
        pdf = next(p for p in pdfs if p.name == 'writer.pdf')
        result, row = convert('pdf-import-to-png', [pdf], ('png',))
        assert row['report']['succeeded'] == 1
        evidence['sourcesUnchanged'] = all(digest(path) == expected for path, expected in originals.items())
        assert evidence['sourcesUnchanged']
        assert not server.converter.active and not server.converter.results
        evidence['status'] = 'PASS'
        (work / 'evidence.json').write_text(json.dumps(evidence, ensure_ascii=False, indent=2))
        print(json.dumps({'status': 'PASS', 'nativeInputCount': len(inputs), 'cases': len(evidence['cases']),
                          'largeBatchBytes': sum(p.stat().st_size for p in big),
                          'pythonTracedPeakBytes': big_row['pythonTracedPeakBytes'],
                          'sourcesUnchanged': evidence['sourcesUnchanged']}, ensure_ascii=False), flush=True)
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=5)


if __name__ == '__main__':
    main()
