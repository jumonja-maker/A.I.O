"""M25/M28/M31 native controls; deterministic fixtures, not historical FULLTEST.

The DBF reader below is deliberately independent of production code. Native
LibreOffice readback is a second reader; Windows GDI enumerates/renders EMFs.
No fixture, executable, installer or downloadable artifact is published.
"""
import argparse
import datetime
from decimal import Decimal
import gzip
import hashlib
import io
import json
import os
from pathlib import Path
import struct
import sys
import threading
import time
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))


def read_dbf(path):
    """Strict, small-fixture dBASE III/IV reader, including native DBT memos."""
    data = path.read_bytes()
    assert len(data) >= 33 and data[0] in (3, 0x83, 0x8b), data[:32].hex()
    count, header, stride = struct.unpack_from('<IHH', data, 4)
    assert header >= 33 and (header - 33) % 32 == 0 and data[header-1] == 13
    assert header + count * stride <= len(data)
    assert data[header+count*stride:] in (b'', b'\x1a')
    fields = []
    for offset in range(32, header-1, 32):
        field = data[offset:offset+32]
        name = field[:11].split(b'\0', 1)[0].decode('utf-8', 'strict')
        fields.append({'name': name, 'type': chr(field[11]),
                       'width': field[16], 'decimal': field[17]})
    assert fields and len({f['name'].upper() for f in fields}) == len(fields), fields
    assert 1 + sum(f['width'] for f in fields) == stride
    rows = []
    for index in range(count):
        record = data[header+index*stride:header+(index+1)*stride]
        assert record[0] == 32, record.hex()
        row, offset = [], 1
        for field in fields:
            raw = record[offset:offset+field['width']]
            offset += field['width']
            kind = field['type']
            if kind == 'C':
                value = raw.rstrip(b' ').decode('utf-8', 'strict')
            elif kind in ('N', 'F'):
                value = str(Decimal(raw.decode('ascii').strip())) if raw.strip() else ''
            elif kind == 'D':
                value = raw.decode('ascii').strip()
                if value:
                    datetime.datetime.strptime(value, '%Y%m%d')
            elif kind == 'L':
                value = raw.decode('ascii')
                assert value in 'TFYN? '
            elif kind == 'M':
                block = int(raw.strip() or b'0')
                if block:
                    memo = path.with_suffix('.dbt').read_bytes()
                    if data[0] == 0x83:
                        start = block * 512
                        end = memo.find(b'\x1a\x1a', start)
                        assert 512 <= start < end <= len(memo), (block, len(memo))
                        value = memo[start:end].decode('utf-8', 'strict')
                    else:
                        block_size = int.from_bytes(memo[20:22], 'big')
                        start = block * block_size
                        marker, length = struct.unpack_from('<II', memo, start)
                        assert marker == 0xffff0800 and length >= 8
                        value = memo[start+8:start+length].decode('utf-8', 'strict')
                else:
                    value = ''
            else:
                raise AssertionError(('unsupported_field', field))
            row.append(value)
        rows.append(row)
    return {'version': data[0], 'ldid': data[29], 'headerBytes': header,
            'recordBytes': stride, 'fields': fields, 'rows': rows}


def emf_records(raw):
    """Physical walk from byte zero, independent of inspect_metafile()."""
    assert raw[40:44] == b' EMF' and struct.unpack_from('<I', raw, 48)[0] == len(raw)
    records, offset = [], 0
    while offset < len(raw):
        kind, size = struct.unpack_from('<II', raw, offset)
        assert size >= 8 and size % 4 == 0 and offset+size <= len(raw)
        records.append(kind)
        offset += size
    assert records[0] == 1 and records[-1] == 14 and records.count(1) == 1
    assert struct.unpack_from('<I', raw, len(raw)-4)[0] >= 20
    return {'physicalRecords': len(records), 'declaredRecords': struct.unpack_from('<I', raw, 52)[0],
            'bytes': len(raw), 'recordTypes': sorted(set(records))}


def gdi_render(path):
    """Independent Windows GDI enumeration and deterministic DIB playback."""
    import ctypes as c
    from ctypes import wintypes as w
    gdi = c.WinDLL('gdi32', use_last_error=True)
    gdi.GetEnhMetaFileW.argtypes = [w.LPCWSTR]
    gdi.GetEnhMetaFileW.restype = w.HANDLE
    gdi.CreateCompatibleDC.argtypes = [w.HDC]
    gdi.CreateCompatibleDC.restype = w.HDC
    gdi.CreateDIBSection.argtypes = [w.HDC, c.c_void_p, w.UINT, c.POINTER(c.c_void_p), w.HANDLE, w.DWORD]
    gdi.CreateDIBSection.restype = w.HBITMAP
    gdi.SelectObject.argtypes = [w.HDC, w.HANDLE]
    gdi.SelectObject.restype = w.HANDLE
    callback_type = c.WINFUNCTYPE(c.c_int, w.HDC, c.c_void_p, c.c_void_p, c.c_int, w.LPARAM)
    gdi.EnumEnhMetaFile.argtypes = [w.HDC, w.HANDLE, callback_type, c.c_void_p, c.POINTER(w.RECT)]
    gdi.PlayEnhMetaFile.argtypes = [w.HDC, w.HANDLE, c.POINTER(w.RECT)]
    gdi.DeleteEnhMetaFile.argtypes = [w.HANDLE]
    gdi.DeleteObject.argtypes = [w.HANDLE]
    gdi.DeleteDC.argtypes = [w.HDC]
    handle, dc, bitmap, old = None, None, None, None
    try:
        handle = gdi.GetEnhMetaFileW(str(path.resolve()))
        assert handle, ('GetEnhMetaFileW', c.get_last_error())
        kinds = []
        def visit(_dc, _table, record, _objects, _param):
            kinds.append(c.cast(record, c.POINTER(w.DWORD))[0])
            return 1
        callback = callback_type(visit)
        rect = w.RECT(0, 0, 256, 256)
        assert gdi.EnumEnhMetaFile(None, handle, callback, None, c.byref(rect))
        dc = gdi.CreateCompatibleDC(None)
        assert dc
        info = struct.pack('<IiiHHIIiiII', 40, 256, -256, 1, 32, 0, 256*256*4, 0, 0, 0, 0)
        buffer = c.create_string_buffer(info)
        bits = c.c_void_p()
        bitmap = gdi.CreateDIBSection(dc, buffer, 0, c.byref(bits), None, 0)
        assert bitmap and bits.value
        old = gdi.SelectObject(dc, bitmap)
        c.memset(bits, 255, 256*256*4)
        assert gdi.PlayEnhMetaFile(dc, handle, c.byref(rect)), c.get_last_error()
        gdi.GdiFlush()
        pixels = c.string_at(bits, 256*256*4)
        assert pixels != b'\xff' * len(pixels), 'empty GDI rendering'
        return {'gdiRecords': len(kinds), 'pixelSha256': hashlib.sha256(pixels).hexdigest()}
    finally:
        if old and dc: gdi.SelectObject(dc, old)
        if bitmap: gdi.DeleteObject(bitmap)
        if dc: gdi.DeleteDC(dc)
        if handle: gdi.DeleteEnhMetaFile(handle)


def main():
    from openpyxl import Workbook, load_workbook
    from app_backend.converter_guard import LibreOfficeAdapter
    from app_backend.libreoffice_engine import convert_path
    from tools.qa.fulltest_vector_pages_native import fixture
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--executable', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=False)
    cases, failures, counter = [], [], 0
    cancel = threading.Event()
    adapter = LibreOfficeAdapter(args.executable)
    def convert(path, kind):
        nonlocal counter
        counter += 1
        proven = adapter.preflight(path, args.output/('probe-%d' % counter), cancel=cancel, timeout=90)
        assert proven.get('state') == 'detected', proven
        adapter.prepare_source(proven)
        return adapter.convert(path, kind, args.output/('task-%d' % counter),
                               cancel=cancel, timeout=120, on_started=lambda: None)
    def run(name, action):
        try:
            result = {'case': name, 'status': 'PASS', **action()}
        except Exception as error:
            result = {'case': name, 'status': 'FAIL', 'error': repr(error),
                      'code': getattr(error, 'code', None), 'diagnostic': getattr(error, 'diagnostic', None)}
            failures.append(name)
        cases.append(result)
        print('DBF_EMF_CASE='+json.dumps(result, ensure_ascii=True), flush=True)
    source = args.output/'cyrillic.xlsx'
    book = Workbook()
    for i in range(2):
        sheet = book.active if i == 0 else book.create_sheet('Second')
        sheet.append(['ID', 'Name', 'Amount', 'Empty', 'Mixed', 'Date', 'Active'])
        sheet.append(['000123', 'Семён Ёлкин', -12.75, None, 42, datetime.date(1993, 3, 30), True])
        sheet.append(['000007', 'Алёна Фёдорова', 0.125, None, '00123', datetime.date(2002, 8, 22), False])
        sheet.append(['000099', 'Ёжик ёлка ' + 'длинная строка ' * 20, 123.5, None, 'текст', None, None])
        sheet['C2'].number_format = sheet['C3'].number_format = sheet['C4'].number_format = '0.000'
    book.save(source)
    book.close()
    literals = args.output/'literal-headers.xlsx'
    book = Workbook()
    sheet = book.active
    names = ['ID,C,8', 'ID', 'Повтор', 'Повтор', 'Long field name', 'X,N,2']
    sheet.append(names)
    sheet.append(['000123', 'A', 'Ёлка', 'ёж', 'long', 42])
    sheet.append(['000007', 'B', 'Юлия', 'Иванов', 'text', '00042'])
    book.save(literals)
    book.close()
    originals = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in (source, literals)}
    with adapter.batch(args.output/'office'):
        def dbf_case(path):
            nonlocal counter
            exported = convert(path, 'dbf')
            dbfs = [p for p in exported if p.suffix == '.dbf']
            assert len(dbfs) == (1 if path.suffix == '.csv' else 2)
            facts = []
            for dbf in dbfs:
                raw = read_dbf(dbf)
                print('DBF_PHYSICAL='+json.dumps(raw, ensure_ascii=True), flush=True)
                assert dbf.with_suffix('.cpg') in exported
                assert dbf.with_suffix('.cpg').read_text().strip() == 'UTF-8'
                assert [r[0] for r in raw['rows']] == ['000123', '000007', '000099']
                assert [Decimal(r[2]) for r in raw['rows']] == [Decimal('-12.75'), Decimal('0.125'), Decimal('123.5')]
                assert [r[1] for r in raw['rows']] == ['Семён Ёлкин', 'Алёна Фёдорова', 'Ёжик ёлка ' + 'длинная строка ' * 20]
                assert [r[4] for r in raw['rows']] == ['42', '00123', 'текст']
                assert all(r[3] == '' for r in raw['rows'])
                if path.suffix != '.csv':
                    assert [f['type'] for f in raw['fields']] == ['C', 'M', 'N', 'C', 'M', 'D', 'L'], raw
                    assert [r[5] for r in raw['rows']] == ['19930330', '20020822', ''], raw
                    assert [r[6] for r in raw['rows'][:2]] == ['T', 'F'], raw
                counter += 1
                imported = convert_path(dbf, 'xlsx', args.output/('readback-%d' % counter),
                    executable=args.executable, import_filter='dBase:76')
                workbook = load_workbook(next(p for p in imported if p.suffix == '.xlsx'), data_only=True)
                rows = list(workbook.active.values)[1:]
                workbook.close()
                assert [r[0] for r in rows] == [r[0] for r in raw['rows']], rows
                assert [r[1] for r in rows] == [r[1] for r in raw['rows']], rows
                assert [Decimal(str(r[2])) for r in rows] == [Decimal(r[2]) for r in raw['rows']], rows
                assert [str(r[4]) for r in rows] == [r[4] for r in raw['rows']], rows
                facts.append({k: v for k, v in raw.items() if k != 'rows'})
            return {'outputs': facts, 'readers': ['independent_physical_reader', 'native_LO_dBase_76_to_XLSX']}
        run('xlsx_utf8_values_and_native_readback', lambda: dbf_case(source))
        # ODS and CSV are created by the same pinned native engine, then
        # exported back through A.I.O's real source admission/UNO route.
        for kind in ('ods', 'csv'):
            def generated_case(kind=kind):
                native = convert(source, kind)
                generated = next(p for p in native if p.suffix == '.'+kind)
                return dbf_case(generated)
            run(kind+'_utf8_values_and_native_readback', generated_case)
        def literal_case():
            exported = convert(literals, 'dbf')
            dbf = next(p for p in exported if p.suffix == '.dbf')
            facts = read_dbf(dbf)
            mappings = [p for p in exported if p.name.endswith('.aio-columns.jsonl')]
            assert mappings, ('literal_headers_have_no_mapping', facts)
            metadata = json.loads(mappings[0].read_text(encoding='utf-8').splitlines()[0])
            assert metadata['originalHeaders'] == names, metadata
            assert metadata.get('renamedFields'), metadata
            assert [r[0] for r in facts['rows']] == ['000123', '000007'], facts
            assert [r[5] for r in facts['rows']] == ['42', '00042'], facts
            assert len({f['name'].upper() for f in facts['fields']}) == len(names)
            return {'fields': facts['fields'], 'mapping': metadata}
        run('literal_duplicate_long_cyrillic_headers', literal_case)
        def literal_only():
            path = args.output/'literal-only.xlsx'
            book = Workbook()
            book.active.append(['ID,C,8'])
            book.active.append(['000123'])
            book.active.append(['000007'])
            book.save(path)
            book.close()
            # Establish provenance independently: the native exporter itself
            # interprets this literal XLSX heading as field metadata.
            native = convert_path(path, 'dbf', args.output/'raw-literal-dbase',
                executable=args.executable, convert_spec='dbf:dBase:76')
            native_facts = read_dbf(next(p for p in native if p.suffix == '.dbf'))
            assert native_facts['fields'][0]['name'] == 'ID', native_facts
            exported = convert(path, 'dbf')
            dbf = next(p for p in exported if p.suffix == '.dbf')
            facts = read_dbf(dbf)
            metadata = json.loads(next(p for p in exported if p.name.endswith('.aio-columns.jsonl'))
                                  .read_text(encoding='utf-8').splitlines()[0])
            assert metadata['originalHeaders'] == ['ID,C,8'], metadata
            assert metadata['fieldNames'] == [facts['fields'][0]['name']], (metadata, facts)
            assert [r[0] for r in facts['rows']] == ['000123', '000007']
            assert 'NATIVE_DBF_FIELD_NAME_MAPPING' in adapter.export_info()['warnings']
            return {'nativeField': native_facts['fields'][0], 'aioField': facts['fields'][0], 'mapping': metadata}
        run('literal_only_header_origin_and_mapping', literal_only)
        def vectors():
            from app_backend.converter_guard.native_metafile import repair_native_emf_counts
            source = args.output/'vectors.fodg'
            source.write_text(fixture('fodg'), encoding='utf-8')
            native = convert_path(source, 'emf', args.output/'raw-native-emf',
                executable=args.executable, convert_spec='emf:draw_emf_Export')
            raw_path = next(p for p in native if p.suffix == '.emf')
            raw_before = raw_path.read_bytes()
            native_before = emf_records(raw_before)
            before_render = gdi_render(raw_path) if os.name == 'nt' else None
            repairs = repair_native_emf_counts([raw_path])
            raw_after = raw_path.read_bytes()
            native_after = emf_records(raw_after)
            assert native_after['declaredRecords'] == native_after['physicalRecords']
            assert raw_before[:52] == raw_after[:52] and raw_before[56:] == raw_after[56:]
            if os.name == 'nt':
                assert before_render == gdi_render(raw_path), (native_before, native_after)
            plain = [p for p in convert(source, 'emf') if p.suffix == '.emf']
            compressed = [p for p in convert(source, 'emz') if p.suffix == '.emz']
            assert len(plain) == len(compressed) == 3
            findings = []
            for index, (emf, emz) in enumerate(zip(sorted(plain), sorted(compressed))):
                raw = emf.read_bytes()
                inner = gzip.decompress(emz.read_bytes())
                facts, gzip_facts = emf_records(raw), emf_records(inner)
                assert facts['declaredRecords'] == facts['physicalRecords'], facts
                assert gzip_facts['declaredRecords'] == gzip_facts['physicalRecords'], gzip_facts
                row = {'page': index+1, 'emf': facts, 'emz': gzip_facts}
                if os.name == 'nt':
                    render = gdi_render(emf)
                    assert render['gdiRecords'] == facts['physicalRecords'], (render, facts)
                    roundtrip = args.output/('gzip-page-%d.emf' % index)
                    roundtrip.write_bytes(inner)
                    gzip_render = gdi_render(roundtrip)
                    assert gzip_render['pixelSha256'] == render['pixelSha256'], (gzip_render, render)
                    row['gdi'] = render
                findings.append(row)
            return {'rawNativeBefore': native_before, 'rawNativeAfter': native_after,
                    'rawNativeGdiBeforeAfter': before_render, 'repairs': repairs, 'pages': findings}
        run('emf_emz_structure_gdi_and_gzip', vectors)
    def save_case():
        from app_backend.converter_guard import ConverterGuard
        guard = ConverterGuard(adapter=LibreOfficeAdapter(args.executable), recovery_parent=args.output/'owned')
        try:
            raw = source.read_bytes()
            job = guard.create([{'id': 'source', 'name': source.name, 'size': len(raw)}], ['dbf'])['jobId']
            guard.upload(job, 'source', io.BytesIO(raw), len(raw))
            guard.start(job)
            deadline = time.monotonic() + 180
            while not (state := guard.status(job))['finished']:
                assert time.monotonic() < deadline, state
                time.sleep(.05)
            assert state['state'] == 'ready', state
            ticket = state['result']['ticket']
            destination = args.output/'saved'
            destination.mkdir()
            saved = guard.save_folder(ticket, destination)
            folder = destination/saved['name']
            manifest = json.loads((folder/'AIO-results-manifest.json').read_text(encoding='utf-8'))
            inventory = {row['path']: row for row in manifest['outputs']}
            dbfs = [name for name in inventory if name.endswith('.dbf')]
            assert len(dbfs) == 2 and manifest['readbackVerified'], manifest
            for name in dbfs:
                assert str(Path(name).with_suffix('.cpg')).replace('\\', '/') in inventory, inventory
                assert str(Path(name).with_suffix('.dbt')).replace('\\', '/') in inventory, inventory
            for name, row in inventory.items():
                assert hashlib.sha256((folder/name).read_bytes()).hexdigest() == row['sha256'], row
            with guard.open_result(ticket) as (stream, name, size):
                assert name.endswith('.zip')
                with ZipFile(stream) as archive:
                    assert archive.testzip() is None
                    zipped = json.loads(archive.read('AIO-results-manifest.json'))
                    assert {r['path'] for r in zipped['outputs']} == set(inventory)
                    for row in zipped['outputs']:
                        assert hashlib.sha256(archive.read(row['path'])).hexdigest() == row['sha256'], row
                    report = json.loads(archive.read('AIO-conversion-report.json'))
                    warnings = report['files'][0]['targets'][0]['warnings']
                    assert warnings, report
            return {'dbfFiles': len(dbfs), 'outputs': len(inventory), 'folderReadback': True,
                    'zipReadback': True, 'reportsAndCompanions': True}
        finally:
            guard.close()
    run('real_native_guard_folder_zip_manifest_cpg_dbt_mapping', save_case)
    assert all(hashlib.sha256(p.read_bytes()).hexdigest() == sha for p, sha in originals.items())
    result = {'status': 'FAIL' if failures else 'PASS', 'cases': cases,
              'sourceUnchanged': True, 'historicalFULLTEST': 'NOT_RUN'}
    (args.output/'evidence.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    if failures:
        raise AssertionError(failures)


if __name__ == '__main__':
    main()
